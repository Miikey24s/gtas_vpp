from pathlib import Path
from shutil import copy2

from docx import Document


WORKING = Path(r"D:\WORK\gtas_vpp\LVTN\NguyenAnNam_DH52201078_working.docx")
CHECKPOINT = Path(r"D:\WORK\gtas_vpp\LVTN\checkpoints\01_chuong1.docx")


REPLACEMENTS = {
    "Trong hoạt động của doanh nghiệp, văn phòng phẩm là nhóm vật tư phục vụ trực tiếp cho công việc hằng ngày của các phòng ban. Nhu cầu sử dụng văn phòng phẩm thường phát sinh theo kỳ, với số lượng và danh mục khác nhau tùy theo đặc thù công việc của từng bộ phận. Khi số lượng người dùng tăng lên và phạm vi quản lý mở rộng, việc tiếp nhận yêu cầu thông qua biểu mẫu rời rạc, bảng tính thủ công hoặc trao đổi qua các kênh liên lạc riêng lẻ sẽ khiến cho công tác tổng hợp trở nên chậm trễ, khó kiểm soát và dễ phát sinh sai sót.":
        "Trong hoạt động của doanh nghiệp, văn phòng phẩm là nhóm vật tư phục vụ trực tiếp cho công việc hằng ngày của các phòng ban. Nhu cầu sử dụng thường phát sinh theo kỳ, với số lượng và danh mục khác nhau tùy theo đặc thù của từng bộ phận. Khi số lượng người dùng và phạm vi quản lý tăng, việc tiếp nhận yêu cầu bằng biểu mẫu rời rạc, bảng tính thủ công hoặc các kênh trao đổi riêng lẻ làm cho công tác tổng hợp trở nên chậm, khó kiểm soát và dễ phát sinh sai sót, trùng lặp hoặc thiếu thông tin.",
    "Quy trình tiếp nhận yêu cầu văn phòng phẩm không chỉ dừng lại ở việc ghi nhận số lượng cần dùng mà còn liên quan đến nhiều yếu tố như thời hạn gửi yêu cầu, trạng thái xử lý đơn, danh mục mặt hàng đang còn, bảng giá áp dụng theo từng thời điểm, quyền truy cập của từng nhóm người dùng và dữ liệu phòng ban/công ty của người yêu cầu.":
        "Quy trình tiếp nhận yêu cầu văn phòng phẩm không chỉ dừng lại ở việc ghi nhận số lượng cần dùng mà còn liên quan đến thời hạn gửi yêu cầu, trạng thái xử lý đơn, danh mục mặt hàng còn hiệu lực, bảng giá áp dụng theo từng thời điểm, quyền truy cập của từng nhóm người dùng và dữ liệu phòng ban, công ty của người yêu cầu.",
    "Bên cạnh đó, trong môi trường doanh nghiệp có nhiều phòng ban sử dụng chung danh mục văn phòng phẩm, yêu cầu đặt ra không chỉ là số hóa thao tác nhập liệu mà còn phải chuẩn hóa toàn bộ quy trình nghiệp vụ từ đăng nhập, tạo đơn, theo dõi lịch sử, phê duyệt đơn bổ sung, quản lý bảng giá đến tổng hợp báo cáo theo từng kỳ.":
        "Trong môi trường doanh nghiệp có nhiều phòng ban sử dụng chung một danh mục văn phòng phẩm, yêu cầu đặt ra không chỉ là số hóa thao tác nhập liệu mà còn phải chuẩn hóa chuỗi nghiệp vụ từ đăng nhập, tạo đơn, theo dõi lịch sử, phê duyệt đơn bổ sung, quản lý bảng giá đến tổng hợp dữ liệu nhu cầu theo từng kỳ.",
    "Xuất phát từ những yêu cầu trên, đề tài lựa chọn nghiên cứu và xây dựng website quản lý văn phòng phẩm Phong Phú theo mô hình tập trung, hỗ trợ xử lý nghiệp vụ theo kỳ, bảo đảm phân quyền người dùng, lưu lịch sử thao tác và tạo nền tảng cho việc triển khai trong môi trường doanh nghiệp.":
        "Xuất phát từ những yêu cầu trên, đề tài lựa chọn phân tích, thiết kế và xây dựng website quản lý văn phòng phẩm Phong Phú theo mô hình tập trung. Hệ thống hỗ trợ xử lý nghiệp vụ theo kỳ, phân quyền người dùng, lưu lịch sử thao tác và tạo nền tảng cho việc triển khai trong môi trường doanh nghiệp.",
    "Mục tiêu tổng quát của luận văn là phân tích, thiết kế và xây dựng hệ thống web hỗ trợ quản lý yêu cầu văn phòng phẩm theo kỳ tại doanh nghiệp, qua đó góp phần chuẩn hóa quy trình tiếp nhận, tổng hợp và xử lý nhu cầu sử dụng văn phòng phẩm giữa các phòng ban.":
        "Mục tiêu tổng quát của luận văn là phân tích, thiết kế, hiện thực và thử nghiệm hệ thống web hỗ trợ quản lý yêu cầu văn phòng phẩm theo kỳ tại doanh nghiệp. Qua đó, đề tài góp phần chuẩn hóa quy trình tiếp nhận, tổng hợp và xử lý nhu cầu sử dụng văn phòng phẩm giữa các phòng ban.",
    "Hệ thống áp dụng quy tắc thời hạn chốt vào 00 giờ 00 ngày 05 hằng tháng. Backend phải là nơi chịu trách nhiệm xác định kỳ hiện tại, kỳ vừa đóng và thời điểm hết hạn để tránh việc Frontend tính sai theo thời gian trên thiết bị người dùng.":
        "Hệ thống áp dụng thời điểm chốt lúc 00 giờ 00 ngày 05 hằng tháng. Backend là nguồn xác định thống nhất kỳ hiện tại, kỳ trước, thời hạn gửi đơn và trạng thái quá hạn; frontend chỉ hiển thị dữ liệu do backend cung cấp để tránh sai lệch do thời gian trên thiết bị người dùng.",
    "Đơn bổ sung được phát sinh sau thời hạn hoặc sau khi kỳ trước đã khép lại nên phải tuân theo quy trình phê duyệt gồm các trạng thái chờ duyệt, duyệt hoặc từ chối. Hệ thống cũng giới hạn số lượng đơn bổ sung và ngăn chặn việc tạo mới khi còn đơn đang chờ duyệt.":
        "Đơn bổ sung được tạo cho kỳ trước sau khi đã qua thời hạn gửi đơn thông thường nhưng trước khi kỳ được đóng. Đơn phải tuân theo luồng phê duyệt với các trạng thái chờ duyệt, đã duyệt hoặc từ chối. Hệ thống giới hạn tối đa ba đơn bổ sung cho mỗi người dùng trong một kỳ và không cho tạo đơn mới khi vẫn còn một đơn bổ sung đang chờ xử lý.",
    "Hệ thống cần kiểm soát quyền ở mức menu, trang, tab và component. Điều này giúp người dùng chỉ nhìn thấy và thao tác đúng phạm vi được cấp, đồng thời backend vẫn kiểm soát các API quan trọng bằng cơ chế phân quyền dựa trên chính sách – Policy-based Authorization.":
        "Hệ thống cần kiểm soát quyền ở mức menu, trang, tab và component để người dùng chỉ nhìn thấy và thao tác trong phạm vi được cấp. Đồng thời, backend kiểm soát các API quan trọng bằng cơ chế phân quyền dựa trên chính sách (policy-based authorization).",
    "Các thao tác tạo đơn, chỉnh sửa, hủy đơn, duyệt đơn bổ sung hoặc đóng kỳ đều ảnh hưởng trực tiếp đến dữ liệu nghiệp vụ. Hệ thống cần kiểm soát quy trình chuyển đổi trạng thái hợp lệ, sử dùng transaction nhằm đảm bảo toàn vẹn dữ liệu và lưu lại lịch sử thao tác.":
        "Các thao tác tạo đơn, chỉnh sửa, hủy đơn, duyệt đơn bổ sung và đóng kỳ đều ảnh hưởng trực tiếp đến dữ liệu nghiệp vụ. Hệ thống cần kiểm soát các chuyển đổi trạng thái hợp lệ, sử dụng transaction để bảo đảm tính toàn vẹn dữ liệu và lưu lại lịch sử thao tác phục vụ truy vết.",
    "Frontend và backend cần phối hợp ổn định, bảo đảm kết nối, xác thực và trao đổi dữ liệu thông suốt; đồng thời hệ thống phải tương thích và vận hành ổn định với Docker Compose, Nginx reverse proxy và SQL Server.":
        "Frontend và backend cần phối hợp ổn định, bảo đảm kết nối, xác thực và trao đổi dữ liệu thông suốt. Bên cạnh đó, cấu hình triển khai phải hỗ trợ Docker Compose, SQL Server và Nginx reverse proxy, đồng thời tách biệt thông tin nhạy cảm theo từng môi trường.",
    "- Xây dựng frontend bằng Blazor Server kết hợp Radzen Blazor, tổ chức giao diện theo các vùng chức năng: Dashboard – Trang chủ, Library – Thư viện, Permission – Phần quyền và Report – Báo cáo.":
        "- Xây dựng frontend bằng Blazor Server kết hợp Radzen Blazor, tổ chức giao diện theo các vùng chức năng Dashboard (nghiệp vụ đơn), Library (dữ liệu nền), Permission (phân quyền) và Report (khung báo cáo đang phát triển).",
    "Phạm vi của đề tài là hệ thống web nội bộ phục vụ quản lý yêu cầu văn phòng phẩm tại doanh nghiệp, không triển khai các chức năng mua bán trực tuyến công khai hoặc thanh toán điện tử. Hệ thống tập trung vào quy trình gửi yêu cầu - tổng hợp - phê duyệt yêu cầu bổ sung - đóng kỳ - quản lý danh mục, vật tư.":
        "Phạm vi của đề tài là hệ thống web nội bộ phục vụ quản lý yêu cầu văn phòng phẩm tại doanh nghiệp. Hệ thống tập trung vào quy trình gửi yêu cầu, tổng hợp nhu cầu, phê duyệt đơn bổ sung, đóng kỳ và quản lý dữ liệu nền. Đề tài không triển khai website bán hàng công khai, thanh toán điện tử, quản lý kho, phát hành đơn mua hàng hoặc tích hợp chính thức với hệ thống nhân sự của doanh nghiệp.",
}


def replace_paragraph_text(paragraph, new_text):
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(new_text)


def main():
    if not WORKING.exists():
        raise FileNotFoundError(WORKING)

    doc = Document(WORKING)
    changed = []
    for paragraph in doc.paragraphs:
        old = " ".join(paragraph.text.split())
        if old in REPLACEMENTS:
            replace_paragraph_text(paragraph, REPLACEMENTS[old])
            changed.append(old[:80])

    expected = len(REPLACEMENTS)
    if len(changed) != expected:
        missing = [key for key in REPLACEMENTS if key[:80] not in changed]
        raise RuntimeError(f"Expected {expected} replacements, made {len(changed)}; missing={missing}")

    chapter2_index = next(
        i for i, p in enumerate(doc.paragraphs)
        if " ".join(p.text.split()) == "Chương 2. PHƯƠNG PHÁP THỰC HIỆN"
    )
    intro = doc.paragraphs[chapter2_index - 1]
    if not intro.text.strip():
        replace_paragraph_text(
            intro,
            "Các kết quả cần đạt và tiêu chí đánh giá tương ứng được trình bày trong bảng sau:",
        )

    table = doc.tables[0]
    table.cell(3, 1).text = "Tổng hợp nhu cầu và đóng kỳ"
    table.cell(3, 2).text = (
        "Hệ thống tổng hợp dữ liệu theo cá nhân, phòng ban và toàn doanh nghiệp; "
        "đóng kỳ bằng bảng giá hợp lệ và chụp cố định đơn giá vào chi tiết đơn."
    )
    table.cell(5, 2).text = (
        "Có kiểm thử backend, frontend helper, UI test Playwright, báo cáo audit giao diện "
        "và cấu hình Docker Compose/Nginx phục vụ triển khai."
    )

    doc.save(WORKING)
    copy2(WORKING, CHECKPOINT)
    print(f"Updated {len(changed)} paragraphs")
    print(f"Working: {WORKING}")
    print(f"Checkpoint: {CHECKPOINT}")


if __name__ == "__main__":
    main()

