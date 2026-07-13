from __future__ import annotations

import copy
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt
from docx.text.paragraph import Paragraph


SOURCE = Path(r"D:\WORK\gtas_vpp\LVTN\checkpoints\04_chuong4_thunghiem.docx")
OUTPUT = Path(r"D:\WORK\gtas_vpp\LVTN\checkpoints\05_chuong5_phuluc_tltk.docx")
WORKING = Path(r"D:\WORK\gtas_vpp\LVTN\NguyenAnNam_DH52201078_working.docx")

TABLE_WIDTH = 9000
NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def paragraph_text(node) -> str:
    return "".join(text_node.text or "" for text_node in node.iter(qn("w:t"))).strip()


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row) -> None:
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def set_cell_margins(cell, top=75, start=95, bottom=75, end=95) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_table_geometry(table, widths: list[int]) -> None:
    assert sum(widths) == TABLE_WIDTH
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def format_run(run, size: float, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        r_fonts.set(qn(attr), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def format_cell(cell, size: float, bold: bool = False, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            format_run(run, size=size, bold=bold)


def add_paragraph(
    doc,
    marker,
    text: str = "",
    style: str = "Normal",
    keep_with_next: bool = False,
    page_break_before: bool = False,
):
    paragraph = doc.add_paragraph(style=style)
    if text:
        paragraph.add_run(text)
    paragraph.paragraph_format.keep_with_next = keep_with_next
    paragraph.paragraph_format.page_break_before = page_break_before
    marker.addprevious(paragraph._p)
    return paragraph


def add_labeled_paragraph(doc, marker, label: str, text: str, indent: bool = True):
    paragraph = doc.add_paragraph(style="Normal")
    if indent:
        paragraph.paragraph_format.left_indent = Pt(18)
        paragraph.paragraph_format.first_line_indent = Pt(-18)
    paragraph.paragraph_format.space_after = Pt(3)
    label_run = paragraph.add_run(label + " ")
    label_run.bold = True
    paragraph.add_run(text)
    marker.addprevious(paragraph._p)
    return paragraph


def add_step(doc, marker, number: int, text: str):
    return add_labeled_paragraph(doc, marker, f"Bước {number}.", text)


def add_caption(doc, marker, number: str, title: str):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    prefix = paragraph.add_run(f"Bảng {number}:")
    format_run(prefix, 11, bold=True, italic=True)
    suffix = paragraph.add_run(f" {title}")
    format_run(suffix, 11, italic=True)
    marker.addprevious(paragraph._p)
    return paragraph


def add_table(doc, marker, headers, rows, widths, font_size=10.0, center_columns=()):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header = table.rows[0]
    for index, value in enumerate(headers):
        header.cells[index].text = value
        set_cell_shading(header.cells[index], "E7E6E6")
        format_cell(header.cells[index], font_size, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_header(header)
    prevent_row_split(header)

    for values in rows:
        row = table.add_row()
        for index, value in enumerate(values):
            row.cells[index].text = value
            alignment = WD_ALIGN_PARAGRAPH.CENTER if index in center_columns else WD_ALIGN_PARAGRAPH.LEFT
            format_cell(row.cells[index], font_size, align=alignment)
        prevent_row_split(row)

    set_table_geometry(table, widths)
    marker.addprevious(table._tbl)
    return table


def next_section_break(body, start_node):
    started = False
    for node in body.iterchildren():
        if node is start_node:
            started = True
            continue
        if started and node.tag == qn("w:p"):
            sect = node.find("./w:pPr/w:sectPr", namespaces={"w": NS_W})
            if sect is not None:
                return node
    raise RuntimeError("Section break not found")


def clear_between(body, start_node, end_node) -> None:
    deleting = False
    for node in list(body.iterchildren()):
        if node is start_node:
            deleting = True
            continue
        if node is end_node:
            break
        if deleting:
            body.remove(node)


def make_text_run(text: str, rpr=None):
    run = OxmlElement("w:r")
    if rpr is not None:
        run.append(copy.deepcopy(rpr))
    text_node = OxmlElement("w:t")
    if text.startswith(" ") or text.endswith(" "):
        text_node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    text_node.text = text
    run.append(text_node)
    return run


def apply_link_style(run_element, underline=True) -> None:
    rpr = run_element.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        run_element.insert(0, rpr)
    color = rpr.find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        rpr.append(color)
    color.set(qn("w:val"), "000000")
    underline_node = rpr.find(qn("w:u"))
    if underline_node is None:
        underline_node = OxmlElement("w:u")
        rpr.append(underline_node)
    underline_node.set(qn("w:val"), "single" if underline else "none")


def internal_hyperlink(text: str, anchor: str, rpr=None):
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    hyperlink.set(qn("w:history"), "1")
    run = make_text_run(text, rpr)
    apply_link_style(run, underline=True)
    hyperlink.append(run)
    return hyperlink


def external_hyperlink(paragraph, text: str, url: str):
    rel_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = make_text_run(text)
    apply_link_style(run, underline=True)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def max_bookmark_id(doc) -> int:
    ids = []
    for node in doc.element.xpath(".//w:bookmarkStart"):
        value = node.get(qn("w:id"))
        if value and value.isdigit():
            ids.append(int(value))
    return max(ids, default=0)


def bookmark_paragraph(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    ppr = paragraph._p.find(qn("w:pPr"))
    insert_at = 1 if ppr is not None else 0
    paragraph._p.insert(insert_at, start)
    paragraph._p.append(end)


def add_reference(doc, marker, number: int, author: str, title: str, source: str, url: str | None, bookmark_id: int, page_break_before: bool = False):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.left_indent = Pt(28)
    paragraph.paragraph_format.first_line_indent = Pt(-28)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.page_break_before = page_break_before
    paragraph.add_run(f"[{number}] {author}, “{title},” {source}.")
    if url:
        paragraph.add_run(" [Trực tuyến]. Địa chỉ: ")
        external_hyperlink(paragraph, url, url)
        paragraph.add_run(". [Truy cập: 13/07/2026].")
    bookmark_paragraph(paragraph, f"ref_{number:02d}", bookmark_id)
    marker.addprevious(paragraph._p)
    return paragraph


def rebuild_citation_paragraph(paragraph, before: str, numbers: list[int], after: str) -> None:
    ppr = paragraph._p.find(qn("w:pPr"))
    for child in list(paragraph._p):
        if child is not ppr:
            paragraph._p.remove(child)
    paragraph._p.append(make_text_run(before))
    for index, number in enumerate(numbers):
        if index:
            paragraph._p.append(make_text_run(", "))
        paragraph._p.append(internal_hyperlink(f"[{number}]", f"ref_{number:02d}"))
    paragraph._p.append(make_text_run(after))


def convert_existing_citation_links(doc) -> int:
    changed = 0
    for hyperlink in doc.element.xpath(".//w:hyperlink[@r:id]"):
        text = "".join(hyperlink.xpath(".//w:t/text()"))
        match = re.fullmatch(r"\[(\d+)\]", text.strip())
        if not match:
            continue
        number = int(match.group(1))
        if not 1 <= number <= 15:
            continue
        hyperlink.attrib.pop(qn("r:id"), None)
        hyperlink.set(qn("w:anchor"), f"ref_{number:02d}")
        hyperlink.set(qn("w:history"), "1")
        for run in hyperlink.findall(qn("w:r")):
            apply_link_style(run, underline=True)
        changed += 1
    return changed


def remove_orphan_hyperlink_relationships(doc) -> int:
    used = {h.get(qn("r:id")) for h in doc.element.xpath(".//w:hyperlink[@r:id]")}
    removed = 0
    for rel_id, rel in list(doc.part.rels.items()):
        if rel.reltype == RT.HYPERLINK and rel_id not in used:
            doc.part.drop_rel(rel_id)
            removed += 1
    return removed


def compact_appendix(doc, body, appendix_node, appendix_end_node) -> None:
    active = False
    for node in body.iterchildren():
        if node is appendix_node:
            active = True
            continue
        if node is appendix_end_node:
            break
        if not active or node.tag != qn("w:p"):
            continue
        paragraph = Paragraph(node, doc._body)
        if paragraph.style.name == "Normal":
            paragraph.paragraph_format.line_spacing = 1.15
            paragraph.paragraph_format.space_after = Pt(1.5)
        elif paragraph.style.name == "Heading 3":
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.keep_with_next = True


def compact_toc_styles(doc) -> None:
    settings = {
        "toc 1": (13.0, 1.0, 0.0),
        "toc 2": (12.0, 1.0, 0.0),
        "toc 3": (11.0, 1.0, 0.0),
    }
    for style_name, (size, line_spacing, space_after) in settings.items():
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        r_fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
            r_fonts.set(qn(attr), "Times New Roman")
        style.paragraph_format.line_spacing = line_spacing
        style.paragraph_format.space_after = Pt(space_after)


doc = Document(SOURCE)
body = doc.element.body
children = list(body.iterchildren())

chapter5 = next(node for node in children if node.tag == qn("w:p") and paragraph_text(node) == "CHƯƠNG 5. KẾT LUẬN")
chapter5_end = next_section_break(body, chapter5)
clear_between(body, chapter5, chapter5_end)
marker = chapter5_end

add_paragraph(
    doc,
    marker,
    "Chương này tổng hợp mức độ hoàn thành của đề tài trên cơ sở các mục tiêu đã nêu ở Chương 1, kết quả hiện thực trong Chương 3 và bằng chứng kiểm thử ở Chương 4. Việc đánh giá tách rõ chức năng đã có, phần mới đạt ở mức đề tài và các nội dung chưa hoàn thiện để tránh đồng nhất bản thử nghiệm với một hệ thống đã sẵn sàng vận hành chính thức.",
)
add_paragraph(doc, marker, "5.1 KẾT QUẢ ĐỐI CHIẾU VỚI MỤC TIÊU", "Heading 2", keep_with_next=True)
add_paragraph(
    doc,
    marker,
    "Sáu mục tiêu chức năng trong mục 1.1.2 và yêu cầu kiểm thử, triển khai trong mục 1.4 được đối chiếu trực tiếp với chức năng đang có trong source. Kết quả tổng hợp được trình bày ở Bảng 5-1.",
)
add_caption(doc, marker, "5-1", "Đối chiếu kết quả thực hiện với mục tiêu đề tài")
add_table(
    doc,
    marker,
    ["STT", "Mục tiêu/tiêu chí", "Kết quả và bằng chứng", "Đánh giá"],
    [
        ["1", "Đăng nhập và sử dụng chức năng theo phân quyền.", "Đã có xác thực, nhóm quyền, ánh xạ trang-component và policy authorization cho các API quan trọng.", "Đạt"],
        ["2", "Tạo, chỉnh sửa, hủy và theo dõi đơn yêu cầu định kỳ.", "Đã có luồng tạo đơn thường, sao chép dữ liệu kỳ trước, sửa/hủy ở trạng thái hợp lệ, lịch sử và quy tắc kỳ theo ngày 5.", "Đạt"],
        ["3", "Tạo và xử lý đơn bổ sung.", "Đơn bổ sung dành cho kỳ trước, bắt buộc lý do, tối đa ba đơn mỗi người dùng/kỳ, không tạo thêm khi còn Pending; quản trị viên có thể duyệt hoặc từ chối.", "Đạt"],
        ["4", "Quản lý tập trung danh mục và bảng giá.", "Đã có nhóm chức năng Library cho danh mục, vật tư, đơn vị tính, nhà cung cấp, phòng ban, giá nhà cung cấp và bảng giá L07; hỗ trợ sao chép và đặt bảng giá mặc định.", "Đạt"],
        ["5", "Tổng hợp nhu cầu theo cá nhân, phòng ban và toàn doanh nghiệp.", "Các tab tổng hợp trên Dashboard đã hiển thị số đơn, số dòng hàng, tổng số lượng và tổng tiền theo phạm vi quyền. Trang Report riêng và xuất file chưa hoàn thiện.", "Đạt một phần"],
        ["6", "Đóng kỳ bằng bảng giá hợp lệ và lưu vết xử lý.", "PeriodSettlementService kiểm tra đơn Pending/giá thiếu, chụp đơn giá vào chi tiết, ghi thời điểm, người đóng kỳ và log nghiệp vụ.", "Đạt"],
        ["7", "Kiểm thử và chuẩn bị triển khai.", "Backend 132/132 test và frontend helper 26/26 test đạt ngày 13/07/2026; solution và UI-test project build thành công; có Docker Compose/Nginx. Playwright E2E chưa chạy lại trong đợt hiện tại.", "Đạt ở mức đề tài"],
    ],
    [600, 2200, 4750, 1450],
    font_size=10.0,
    center_columns=(0, 3),
)
add_paragraph(
    doc,
    marker,
    "Như vậy, các luồng nghiệp vụ cốt lõi về đơn yêu cầu, đơn bổ sung, dữ liệu nền, phân quyền và đóng kỳ đã được hiện thực. Phần tổng hợp đã sử dụng được trong Dashboard nhưng mục tiêu báo cáo chỉ đạt một phần vì trang Report vẫn là giao diện chờ và chưa có xuất Excel/PDF. Kết quả kiểm thử hiện tại đủ chứng minh mức độ hoàn thành của đề tài, chưa thay thế đánh giá vận hành trong môi trường doanh nghiệp.",
)

add_paragraph(doc, marker, "5.2 CÁC VẤN ĐỀ CÒN TỒN ĐỌNG", "Heading 2", keep_with_next=True)
add_labeled_paragraph(doc, marker, "Báo cáo và xuất dữ liệu:", "Trang /report mới hiển thị số liệu chờ, nút xuất báo cáo đang bị vô hiệu hóa. Các tổng hợp thực tế hiện nằm tại Dashboard; chưa có biểu mẫu Excel/PDF theo chuẩn doanh nghiệp.")
add_labeled_paragraph(doc, marker, "Bằng chứng kiểm thử giao diện:", "Project Playwright build thành công nhưng chưa chạy E2E ngày 13/07/2026 vì cần frontend, backend, tài khoản và dữ liệu test đồng bộ. Kết quả 54/54 lượt audit là snapshot ngày 28/05/2026, không phải lần chạy hiện tại.")
add_labeled_paragraph(doc, marker, "Hiệu năng và môi trường dữ liệu:", "Kiểm thử backend chủ yếu dùng EF Core InMemory/SQLite; chưa có kiểm thử tải dài hạn và bộ integration test đầy đủ trên SQL Server với dữ liệu gần quy mô thực tế.")
add_labeled_paragraph(doc, marker, "An toàn phụ thuộc:", "Quá trình restore backend test còn cảnh báo NU1903 đối với Microsoft.OpenApi 2.4.1 và SQLitePCLRaw.lib.e_sqlite3 2.1.11. Cần nâng cấp phiên bản tương thích và chạy lại toàn bộ test trước khi triển khai.")
add_labeled_paragraph(doc, marker, "Tích hợp và thông báo:", "Dữ liệu người dùng, công ty, phòng ban vẫn nằm trong phạm vi đề tài; chưa tích hợp chính thức với hệ thống nhân sự/danh bạ nội bộ và chưa có email hoặc notification chủ động cho hạn gửi đơn, kết quả duyệt và đóng kỳ.")
add_labeled_paragraph(doc, marker, "Quy trình nâng cao:", "Phê duyệt nhiều cấp, hạn mức ngân sách theo phòng ban, so sánh nhà cung cấp và quản lý kho không thuộc phạm vi hiện tại.")

add_paragraph(doc, marker, "5.3 HƯỚNG PHÁT TRIỂN", "Heading 2", keep_with_next=True)
add_labeled_paragraph(doc, marker, "Hoàn thiện báo cáo:", "Kết nối trang Report với dữ liệu tổng hợp, bổ sung bộ lọc kỳ/phòng ban/trạng thái và xuất Excel/PDF cho tổng hợp kỳ, phòng ban, toàn doanh nghiệp và danh sách mặt hàng cần mua.")
add_labeled_paragraph(doc, marker, "Nâng độ tin cậy kiểm thử:", "Xây dựng môi trường test cố định, chạy Playwright E2E và UI audit trong CI; bổ sung integration test trên SQL Server, kiểm thử tải, dữ liệu lớn và các tình huống đồng thời.")
add_labeled_paragraph(doc, marker, "Gia cố bảo mật và vận hành:", "Nâng cấp các package có cảnh báo, quản lý secrets theo môi trường, bổ sung giám sát log/cảnh báo, quy trình sao lưu và hướng dẫn rollback khi migration gặp lỗi.")
add_labeled_paragraph(doc, marker, "Tích hợp và mở rộng nghiệp vụ:", "Đồng bộ người dùng/phòng ban với hệ thống nhân sự, bổ sung thông báo; sau đó phát triển phân tích chi phí, phê duyệt nhiều cấp, hạn mức ngân sách và so sánh giá theo nhu cầu thực tế.")
add_paragraph(doc, marker, "", "Normal")


children = list(body.iterchildren())
appendix = next(node for node in children if node.tag == qn("w:p") and paragraph_text(node) == "PHỤ LỤC")
appendix_end = next_section_break(body, appendix)
clear_between(body, appendix, appendix_end)
marker = appendix_end

add_paragraph(doc, marker, "PHỤ LỤC A. HƯỚNG DẪN DÀNH CHO NHÂN VIÊN", "Heading 2", keep_with_next=True)
add_paragraph(doc, marker, "A.1 Điều kiện sử dụng và đăng nhập", "Heading 3", keep_with_next=True)
add_labeled_paragraph(doc, marker, "Điều kiện:", "Người dùng có tài khoản hợp lệ, được gán phòng ban/công ty và có quyền REQUEST_ORDER. Các tab lịch sử, danh mục hoặc tổng hợp chỉ xuất hiện khi tài khoản có quyền tương ứng.")
add_step(doc, marker, 1, "Truy cập website GTAS VPP và chọn đúng môi trường được cấu hình.")
add_step(doc, marker, 2, "Nhập tên đăng nhập, mật khẩu và thực hiện đăng nhập. Nếu thông tin sai hoặc tài khoản không có quyền, hệ thống hiển thị thông báo và không mở chức năng bị hạn chế.")
add_step(doc, marker, 3, "Tại Dashboard, kiểm tra kỳ hiện tại, thời hạn gửi đơn và các tab được cấp cho tài khoản.")

add_paragraph(doc, marker, "A.2 Tạo đơn yêu cầu thông thường", "Heading 3", keep_with_next=True)
add_step(doc, marker, 1, "Tại Dashboard, chọn tạo đơn mới. Hệ thống lấy kỳ hiện tại từ backend; người dùng không tự chọn một kỳ bất kỳ.")
add_step(doc, marker, 2, "Ở bước Chọn mặt hàng, tìm theo mã hoặc tên, thêm mặt hàng và nhập số lượng lớn hơn 0. Không thêm trùng cùng một mặt hàng trong đơn.")
add_step(doc, marker, 3, "Chuyển sang bước Kiểm tra và gửi, rà lại danh sách chi tiết, số lượng và ghi chú. Đơn phải có ít nhất một mặt hàng hợp lệ.")
add_step(doc, marker, 4, "Nhấn gửi đơn. Đơn thường hợp lệ được lưu ở trạng thái Submitted và xuất hiện trong danh sách đơn của tôi.")
add_labeled_paragraph(doc, marker, "Lưu ý:", "Mỗi người dùng chỉ có tối đa một đơn thường trong một kỳ. Khi đã quá hạn hoặc kỳ đã đóng, backend từ chối tạo đơn mới dù giao diện vẫn đang mở từ trước.")

add_paragraph(doc, marker, "A.3 Sao chép, chỉnh sửa, hủy và xem lịch sử", "Heading 3", keep_with_next=True)
add_labeled_paragraph(doc, marker, "Sao chép kỳ trước:", "Chọn chức năng sao chép để nạp lại danh sách mặt hàng từ đơn gần nhất, sau đó kiểm tra số lượng và gửi như một đơn mới của kỳ hiện tại.")
add_labeled_paragraph(doc, marker, "Chỉnh sửa:", "Mở đơn ở trạng thái Submitted hoặc Pending, thay đổi chi tiết hợp lệ và lưu lại. Đơn Approved, Rejected hoặc Cancelled không được chỉnh sửa.")
add_labeled_paragraph(doc, marker, "Hủy đơn:", "Chọn hủy và xác nhận. Đơn Submitted hoặc Pending được chuyển sang Cancelled; thao tác không xóa vật lý bản ghi khỏi lịch sử.")
add_labeled_paragraph(doc, marker, "Xem lịch sử:", "Dùng bộ lọc năm, tháng và trạng thái để tra cứu đơn; mở dòng đơn để xem chi tiết mặt hàng, số lượng và thông tin xử lý.")

add_paragraph(doc, marker, "A.4 Tạo đơn bổ sung", "Heading 3", keep_with_next=True)
add_step(doc, marker, 1, "Chỉ tạo đơn bổ sung cho kỳ trước sau khi đã qua thời hạn đơn thường và trước khi kỳ đó được đóng.")
add_step(doc, marker, 2, "Chọn tạo đơn bổ sung, thêm mặt hàng, nhập số lượng và bắt buộc nêu lý do phát sinh.")
add_step(doc, marker, 3, "Gửi đơn. Đơn bổ sung được lưu ở trạng thái Pending để chờ quản trị viên duyệt hoặc từ chối.")
add_labeled_paragraph(doc, marker, "Giới hạn:", "Mỗi người dùng có tối đa ba đơn bổ sung trong một kỳ và không thể tạo đơn bổ sung mới khi vẫn còn một đơn Pending. Kỳ đã đóng không nhận thêm đơn.")

add_paragraph(doc, marker, "PHỤ LỤC B. HƯỚNG DẪN DÀNH CHO QUẢN TRỊ VIÊN", "Heading 2", keep_with_next=True)
add_paragraph(doc, marker, "B.1 Duyệt hoặc từ chối đơn bổ sung", "Heading 3", keep_with_next=True)
add_labeled_paragraph(doc, marker, "Điều kiện:", "Tài khoản có quyền REQUEST_ADMIN_APPROVAL.")
add_step(doc, marker, 1, "Mở Dashboard và chọn danh sách đơn bổ sung đang chờ duyệt.")
add_step(doc, marker, 2, "Mở chi tiết từng đơn, kiểm tra kỳ, người yêu cầu, mặt hàng, số lượng và lý do phát sinh.")
add_step(doc, marker, 3, "Chọn Duyệt để chuyển Pending thành Approved; hoặc chọn Từ chối, nhập lý do khi cần và chuyển đơn thành Rejected.")
add_labeled_paragraph(doc, marker, "Kiểm soát trạng thái:", "Chỉ đơn Pending được duyệt hoặc từ chối. Hệ thống không cho thực hiện lại thao tác trên đơn đã xử lý.")

add_paragraph(doc, marker, "B.2 Đóng kỳ", "Heading 3", keep_with_next=True)
add_labeled_paragraph(doc, marker, "Điều kiện:", "Tài khoản có quyền PERIOD_SETTLE; kỳ chưa đóng; không còn đơn bổ sung Pending; bảng giá được chọn hoặc bảng giá mặc định có đủ giá cho tất cả mặt hàng phát sinh.")
add_step(doc, marker, 1, "Mở vùng thao tác kỳ, chọn kỳ cần đóng và bảng giá áp dụng. Có thể để trống lựa chọn để hệ thống dùng bảng giá mặc định.")
add_step(doc, marker, 2, "Tải lại dữ liệu và kiểm tra số đơn, số dòng hàng, tổng số lượng, tổng tiền cùng cảnh báo giá thiếu hoặc đơn Pending.")
add_step(doc, marker, 3, "Khi nút Đóng kỳ được bật, thực hiện đóng kỳ và xác nhận. Hệ thống chụp đơn giá vào chi tiết đơn, ghi thời điểm, người thực hiện và bảng giá đã dùng.")
add_step(doc, marker, 4, "Mở lại tổng hợp kỳ để đối chiếu kết quả. Lần gọi đóng kỳ sau phải giữ kết quả nhất quán, không chụp giá lặp.")
add_labeled_paragraph(doc, marker, "Khi không đóng được:", "Xử lý hết đơn Pending; kiểm tra bảng giá mặc định/lựa chọn và bổ sung giá còn thiếu; tải lại trạng thái kỳ rồi thao tác lại. Không sửa trực tiếp dữ liệu đã chụp giá trong cơ sở dữ liệu.")

add_paragraph(doc, marker, "B.3 Kiểm tra tổng hợp và giới hạn báo cáo", "Heading 3", keep_with_next=True)
add_labeled_paragraph(doc, marker, "Tổng hợp phòng ban:", "Chọn năm, tháng và trạng thái để đối chiếu các đơn thuộc phòng ban trong phạm vi quyền.")
add_labeled_paragraph(doc, marker, "Tổng hợp toàn doanh nghiệp:", "Tài khoản có REQUEST_ALL_ORDERS_SUMMARY có thể xem số đơn, số dòng hàng, tổng số lượng, tổng tiền và chi tiết theo bộ lọc.")
add_paragraph(doc, marker, "", "Normal")
compact_appendix(doc, body, appendix, appendix_end)


children = list(body.iterchildren())
references = next(node for node in children if node.tag == qn("w:p") and paragraph_text(node) == "TÀI LIỆU THAM KHẢO")
final_sectpr = body.sectPr
clear_between(body, references, final_sectpr)
marker = final_sectpr

reference_data = [
    (1, "Microsoft", "ASP.NET Core documentation", "Microsoft Learn", "https://learn.microsoft.com/en-us/aspnet/core/?view=aspnetcore-10.0"),
    (2, "Microsoft", "Entity Framework Core documentation", "Microsoft Learn", "https://learn.microsoft.com/en-us/ef/core/"),
    (3, "Microsoft", "SQL Server technical documentation", "Microsoft Learn", "https://learn.microsoft.com/en-us/sql/sql-server/?view=sql-server-ver17"),
    (4, "Microsoft", "ASP.NET Core Blazor", "Microsoft Learn", "https://learn.microsoft.com/en-us/aspnet/core/blazor/?view=aspnetcore-10.0"),
    (5, "Radzen", "Radzen Blazor Components", "Radzen", "https://blazor.radzen.com/"),
    (6, "Mapster Project", "Mapster", "GitHub", "https://github.com/MapsterMapper/Mapster"),
    (7, "Serilog Project", "Serilog - structured logging for .NET", "Serilog", "https://serilog.net/"),
    (8, "Docker", "Docker Compose", "Docker Docs", "https://docs.docker.com/compose/"),
    (9, "Microsoft", "Playwright for .NET", "Playwright", "https://playwright.dev/dotnet/"),
    (10, "Nguyễn An Nam", "Mã nguồn đề tài GTAS VPP", "workspace luận văn gồm backend, frontend, shared DTO, test và hồ sơ audit, 2026", None),
    (11, "Google Workspace Learning Center", "Share & collaborate on a spreadsheet", "Google Support", "https://support.google.com/a/users/answer/13309904?hl=en"),
    (12, "Microsoft Support", "Basic tasks in Excel for the web", "Microsoft Support", "https://support.microsoft.com/en-US/Excel/basic-tasks-in-excel-for-the-web"),
    (13, "Odoo", "Purchase", "Odoo 18.0 documentation", "https://www.odoo.com/documentation/18.0/applications/inventory_and_mrp/purchase.html"),
    (14, "Atlassian", "Get started with service requests in Jira Service Management", "Atlassian", "https://www.atlassian.com/software/jira/service-management/product-guide/getting-started/service-request-management"),
    (15, "Zoho Creator", "Approval workflows", "Zoho Creator", "https://www.zoho.com/creator/approval-workflow/"),
]

bookmark_id = max_bookmark_id(doc) + 1
for number, author, title, source, url in reference_data:
    add_reference(doc, marker, number, author, title, source, url, bookmark_id, page_break_before=(number == 10))
    bookmark_id += 1


for paragraph in doc.paragraphs:
    text = paragraph.text
    if text.startswith("Tài liệu chính thức cho thấy Google Sheets"):
        before, after = text.split("[11]-[15]")
        rebuild_citation_paragraph(paragraph, before, [11, 12, 13, 14, 15], after)
    elif text.startswith("Hệ thống GTAS VPP được hiện thực trên nền tảng .NET 10"):
        before, after = text.split("[1]-[9]")
        rebuild_citation_paragraph(paragraph, before, list(range(1, 10)), after)

converted_citations = convert_existing_citation_links(doc)
removed_hyperlinks = remove_orphan_hyperlink_relationships(doc)
compact_toc_styles(doc)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
shutil.copy2(OUTPUT, WORKING)
print(f"saved={OUTPUT}")
print(f"converted_citation_links={converted_citations}")
print(f"removed_orphan_hyperlink_relationships={removed_hyperlinks}")
