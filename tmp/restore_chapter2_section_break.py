from copy import deepcopy
from pathlib import Path
from shutil import copy2

from docx import Document


ORIGINAL = Path(r"D:\WORK\gtas_vpp\LVTN\NguyenAnNam_DH52201078.docx")
WORKING = Path(r"D:\WORK\gtas_vpp\LVTN\NguyenAnNam_DH52201078_working.docx")
CHECKPOINT = Path(r"D:\WORK\gtas_vpp\LVTN\checkpoints\01_chuong1.docx")
CHAPTER2 = "Chương 2. PHƯƠNG PHÁP THỰC HIỆN"


def find_paragraph(doc, text):
    return next(p for p in doc.paragraphs if " ".join(p.text.split()) == text)


def main():
    original = Document(ORIGINAL)
    source_heading = find_paragraph(original, CHAPTER2)
    source_break = source_heading._p.getprevious()
    if source_break is None or not source_break.xpath("./w:pPr/w:sectPr"):
        raise RuntimeError("Original paragraph before Chapter 2 does not contain a section break.")

    working = Document(WORKING)
    target_heading = find_paragraph(working, CHAPTER2)
    previous = target_heading._p.getprevious()
    if previous is not None and previous.xpath("./w:pPr/w:sectPr"):
        print("Section break already exists; no insertion required.")
    else:
        target_heading._p.addprevious(deepcopy(source_break))
        print("Restored the original section-break paragraph before Chapter 2.")

    working.save(WORKING)
    copy2(WORKING, CHECKPOINT)


if __name__ == "__main__":
    main()

