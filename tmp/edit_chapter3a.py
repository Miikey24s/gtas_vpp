from copy import deepcopy
from pathlib import Path
from shutil import copyfile
from zipfile import ZIP_DEFLATED, ZipFile
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree


ROOT = Path(r"D:\WORK\gtas_vpp")
INPUT = ROOT / "LVTN" / "checkpoints" / "02C_chuong2_sodotongquat.docx"
OUTPUT = ROOT / "LVTN" / "checkpoints" / "03A_chuong3_mohinhdulieu.docx"
WORKING = ROOT / "LVTN" / "NguyenAnNam_DH52201078_working.docx"
DIAGRAM_DIR = ROOT / "LVTN" / "diagrams" / "ch03"

DIAGRAMS = [
    {
        "key": "conceptual",
        "png": DIAGRAM_DIR / "data-conceptual.png",
        "svg": DIAGRAM_DIR / "data-conceptual.svg",
        "media": "data-conceptual.svg",
        "caption": "Hình 3-1: Mô hình dữ liệu ý niệm của hệ thống GTAS VPP",
        "alt": "Mô hình dữ liệu ý niệm của hệ thống GTAS VPP",
        "width": 6.20,
    },
    {
        "key": "authorization",
        "png": DIAGRAM_DIR / "erd-authorization.png",
        "svg": DIAGRAM_DIR / "erd-authorization.svg",
        "media": "erd-authorization.svg",
        "caption": "Hình 3-2: ERD luận lý nhóm tổ chức và phân quyền",
        "alt": "ERD luận lý nhóm tổ chức và phân quyền",
        "width": 6.15,
    },
    {
        "key": "catalog",
        "png": DIAGRAM_DIR / "erd-catalog-pricing.png",
        "svg": DIAGRAM_DIR / "erd-catalog-pricing.svg",
        "media": "erd-catalog-pricing.svg",
        "caption": "Hình 3-3: ERD luận lý nhóm danh mục và bảng giá",
        "alt": "ERD luận lý nhóm danh mục và bảng giá",
        "width": 6.15,
    },
    {
        "key": "request",
        "png": DIAGRAM_DIR / "erd-request-log.png",
        "svg": DIAGRAM_DIR / "erd-request-log.svg",
        "media": "erd-request-log.svg",
        "caption": "Hình 3-4: ERD luận lý nhóm đơn yêu cầu và nhật ký",
        "alt": "ERD luận lý nhóm đơn yêu cầu và nhật ký",
        "width": 6.15,
    },
]

for diagram in DIAGRAMS:
    if not diagram["png"].exists() or not diagram["svg"].exists():
        raise FileNotFoundError(diagram)


def p_text(paragraph):
    return paragraph.text.strip()


def set_run_font(run, size=13, bold=None, italic=None, underline=None):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Times New Roman")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if underline is not None:
        run.underline = underline


def set_keep_next(paragraph, enabled=True):
    paragraph.paragraph_format.keep_with_next = enabled


def replace_plain(paragraph, text):
    p_pr = paragraph._p.pPr
    for child in list(paragraph._p):
        if child is not p_pr:
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    return run


def insert_after(paragraph, text):
    new_p = OxmlElement("w:p")
    if paragraph._p.pPr is not None:
        new_p.append(deepcopy(paragraph._p.pPr))
    new_r = OxmlElement("w:r")
    new_t = OxmlElement("w:t")
    new_t.text = text
    new_r.append(new_t)
    new_p.append(new_r)
    paragraph._p.addnext(new_p)
    from docx.text.paragraph import Paragraph

    return Paragraph(new_p, paragraph._parent)


def renumber_existing_chapter3_figures(doc):
    pattern = re.compile(r"Hình 3-(\d+)")

    def repl(match):
        number = int(match.group(1))
        return f"Hình 3-{number + 3}" if number >= 2 else match.group(0)

    for text_node in doc.element.body.xpath(".//w:t"):
        if text_node.text and "Hình 3-" in text_node.text:
            text_node.text = pattern.sub(repl, text_node.text)


doc = Document(INPUT)
renumber_existing_chapter3_figures(doc)

# Rebuild the static list of figures around the old Hình 3-1 entry.
list_entry = next(
    (p for p in doc.paragraphs if p_text(p).startswith("Hình 3-1: Sơ đồ lớp/mô hình dữ liệu chính")),
    None,
)
if list_entry is None:
    raise RuntimeError("Cannot locate list-of-figures entry Hình 3-1")
replace_plain(list_entry, DIAGRAMS[0]["caption"])
current = list_entry
for diagram in DIAGRAMS[1:]:
    current = insert_after(current, diagram["caption"])

heading_31 = next((p for p in doc.paragraphs if p_text(p) == "3.1 MÔ HÌNH DỮ LIỆU"), None)
heading_32 = next((p for p in doc.paragraphs if p_text(p) == "3.2 MÔ HÌNH XỬ LÝ"), None)
if heading_31 is None or heading_32 is None:
    raise RuntimeError("Cannot locate Chapter 3 data-model boundaries")
heading_32.paragraph_format.page_break_before = True

# Remove the old all-in-one ERD, caption, paragraph, and summary table.
node = heading_31._p.getnext()
while node is not None and node is not heading_32._p:
    next_node = node.getnext()
    node.getparent().remove(node)
    node = next_node

anchor = heading_32._p
new_picture_rids = {}


def move_before_anchor(element):
    anchor.addprevious(element)


def add_paragraph(text="", style=None, *, justify=False, page_break=False, keep_next=False, first_indent=True):
    paragraph = doc.add_paragraph(style=style)
    if text:
        run = paragraph.add_run(text)
        if style not in ("Heading 1", "Heading 2", "Heading 3", "Heading 4"):
            set_run_font(run, 13)
    if justify:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.line_spacing = 1.5
        paragraph.paragraph_format.space_after = Pt(0)
        if first_indent:
            paragraph.paragraph_format.first_line_indent = Inches(0.5)
    if page_break:
        paragraph.paragraph_format.page_break_before = True
    if keep_next:
        set_keep_next(paragraph)
    move_before_anchor(paragraph._p)
    return paragraph


def add_caption(text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_together = True
    label, description = text.split(":", 1)
    label_run = paragraph.add_run(label + ":")
    set_run_font(label_run, 12, bold=True, italic=True, underline=True)
    description_run = paragraph.add_run(description)
    set_run_font(description_run, 12)
    move_before_anchor(paragraph._p)
    return paragraph


def add_figure(diagram):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    shape = run.add_picture(str(diagram["png"]), width=Inches(diagram["width"]))
    drawing = paragraph._p.xpath(".//w:drawing")[0]
    blip = drawing.xpath(".//a:blip")[0]
    rid = blip.get(qn("r:embed"))
    new_picture_rids[diagram["key"]] = rid
    for doc_pr in drawing.xpath(".//wp:docPr"):
        doc_pr.set("name", diagram["alt"])
        doc_pr.set("descr", diagram["alt"])
    for c_nv_pr in drawing.xpath(".//pic:cNvPr"):
        c_nv_pr.set("name", diagram["alt"])
        c_nv_pr.set("descr", diagram["alt"])
    move_before_anchor(paragraph._p)
    add_caption(diagram["caption"])


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_physical_table(number, title, rows, *, page_break=False):
    title_p = add_paragraph(f"Bảng 3-{number}: {title}", page_break=page_break, keep_next=True, first_indent=False)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title_p.runs:
        set_run_font(run, 12, bold=True, italic=True)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ("Đối tượng", "Khóa chính", "Khóa ngoại / tham chiếu", "Trường và ràng buộc chính")
    for index, (cell, text) in enumerate(zip(table.rows[0].cells, headers)):
        cell.text = text
        set_cell_shading(cell, "E6E6E6")
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            set_run_font(run, 10, bold=True)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:tblHeader"))
    for values in rows:
        cells = table.add_row().cells
        for col_index, (cell, value) in enumerate(zip(cells, values)):
            cell.text = value
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_index < 2 else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            for run in paragraph.runs:
                set_run_font(run, 9.5, bold=(col_index == 0))
    set_table_widths(table, [1900, 1450, 2300, 3422])
    move_before_anchor(table._tbl)
    spacer = add_paragraph("", first_indent=False)
    spacer.paragraph_format.space_after = Pt(0)


# 3.1 introduction and conceptual model.
add_paragraph(
    "Mô hình dữ liệu được trình bày theo ba mức nhằm tách biệt góc nhìn nghiệp vụ, cấu trúc quan hệ và cách cài đặt trên SQL Server. Nội dung được đối chiếu với entity, cấu hình Entity Framework Core và migration snapshot hiện tại của project.",
    justify=True,
)
add_paragraph("3.1.1 Mô hình dữ liệu ý niệm", style="Heading 3", keep_next=True)
add_paragraph(
    "Ở mức ý niệm, dữ liệu được chia thành ba miền: tổ chức - phân quyền, danh mục - bảng giá và đơn yêu cầu. Người dùng thuộc phạm vi tổ chức, được gán nhóm quyền; đơn yêu cầu chứa nhiều chi tiết vật tư và phát sinh nhật ký thao tác; giá vật tư được xác định theo nhà cung cấp và bảng giá.",
    justify=True,
)
add_figure(DIAGRAMS[0])

# Logical model, split into three readable ERDs.
add_paragraph("3.1.2 Mô hình dữ liệu luận lý", style="Heading 3", page_break=True, keep_next=True)
add_paragraph(
    "Mô hình luận lý giữ nguyên tên bảng và quan hệ trong source nhưng chỉ hiển thị khóa cùng các thuộc tính quan trọng. Sơ đồ được chia theo miền dữ liệu để chữ đủ lớn khi in trên khổ A4.",
    justify=True,
)
add_paragraph("3.1.2.1 Nhóm tổ chức và phân quyền", style="Heading 4", keep_next=True)
add_paragraph(
    "Quyền được cấu thành từ trang, component và ánh xạ trang - component; nhóm quyền liên kết với ánh xạ này theo phạm vi công ty. P04_UserGroup gán người dùng vào nhóm và đơn vị tổ chức. UserId tham chiếu dữ liệu từ view v_Users nhưng chưa có khóa ngoại vật lý; ParentGroupId và ParentId cũng chưa được cấu hình thành quan hệ tự tham chiếu.",
    justify=True,
)
add_figure(DIAGRAMS[1])

add_paragraph("3.1.2.2 Nhóm danh mục và bảng giá", style="Heading 4", page_break=True, keep_next=True)
add_paragraph(
    "Vật tư thuộc một loại và một đơn vị tính. Giá được lưu tại L06_VPPSupplierMapping theo bộ ba vật tư - nhà cung cấp - bảng giá. Hai filtered unique index bảo đảm chỉ có một bảng giá mặc định và chỉ một nhà cung cấp mặc định cho mỗi vật tư trong từng bảng giá khi bản ghi chưa bị xóa mềm.",
    justify=True,
)
add_figure(DIAGRAMS[2])

add_paragraph("3.1.2.3 Nhóm đơn yêu cầu và nhật ký", style="Heading 4", page_break=True, keep_next=True)
add_paragraph(
    "VPP01_RequestHeader lưu kỳ, trạng thái và thông tin duyệt; VPP02_RequestDetail lưu vật tư, số lượng và đơn giá chụp; VPP03_Log lưu lịch sử thao tác dạng JSON. Đơn đã đóng kỳ có thể tham chiếu bảng giá được sử dụng. Các mã người dùng, phòng ban và công ty hiện là giá trị tham chiếu, không phải khóa ngoại trong database.",
    justify=True,
)
add_figure(DIAGRAMS[3])

# Physical model and database constraints.
add_paragraph("3.1.3 Mô hình dữ liệu vật lý", style="Heading 3", page_break=True, keep_next=True)
add_paragraph(
    "Migration hiện tạo 17 bảng. Phần lớn entity kế thừa BaseModel với Id kiểu uniqueidentifier, Description nvarchar(500), thông tin tạo/cập nhật và cờ IsDeleted. Ba bảng ánh xạ P05, P06 và nhật ký VPP03 có cấu trúc riêng. Hai view v_Users và v_WFXCompany được ánh xạ keyless để đọc dữ liệu tích hợp, không thuộc migration của hệ thống.",
    justify=True,
)

auth_rows = [
    ("P01_Page", "Id\nuniqueidentifier", "-", "PageCode nvarchar(50) NOT NULL; PageName nvarchar(250); Type nvarchar(50) NOT NULL."),
    ("P02_Group", "Id\nuniqueidentifier", "-", "GroupName nvarchar(50) NOT NULL; ParentGroupId nullable, chưa có FK tự tham chiếu."),
    ("P03_Component", "Id\nuniqueidentifier", "-", "ComponentCode nvarchar(50) và ComponentName nvarchar(150), đều NOT NULL."),
    ("P04_UserGroup", "Id\nuniqueidentifier", "P02_GroupId; LEX02_CompanyDepartmentLocationId", "UserId int tham chiếu view người dùng nhưng không có FK; các quan hệ cấu hình DeleteBehavior.Restrict."),
    ("P05_PageComponent-\nMapping", "Id\nuniqueidentifier", "P01_PageId; P03_ComponentId", "Bảng nối trang - component; hai quan hệ dùng Restrict."),
    ("P06_GroupPage-\nComponentMapping", "P02_GroupId + P05_PageComponentMappingId + MemberCompanyCode", "P02_GroupId; P05_PageComponentMappingId", "MemberCompanyCode bigint; IsEnable và IsVisible kiểu bit; khóa chính ghép ba cột."),
    ("LEX02_Company-\nDepartmentLocation", "Id\nuniqueidentifier", "-", "LEX02Type nvarchar(max) NOT NULL; ParentId nullable, chưa có FK tự tham chiếu."),
    ("v_Users (view)", "Không có khóa", "-", "View keyless dùng đọc UserId, tài khoản và họ tên; không do migration tạo."),
]
add_physical_table(1, "Mô hình vật lý nhóm tổ chức và phân quyền", auth_rows)

catalog_rows = [
    ("L01_Class", "Id\nuniqueidentifier", "-", "ClassCode nvarchar(max); ClassName và ClassModul nvarchar(200)."),
    ("L02_ClassDetail", "Id\nuniqueidentifier", "ClassId -> L01_Class", "ClassId nullable; mã/giá trị chi tiết nvarchar(max) NOT NULL; Sort int."),
    ("L03_VPP-\nCategory", "Id\nuniqueidentifier", "-", "Mã và tên loại vật tư nvarchar(max)."),
    ("L04_VPP", "Id\nuniqueidentifier", "UOMId -> L02; VPPCategoryId -> L03", "Mã/tên vật tư nvarchar(max); hai FK bắt buộc và dùng Restrict."),
    ("L05_VPP-\nSupplier", "Id\nuniqueidentifier", "-", "Tên, tên ngắn và địa chỉ nhà cung cấp dùng nvarchar(max)."),
    ("L06_VPPSupplier-\nMapping", "Id\nuniqueidentifier", "L04_VPPId; L05_VPPSupplierId; L07_PriceListId", "Price bigint; IsDefault bit; filtered unique index theo PriceListId + VPPId khi IsDefault = 1 và IsDeleted = 0."),
    ("L07_PriceList", "Id\nuniqueidentifier", "-", "PriceListCode nvarchar(50); PriceListName nvarchar(200); IsDefault bit; filtered unique index chỉ cho một mặc định."),
]
add_physical_table(2, "Mô hình vật lý nhóm danh mục và bảng giá", catalog_rows, page_break=True)

request_rows = [
    ("VPP01_Request-\nHeader", "Id\nuniqueidentifier", "SettledByPriceListId -> L07_PriceList", "VPPCode nvarchar(64) unique; Y/M/Status int; IsAdditionalOrder bit; RowVersion rowversion; unique index một đơn thường/người/kỳ."),
    ("VPP02_Request-\nDetail", "Id\nuniqueidentifier", "VPP01_RequestHeaderId; VPPId -> L04_VPP", "Qty int; CurrentSinglePrice bigint. Header dùng Restrict; quan hệ VPP dùng Cascade theo snapshot."),
    ("VPP03_Log", "Id\nuniqueidentifier", "VPP01_RequestHeaderId", "LogTitle nvarchar(255); LogJS nvarchar(max); LogDate datetime2 mặc định GETDATE(); Id mặc định NEWID()."),
    ("v_WFXCompany (view)", "Không có khóa", "-", "View keyless dùng đọc thông tin công ty tích hợp; không do migration tạo."),
]
add_physical_table(3, "Mô hình vật lý nhóm đơn yêu cầu và nhật ký", request_rows, page_break=True)
add_paragraph(
    "Các trạng thái đơn được lưu dưới dạng int: Submitted = 1, Cancelled = 4, Pending = 6, Approved = 7 và Rejected = 8. Ngoại trừ quan hệ VPP02_RequestDetail - L04_VPP đang là Cascade trong snapshot, các quan hệ chính còn lại dùng Restrict để hạn chế xóa dây chuyền; dữ liệu nghiệp vụ chủ yếu được loại bỏ bằng cờ IsDeleted.",
    justify=True,
)

doc.save(OUTPUT)

# Add SVG alternatives to the four newly inserted PNG drawings and remove the obsolete old ERD media.
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL = "http://schemas.openxmlformats.org/package/2006/relationships"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"
ASVG = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"


def ns(namespace, tag):
    return f"{{{namespace}}}{tag}"


def next_rid(rels_root):
    values = []
    for rel in rels_root:
        rid = rel.get("Id", "")
        if rid.startswith("rId") and rid[3:].isdigit():
            values.append(int(rid[3:]))
    return f"rId{max(values, default=0) + 1}"


with ZipFile(OUTPUT, "r") as source:
    document_root = etree.fromstring(source.read("word/document.xml"))
    rels_root = etree.fromstring(source.read("word/_rels/document.xml.rels"))
    content_types = etree.fromstring(source.read("[Content_Types].xml"))

    if not content_types.xpath('./ct:Default[@Extension="svg"]', namespaces={"ct": CT}):
        default = etree.SubElement(content_types, ns(CT, "Default"))
        default.set("Extension", "svg")
        default.set("ContentType", "image/svg+xml")

    for diagram in DIAGRAMS:
        svg_rid = next_rid(rels_root)
        rel = etree.Element(ns(REL, "Relationship"))
        rel.set("Id", svg_rid)
        rel.set("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image")
        rel.set("Target", f"media/{diagram['media']}")
        rels_root.append(rel)

        blips = document_root.xpath(
            f'//a:blip[@r:embed="{new_picture_rids[diagram["key"]]}"]',
            namespaces={"a": A, "r": R},
        )
        if len(blips) != 1:
            raise RuntimeError(f"Expected one new PNG blip for {diagram['key']}, got {len(blips)}")
        blip = blips[0]
        ext_list = blip.find(ns(A, "extLst"))
        if ext_list is None:
            ext_list = etree.SubElement(blip, ns(A, "extLst"))
        ext = etree.SubElement(ext_list, ns(A, "ext"))
        ext.set("uri", "{96DAC541-7B7A-43D3-8B79-37D633B846F1}")
        svg_blip = etree.SubElement(ext, ns(ASVG, "svgBlip"), nsmap={"asvg": ASVG})
        svg_blip.set(ns(R, "embed"), svg_rid)

    referenced_rids = set(document_root.xpath("//@r:embed", namespaces={"r": R}))
    obsolete_targets = set()
    for rel in list(rels_root):
        if rel.get("Id") not in referenced_rids and rel.get("Target") == "media/image4.png":
            obsolete_targets.add("word/media/image4.png")
            rels_root.remove(rel)

    document_xml = etree.tostring(document_root, xml_declaration=True, encoding="UTF-8", standalone=True)
    rels_xml = etree.tostring(rels_root, xml_declaration=True, encoding="UTF-8", standalone=True)
    content_types_xml = etree.tostring(content_types, xml_declaration=True, encoding="UTF-8", standalone=True)

    temp = OUTPUT.with_suffix(".tmp.docx")
    with ZipFile(temp, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            if item.filename in obsolete_targets:
                continue
            data = source.read(item.filename)
            if item.filename == "word/document.xml":
                data = document_xml
            elif item.filename == "word/_rels/document.xml.rels":
                data = rels_xml
            elif item.filename == "[Content_Types].xml":
                data = content_types_xml
            target.writestr(item, data)
        existing = set(source.namelist())
        for diagram in DIAGRAMS:
            name = f"word/media/{diagram['media']}"
            if name not in existing:
                target.writestr(name, diagram["svg"].read_bytes())

temp.replace(OUTPUT)
copyfile(OUTPUT, WORKING)
print(f"checkpoint={OUTPUT}")
print(f"working={WORKING}")
