from __future__ import annotations

import json
import re
import sys
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn


sys.stdout.reconfigure(encoding="utf-8")
DOCX = Path(sys.argv[1])
doc = Document(DOCX)


def pt(value):
    return round(value.pt, 2) if value is not None else None


def style_record(name: str):
    style = doc.styles[name]
    pf = style.paragraph_format
    font = style.font
    return {
        "style_id": style.style_id,
        "font": font.name,
        "size_pt": pt(font.size),
        "bold": font.bold,
        "italic": font.italic,
        "underline": font.underline,
        "color": str(font.color.rgb) if font.color is not None and font.color.rgb is not None else None,
        "alignment": str(pf.alignment),
        "space_before_pt": pt(pf.space_before),
        "space_after_pt": pt(pf.space_after),
        "line_spacing": pf.line_spacing if isinstance(pf.line_spacing, float) else pt(pf.line_spacing),
        "keep_with_next": pf.keep_with_next,
        "page_break_before": pf.page_break_before,
    }


styles = {}
for name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Heading 4", "TOC 1", "TOC 2", "TOC 3"):
    if name in [style.name for style in doc.styles if style.type == WD_STYLE_TYPE.PARAGRAPH]:
        styles[name] = style_record(name)

chapter_start = next(
    index for index, paragraph in enumerate(doc.paragraphs)
    if paragraph.style.style_id == "Heading1" and paragraph.text.upper().startswith("CHƯƠNG 1.")
)

body_font_sizes = Counter()
body_fonts = Counter()
body_colors = Counter()
heading_direct_sizes = Counter()
for paragraph in doc.paragraphs[chapter_start:]:
    for run in paragraph.runs:
        if run.text.strip():
            body_font_sizes[pt(run.font.size)] += len(run.text)
            body_fonts[run.font.name] += len(run.text)
            color = str(run.font.color.rgb) if run.font.color is not None and run.font.color.rgb is not None else None
            body_colors[color] += len(run.text)
            if paragraph.style.style_id in {"Heading1", "Heading2", "Heading3", "Heading4"}:
                heading_direct_sizes[(paragraph.style.style_id, pt(run.font.size))] += len(run.text)

table_font_sizes = Counter()
table_fonts = Counter()
table_para_spacing = Counter()
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                pf = paragraph.paragraph_format
                line = pf.line_spacing if isinstance(pf.line_spacing, float) else pt(pf.line_spacing)
                table_para_spacing[(pt(pf.space_after), line)] += 1
                for run in paragraph.runs:
                    if run.text.strip():
                        table_font_sizes[pt(run.font.size)] += len(run.text)
                        table_fonts[run.font.name] += len(run.text)

caption_records = []
caption_re = re.compile(r"^(Hình|Bảng)\s+\d+-\d+:")
for index, paragraph in enumerate(doc.paragraphs):
    if not caption_re.match(paragraph.text.strip()):
        continue
    caption_records.append({
        "index": index,
        "text": paragraph.text,
        "alignment": str(paragraph.alignment),
        "runs": [
            {
                "text": run.text,
                "font": run.font.name,
                "size": pt(run.font.size),
                "bold": run.bold,
                "italic": run.italic,
                "underline": run.underline,
                "color": str(run.font.color.rgb) if run.font.color is not None and run.font.color.rgb is not None else None,
            }
            for run in paragraph.runs
            if run.text
        ],
    })

headers = []
footers = []
for section_index, section in enumerate(doc.sections, 1):
    for kind, header in (("default", section.header), ("first", section.first_page_header), ("even", section.even_page_header)):
        text = " | ".join(p.text for p in header.paragraphs if p.text.strip())
        headers.append({"section": section_index, "kind": kind, "text": text})
    for kind, footer in (("default", section.footer), ("first", section.first_page_footer), ("even", section.even_page_footer)):
        text = " | ".join(p.text for p in footer.paragraphs if p.text.strip())
        footers.append({"section": section_index, "kind": kind, "text": text})

result = {
    "document": str(DOCX),
    "paragraphs": len(doc.paragraphs),
    "tables": len(doc.tables),
    "sections": len(doc.sections),
    "chapter_start": chapter_start,
    "styles": styles,
    "body_direct_font_sizes_by_char": {str(k): v for k, v in body_font_sizes.most_common()},
    "body_direct_fonts_by_char": {str(k): v for k, v in body_fonts.most_common()},
    "body_direct_colors_by_char": {str(k): v for k, v in body_colors.most_common()},
    "heading_direct_sizes_by_char": {str(k): v for k, v in heading_direct_sizes.most_common()},
    "table_direct_font_sizes_by_char": {str(k): v for k, v in table_font_sizes.most_common()},
    "table_direct_fonts_by_char": {str(k): v for k, v in table_fonts.most_common()},
    "table_paragraph_spacing_counts": {str(k): v for k, v in table_para_spacing.most_common()},
    "caption_count": len(caption_records),
    "caption_examples": caption_records[:8] + caption_records[-5:],
    "headers": headers,
    "footers": footers,
}

print(json.dumps(result, ensure_ascii=False, indent=2))

