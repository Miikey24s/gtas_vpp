from pathlib import Path
from shutil import copy2

from docx import Document


WORKING = Path(r"D:\WORK\gtas_vpp\LVTN\NguyenAnNam_DH52201078_working.docx")
CHECKPOINT = Path(r"D:\WORK\gtas_vpp\LVTN\checkpoints\01_chuong1.docx")
INTRO_TEXT = "Các kết quả cần đạt và tiêu chí đánh giá tương ứng được trình bày trong bảng sau:"


def main():
    doc = Document(WORKING)
    intro = next(p for p in doc.paragraphs if " ".join(p.text.split()) == INTRO_TEXT)
    table = doc.tables[0]
    table._element.addprevious(intro._p)
    doc.save(WORKING)
    copy2(WORKING, CHECKPOINT)
    print("Moved the section 1.4 introduction immediately before Table 1.")


if __name__ == "__main__":
    main()

