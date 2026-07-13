from copy import deepcopy
from pathlib import Path
from shutil import copyfile

from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(r"D:\WORK\gtas_vpp")
INPUT = ROOT / "LVTN" / "checkpoints" / "03C_chuong3_sequence_activity.docx"
OUTPUT = ROOT / "LVTN" / "checkpoints" / "03D_chuong3_giaodien_baobieu.docx"
WORKING = ROOT / "LVTN" / "NguyenAnNam_DH52201078_working.docx"
IMAGE_DIR = ROOT / "LVTN" / "screenshots" / "ch03"


RENAMES = {
    "3.3.1.2 Dashboard tổng quan": "3.3.1.2 Dashboard đơn hàng cá nhân",
    "3.3.2.3 Duyệt đơn bổ sung": "3.3.2.3 Thao tác kỳ và đơn bổ sung chờ duyệt",
    "3.3.4 Giao diện báo cáo": "3.3.4 Giao diện tổng hợp và báo biểu",
    "3.3.4.1 Báo cáo tổng hợp theo kỳ": "3.3.4.1 Tổng hợp toàn doanh nghiệp theo kỳ",
    "Hình 3-27: Giao diện dashboard tổng quan": "Hình 3-27: Giao diện dashboard đơn hàng cá nhân",
    "Hình 3-30: Giao diện duyệt đơn bổ sung": "Hình 3-30: Giao diện thao tác kỳ và đơn bổ sung chờ duyệt",
    "Hình 3-34: Giao diện báo cáo tổng hợp theo kỳ": "Hình 3-34: Giao diện tổng hợp toàn doanh nghiệp theo kỳ",
}


SCREENS = [
    {
        "heading": "3.3.1.1 Đăng nhập hệ thống",
        "caption": "Hình 3-26: Giao diện đăng nhập hệ thống",
        "file": "ui-login.png",
        "alt": "Giao diện đăng nhập GTAS VPP",
        "intro": "Màn hình đăng nhập chia bố cục thành vùng nhận diện PPJ Group và biểu mẫu xác thực. Người dùng nhập tên đăng nhập, mật khẩu; lựa chọn môi trường chỉ xuất hiện khi hệ thống chạy ở chế độ phát triển.",
    },
    {
        "heading": "3.3.1.2 Dashboard đơn hàng cá nhân",
        "caption": "Hình 3-27: Giao diện dashboard đơn hàng cá nhân",
        "file": "ui-dashboard-my-orders.png",
        "alt": "Dashboard đơn hàng cá nhân",
        "intro": "Dashboard cá nhân hiển thị kỳ đặt hàng hiện tại, hạn gửi đơn, tổng dòng hàng, tổng số lượng và các đơn của người dùng. Từ đây người dùng có thể tạo đơn mới, sao chép kỳ trước hoặc tạo đơn bổ sung khi đủ điều kiện.",
    },
    {
        "heading": "3.3.2.1 Danh sách đơn yêu cầu",
        "caption": "Hình 3-28: Giao diện danh sách đơn yêu cầu",
        "file": "ui-order-history.png",
        "alt": "Giao diện lịch sử và danh sách đơn yêu cầu",
        "intro": "Tab lịch sử cho phép lọc theo năm, tháng và nhiều trạng thái; các thẻ KPI tổng hợp số đơn, số dòng hàng và tổng số lượng. Người dùng có thể mở từng dòng để xem chi tiết đơn.",
    },
    {
        "heading": "3.3.2.2 Tạo đơn yêu cầu văn phòng phẩm",
        "caption": "Hình 3-29: Giao diện tạo đơn yêu cầu văn phòng phẩm",
        "file": "ui-order-create.png",
        "alt": "Giao diện tạo đơn yêu cầu văn phòng phẩm",
        "intro": "Màn hình tạo đơn dùng wizard hai bước. Bước đầu cho phép tìm kiếm và chọn vật tư từ danh mục; vùng bên phải hiển thị các mặt hàng đã chọn trước khi người dùng chuyển sang bước rà soát và gửi đơn.",
    },
    {
        "heading": "3.3.2.3 Thao tác kỳ và đơn bổ sung chờ duyệt",
        "caption": "Hình 3-30: Giao diện thao tác kỳ và đơn bổ sung chờ duyệt",
        "file": "ui-period-operations.png",
        "alt": "Giao diện thao tác kỳ và đơn bổ sung chờ duyệt",
        "intro": "Vùng thao tác kỳ gồm hai tab: tổng kết kỳ/đóng kỳ và đơn bổ sung chờ duyệt. Màn hình hiển thị kỳ, bảng giá, cảnh báo số đơn Pending và các chỉ số tổng hợp trước khi người quản trị thực hiện thao tác.",
    },
    {
        "heading": "3.3.3.1 Quản lý bảng giá",
        "caption": "Hình 3-31: Giao diện quản lý bảng giá",
        "file": "ui-price-lists.png",
        "alt": "Giao diện quản lý danh sách bảng giá",
        "intro": "Tab bảng giá cho phép thêm, sửa, sao chép, đặt mặc định hoặc xóa theo ràng buộc dịch vụ. Cột số đơn liên quan giúp người quản trị nhận biết mức độ sử dụng của từng bảng giá.",
    },
    {
        "heading": "3.3.3.2 Quản lý danh mục mặt hàng",
        "caption": "Hình 3-32: Giao diện quản lý danh mục mặt hàng",
        "file": "ui-library-items.png",
        "alt": "Giao diện quản lý danh mục mặt hàng",
        "intro": "Danh mục mặt hàng sử dụng lưới dữ liệu có lọc, phân trang và lựa chọn cột. Mỗi bản ghi liên kết mã hàng, tên hàng, đơn vị tính, nhóm hàng, nhà cung cấp và giá để phục vụ lập đơn.",
    },
    {
        "heading": "3.3.3.3 Phân quyền theo nhóm",
        "caption": "Hình 3-33: Giao diện phân quyền theo nhóm",
        "file": "ui-permission-groups.png",
        "alt": "Giao diện phân quyền theo nhóm và trang",
        "intro": "Màn hình phân quyền hiển thị danh sách nhóm và vùng cấu hình quyền nhóm/trang. Người quản trị chọn nhóm, sau đó cập nhật quan hệ nhóm cha và các mapping component bằng IsVisible, IsEnable.",
    },
    {
        "heading": "3.3.4.1 Tổng hợp toàn doanh nghiệp theo kỳ",
        "caption": "Hình 3-34: Giao diện tổng hợp toàn doanh nghiệp theo kỳ",
        "file": "ui-all-orders-summary.png",
        "alt": "Giao diện tổng hợp toàn doanh nghiệp theo kỳ",
        "intro": "Màn hình tổng hợp toàn doanh nghiệp cung cấp KPI tổng số đơn, tổng dòng hàng, tổng số lượng và danh sách đơn theo phòng ban, người tạo, kỳ và trạng thái. Đây là dữ liệu tổng hợp đã hoạt động, được dùng thay cho trang Report chưa hoàn thiện.",
    },
]


def set_run_font(run, size=11, bold=None, italic=None, underline=None):
    run.font.name = "Times New Roman"
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rfonts.set(qn(key), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if underline is not None:
        run.underline = underline


def replace_plain(paragraph, text, *, size=None):
    old_rpr = None
    if paragraph.runs and paragraph.runs[0]._r.rPr is not None:
        old_rpr = deepcopy(paragraph.runs[0]._r.rPr)
    p_pr = paragraph._p.pPr
    for child in list(paragraph._p):
        if child is not p_pr:
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if old_rpr is not None:
        run._r.insert(0, old_rpr)
    if size is not None:
        set_run_font(run, size=size)


def paragraph_after(doc, paragraph, exact_text):
    seen = False
    for candidate in doc.paragraphs:
        if candidate._p is paragraph._p:
            seen = True
            continue
        if seen and candidate.text.strip() == exact_text:
            return candidate
    raise RuntimeError(f"Cannot find paragraph after heading: {exact_text}")


def format_body(paragraph, text, *, italic=False, size=11, keep_with_next=False):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Inches(0.39)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.keep_with_next = keep_with_next
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(text)
    set_run_font(run, size=size, italic=italic)


def format_caption(paragraph, text):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = False
    p_pr = paragraph._p.pPr
    for child in list(paragraph._p):
        if child is not p_pr:
            paragraph._p.remove(child)
    label, description = text.split(":", 1)
    label_run = paragraph.add_run(label + ":")
    set_run_font(label_run, 12, bold=True, italic=True, underline=True)
    desc_run = paragraph.add_run(description)
    set_run_font(desc_run, 12)


doc = Document(INPUT)

# Rename headings and captions that previously overstated or mismatched the actual screen.
for paragraph in doc.paragraphs:
    text = paragraph.text.strip()
    if text in RENAMES:
        replace_plain(paragraph, RENAMES[text], size=12 if text.startswith("Hình ") else None)

# Make the section introduction explicit about the source and status of the screenshots.
section_intro = next(
    p for p in doc.paragraphs
    if p.text.strip().startswith("Giao diện hệ thống được tổ chức theo mô hình ứng dụng quản trị nội bộ")
)
replace_plain(
    section_intro,
    "Giao diện hệ thống được tổ chức theo mô hình ứng dụng quản trị nội bộ, ưu tiên thao tác lặp lại, lọc dữ liệu, phân trang và hiển thị trạng thái. Các vùng đã triển khai tập trung ở Dashboard, Library và Permission; trang Report được giữ như khung mở rộng nhưng chưa có dữ liệu báo cáo hoàn chỉnh.",
)
for run in section_intro.runs:
    set_run_font(run, 11)
section_intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
section_intro.paragraph_format.first_line_indent = Inches(0.39)
section_intro.paragraph_format.line_spacing = 1.15

source_note = doc.add_paragraph()
format_body(
    source_note,
    "Các ảnh nghiệp vụ và quản trị sử dụng snapshot kiểm thử giao diện ngày 28/05/2026 trên môi trường SERVER TEST; riêng màn hình đăng nhập được chụp lại từ frontend hiện tại ngày 13/07/2026. Dữ liệu hiển thị trong hình là dữ liệu kiểm thử.",
    italic=True,
    size=10,
)
section_intro._p.addnext(source_note._p)

# Clarify why the implemented dashboard summary is presented instead of the unfinished Report page.
report_heading = next(p for p in doc.paragraphs if p.text.strip() == "3.3.4 Giao diện tổng hợp và báo biểu")
report_status = doc.add_paragraph()
format_body(
    report_status,
    "Hệ thống đã có các vùng tổng hợp cá nhân, phòng ban và toàn doanh nghiệp trên Dashboard. Trang /report hiện chỉ hiển thị các thẻ chỉ số rỗng, cảnh báo nội dung chưa được cấu hình và nút xuất bị vô hiệu hóa; vì vậy trang này không được mô tả như một chức năng báo cáo đã hoàn thành.",
)
report_heading._p.addnext(report_status._p)

for screen in SCREENS:
    image_path = IMAGE_DIR / screen["file"]
    if not image_path.exists():
        raise FileNotFoundError(image_path)

    heading = next((p for p in doc.paragraphs if p.text.strip() == screen["heading"]), None)
    if heading is None:
        raise RuntimeError(f"Cannot locate heading: {screen['heading']}")
    caption = paragraph_after(doc, heading, screen["caption"])

    node = heading._p.getnext()
    while node is not None and node is not caption._p:
        next_node = node.getnext()
        node.getparent().remove(node)
        node = next_node

    intro = doc.add_paragraph()
    format_body(intro, screen["intro"], keep_with_next=True)

    picture = doc.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.space_before = Pt(2)
    picture.paragraph_format.space_after = Pt(0)
    picture.paragraph_format.keep_with_next = True
    picture.paragraph_format.keep_together = True

    with Image.open(image_path) as image:
        ratio = image.width / image.height
    max_width = 6.15
    max_height = 5.25
    if ratio >= max_width / max_height:
        width = max_width
        height = width / ratio
    else:
        height = max_height
        width = height * ratio

    run = picture.add_run()
    run.add_picture(str(image_path), width=Inches(width), height=Inches(height))
    drawing = picture._p.xpath(".//w:drawing")[0]
    for doc_pr in drawing.xpath(".//wp:docPr"):
        doc_pr.set("name", screen["alt"])
        doc_pr.set("descr", screen["alt"])
    for c_nv_pr in drawing.xpath(".//pic:cNvPr"):
        c_nv_pr.set("name", screen["alt"])
        c_nv_pr.set("descr", screen["alt"])

    heading.paragraph_format.keep_with_next = True
    heading._p.addnext(intro._p)
    intro._p.addnext(picture._p)
    format_caption(caption, screen["caption"])

doc.save(OUTPUT)
working_temp = WORKING.with_suffix(".tmp.docx")
copyfile(OUTPUT, working_temp)
working_temp.replace(WORKING)
print(f"checkpoint={OUTPUT}")
print(f"working={WORKING}")

