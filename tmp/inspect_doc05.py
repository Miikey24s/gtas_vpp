from pathlib import Path
import re
import zipfile
from lxml import etree
from docx import Document
import sys

sys.stdout.reconfigure(encoding="utf-8")

DOCX = Path(r"D:\WORK\gtas_vpp\LVTN\checkpoints\04_chuong4_thunghiem.docx")
doc = Document(DOCX)

print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} sections={len(doc.sections)}")

start = None
for i, p in enumerate(doc.paragraphs):
    if "CHƯƠNG 5" in p.text.upper() and p.style.name == "Heading 1":
        start = i
        break
print(f"chapter5_start={start}")
for i in range(max(0, (start or 0) - 3), len(doc.paragraphs)):
    p = doc.paragraphs[i]
    text = " ".join(p.text.split())
    if text or i >= (start or 0):
        print(f"P{i:04d}\t{p.style.name!r}\t{text}")

print("\n=== CHAPTER 1 OBJECTIVES AND SCOPE ===")
for i, p in enumerate(doc.paragraphs):
    if 120 <= i <= 180:
        text = " ".join(p.text.split())
        if text:
            print(f"O{i:04d}\t{p.style.name!r}\t{text}")

print("\n=== IN-TEXT CITATION PARAGRAPHS ===")
pat = re.compile(r"\[(?:\d+)(?:\s*[-,]\s*\d+)*\]")
for i, p in enumerate(doc.paragraphs):
    if pat.search(p.text):
        print(f"C{i:04d}\t{p.style.name!r}\t{' '.join(p.text.split())}")

print("\n=== TABLE TEXT NEAR END ===")
for ti, table in enumerate(doc.tables):
    combined = " | ".join(" / ".join(c.text.replace("\n", " ") for c in row.cells) for row in table.rows)
    if ti >= len(doc.tables) - 6:
        print(f"T{ti:02d}\t{combined[:900]}")

with zipfile.ZipFile(DOCX) as z:
    xml = etree.fromstring(z.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
          "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
    bookmarks = [(b.get(f"{{{ns['w']}}}name"), b.get(f"{{{ns['w']}}}id")) for b in xml.xpath(".//w:bookmarkStart", namespaces=ns)]
    links = xml.xpath(".//w:hyperlink", namespaces=ns)
    internal = [h.get(f"{{{ns['w']}}}anchor") for h in links if h.get(f"{{{ns['w']}}}anchor")]
    external = [h.get(f"{{{ns['r']}}}id") for h in links if h.get(f"{{{ns['r']}}}id")]
    print(f"\nbookmarks={len(bookmarks)} hyperlinks={len(links)} internal={len(internal)} external={len(external)}")
    print("bookmark_names_tail=", [n for n, _ in bookmarks[-30:]])
    print("internal_tail=", internal[-30:])
