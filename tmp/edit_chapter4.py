from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


SOURCE = Path(r"D:\WORK\gtas_vpp\LVTN\checkpoints\03D_chuong3_giaodien_baobieu.docx")
OUTPUT = Path(r"D:\WORK\gtas_vpp\LVTN\checkpoints\04_chuong4_thunghiem.docx")
WORKING = Path(r"D:\WORK\gtas_vpp\LVTN\NguyenAnNam_DH52201078_working.docx")

TABLE_WIDTH = 9000
MATRIX_WIDTHS = [800, 1500, 2000, 1850, 1900, 950]


def paragraph_text(node) -> str:
    return "".join(text_node.text or "" for text_node in node.iter(qn("w:t"))).strip()


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_table_geometry(table, widths: list[int]) -> None:
    assert sum(widths) == TABLE_WIDTH
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH))
    tbl_w.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def format_run(run, size: float, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Times New Roman")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def format_cell(cell, size: float, bold: bool = False, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            format_run(run, size=size, bold=bold)


def add_paragraph(
    doc,
    marker,
    text: str = "",
    style: str = "Normal",
    keep_with_next: bool = False,
    page_break_before: bool = False,
):
    paragraph = doc.add_paragraph(style=style)
    paragraph.add_run(text)
    paragraph.paragraph_format.keep_with_next = keep_with_next
    paragraph.paragraph_format.page_break_before = page_break_before
    marker.addprevious(paragraph._p)
    return paragraph


def add_caption(doc, marker, number: str, title: str, page_break_before: bool = False):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.page_break_before = page_break_before
    prefix = paragraph.add_run(f"Bảng {number}:")
    format_run(prefix, 11, bold=True, italic=True)
    suffix = paragraph.add_run(f" {title}")
    format_run(suffix, 11, italic=True)
    marker.addprevious(paragraph._p)
    return paragraph


def add_table(doc, marker, headers, rows, widths, font_size=9.5, center_columns=()):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header = table.rows[0]
    for index, value in enumerate(headers):
        header.cells[index].text = value
        set_cell_shading(header.cells[index], "E7E6E6")
        format_cell(header.cells[index], font_size, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_header(header)
    prevent_row_split(header)

    for values in rows:
        row = table.add_row()
        for index, value in enumerate(values):
            row.cells[index].text = value
            alignment = WD_ALIGN_PARAGRAPH.CENTER if index in center_columns else WD_ALIGN_PARAGRAPH.LEFT
            format_cell(row.cells[index], font_size, align=alignment)
        prevent_row_split(row)

    set_table_geometry(table, widths)
    marker.addprevious(table._tbl)
    return table


def add_labeled_paragraph(doc, marker, label: str, text: str):
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.paragraph_format.left_indent = Pt(18)
    paragraph.paragraph_format.first_line_indent = Pt(-18)
    paragraph.paragraph_format.space_after = Pt(3)
    label_run = paragraph.add_run(label + " ")
    label_run.bold = True
    paragraph.add_run(text)
    marker.addprevious(paragraph._p)
    return paragraph


doc = Document(SOURCE)
body = doc.element.body
children = list(body.iterchildren())

chapter4 = next(node for node in children if node.tag == qn("w:p") and paragraph_text(node) == "CHƯƠNG 4. THỬ NGHIỆM")
start_index = children.index(chapter4)
section_end = next(
    node
    for node in children[start_index + 1 :]
    if node.tag == qn("w:p")
    and node.find("./w:pPr/w:sectPr", namespaces={"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}) is not None
)

for node in list(body.iterchildren()):
    if node is chapter4:
        deleting = True
        continue
    if node is section_end:
        break
    if "deleting" in locals() and deleting:
        body.remove(node)

marker = section_end

add_paragraph(doc, marker, "4.1 CÁC KỊCH BẢN THỬ NGHIỆM", "Heading 2", keep_with_next=True)
add_paragraph(
    doc,
    marker,
    "Mục tiêu thử nghiệm là kiểm chứng các quy tắc nghiệp vụ có rủi ro cao, khả năng xử lý lỗi và tính ổn định của các hàm hỗ trợ giao diện. Mỗi trường hợp được mô tả bằng mã, tiền điều kiện, bước thực hiện, kết quả mong đợi, kết quả thực tế và trạng thái để có thể truy vết về test tự động hoặc báo cáo audit tương ứng.",
)

add_paragraph(doc, marker, "4.1.1 Phạm vi và phương pháp ghi nhận", "Heading 3", keep_with_next=True)
add_paragraph(
    doc,
    marker,
    "Kết quả hiện tại được lấy từ các lệnh kiểm thử và build chạy ngày 13/07/2026 trên .NET SDK 10.0.301. Riêng audit giao diện là snapshot độc lập ngày 28/05/2026; số liệu này chỉ chứng minh trạng thái của phiên bản tại thời điểm audit, không được xem là kết quả chạy lại hiện tại.",
)
add_caption(doc, marker, "4-1", "Phạm vi và bằng chứng kiểm thử")
add_table(
    doc,
    marker,
    ["Tầng kiểm thử", "Công cụ", "Phạm vi", "Bằng chứng"],
    [
        ["Backend unit/integration", "xUnit, Moq, FluentAssertions, EF Core InMemory/SQLite", "Quy tắc kỳ, trạng thái đơn, tạo đơn, bảng giá, đóng kỳ, controller và middleware.", "Tệp TRX sinh ngày 13/07/2026."],
        ["Frontend helper", "xUnit", "Chuẩn hóa URL API, quy tắc bỏ qua chứng chỉ trong Development và định dạng ngày giờ.", "Tệp TRX sinh ngày 13/07/2026."],
        ["UI end-to-end", "Playwright for .NET", "Project có test đăng nhập, tạo đơn, lịch sử, tổng hợp, danh mục và phân quyền.", "Project build thành công; không chạy E2E trong đợt này vì cần môi trường FE/BE và dữ liệu test đã cấu hình."],
        ["UI audit snapshot", "Playwright audit runner", "18 route xác thực trên mobile, tablet và desktop.", "Snapshot 28/05/2026, tổng cộng 54 lượt."],
    ],
    [1500, 2050, 3200, 2250],
    font_size=10,
)

add_paragraph(doc, marker, "4.1.2 Ma trận test case đại diện", "Heading 3", keep_with_next=True)
add_paragraph(
    doc,
    marker,
    "Ma trận dưới đây chọn các trường hợp đại diện cho luồng chính, biên dữ liệu và nhánh lỗi. Cột kết quả thực tế ghi theo nhóm test đã chạy; một test case có thể được bao phủ bởi nhiều test tham số hóa.",
)

add_caption(doc, marker, "4-2", "Ma trận kiểm thử xác thực, kỳ và đơn yêu cầu", page_break_before=True)
add_table(
    doc,
    marker,
    ["Mã", "Tiền điều kiện", "Bước thực hiện", "Kết quả mong đợi", "Kết quả thực tế", "TT"],
    [
        ["TC-01", "Có tài khoản đúng, sai và môi trường Live-only.", "Gọi API đăng nhập lần lượt với ba bộ dữ liệu.", "Đúng thông tin trả token; sai thông tin trả 401; môi trường không hợp lệ trả 400.", "3/3 test xác thực đạt.", "Đạt"],
        ["TC-02", "Mốc thời gian trước, đúng và sau 00:00 ngày 5.", "Tính kỳ hiện tại, kỳ trước, deadline và kiểm tra đã quá hạn.", "Chuyển kỳ đúng biên ngày 5; deadline là ngày 5 của tháng kế tiếp.", "18/18 test tính kỳ đạt.", "Đạt"],
        ["TC-03", "Danh sách vật tư null/rỗng, số lượng không dương hoặc trùng mã.", "Gọi validation trước khi tạo đơn.", "Từ chối dữ liệu không hợp lệ và không ghi đơn.", "4 trường hợp validation đạt.", "Đạt"],
        ["TC-04", "Kỳ còn mở và dữ liệu vật tư hợp lệ.", "Tạo một đơn thường và một đơn bổ sung.", "Đơn thường ở Submitted; đơn bổ sung ở Pending.", "2 test tạo đơn đạt đúng trạng thái.", "Đạt"],
        ["TC-05", "Hai yêu cầu tạo đơn được gửi gần đồng thời.", "Chạy cặp regular/regular, regular/additional và regular khác kỳ.", "Chỉ một regular cùng kỳ thành công; hai loại khác nhau hoặc khác kỳ được phép.", "3/3 race-condition tests đạt.", "Đạt"],
    ],
    MATRIX_WIDTHS,
    font_size=9.5,
    center_columns=(0, 5),
)

add_caption(doc, marker, "4-3", "Ma trận kiểm thử trạng thái, bảng giá và đóng kỳ")
add_table(
    doc,
    marker,
    ["Mã", "Tiền điều kiện", "Bước thực hiện", "Kết quả mong đợi", "Kết quả thực tế", "TT"],
    [
        ["TC-06", "Đơn ở Submitted, Pending, Approved, Rejected hoặc Cancelled.", "Thực hiện Update, Cancel, Approve và Reject theo từng trạng thái.", "Chỉ các chuyển trạng thái hợp lệ thành công; tổ hợp sai phát sinh lỗi nghiệp vụ.", "20/20 test chuyển trạng thái đạt.", "Đạt"],
        ["TC-07", "Có nhiều bảng giá/giá nhà cung cấp và một bản ghi mặc định.", "Đặt mặc định, sao chép, cập nhật và thử xóa bản mặc định hoặc đang được dùng.", "Luôn còn một mặc định; clone tạo ID mới; thao tác xóa không hợp lệ bị chặn.", "13/13 test bảng giá và giá nhà cung cấp đạt.", "Đạt"],
        ["TC-08", "Không còn đơn bổ sung Pending; bảng giá đủ giá.", "Đóng kỳ và gọi lại settlement.", "Chụp giá từng dòng, đóng dấu header, ghi log; lần gọi sau vẫn nhất quán.", "Các nhánh happy path, log và idempotent đạt.", "Đạt"],
        ["TC-09", "Còn đơn Pending hoặc thiếu giá cho một vật tư.", "Thử đóng kỳ ở từng điều kiện lỗi.", "Không hoàn tất settlement; trả lỗi nghiệp vụ và không để dữ liệu dở dang.", "Test blocked-by-pending và missing-price đều đạt.", "Đạt"],
        ["TC-10", "Thiếu UserID hoặc service phát sinh exception đã biết/không biết.", "Gọi controller và middleware với từng loại lỗi.", "Trả 401 hoặc Problem Details 400/403/404/500 tương ứng.", "1 test claim và 4 middleware mapping đạt.", "Đạt"],
    ],
    MATRIX_WIDTHS,
    font_size=9.5,
    center_columns=(0, 5),
)

add_caption(doc, marker, "4-4", "Ma trận kiểm thử frontend và audit giao diện")
add_table(
    doc,
    marker,
    ["Mã", "Tiền điều kiện", "Bước thực hiện", "Kết quả mong đợi", "Kết quả thực tế", "TT"],
    [
        ["TC-11", "Có URL localhost, wildcard, container và cấu hình thiếu.", "Resolve URL và kiểm tra quy tắc bypass chứng chỉ.", "Chuẩn hóa đúng host; chỉ bypass HTTPS localhost trong Development; thiếu URL phát sinh lỗi.", "12/12 test URL đạt.", "Đạt"],
        ["TC-12", "Có giá trị ngày giờ hợp lệ, nullable và culture khác nhau.", "Định dạng ngày ngắn, dài, giờ, tháng/năm và null.", "Hiển thị dd/MM/yyyy, HH:mm ổn định; null thành dấu gạch ngang.", "14/14 test định dạng đạt.", "Đạt"],
        ["TC-13", "Frontend tại commit và dữ liệu SERVER TEST của đợt audit.", "Mở 18 route trên 3 viewport; kiểm tra lỗi Blazor, response, console và tràn ngang.", "Mỗi lượt không có lỗi chặn theo tiêu chí audit.", "54/54 lượt OK ngày 28/05/2026; đây là snapshot, chưa chạy lại hiện tại.", "Đạt tại snapshot"],
    ],
    MATRIX_WIDTHS,
    font_size=9.5,
    center_columns=(0, 5),
)

add_paragraph(doc, marker, "4.2 KẾT QUẢ THỬ NGHIỆM CÁC KỊCH BẢN", "Heading 2", keep_with_next=True)
result_note = add_paragraph(
    doc,
    marker,
    "Hai bộ test tự động hiện tại có tổng cộng 158 test và đều đạt. Ngoài ra, solution chính và project UI test đều build thành công. Bảng kết quả tách rõ số liệu chạy ngày 13/07/2026 với snapshot audit ngày 28/05/2026 để tránh hiểu nhầm 54 lượt audit là kết quả hiện tại.",
)
add_caption(doc, marker, "4-5", "Tổng hợp kết quả kiểm thử và build")
add_table(
    doc,
    marker,
    ["Hạng mục", "Tổng", "Đạt", "Không đạt / bỏ qua", "Thời gian", "Mốc và ghi chú"],
    [
        ["Backend test", "132", "132", "0 / 0", "4 giây", "Chạy 13/07/2026; có TRX."],
        ["Frontend helper test", "26", "26", "0 / 0", "585 ms", "Chạy 13/07/2026; có TRX."],
        ["Build solution chính", "-", "0 lỗi", "0 cảnh báo", "10,54 giây", "Build 13/07/2026; solution không gồm project test."],
        ["Build project UI test", "-", "0 lỗi", "0 cảnh báo", "8,07 giây", "Build 13/07/2026; không chạy E2E."],
        ["UI audit", "54", "54 OK", "0 lỗi audit", "-", "Snapshot 28/05/2026: 18 route × 3 viewport."],
    ],
    [1800, 650, 850, 1350, 950, 3400],
    font_size=10,
    center_columns=(1, 2, 3, 4),
)
add_paragraph(
    doc,
    marker,
    "Kết quả build hiện tại không còn xuất hiện tham chiếu project AI bị thiếu. Tuy nhiên, điều này không thay thế kiểm thử UI end-to-end: các test Playwright cần backend, frontend, tài khoản và dữ liệu test đồng bộ, nên chỉ được ghi nhận là đã build thành công trong đợt này.",
)
result_note.paragraph_format.space_before = Pt(6)

add_paragraph(
    doc,
    marker,
    "4.3 XỬ LÝ NGOẠI LỆ VÀ TỒN ĐỌNG KỸ THUẬT",
    "Heading 2",
    keep_with_next=True,
    page_break_before=True,
)
add_paragraph(doc, marker, "4.3.1 Xử lý các trường hợp ngoại lệ", "Heading 3", keep_with_next=True)
add_labeled_paragraph(doc, marker, "Xác thực và môi trường:", "Sai thông tin đăng nhập trả 401; lựa chọn môi trường không hợp lệ trong chế độ Live-only trả 400; thiếu UserID trong claim trả Unauthorized thay vì tiếp tục xử lý.")
add_labeled_paragraph(doc, marker, "Dữ liệu đơn không hợp lệ:", "Danh sách vật tư null/rỗng, số lượng không dương hoặc trùng mặt hàng bị từ chối trước khi ghi dữ liệu.")
add_labeled_paragraph(doc, marker, "Trùng đơn và đồng thời:", "Hai yêu cầu tạo đơn thường cùng kỳ chỉ một yêu cầu thành công; đơn khác kỳ hoặc cặp đơn thường/đơn bổ sung được phép theo quy tắc nghiệp vụ.")
add_labeled_paragraph(doc, marker, "Đơn bổ sung và đóng kỳ:", "Hệ thống chặn settlement khi còn đơn Pending hoặc bảng giá thiếu mặt hàng; giao dịch không được xác nhận khi validation thất bại.")
add_labeled_paragraph(doc, marker, "Ánh xạ lỗi API:", "Middleware chuyển InvalidOperationException, UnauthorizedAccessException, KeyNotFoundException và lỗi không xác định thành Problem Details với mã 400, 403, 404 và 500 tương ứng.")

add_paragraph(doc, marker, "4.3.2 Tồn đọng kỹ thuật", "Heading 3", keep_with_next=True)
add_labeled_paragraph(doc, marker, "Cảnh báo NU1903:", "Khi restore project backend test, .NET báo Microsoft.OpenApi 2.4.1 và SQLitePCLRaw.lib.e_sqlite3 2.1.11 có lỗ hổng mức nghiêm trọng cao. Cần đánh giá phiên bản nâng cấp tương thích, chạy lại toàn bộ test và kiểm tra migration trước khi triển khai.")
add_labeled_paragraph(doc, marker, "Giới hạn bằng chứng UI:", "Kết quả 54/54 là snapshot ngày 28/05/2026, không phải lần chạy ngày 13/07/2026. Cần chạy lại audit và Playwright E2E trên bản sắp bàn giao với môi trường test cố định.")
add_labeled_paragraph(doc, marker, "Phạm vi dữ liệu:", "Backend test hiện chủ yếu dùng EF Core InMemory/SQLite; chưa có kiểm thử tải dài hạn và bộ integration test đầy đủ trên SQL Server với dữ liệu gần quy mô thực tế.")
add_labeled_paragraph(doc, marker, "Chức năng chưa hoàn thiện:", "Trang Report và xuất Excel/PDF chưa được tính là chức năng đã đạt; các test hiện tại chỉ chứng minh dashboard tổng hợp và các helper liên quan.")

add_paragraph(doc, marker, "", "Normal")

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
shutil.copy2(OUTPUT, WORKING)
print(OUTPUT)
