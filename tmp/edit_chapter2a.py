from copy import deepcopy
from pathlib import Path
from shutil import copy2
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from lxml import etree


ROOT = Path(r"D:\WORK\gtas_vpp")
WORKING = ROOT / "LVTN" / "NguyenAnNam_DH52201078_working.docx"
CHECKPOINT = ROOT / "LVTN" / "checkpoints" / "02A_chuong2_congnghe.docx"
STAGE_TEXT = ROOT / "tmp" / "02a_text.docx"
STAGE_PATCHED = ROOT / "tmp" / "02a_patched.docx"
SVG = ROOT / "LVTN" / "diagrams" / "ch02" / "architecture-overview.svg"
PNG = ROOT / "tmp" / "ch02_arch" / "architecture-overview.png"

P21_OLD = (
    "Để xác định hướng xây dựng hệ thống, đề tài khảo sát một số hình thức và hệ thống "
    "thường gặp trong quản lý yêu cầu nội bộ của doanh nghiệp."
)
P21_NEW = (
    "Để xác định hướng xây dựng hệ thống, đề tài khảo sát bốn nhóm giải pháp thường gặp "
    "trong quản lý yêu cầu nội bộ: biểu mẫu hoặc email, bảng tính dùng chung, nền tảng "
    "mua sắm tổng quát và nền tảng quản lý yêu cầu hoặc quy trình."
)
P21_EVIDENCE = (
    "Tài liệu chính thức cho thấy Google Sheets và Excel for the web hỗ trợ chia sẻ, "
    "đồng biên soạn; Odoo Purchase hỗ trợ quản lý báo giá và đơn mua hàng; Jira Service "
    "Management và Zoho Creator hỗ trợ tiếp nhận yêu cầu, luồng xử lý hoặc phê duyệt "
    "[11]-[15]. Bảng so sánh dưới đây đánh giá các nhóm giải pháp theo mức độ phù hợp "
    "với nghiệp vụ riêng của GTAS VPP, gồm quy tắc ngày 5, đơn bổ sung, chụp giá và "
    "phân quyền chi tiết."
)

P22_INTRO_1 = (
    "Hệ thống GTAS VPP được hiện thực trên nền tảng .NET 10, gồm frontend Blazor Server, "
    "backend ASP.NET Core Web API và cơ sở dữ liệu SQL Server. Các lựa chọn công nghệ "
    "trong mục này được đối chiếu trực tiếp với project, package reference và cấu hình "
    "khởi động của hệ thống [1]-[9]."
)
P22_INTRO_2 = (
    "Trong mô hình triển khai, Nginx tiếp nhận kết nối HTTPS/WSS và định tuyến giao diện, "
    "SignalR hoặc API đến dịch vụ tương ứng. Frontend duy trì phiên đăng nhập bằng cookie; "
    "backend xác thực JWT, áp dụng policy authorization và truy cập SQL Server qua Entity "
    "Framework Core. Dịch vụ migrator thực hiện migration và seed dữ liệu trước khi backend "
    "phục vụ yêu cầu."
)

REFERENCES = [
    '[11] Google Workspace Learning Center, "Share & collaborate on a spreadsheet." Available: https://support.google.com/a/users/answer/13309904',
    '[12] Microsoft Support, "Basic tasks in Excel for the web." Available: https://support.microsoft.com/en-US/Excel/basic-tasks-in-excel-for-the-web',
    '[13] Odoo, "Purchase - Odoo 18.0 documentation." Available: https://www.odoo.com/documentation/18.0/applications/inventory_and_mrp/purchase.html',
    '[14] Atlassian, "Get started with service requests in Jira Service Management." Available: https://www.atlassian.com/software/jira/service-management/product-guide/getting-started/service-request-management',
    '[15] Zoho Creator, "Approval workflows." Available: https://www.zoho.com/creator/approval-workflow/',
]


def normalized(text):
    return " ".join(text.split())


def set_paragraph_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def append_reference(doc, text):
    if any(normalized(p.text) == normalized(text) for p in doc.paragraphs):
        return
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.add_run(text)


def edit_text_and_tables():
    doc = Document(WORKING)

    p21 = next(p for p in doc.paragraphs if normalized(p.text) == P21_OLD)
    set_paragraph_text(p21, P21_NEW)
    evidence = doc.add_paragraph(P21_EVIDENCE, style="Normal")
    p21._p.addnext(evidence._p)

    table = doc.tables[1]
    table.cell(2, 1).text = "Bảng tính dùng chung (Google Sheets, Excel for the web) [11], [12]"
    table.cell(2, 2).text = (
        "Dễ tiếp cận; hỗ trợ chia sẻ, đồng biên soạn, lọc và tổng hợp dữ liệu cơ bản."
    )
    table.cell(2, 3).text = (
        "Không cung cấp sẵn quy tắc kỳ/ngày 5, luồng trạng thái, chụp giá hoặc phân quyền "
        "ở mức component/API; các ràng buộc phụ thuộc vào cách tổ chức bảng và công thức."
    )
    table.cell(3, 1).text = "Nền tảng mua sắm tổng quát (Odoo Purchase) [13]"
    table.cell(3, 2).text = (
        "Hỗ trợ nhà cung cấp, yêu cầu báo giá, đơn mua hàng, thỏa thuận mua và báo cáo mua sắm."
    )
    table.cell(3, 3).text = (
        "Phạm vi rộng hơn nhu cầu của đề tài; cần tùy biến thêm quy tắc kỳ/ngày 5, "
        "đơn bổ sung và cơ chế chụp giá của GTAS VPP."
    )
    table.cell(4, 1).text = (
        "Nền tảng quản lý yêu cầu/quy trình (Jira Service Management, Zoho Creator) [14], [15]"
    )
    table.cell(4, 2).text = (
        "Có cổng tiếp nhận, biểu mẫu, trạng thái, luồng phê duyệt, phân quyền, tự động hóa "
        "và báo cáo."
    )
    table.cell(4, 3).text = (
        "Cần thiết kế hoặc cấu hình thêm mô hình dữ liệu, quy tắc ngày 5, giới hạn đơn bổ sung "
        "và chụp giá theo nghiệp vụ riêng của doanh nghiệp."
    )

    paragraphs = doc.paragraphs
    heading_index = next(
        i for i, p in enumerate(paragraphs)
        if normalized(p.text) == "2.2 CÔNG NGHỆ SỬ DỤNG"
    )
    picture_paragraph = next(
        p for p in paragraphs[heading_index + 1 :]
        if p._p.xpath(".//w:drawing")
    )
    intro1 = doc.add_paragraph(P22_INTRO_1, style="Normal")
    intro2 = doc.add_paragraph(P22_INTRO_2, style="Normal")
    picture_paragraph._p.addprevious(intro1._p)
    picture_paragraph._p.addprevious(intro2._p)

    caption = next(p for p in doc.paragraphs if normalized(p.text) == "Hình 2-1: Kiến trúc tổng thể hệ thống GTAS VPP")
    for run in caption.runs:
        run.text = run.text.replace(
            "Kiến trúc tổng thể hệ thống GTAS VPP",
            "Kiến trúc triển khai tổng thể hệ thống GTAS VPP",
        )

    tech = doc.tables[2]
    tech.cell(1, 1).text = "ASP.NET Core Web API (.NET 10) [1]"
    tech.cell(2, 1).text = "Blazor Server, Radzen Blazor [4], [5]"
    tech.cell(3, 1).text = "SQL Server, Entity Framework Core [2], [3]"
    tech.cell(4, 1).text = "JWT Bearer, cookie frontend, policy-based authorization [1], [4]"
    tech.cell(5, 1).text = "Mapster [6]"
    tech.cell(6, 1).text = "Serilog, bảng VPP03_Log [7]"
    tech.cell(7, 1).text = "Docker Compose, Nginx reverse proxy [8]"
    tech.cell(7, 2).text = (
        "Đóng gói SQL Server, migrator, backend và frontend; Nginx định tuyến /api, /_blazor, "
        "thực hiện reverse proxy và SSL termination."
    )
    tech.cell(8, 1).text = "xUnit, Playwright, FluentAssertions, Moq, EF Core InMemory/SQLite [9]"

    for reference in REFERENCES:
        append_reference(doc, reference)

    doc.save(STAGE_TEXT)


def next_relationship_id(root):
    values = []
    for element in root:
        rid = element.get("Id", "")
        if rid.startswith("rId") and rid[3:].isdigit():
            values.append(int(rid[3:]))
    return f"rId{max(values, default=0) + 1}"


def patch_svg_package():
    ns = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
        "pic": "http://schemas.openxmlformats.org/drawingml/2006/picture",
    }
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    ct_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    asvg_ns = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"

    with ZipFile(STAGE_TEXT, "r") as source:
        document_xml = etree.fromstring(source.read("word/document.xml"))
        rels_xml = etree.fromstring(source.read("word/_rels/document.xml.rels"))
        content_types = etree.fromstring(source.read("[Content_Types].xml"))

        svg_rid = next_relationship_id(rels_xml)
        relationship = etree.Element(f"{{{rel_ns}}}Relationship")
        relationship.set("Id", svg_rid)
        relationship.set(
            "Type",
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image",
        )
        relationship.set("Target", "media/architecture-overview.svg")
        rels_xml.append(relationship)

        if not content_types.xpath(
            './ct:Default[@Extension="svg"]',
            namespaces={"ct": ct_ns},
        ):
            default = etree.Element(f"{{{ct_ns}}}Default")
            default.set("Extension", "svg")
            default.set("ContentType", "image/svg+xml")
            content_types.append(default)

        blip = document_xml.xpath(
            '//a:blip[@r:embed="rId13"]',
            namespaces=ns,
        )[0]
        ext_list = blip.find(f"{{{ns['a']}}}extLst")
        if ext_list is None:
            ext_list = etree.SubElement(blip, f"{{{ns['a']}}}extLst")
        svg_ext = etree.SubElement(ext_list, f"{{{ns['a']}}}ext")
        svg_ext.set("uri", "{96DAC541-7B7A-43D3-8B79-37D633B846F1}")
        svg_blip = etree.SubElement(svg_ext, f"{{{asvg_ns}}}svgBlip")
        svg_blip.set(f"{{{ns['r']}}}embed", svg_rid)

        drawing = blip.xpath("ancestor::w:drawing[1]", namespaces=ns)[0]
        width_in = 5.25
        height_in = width_in * 888 / 866
        cx = str(round(width_in * 914400))
        cy = str(round(height_in * 914400))
        for extent in drawing.xpath(".//wp:extent | .//a:xfrm/a:ext", namespaces=ns):
            extent.set("cx", cx)
            extent.set("cy", cy)
        doc_pr = drawing.xpath(".//wp:docPr", namespaces=ns)[0]
        doc_pr.set("descr", "Kiến trúc triển khai tổng thể hệ thống GTAS VPP")
        doc_pr.set("title", "Hình 2-1")

        with ZipFile(STAGE_PATCHED, "w", ZIP_DEFLATED) as target:
            names = set(source.namelist())
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == "word/document.xml":
                    data = etree.tostring(
                        document_xml,
                        xml_declaration=True,
                        encoding="UTF-8",
                        standalone="yes",
                    )
                elif item.filename == "word/_rels/document.xml.rels":
                    data = etree.tostring(
                        rels_xml,
                        xml_declaration=True,
                        encoding="UTF-8",
                        standalone="yes",
                    )
                elif item.filename == "[Content_Types].xml":
                    data = etree.tostring(
                        content_types,
                        xml_declaration=True,
                        encoding="UTF-8",
                        standalone="yes",
                    )
                elif item.filename == "word/media/image1.png":
                    data = PNG.read_bytes()
                target.writestr(item, data)
            if "word/media/architecture-overview.svg" not in names:
                target.writestr(
                    "word/media/architecture-overview.svg",
                    SVG.read_bytes(),
                )

    copy2(STAGE_PATCHED, WORKING)
    copy2(STAGE_PATCHED, CHECKPOINT)


def main():
    for path in (WORKING, SVG, PNG):
        if not path.exists():
            raise FileNotFoundError(path)
    edit_text_and_tables()
    patch_svg_package()
    print(f"Working: {WORKING}")
    print(f"Checkpoint: {CHECKPOINT}")
    print(f"PlantUML: {SVG.with_suffix('.puml')}")
    print(f"SVG: {SVG}")


if __name__ == "__main__":
    main()
