from pathlib import Path
from shutil import copyfile
from zipfile import ZIP_DEFLATED, ZipFile

from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree


ROOT = Path(r"D:\WORK\gtas_vpp")
INPUT = ROOT / "LVTN" / "checkpoints" / "03B_chuong3_usecase_chitiet.docx"
OUTPUT = ROOT / "LVTN" / "checkpoints" / "03C_chuong3_sequence_activity.docx"
WORKING = ROOT / "LVTN" / "NguyenAnNam_DH52201078_working.docx"
DIAGRAM_DIR = ROOT / "LVTN" / "diagrams" / "ch03"


DIAGRAMS = [
    {
        "key": "seq_login",
        "file": "sequence-login-permissions",
        "heading": "3.2.2.1 Đăng nhập và tải quyền",
        "caption": "Hình 3-15: Sơ đồ tuần tự chức năng đăng nhập và tải quyền",
        "alt": "Sơ đồ tuần tự đăng nhập và tải quyền",
        "intro": "Luồng đăng nhập gồm hai giai đoạn: backend xác thực và phát JWT; frontend dùng ticket một lần để tạo cookie, sau đó PermissionState tải quyền page/component trước khi điều hướng. Nhánh sai thông tin hoặc lỗi kết nối không tạo phiên đăng nhập.",
    },
    {
        "key": "seq_regular",
        "file": "sequence-create-regular-request",
        "heading": "3.2.2.2 Tạo đơn yêu cầu thông thường",
        "caption": "Hình 3-16: Sơ đồ tuần tự chức năng tạo đơn yêu cầu thông thường",
        "alt": "Sơ đồ tuần tự tạo đơn yêu cầu thông thường",
        "intro": "Frontend lấy kỳ hiện tại do backend xác định rồi gửi đơn. VPPRequestService kiểm tra deadline ngày 5, một đơn thường trong kỳ, trạng thái settlement và vật tư còn hiệu lực trước khi lưu header, details, giá hiện tại và log CREATE trong cùng transaction.",
    },
    {
        "key": "seq_additional",
        "file": "sequence-create-additional-request",
        "heading": "3.2.2.3 Tạo đơn bổ sung",
        "caption": "Hình 3-17: Sơ đồ tuần tự chức năng tạo đơn bổ sung",
        "alt": "Sơ đồ tuần tự tạo đơn bổ sung",
        "intro": "Đơn bổ sung chỉ nhắm đến kỳ liền trước. Dịch vụ kiểm tra giới hạn tối đa ba đơn, không còn đơn bổ sung Pending khác và kỳ chưa settlement; đơn hợp lệ được tạo ở trạng thái Pending để chờ phê duyệt.",
    },
    {
        "key": "seq_approval",
        "file": "sequence-approve-additional-request",
        "heading": "3.2.2.4 Duyệt đơn bổ sung",
        "caption": "Hình 3-18: Sơ đồ tuần tự chức năng duyệt đơn bổ sung",
        "alt": "Sơ đồ tuần tự duyệt hoặc từ chối đơn bổ sung",
        "intro": "Tab_AdminApproval tải danh sách Pending và gửi quyết định sau khi người quản trị xác nhận. State machine chỉ cho phép chuyển từ Pending sang Approved hoặc Rejected; người xử lý, thời điểm, lý do và log tương ứng được ghi trong transaction.",
    },
    {
        "key": "seq_settle",
        "file": "sequence-settle-period",
        "heading": "3.2.2.5 Đóng kỳ và chụp giá",
        "caption": "Hình 3-19: Sơ đồ tuần tự chức năng đóng kỳ và chụp giá",
        "alt": "Sơ đồ tuần tự đóng kỳ và chụp giá",
        "intro": "PeriodSettlementService kiểm tra bảng giá, đơn bổ sung Pending và độ phủ giá trước khi cập nhật. Việc chụp CurrentSinglePrice, đánh dấu settlement và ghi PERIOD_SETTLED được thực hiện nguyên tử; bất kỳ lỗi nào cũng rollback toàn bộ kỳ.",
    },
    {
        "key": "seq_pricing",
        "file": "sequence-manage-pricing",
        "heading": "3.2.2.6 Quản lý bảng giá",
        "caption": "Hình 3-20: Sơ đồ tuần tự chức năng quản lý bảng giá",
        "alt": "Sơ đồ tuần tự quản lý bảng giá và ánh xạ giá",
        "intro": "Luồng minh họa thao tác thêm hoặc sửa giá vật tư theo nhà cung cấp trong một bảng giá. VPPPriceService kiểm tra giá không âm, các khóa tham chiếu và ràng buộc giá mặc định duy nhất trước khi commit.",
    },
    {
        "key": "act_regular",
        "file": "activity-create-regular-request",
        "heading": "3.2.3.1 Tạo đơn yêu cầu thông thường",
        "caption": "Hình 3-21: Sơ đồ hoạt động quy trình tạo đơn yêu cầu thông thường",
        "alt": "Sơ đồ hoạt động tạo đơn yêu cầu thông thường",
        "intro": "Quy trình hoạt động nhấn mạnh các điểm chặn nghiệp vụ trước khi tạo đơn: hạn ngày 5, đơn trùng, dữ liệu chi tiết và trạng thái đóng kỳ.",
    },
    {
        "key": "act_additional",
        "file": "activity-additional-request",
        "heading": "3.2.3.2 Xử lý đơn bổ sung",
        "caption": "Hình 3-22: Sơ đồ hoạt động quy trình xử lý đơn bổ sung",
        "alt": "Sơ đồ hoạt động xử lý đơn bổ sung",
        "intro": "Quy trình kết hợp hai vai trò: nhân viên lập đơn Pending và quản trị viên quyết định phê duyệt hoặc từ chối. Mỗi nhánh kết thúc bằng trạng thái và nhật ký có thể truy vết.",
    },
    {
        "key": "act_settle",
        "file": "activity-settle-period",
        "heading": "3.2.3.3 Đóng kỳ",
        "caption": "Hình 3-23: Sơ đồ hoạt động quy trình đóng kỳ",
        "alt": "Sơ đồ hoạt động đóng kỳ",
        "intro": "Đóng kỳ chỉ hoàn tất khi không còn đơn bổ sung chờ duyệt và bảng giá bao phủ toàn bộ vật tư của các đơn Submitted/Approved. Các bước cập nhật được bao trong một transaction.",
    },
    {
        "key": "act_pricing",
        "file": "activity-manage-pricing",
        "heading": "3.2.3.4 Quản lý bảng giá",
        "caption": "Hình 3-24: Sơ đồ hoạt động quy trình quản lý bảng giá",
        "alt": "Sơ đồ hoạt động quản lý bảng giá",
        "intro": "Quy trình quản lý giá thể hiện việc kiểm tra tham chiếu, xử lý cờ mặc định và rollback khi vi phạm ràng buộc duy nhất; không giả định bảng giá bị khóa sau settlement vì source hiện không có quy tắc đó.",
    },
    {
        "key": "act_permissions",
        "file": "activity-manage-permissions",
        "heading": "3.2.3.5 Phân quyền",
        "caption": "Hình 3-25: Sơ đồ hoạt động quy trình phân quyền",
        "alt": "Sơ đồ hoạt động phân quyền",
        "intro": "Quy trình phân quyền bao gồm kiểm tra quan hệ nhóm cha và cập nhật mapping trang/component. Hệ thống từ chối quan hệ tạo vòng lặp hoặc cây nhóm có độ sâu từ 10 cấp, sau đó làm mới PermissionState để giao diện áp dụng IsVisible/IsEnable mới.",
    },
]


for diagram in DIAGRAMS:
    diagram["png"] = DIAGRAM_DIR / f"{diagram['file']}.png"
    diagram["svg"] = DIAGRAM_DIR / f"{diagram['file']}.svg"
    diagram["media"] = f"{diagram['file']}.svg"
    if not diagram["png"].exists() or not diagram["svg"].exists():
        raise FileNotFoundError(diagram["file"])


def set_run_font(run, size=11, bold=None, italic=None, underline=None):
    run.font.name = "Times New Roman"
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rfonts.set(qn(key), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if underline is not None:
        run.underline = underline


def paragraph_after(doc, paragraph, exact_text):
    seen = False
    for candidate in doc.paragraphs:
        if candidate._p is paragraph._p:
            seen = True
            continue
        if seen and candidate.text.strip() == exact_text:
            return candidate
    raise RuntimeError(f"Cannot find paragraph after heading: {exact_text}")


def format_intro(paragraph, text):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = Inches(0.39)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(text)
    set_run_font(run, 11)


def format_caption(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.keep_with_next = False
    for run in paragraph.runs:
        set_run_font(run, 12)


doc = Document(INPUT)
new_picture_rids = {}

for diagram in DIAGRAMS:
    heading = next((p for p in doc.paragraphs if p.text.strip() == diagram["heading"]), None)
    if heading is None:
        raise RuntimeError(f"Cannot locate heading: {diagram['heading']}")
    caption = paragraph_after(doc, heading, diagram["caption"])

    node = heading._p.getnext()
    while node is not None and node is not caption._p:
        next_node = node.getnext()
        node.getparent().remove(node)
        node = next_node

    intro = doc.add_paragraph()
    format_intro(intro, diagram["intro"])

    picture = doc.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.space_before = Pt(2)
    picture.paragraph_format.space_after = Pt(0)
    picture.paragraph_format.keep_with_next = True
    picture.paragraph_format.keep_together = True

    with Image.open(diagram["png"]) as image:
        ratio = image.width / image.height
    max_width = 6.15
    max_height = 5.40
    if ratio >= max_width / max_height:
        width = max_width
        height = width / ratio
    else:
        height = max_height
        width = height * ratio

    run = picture.add_run()
    run.add_picture(str(diagram["png"]), width=Inches(width), height=Inches(height))
    drawing = picture._p.xpath(".//w:drawing")[0]
    blip = drawing.xpath(".//a:blip")[0]
    rid = blip.get(qn("r:embed"))
    new_picture_rids[diagram["key"]] = rid
    for doc_pr in drawing.xpath(".//wp:docPr"):
        doc_pr.set("name", diagram["alt"])
        doc_pr.set("descr", diagram["alt"])
    for c_nv_pr in drawing.xpath(".//pic:cNvPr"):
        c_nv_pr.set("name", diagram["alt"])
        c_nv_pr.set("descr", diagram["alt"])

    heading.paragraph_format.keep_with_next = True
    heading._p.addnext(intro._p)
    intro._p.addnext(picture._p)
    format_caption(caption)

doc.save(OUTPUT)


# Attach each SVG as the preferred Office image while keeping its PNG fallback.
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL = "http://schemas.openxmlformats.org/package/2006/relationships"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"
ASVG = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"


def ns(namespace, tag):
    return f"{{{namespace}}}{tag}"


def next_rid(rels_root):
    values = []
    for rel in rels_root:
        rid = rel.get("Id", "")
        if rid.startswith("rId") and rid[3:].isdigit():
            values.append(int(rid[3:]))
    return f"rId{max(values, default=0) + 1}"


with ZipFile(OUTPUT, "r") as source:
    document_root = etree.fromstring(source.read("word/document.xml"))
    rels_root = etree.fromstring(source.read("word/_rels/document.xml.rels"))
    content_types = etree.fromstring(source.read("[Content_Types].xml"))

    if not content_types.xpath('./ct:Default[@Extension="svg"]', namespaces={"ct": CT}):
        default = etree.SubElement(content_types, ns(CT, "Default"))
        default.set("Extension", "svg")
        default.set("ContentType", "image/svg+xml")

    for diagram in DIAGRAMS:
        svg_rid = next_rid(rels_root)
        rel = etree.Element(ns(REL, "Relationship"))
        rel.set("Id", svg_rid)
        rel.set("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image")
        rel.set("Target", f"media/{diagram['media']}")
        rels_root.append(rel)

        blips = document_root.xpath(
            f'//a:blip[@r:embed="{new_picture_rids[diagram["key"]]}"]',
            namespaces={"a": A, "r": R},
        )
        if len(blips) != 1:
            raise RuntimeError(f"Expected one PNG blip for {diagram['key']}, got {len(blips)}")
        blip = blips[0]
        ext_list = blip.find(ns(A, "extLst"))
        if ext_list is None:
            ext_list = etree.SubElement(blip, ns(A, "extLst"))
        ext = etree.SubElement(ext_list, ns(A, "ext"))
        ext.set("uri", "{96DAC541-7B7A-43D3-8B79-37D633B846F1}")
        svg_blip = etree.SubElement(ext, ns(ASVG, "svgBlip"), nsmap={"asvg": ASVG})
        svg_blip.set(ns(R, "embed"), svg_rid)

    document_xml = etree.tostring(document_root, xml_declaration=True, encoding="UTF-8", standalone=True)
    rels_xml = etree.tostring(rels_root, xml_declaration=True, encoding="UTF-8", standalone=True)
    content_types_xml = etree.tostring(content_types, xml_declaration=True, encoding="UTF-8", standalone=True)

    temp = OUTPUT.with_suffix(".tmp.docx")
    with ZipFile(temp, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "word/document.xml":
                data = document_xml
            elif item.filename == "word/_rels/document.xml.rels":
                data = rels_xml
            elif item.filename == "[Content_Types].xml":
                data = content_types_xml
            target.writestr(item, data)

        existing = set(source.namelist())
        for diagram in DIAGRAMS:
            name = f"word/media/{diagram['media']}"
            if name not in existing:
                target.writestr(name, diagram["svg"].read_bytes())

temp.replace(OUTPUT)
working_temp = WORKING.with_suffix(".tmp.docx")
copyfile(OUTPUT, working_temp)
working_temp.replace(WORKING)
print(f"checkpoint={OUTPUT}")
print(f"working={WORKING}")

