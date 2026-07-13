from __future__ import annotations

import argparse
import os
import posixpath
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Pt
from docx.text.paragraph import Paragraph
from lxml import etree


FONT = "Times New Roman"
NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
}

HERE = Path(__file__).resolve().parent
LVTN_ROOT = HERE.parent

DIAGRAMS = {
    "2-1": LVTN_ROOT / "diagrams" / "ch02" / "architecture-overview",
    "2-2": LVTN_ROOT / "diagrams" / "ch02" / "functional-decomposition",
    "2-3": LVTN_ROOT / "diagrams" / "ch02" / "use-case-overview",
    "3-1": LVTN_ROOT / "diagrams" / "ch03" / "data-conceptual",
    "3-4": LVTN_ROOT / "diagrams" / "ch03" / "erd-request-log",
    "3-13": LVTN_ROOT / "diagrams" / "ch03" / "use-case-manage-permissions",
    "3-15": LVTN_ROOT / "diagrams" / "ch03" / "sequence-login-permissions",
    "3-25": LVTN_ROOT / "diagrams" / "ch03" / "activity-manage-permissions",
}


def normalize(value: str) -> str:
    return " ".join(value.split())


def set_run_font(run, size: float = 13) -> None:
    run.font.name = FONT
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for name in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{name}"), FONT)


def clear_paragraph(paragraph: Paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def replace_by_prefix(doc: Document, prefix: str, replacement: str) -> None:
    matches = [paragraph for paragraph in doc.paragraphs if normalize(paragraph.text).startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph starting with {prefix!r}, found {len(matches)}")
    paragraph = matches[0]
    clear_paragraph(paragraph)
    set_run_font(paragraph.add_run(replacement))


def find_by_prefix(doc: Document, prefix: str) -> Paragraph:
    matches = [paragraph for paragraph in doc.paragraphs if normalize(paragraph.text).startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph starting with {prefix!r}, found {len(matches)}")
    return matches[0]


def insert_after(paragraph: Paragraph) -> Paragraph:
    element = OxmlElement("w:p")
    paragraph._p.addnext(element)
    created = Paragraph(element, paragraph._parent)
    created.style = paragraph._parent.part.document.styles["Normal"]
    return created


def add_internal_hyperlink(paragraph: Paragraph, text: str, anchor: str) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    hyperlink.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    rpr.append(style)
    fonts = OxmlElement("w:rFonts")
    for name in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{name}"), FONT)
    rpr.append(fonts)
    run.append(rpr)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_external_hyperlink(paragraph: Paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    rpr.append(style)
    fonts = OxmlElement("w:rFonts")
    for name in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{name}"), FONT)
    rpr.append(fonts)
    run.append(rpr)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def next_bookmark_id(doc: Document) -> int:
    values = []
    for node in doc._element.xpath(".//w:bookmarkStart"):
        value = node.get(qn("w:id"))
        if value and value.isdigit():
            values.append(int(value))
    return max(values, default=0) + 1


def bookmark_paragraph(paragraph: Paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    insert_at = 1 if len(paragraph._p) and paragraph._p[0].tag == qn("w:pPr") else 0
    paragraph._p.insert(insert_at, start)
    paragraph._p.append(end)


def add_ai_sources(doc: Document) -> None:
    if any(normalize(paragraph.text).startswith("[16]") for paragraph in doc.paragraphs):
        return

    technology_anchor = find_by_prefix(doc, "Trong mô hình triển khai, Nginx")
    paragraph = insert_after(technology_anchor)
    set_run_font(paragraph.add_run(
        "Module nhận định báo cáo là tính năng tùy chọn. Backend gọi OpenAI Responses API khi người dùng chủ động yêu cầu, "
        "ràng buộc kết quả bằng JSON Schema, đặt store=false, chỉ gửi chỉ số tổng hợp và dùng quy tắc xác định khi tính năng "
        "tắt, thiếu khóa, quá thời gian hoặc API lỗi "
    ))
    for index in range(16, 19):
        add_internal_hyperlink(paragraph, f"[{index}]", f"ref_{index:02d}")
        if index < 18:
            set_run_font(paragraph.add_run(", "))
    set_run_font(paragraph.add_run("."))

    sources = [
        (
            16,
            "OpenAI, “Responses API,” OpenAI Developer Documentation. [Trực tuyến]. Địa chỉ: ",
            "https://developers.openai.com/api/docs/guides/text",
        ),
        (
            17,
            "OpenAI, “Structured Outputs,” OpenAI Developer Documentation. [Trực tuyến]. Địa chỉ: ",
            "https://developers.openai.com/api/docs/guides/structured-outputs",
        ),
        (
            18,
            "OpenAI, “Production best practices,” OpenAI Developer Documentation. [Trực tuyến]. Địa chỉ: ",
            "https://developers.openai.com/api/docs/guides/production-best-practices",
        ),
    ]
    bookmark_id = next_bookmark_id(doc)
    for number, lead, url in sources:
        reference = doc.add_paragraph(style="Normal")
        set_run_font(reference.add_run(f"[{number}] {lead}"))
        add_external_hyperlink(reference, url, url)
        set_run_font(reference.add_run(". [Truy cập: 13/07/2026]."))
        bookmark_paragraph(reference, f"ref_{number:02d}", bookmark_id)
        bookmark_id += 1


def add_appendix_guidance(doc: Document) -> None:
    headings = [
        paragraph for paragraph in doc.paragraphs
        if normalize(paragraph.text).startswith("B.3 Kiểm tra tổng hợp")
        and paragraph.style.style_id == "Heading3"
    ]
    if len(headings) != 1:
        raise ValueError(f"Expected one B.3 Heading 3 paragraph, found {len(headings)}")
    heading = headings[0]
    clear_paragraph(heading)
    set_run_font(heading.add_run("B.3 Báo cáo, xuất dữ liệu và thông báo"))

    replace_by_prefix(
        doc,
        "Tổng hợp phòng ban:",
        "Báo cáo theo phạm vi: Mở trang Report, chọn phạm vi cá nhân, phòng ban hoặc toàn doanh nghiệp mà tài khoản được cấp; "
        "chọn năm/tháng và áp dụng bộ lọc để xem KPI, xu hướng theo kỳ, phân bố trạng thái, phòng ban và mặt hàng nổi bật.",
    )
    replace_by_prefix(
        doc,
        "Tổng hợp toàn doanh nghiệp:",
        "Xuất dữ liệu: Tài khoản có REPORT_EXPORT chọn Xuất báo cáo để tải CSV UTF-8 trong đúng phạm vi quyền. "
        "Nếu dữ liệu quá lớn, thu hẹp năm hoặc tháng trước khi xuất.",
    )
    export_paragraph = find_by_prefix(doc, "Xuất dữ liệu:")
    notification = insert_after(export_paragraph)
    set_run_font(notification.add_run(
        "Thông báo: Chọn biểu tượng chuông để xem hộp thư, số chưa đọc, đánh dấu từng mục hoặc tất cả là đã đọc. "
        "Đơn bổ sung mới và kết quả duyệt được phát theo thời gian thực; chọn mục thông báo để mở chức năng liên quan."
    ))
    insight = insert_after(notification)
    set_run_font(insight.add_run(
        "Nhận định báo cáo: Chọn Tạo nhận định khi báo cáo có dữ liệu. Hệ thống luôn có kết quả theo quy tắc; "
        "nhận định AI chỉ xuất hiện khi backend được bật và cấu hình khóa hợp lệ, không thay thế quyết định nghiệp vụ."
    ))


def update_paragraphs(doc: Document) -> None:
    replacements = {
        "- Xây dựng frontend bằng Blazor Server":
            "- Xây dựng frontend bằng Blazor Server kết hợp Radzen Blazor, tổ chức giao diện theo các vùng Dashboard, "
            "Library, Permission và Report; tích hợp báo cáo theo phạm vi, xuất CSV, nhận định tùy chọn và trung tâm thông báo.",
        "Trong mô hình triển khai, Nginx":
            "Trong mô hình triển khai, Nginx tiếp nhận HTTPS/WSS và định tuyến giao diện, Blazor SignalR, API cùng các hub "
            "quyền/thông báo. Frontend duy trì phiên bằng cookie và gọi backend bằng JWT; backend áp dụng action policy, "
            "truy cập SQL Server qua Entity Framework Core, còn migrator chạy migration và seed trước khi phục vụ.",
        "Lưu ý về phạm vi phân quyền:":
            "Phân quyền được thực hiện ở hai lớp. Frontend dùng snapshot trang/component để ẩn hoặc khóa giao diện; backend "
            "đọc mapping hiện hành ở mỗi action có policy nên quyền vừa bị thu hồi bị từ chối ngay ở lần gọi API kế tiếp. "
            "SignalR phát PermissionsChanged theo nhóm/người dùng để phiên đang mở tải lại snapshot mà không cần đăng nhập lại.",
        "- Backend đồng thời trả về danh sách quyền":
            "- Backend trả dữ liệu đăng nhập và access token; sau khi frontend tạo cookie, PermissionState gọi "
            "/api/Auth/me/permissions để lấy một snapshot gồm trang, component, action hiệu lực, nhóm, công ty và version.",
        "- Frontend nhận token JWT":
            "- Frontend dùng ticket một lần để đổi token đăng nhập thành cookie, tải snapshot quyền và khởi tạo kết nối "
            "SignalR /hubs/permissions bằng access token.",
        "- Giao diện tự động ẩn hoặc hiển thị":
            "- Giao diện tự động ẩn/hiển thị và bật/tắt menu, trang, tab, component theo IsVisible/IsEnable; "
            "backend vẫn kiểm tra mã action tương ứng, không tin cậy việc ẩn nút ở frontend.",
        "- Hệ thống điều hướng đến trang đầu tiên":
            "- Hệ thống điều hướng đến tuyến đầu tiên được phép. Khi quản trị viên đổi quyền, phiên đang mở nhận sự kiện, "
            "tải lại snapshot và rời chức năng nếu quyền truy cập vừa bị thu hồi.",
        "- Hệ thống tạo đầu đơn với mã tự động":
            "- Hệ thống tạo đầu đơn với mã tự động dạng VPP-YYYYMM-GUID, trạng thái Submitted, DepartmentCode và thời điểm gửi. "
            "DepartmentCode và MemberCompanyCode được lấy từ JWT; thiếu phạm vi tổ chức thì API từ chối thao tác.",
        "- Người dùng được cấu hình hiển thị và kích hoạt tab \"Duyệt đơn bổ sung\"":
            "- Người dùng có quyền REQUEST_APPROVE được hiển thị và kích hoạt tab \"Duyệt đơn bổ sung\" để xem danh sách "
            "các đơn Pending. API duyệt yêu cầu policy REQUEST_APPROVE; API từ chối yêu cầu REQUEST_REJECT, nên việc ẩn nút "
            "ở frontend không thay thế kiểm tra quyền tại backend.",
        "- Người dùng được hiển thị tab \"Đóng kỳ\"":
            "- Người dùng có quyền PERIOD_SETTLE được hiển thị tab \"Đóng kỳ\" để chọn năm, tháng và bảng giá áp dụng. "
            "Cả các endpoint xem trạng thái, xem trước và thực hiện đóng kỳ đều yêu cầu policy PERIOD_SETTLE; service tiếp tục "
            "kiểm tra bảng giá, đơn bổ sung Pending và dữ liệu đơn trước khi ghi nhận.",
        "- Phòng ban không còn hoạt động được vô hiệu hóa hoặc kích hoạt lại bằng PATCH IsDeleted.":
            "- Phòng ban không còn hoạt động được vô hiệu hóa hoặc kích hoạt lại bằng PATCH IsDeleted. Endpoint nằm trong "
            "nhóm thao tác ghi của LibraryController và yêu cầu policy LIBRARY_MANAGE tại backend.",
        "- Tổng hợp phòng ban: frontend mặc định":
            "- Tổng hợp phòng ban và báo cáo phòng ban lấy DepartmentCode từ claim; service giới hạn truy vấn theo công ty và "
            "phòng ban của phiên, đồng thời yêu cầu REQUEST_VIEW_DEPARTMENT hoặc REPORT_VIEW_DEPARTMENT.",
        "- Tổng hợp toàn doanh nghiệp: endpoint all-orders":
            "- Tổng hợp/báo cáo toàn doanh nghiệp vẫn giới hạn theo MemberCompanyCode trong claim và yêu cầu action "
            "REQUEST_VIEW_ALL hoặc REPORT_VIEW_ALL; người dùng chỉ có quyền phòng ban không thể đổi query để mở rộng phạm vi.",
        "- Endpoint xóa nhóm quyền hiện gọi":
            "- Các thao tác tạo, sửa, xóa nhóm/mapping/người dùng yêu cầu PERMISSION_MANAGE. Endpoint xóa nhóm từ chối khi "
            "nhóm còn người dùng hoạt động; hệ thống cũng chặn xóa quyền quản trị phân quyền cuối cùng của công ty.",
        "- Khi quản trị viên thay đổi cấu hình":
            "- Khi quản trị viên thay đổi cấu hình, backend lưu thời điểm/người cập nhật và phát sự kiện theo nhóm. "
            "Các phiên liên quan làm mới quyền gần như tức thời; action API tiếp theo luôn đối chiếu mapping mới.",
        "- Thay đổi nhóm quyền sẽ ảnh hưởng trực tiếp":
            "- Khi đổi nhóm hoặc phòng ban, backend thông báo cả nhóm cũ, nhóm mới và người dùng. Phiên đang mở tải lại snapshot; "
            "quyền giao diện và quyền action thay đổi mà không cần đăng nhập lại.",
        "- Người dùng nhấn vào đơn để xem đầu đơn":
            "- Người dùng nhấn vào đơn để xem đầu đơn, mặt hàng, số lượng, đơn giá, ghi chú và thời điểm xử lý. "
            "Controller kiểm tra quyền xem theo phạm vi sở hữu/phòng ban/toàn công ty trước khi trả dữ liệu; sửa/hủy còn yêu cầu "
            "quyền action và đơn thuộc chính người dùng trong cùng công ty.",
        "Lưu ý: Các tác nhân trên phản ánh vai trò nghiệp vụ":
            "Lưu ý: Tác nhân thể hiện vai trò nghiệp vụ. Quyền thực tế được cấu hình bằng page-component mapping và mở rộng "
            "sang các action tương thích; backend cưỡng chế policy cho nghiệp vụ quan trọng, còn SignalR đồng bộ thay đổi tới phiên đang mở.",
        "VPP01_RequestHeader lưu kỳ":
            "VPP01_RequestHeader lưu kỳ, trạng thái và thông tin duyệt; VPP02_RequestDetail lưu vật tư, số lượng và đơn giá chụp; "
            "VPP03_Log lưu lịch sử thao tác dạng JSON. N01_Notification lưu hộp thư theo người dùng/công ty, loại, nội dung, tuyến mở, "
            "CorrelationId, thời điểm tạo và đọc. Các mã tổ chức/người dùng là tham chiếu logic, không phải khóa ngoại vật lý.",
        "Migration hiện tạo 17 bảng.":
            "Migration hiện tạo 18 bảng, gồm N01_Notification được bổ sung cho hộp thư bền vững. Phần lớn entity kế thừa "
            "BaseModel; các bảng ánh xạ P05, P06, VPP03_Log và N01_Notification có cấu trúc riêng. Hai view v_Users và "
            "v_WFXCompany được ánh xạ keyless để đọc dữ liệu tích hợp, không thuộc migration của hệ thống.",
        "Ở mức thiết kế xử lý, các luồng quan trọng":
            "Ở mức thiết kế xử lý, các luồng chính đi qua giao diện Blazor, controller API và service. Controller đọc claim, "
            "policy handler kiểm tra action theo mapping hiện hành; service kiểm tra nghiệp vụ, giao dịch dữ liệu và ghi log. "
            "SignalR chỉ thông báo thay đổi để frontend tải lại, không thay thế kiểm tra quyền ở API.",
        "Luồng đăng nhập gồm hai giai đoạn:":
            "Luồng đăng nhập gồm hai giai đoạn: backend xác thực và phát JWT; frontend dùng ticket một lần để tạo cookie, gọi một "
            "endpoint lấy PermissionSnapshot rồi kết nối hub quyền. Nhánh sai thông tin/lỗi kết nối không tạo phiên; thay đổi quyền "
            "sau đó được đồng bộ qua PermissionsChanged.",
        "Tab_AdminApproval tải danh sách Pending":
            "Tab_AdminApproval tải danh sách Pending và gửi quyết định sau khi người quản trị xác nhận. Endpoint duyệt yêu cầu "
            "REQUEST_APPROVE, endpoint từ chối yêu cầu REQUEST_REJECT; state machine chỉ cho phép chuyển từ Pending sang Approved "
            "hoặc Rejected. Người xử lý, thời điểm, lý do và log tương ứng được ghi trong transaction.",
        "Quy trình phân quyền bao gồm":
            "Quy trình phân quyền kiểm tra quan hệ nhóm cha và cập nhật mapping trang/component. Hệ thống chặn vòng lặp hoặc cây từ "
            "10 cấp, ghi audit, phát PermissionsChanged qua SignalR; frontend tải lại snapshot và API tiếp tục kiểm tra policy theo action.",
        "Giao diện hệ thống được tổ chức":
            "Giao diện hệ thống được tổ chức như ứng dụng quản trị nội bộ, dùng chung token thiết kế, kiểu chữ, khoảng cách, trạng thái "
            "rỗng/loading/error và responsive breakpoint. Dashboard, Library, Permission, Report cùng trung tâm thông báo dùng một "
            "bố cục/điều hướng nhất quán.",
        "Hệ thống đã có các vùng tổng hợp cá nhân":
            "Hệ thống có Dashboard tổng hợp cá nhân/phòng ban/toàn công ty và trang Report riêng. Report cho phép chọn phạm vi được cấp, "
            "năm/tháng, hiển thị KPI, xu hướng, trạng thái, phòng ban, mặt hàng nổi bật và xuất CSV khi có REPORT_EXPORT.",
        "Màn hình tổng hợp toàn doanh nghiệp cung cấp":
            "Màn hình tổng hợp toàn doanh nghiệp cung cấp KPI và danh sách đơn theo phòng ban, người tạo, kỳ, trạng thái. "
            "Trang Report tái sử dụng dữ liệu đã giới hạn theo claim để tạo biểu đồ và file CSV, không nhận mã phòng ban/công ty tùy ý từ client.",
        "Trong phạm vi đề tài, hệ thống báo biểu":
            "Hệ thống báo biểu đã có tổng hợp cá nhân/phòng ban/toàn công ty, KPI, xu hướng theo kỳ, phân bố trạng thái, top vật tư, "
            "phân bố phòng ban và xuất CSV UTF-8 (tối đa 50.000 dòng, chống CSV formula injection). Nhận định báo cáo chạy theo quy tắc; "
            "OpenAI là tùy chọn, mặc định tắt, chỉ nhận số liệu tổng hợp và có fallback khi lỗi.",
        "Hai bộ test tự động hiện tại":
            "Hai bộ test tự động hiện tại có tổng cộng 176 test (147 backend và 29 frontend), đều đạt ngày 13/07/2026. "
            "Solution Release build sạch; 12 UI test được discovery thành công nhưng chưa chạy E2E do thiếu môi trường/tài khoản/dữ liệu test cố định. "
            "Snapshot audit 54/54 ngày 28/05/2026 được ghi riêng, không tính vào kết quả hiện tại.",
        "Kết quả build hiện tại không còn xuất hiện":
            "Release build hiện tại đạt 0 lỗi, 0 cảnh báo; audit NuGet không phát hiện package dễ tổn thương theo các nguồn cấu hình. "
            "Các project xUnit v3 và Playwright đều build/discovery được, nhưng E2E xác thực vẫn cần môi trường tích hợp có dữ liệu kiểm thử.",
        "Cảnh báo NU1903:":
            "Phụ thuộc ngoài và AI: cảnh báo package cũ đã được xử lý và audit hiện không còn kết quả dễ tổn thương. Tuy vậy, "
            "nhận định OpenAI phụ thuộc dịch vụ ngoài, chi phí và khóa bí mật nên mặc định tắt, giới hạn 5 yêu cầu/10 phút/người dùng và luôn có fallback.",
        "Chức năng chưa hoàn thiện:":
            "Giới hạn báo cáo: CSV đã hoạt động nhưng chưa có mẫu Excel/PDF theo biểu mẫu doanh nghiệp, lịch gửi định kỳ hoặc lưu phiên bản báo cáo. "
            "Nhận định AI là hỗ trợ đọc số liệu, không tham gia phê duyệt hay thay đổi dữ liệu nghiệp vụ.",
        "Như vậy, các luồng nghiệp vụ cốt lõi":
            "Như vậy, các luồng nghiệp vụ cốt lõi, phân quyền theo action, thông báo bền vững và báo cáo theo phạm vi đã được hiện thực. "
            "Report có KPI, biểu đồ, top vật tư, CSV và nhận định tùy chọn; kết quả kiểm thử đủ chứng minh mức đề tài nhưng chưa thay thế "
            "kiểm thử tải, E2E tích hợp và đánh giá vận hành chính thức.",
        "Báo cáo và xuất dữ liệu:":
            "Báo cáo và xuất dữ liệu: Trang Report và CSV đã hoàn thành ở mức đề tài. Phần còn lại là mẫu Excel/PDF chính thức, "
            "lịch gửi, lưu snapshot báo cáo và kiểm thử với dữ liệu gần quy mô thật.",
        "An toàn phụ thuộc:":
            "Vận hành AI và secret: Không commit API key; production phải dùng secret manager, theo dõi chi phí/timeout/rate limit và "
            "giữ REPORT_INSIGHTS_ENABLED=false ở môi trường không cần AI.",
        "Tích hợp và thông báo:":
            "Tích hợp và thông báo: Hộp thư trong ứng dụng đã lưu bền vững và cập nhật qua SignalR cho đơn bổ sung/kết quả duyệt. "
            "Chưa có email, push notification, nhắc hạn gửi đơn/đóng kỳ hoặc tích hợp chính thức với hệ thống nhân sự/danh bạ.",
        "Hoàn thiện báo cáo:":
            "Mở rộng báo cáo: Bổ sung mẫu Excel/PDF theo chuẩn doanh nghiệp, lịch gửi, lưu snapshot, so sánh kỳ và dashboard ngân sách; "
            "đánh giá nhận định AI bằng bộ dữ liệu/prompt test trước khi bật production.",
        "Gia cố bảo mật và vận hành:":
            "Gia cố bảo mật và vận hành: Duy trì audit dependency/secret trong CI, quản lý secret theo môi trường, bổ sung giám sát log, "
            "cảnh báo, sao lưu và hướng dẫn rollback migration; kiểm thử định kỳ việc thu hồi quyền trên phiên đang mở.",
        "Tích hợp và mở rộng nghiệp vụ:":
            "Tích hợp và mở rộng nghiệp vụ: Đồng bộ người dùng/phòng ban với hệ thống nhân sự; mở rộng thông báo sang email/push và nhắc hạn; "
            "sau đó cân nhắc phê duyệt nhiều cấp, hạn mức ngân sách, so sánh nhà cung cấp và quản lý kho theo nhu cầu thực tế.",
        "Điều kiện: Người dùng có tài khoản hợp lệ":
            "Điều kiện: Người dùng có tài khoản hợp lệ, được gán đúng một nhóm quyền hoạt động và có phạm vi công ty/phòng ban. "
            "Menu, tab, action, Report và nút xuất chỉ xuất hiện khi snapshot quyền hiện tại cho phép.",
        "Điều kiện: Tài khoản có quyền REQUEST_ADMIN_APPROVAL.":
            "Điều kiện: Tài khoản có quyền giao diện xử lý đơn bổ sung; action duyệt cần REQUEST_APPROVE và action từ chối cần REQUEST_REJECT.",
    }
    for prefix, replacement in replacements.items():
        replace_by_prefix(doc, prefix, replacement)


def set_cell(cell, text: str) -> None:
    cell.text = text


def update_tables(doc: Document) -> None:
    technology = doc.tables[2]
    if not any("SignalR" in cell.text for row in technology.rows for cell in row.cells):
        row = technology.add_row().cells
        set_cell(row[0], "Thời gian thực")
        set_cell(row[1], "ASP.NET Core SignalR")
        set_cell(row[2], "Đồng bộ thay đổi quyền và hộp thư thông báo tới phiên đang mở.")
        row = technology.add_row().cells
        set_cell(row[0], "Nhận định tùy chọn")
        set_cell(row[1], "OpenAI Responses API [16], [17], [18]")
        set_cell(row[2], "Phân tích số liệu tổng hợp theo yêu cầu; mặc định tắt và có fallback quy tắc.")

    actors = doc.tables[3]
    set_cell(actors.rows[1].cells[2], "Đăng nhập; xem danh mục; tạo/sửa/hủy đơn; xem lịch sử; tạo đơn bổ sung; xem báo cáo cá nhân và hộp thư thông báo.")
    set_cell(actors.rows[2].cells[2], "Xem Dashboard/Report phòng ban trong phạm vi claim và nhận thông báo theo quyền.")
    set_cell(actors.rows[3].cells[2], "Quản lý dữ liệu nền, bảng giá, đơn bổ sung, đóng kỳ, phân quyền; xem báo cáo toàn công ty và xuất CSV.")
    set_cell(actors.rows[4].cells[2], "Tính kỳ, kiểm tra ràng buộc, chụp giá, ghi log, phát sự kiện quyền và thông báo.")

    physical = doc.tables[6]
    if not any("N01_Notification" in cell.text for row in physical.rows for cell in row.cells):
        row = physical.add_row().cells
        set_cell(row[0], "N01_Notification")
        set_cell(row[1], "Id uniqueidentifier")
        set_cell(row[2], "UserId; MemberCompanyCode; CorrelationId là tham chiếu logic, không FK")
        set_cell(row[3], "Type nvarchar(50); Title nvarchar(150); Message nvarchar(1000); Route nvarchar(250); CreatedAt/ReadAt datetime2; hai index hộp thư.")

    login = doc.tables[7]
    set_cell(login.rows[2].cells[1], "Xác thực qua stored procedure, tạo cookie, tải PermissionSnapshot và kết nối SignalR để giao diện cùng API phản ánh quyền hiện hành.")
    set_cell(login.rows[4].cells[1], "Thành công: JWT 24 giờ, cookie và snapshot quyền được tạo; hub quyền được kết nối. Thất bại: không tạo phiên.")
    set_cell(login.rows[5].cells[1], "Nhập thông tin; frontend gọi login; backend xác thực và phát JWT; frontend đổi ticket thành cookie, gọi /api/Auth/me/permissions, kết nối /hubs/permissions rồi mở tuyến đầu tiên được phép.")
    set_cell(login.rows[6].cells[1], "Sai dữ liệu/kết nối: hiển thị lỗi. Không có quyền: ẩn/chặn chức năng. Quyền bị thu hồi: nhận PermissionsChanged, tải lại snapshot; API action kế tiếp trả 403.")

    approval = doc.tables[11]
    set_cell(approval.rows[2].cells[1], "Cho phép người được cấp action REQUEST_APPROVE xem và duyệt đơn bổ sung Pending; action REQUEST_REJECT được kiểm tra riêng khi từ chối.")
    set_cell(approval.rows[3].cells[1], "Người dùng đã đăng nhập; giao diện có quyền xử lý đơn bổ sung; backend yêu cầu REQUEST_APPROVE khi duyệt hoặc REQUEST_REJECT khi từ chối. Đơn phải tồn tại, chưa xóa mềm, là đơn bổ sung và đang Pending.")
    set_cell(approval.rows[5].cells[1], "Quản trị viên mở danh sách Pending, xem chi tiết và chọn duyệt hoặc từ chối. Controller kiểm tra action policy tương ứng; service khóa thao tác trong transaction, kiểm tra chuyển trạng thái, ghi người và thời điểm xử lý, tạo nhật ký và thông báo rồi commit.")
    set_cell(approval.rows[6].cells[1], "Đơn không tồn tại, không phải đơn bổ sung hoặc không còn Pending: rollback và không cập nhật. Thiếu action tương ứng: API trả 403 kể cả khi người dùng gọi trực tiếp; frontend đồng thời ẩn hoặc khóa thao tác.")

    catalog = doc.tables[12]
    set_cell(catalog.rows[3].cells[1], "Người dùng đã đăng nhập; đọc danh mục cần LIBRARY_VIEW, còn tạo, sửa, xóa mềm hoặc khôi phục cần LIBRARY_MANAGE. Dữ liệu liên quan cần thiết đã tồn tại khi tạo ánh xạ.")
    set_cell(catalog.rows[5].cells[1], "Quản trị viên chọn tab danh mục, tải danh sách, thêm mới hoặc chỉnh sửa rồi lưu. Controller kiểm tra LIBRARY_VIEW/LIBRARY_MANAGE theo loại thao tác; service kiểm tra trường bắt buộc và quan hệ, sau đó thực hiện xóa mềm hoặc khôi phục khi hợp lệ.")
    set_cell(catalog.rows[6].cells[1], "Thiếu dữ liệu, khóa không tồn tại, bản ghi trùng hoặc vi phạm quan hệ: API trả lỗi và giao diện giữ dữ liệu để sửa. Thiếu action: API trả 403; quyền page/component chỉ điều khiển khả năng nhìn thấy và kích hoạt giao diện.")

    price = doc.tables[13]
    set_cell(price.rows[3].cells[1], "Người dùng đã đăng nhập; xem bảng giá cần LIBRARY_VIEW, còn tạo, sửa, sao chép, đặt mặc định hoặc xóa cần LIBRARY_MANAGE. Vật tư và nhà cung cấp liên quan phải đang hoạt động.")
    set_cell(price.rows[5].cells[1], "Quản trị viên tạo hoặc chọn bảng giá, cập nhật thông tin, sao chép nếu cần; sau đó quản lý giá theo vật tư và nhà cung cấp. Backend kiểm tra LIBRARY_MANAGE cho thao tác ghi, thực thể liên quan, giá không âm và ràng buộc duy nhất trước khi commit.")
    set_cell(price.rows[6].cells[1], "Không tìm thấy dữ liệu, giá âm hoặc vi phạm ràng buộc mặc định duy nhất: rollback và thông báo lỗi. Thiếu LIBRARY_MANAGE: API trả 403. Không được xóa mềm bảng giá mặc định; bảng giá có chi tiết chỉ xóa khi đáp ứng quy tắc dịch vụ.")

    settlement = doc.tables[14]
    set_cell(settlement.rows[3].cells[1], "Người dùng đã đăng nhập và có PERIOD_SETTLE; năm/tháng hợp lệ; có bảng giá áp dụng; không còn đơn bổ sung Pending và bảng giá có giá cho toàn bộ vật tư thuộc các đơn Submitted/Approved trong kỳ.")
    set_cell(settlement.rows[5].cells[1], "Quản trị viên chọn kỳ và bảng giá. Controller kiểm tra PERIOD_SETTLE cho xem trạng thái, xem trước và thực hiện đóng kỳ; service mở transaction, kiểm tra đơn Pending và đủ giá, chụp giá, đánh dấu settlement, ghi log rồi commit.")
    set_cell(settlement.rows[6].cells[1], "Thiếu PERIOD_SETTLE: API trả 403. Không có bảng giá, còn đơn Pending, thiếu giá vật tư hoặc kỳ không hợp lệ: rollback toàn bộ transaction; không để lại dữ liệu chụp giá một phần.")

    permission = doc.tables[15]
    set_cell(permission.rows[2].cells[1], "Quản lý nhóm quyền, gán người dùng vào nhóm/phòng ban và cấu hình page/component cùng action. Page/component điều khiển giao diện; action policy được backend kiểm tra cho từng thao tác quan trọng.")
    set_cell(permission.rows[3].cells[1], "Người dùng đã đăng nhập; đọc cấu hình cần PERMISSION_VIEW, còn tạo, sửa, xóa nhóm/mapping/người dùng cần PERMISSION_MANAGE. Giao diện Permission và component tương ứng phải được bật.")
    set_cell(permission.rows[4].cells[1], "Cấu hình được lưu và audit; backend phát PermissionsChanged theo nhóm/người dùng. Phiên đang mở tải lại snapshot, còn API action kế tiếp đối chiếu mapping mới ngay cả khi frontend chưa kịp cập nhật.")
    set_cell(permission.rows[5].cells[1], "Quản trị viên cập nhật nhóm cha, phòng ban, page/component hoặc action mapping. Backend kiểm tra PERMISSION_MANAGE, tính hợp lệ và quyền quản trị cuối cùng, lưu thay đổi rồi phát sự kiện SignalR để các phiên liên quan tải lại.")
    set_cell(permission.rows[6].cells[1], "Nhóm/người dùng/ánh xạ không tồn tại, tạo vòng lặp, cây từ 10 cấp hoặc làm mất quản trị phân quyền cuối cùng: từ chối cập nhật. Thiếu PERMISSION_MANAGE: API trả 403; frontend ẩn hoặc khóa chức năng.")

    summary = doc.tables[16]
    set_cell(summary.rows[0].cells[1], "Xem dashboard, báo cáo và dữ liệu tổng hợp")
    set_cell(summary.rows[2].cells[1], "Xem Dashboard và Report theo phạm vi cá nhân, phòng ban hoặc toàn công ty; xuất CSV và tạo nhận định khi được phép.")
    set_cell(summary.rows[3].cells[1], "Người dùng đã đăng nhập; có action xem phạm vi tương ứng. Xuất file cần REPORT_EXPORT.")
    set_cell(summary.rows[4].cells[1], "KPI, biểu đồ xu hướng/trạng thái/phòng ban, top vật tư và bảng tổng hợp hiển thị đúng claim; CSV chỉ chứa phạm vi được cấp.")
    set_cell(summary.rows[5].cells[1], "Chọn phạm vi, năm/tháng; frontend gọi ReportsController; backend đối chiếu action và claim rồi tổng hợp. Người dùng có thể xuất CSV hoặc yêu cầu nhận định quy tắc/AI tùy cấu hình.")
    set_cell(summary.rows[6].cells[1], "Không dữ liệu: trạng thái rỗng. Không quyền: 403/ẩn lựa chọn. AI tắt, thiếu khóa, timeout hoặc lỗi: trả nhận định quy tắc; lỗi tải báo cáo hiển thị thông báo.")

    layers = doc.tables[17]
    set_cell(layers.rows[1].cells[2], "Nghiệp vụ đơn/kỳ/bảng giá; action permission; phạm vi báo cáo/CSV; notification; AI fallback/structured output; controller và middleware.")
    set_cell(layers.rows[1].cells[3], "147/147 test đạt ngày 13/07/2026.")
    set_cell(layers.rows[2].cells[2], "URL/certificate, định dạng ngày giờ, localization parity, route và dịch vụ giao diện.")
    set_cell(layers.rows[2].cells[3], "29/29 test đạt ngày 13/07/2026.")
    set_cell(layers.rows[3].cells[2], "12 test cho đăng nhập, responsive/a11y, tạo đơn, lịch sử, tổng hợp, danh mục và thu hồi quyền tức thời.")
    set_cell(layers.rows[3].cells[3], "Discovery thành công; chưa chạy E2E vì thiếu môi trường FE/BE, tài khoản và dữ liệu test cố định.")

    cases = doc.tables[20]
    if not any(cell.text == "TC-14" for row in cases.rows for cell in row.cells):
        rows = [
            ("TC-14", "Có mapping action đang bật và phiên người dùng đang mở.", "Tắt quyền, kiểm tra snapshot/action tiếp theo và discovery test UI.", "Frontend làm mới; API từ chối quyền đã thu hồi mà không cần đăng nhập lại.", "PermissionService test đạt; test toggle realtime được discovery.", "Đạt"),
            ("TC-15", "Có dữ liệu đơn, quyền báo cáo và người nhận thông báo.", "Tổng hợp/xuất CSV; phát, đọc và đánh dấu thông báo.", "Đúng phạm vi claim; CSV chống formula injection; hộp thư bền vững và đúng người nhận.", "Các test report/notification đạt trong bộ 147 backend test.", "Đạt"),
            ("TC-16", "AI tắt; hoặc bật với HTTP giả lập thành công/lỗi 429.", "Tạo nhận định trong ba cấu hình.", "Không gọi API khi tắt; structured output khi thành công; fallback quy tắc khi lỗi.", "3/3 test ReportInsightService đạt; không dùng khóa/API thật.", "Đạt"),
        ]
        for values in rows:
            cells = cases.add_row().cells
            for cell, value in zip(cells, values):
                set_cell(cell, value)

    results = doc.tables[21]
    values = [
        ("Backend test", "147", "147", "0 / 0", "1 giây", "Chạy Release ngày 13/07/2026."),
        ("Frontend test", "29", "29", "0 / 0", "289 ms", "Chạy Release ngày 13/07/2026."),
        ("Build solution Release", "-", "0 lỗi", "0 cảnh báo", "3,25 giây", "Build toàn solution sau nâng cấp."),
        ("UI test discovery", "12", "12 được phát hiện", "Chưa chạy", "2,9 giây", "Cần môi trường/tài khoản/dữ liệu test để chạy E2E."),
        ("UI audit", "54", "54 OK", "0 lỗi audit", "-", "Snapshot 28/05/2026: 18 route × 3 viewport."),
    ]
    for row, row_values in zip(results.rows[1:], values):
        for cell, value in zip(row.cells, row_values):
            set_cell(cell, value)
    if not any("NuGet vulnerability audit" in cell.text for row in results.rows for cell in row.cells):
        row = results.add_row().cells
        for cell, value in zip(row, ("NuGet vulnerability audit", "11 project", "Không phát hiện", "0", "12,9 giây", "Dùng nguồn NuGet hiện được cấu hình ngày 13/07/2026.")):
            set_cell(cell, value)

    goals = doc.tables[22]
    set_cell(goals.rows[5].cells[2], "Dashboard và Report hiển thị KPI/biểu đồ theo cá nhân, phòng ban, toàn công ty; hỗ trợ xuất CSV và nhận định tùy chọn theo phạm vi quyền.")
    set_cell(goals.rows[5].cells[3], "Đạt")
    set_cell(goals.rows[7].cells[2], "Backend 147/147 và frontend 29/29 test đạt; solution Release 0 lỗi/0 cảnh báo; 12 UI test discovery thành công; có Docker Compose/Nginx. E2E chưa chạy lại.")


def resolve_target(target: str) -> str:
    target = target.replace("\\", "/")
    return posixpath.normpath(posixpath.join("word", target)).lstrip("/")


def replace_diagram_media(path: Path) -> dict[str, list[str]]:
    with zipfile.ZipFile(path, "r") as archive:
        members = {item.filename: archive.read(item.filename) for item in archive.infolist()}
        infos = {item.filename: item for item in archive.infolist()}

    parser = etree.XMLParser(remove_blank_text=False)
    document = etree.fromstring(members["word/document.xml"], parser)
    relationships = etree.fromstring(members["word/_rels/document.xml.rels"], parser)
    relation_targets = {
        rel.get("Id"): resolve_target(rel.get("Target", ""))
        for rel in relationships.findall("pr:Relationship", NS)
    }

    replacements: dict[str, list[str]] = {}
    last_embeds: list[str] = []
    for paragraph in document.xpath("./w:body/w:p", namespaces=NS):
        embeds = paragraph.xpath(".//*[@r:embed]/@r:embed", namespaces=NS)
        if embeds:
            last_embeds = list(dict.fromkeys(embeds))
            continue
        text = "".join(paragraph.xpath(".//w:t/text()", namespaces=NS)).strip()
        if not text.startswith("Hình ") or not last_embeds:
            continue
        figure_id = text.split(":", 1)[0].removeprefix("Hình ").strip()
        source_base = DIAGRAMS.get(figure_id)
        if source_base is None:
            last_embeds = []
            continue
        touched = []
        for relationship_id in last_embeds:
            target = relation_targets.get(relationship_id)
            if not target:
                continue
            suffix = Path(target).suffix.lower()
            source = source_base.with_suffix(suffix)
            if suffix not in {".svg", ".png"} or not source.exists():
                continue
            members[target] = source.read_bytes()
            touched.append(target)
        replacements[figure_id] = touched
        last_embeds = []

    missing = sorted(set(DIAGRAMS) - set(replacements))
    if missing:
        raise ValueError(f"Could not locate embedded media for figures: {missing}")
    incomplete = {key: value for key, value in replacements.items() if len(value) != 2}
    if incomplete:
        raise ValueError(f"Expected SVG and PNG fallback for each diagram: {incomplete}")

    fd, temp_name = tempfile.mkstemp(suffix=".docx", dir=path.parent)
    os.close(fd)
    temp_path = Path(temp_name)
    try:
        with zipfile.ZipFile(temp_path, "w") as output:
            for name, data in members.items():
                output.writestr(infos[name], data)
        with zipfile.ZipFile(temp_path, "r") as test_archive:
            bad = test_archive.testzip()
            if bad:
                raise ValueError(f"Invalid DOCX member after media replacement: {bad}")
        os.replace(temp_path, path)
    finally:
        temp_path.unlink(missing_ok=True)
    return replacements


def build(source: Path, output: Path) -> None:
    doc = Document(source)
    update_paragraphs(doc)
    update_tables(doc)
    add_ai_sources(doc)
    add_appendix_guidance(doc)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    replaced = replace_diagram_media(output)
    print(f"OUTPUT={output}")
    print(f"PARAGRAPHS={len(doc.paragraphs)} TABLES={len(doc.tables)}")
    print(f"UPDATED_DIAGRAMS={len(replaced)}")
    for figure_id, targets in replaced.items():
        print(f"- Hình {figure_id}: {', '.join(targets)}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build(args.source.resolve(), args.output.resolve())


if __name__ == "__main__":
    main()
