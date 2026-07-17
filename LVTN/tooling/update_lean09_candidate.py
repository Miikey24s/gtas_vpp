from __future__ import annotations

import argparse
import os
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PR = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"w": W, "r": R, "pr": PR}


def tidy(value: str) -> str:
    return " ".join(value.split())


def replace_paragraph_text(doc: Document, old: str, new: str) -> int:
    changed = 0
    for paragraph in doc.paragraphs:
        if tidy(paragraph.text) != old:
            continue
        if paragraph._p.xpath(".//w:bookmarkStart | .//w:fldChar | .//w:hyperlink"):
            text_nodes = paragraph._p.xpath(".//w:t")
            combined = "".join(node.text or "" for node in text_nodes)
            if old not in combined:
                raise ValueError(f"Could not preserve markup while replacing {old!r}")
            replaced = False
            remaining = combined
            for node in text_nodes:
                if not replaced:
                    node.text = new
                    replaced = True
                else:
                    node.text = ""
            if not replaced:
                raise ValueError(f"No text node found while replacing {old!r}")
        else:
            paragraph.text = new
        changed += 1
    return changed


def replace_paragraph_prefix(doc: Document, prefix: str, new: str) -> int:
    changed = 0
    for paragraph in doc.paragraphs:
        if not tidy(paragraph.text).startswith(prefix):
            continue
        paragraph.text = new
        changed += 1
    return changed


def set_cell(table, row: int, col: int, value: str) -> None:
    table.rows[row].cells[col].text = value


def update_docx_text(path: Path) -> dict[str, int]:
    doc = Document(path)
    changed: dict[str, int] = {}

    paragraph_updates = {
        "Màn hình đăng nhập chia bố cục thành vùng nhận diện PPJ Group và biểu mẫu xác thực. Người dùng nhập tên đăng nhập, mật khẩu; lựa chọn môi trường chỉ xuất hiện khi hệ thống chạy ở chế độ phát triển.":
            "Màn hình đăng nhập chia bố cục thành vùng nhận diện GTAS VPP và biểu mẫu xác thực. Người dùng nhập tên đăng nhập, mật khẩu; lựa chọn môi trường chỉ xuất hiện khi hệ thống chạy ở chế độ phát triển.",
        "VPP01_RequestHeader lưu kỳ, trạng thái và thông tin duyệt; VPP02_RequestDetail lưu vật tư, số lượng và đơn giá chụp; VPP03_Log lưu lịch sử thao tác dạng JSON. N01_Notification lưu hộp thư theo người dùng/công ty, loại, nội dung, tuyến mở, CorrelationId, thời điểm tạo và đọc. Các mã tổ chức/người dùng là tham chiếu logic, không phải khóa ngoại vật lý.":
            "VPP00_Period lưu biên kỳ Việt Nam từ ngày 05 đến trước ngày 05 tháng sau và state machine Open → SubmissionClosed → Pricing → Settled. VPP01_RequestHeader và VPP02_RequestDetail lưu series/revision, lineage, đơn thường/đơn bổ sung, rowversion và giá tại thời điểm hợp lệ; VPP03_Log lưu audit có actor, action, correlation và lý do. VPP04_Settlement, VPP05_SettlementItem, VPP06_SettlementCharge và VPP07_SettlementAllocation lưu snapshot bất biến của nhà cung cấp, giá net/VAT/gross, phí và phân bổ; N01_Notification và N02_EmailOutbox phục vụ hộp thư bền vững và email sandbox.",
        "Migration hiện tạo 18 bảng, gồm N01_Notification được bổ sung cho hộp thư bền vững. Phần lớn entity kế thừa BaseModel; các bảng ánh xạ P05, P06, VPP03_Log và N01_Notification có cấu trúc riêng. Hai view v_Users và v_WFXCompany được ánh xạ keyless để đọc dữ liệu tích hợp, không thuộc migration của hệ thống.":
            "Các migration hiện tại tạo thêm aggregate kỳ, revision đơn, catalog/bảng giá, snapshot settlement VPP04–VPP07, N01_Notification và N02_EmailOutbox. Các migration settlement/outbox là forward-only; rollback dùng backup/restore hoặc forward correction. Hai view v_Users và v_WFXCompany được ánh xạ keyless để đọc dữ liệu tích hợp, không thuộc migration của hệ thống.",
        "3.2.2.5 Đóng kỳ và chụp giá": "3.2.2.5 Xác nhận settlement bất biến và hiệu chỉnh revision",
        "PeriodSettlementService kiểm tra bảng giá, đơn bổ sung Pending và độ phủ giá trước khi cập nhật. Việc chụp CurrentSinglePrice, đánh dấu settlement và ghi PERIOD_SETTLED được thực hiện nguyên tử; bất kỳ lỗi nào cũng rollback toàn bộ kỳ.":
            "PeriodSettlementService trước hết tạo preview có input hash, blocker, nhà cung cấp chính và phân bổ toàn công ty. Lệnh confirm kiểm tra lại hash/idempotency, chụp VPP04–VPP07 (supplier, item price net/VAT/gross, charge, allocation), reconcile tổng tiền rồi chuyển Pricing → Settled trong transaction. Correction không ghi đè snapshot cũ mà tạo revision mới, bắt buộc reason và four-eyes.",
        "Hình 3-19: Sơ đồ tuần tự chức năng đóng kỳ và chụp giá": "Hình 3-19: Sơ đồ tuần tự xác nhận settlement bất biến và hiệu chỉnh revision",
        "3.2.3.3 Đóng kỳ": "3.2.3.3 Xác nhận settlement và hiệu chỉnh",
        "Đóng kỳ chỉ hoàn tất khi không còn đơn bổ sung chờ duyệt và bảng giá bao phủ toàn bộ vật tư của các đơn Submitted/Approved. Các bước cập nhật được bao trong một transaction.":
            "Preview phải hết blocker và reconcile được tổng nhu cầu/phân bổ. Confirm tạo snapshot bất biến VPP04–VPP07, bảo vệ idempotency và chuyển kỳ sang Settled; correction tạo revision mới, giữ nguyên dữ liệu revision cũ và yêu cầu four-eyes.",
        "Hình 3-23: Sơ đồ hoạt động quy trình đóng kỳ": "Hình 3-23: Sơ đồ hoạt động quy trình xác nhận settlement và hiệu chỉnh",
        "Hệ thống có Dashboard tổng hợp cá nhân/phòng ban/toàn công ty và trang Report riêng. Report cho phép chọn phạm vi được cấp, năm/tháng, hiển thị KPI, xu hướng, trạng thái, phòng ban, mặt hàng nổi bật và xuất CSV khi có REPORT_EXPORT.":
            "Hệ thống có Dashboard tổng hợp cá nhân/phòng ban/toàn công ty và trang Report riêng. Report giới hạn theo claim, đối chiếu allocation snapshot của settlement đã Settled, hiển thị KPI/xu hướng/trạng thái/phòng ban/top vật tư, xuất CSV và tải workbook XLSX gồm Summary, Items, Departments, Trend và TopProducts khi có REPORT_EXPORT.",
        "Hệ thống báo biểu đã có tổng hợp cá nhân/phòng ban/toàn công ty, KPI, xu hướng theo kỳ, phân bố trạng thái, top vật tư, phân bố phòng ban và xuất CSV UTF-8 (tối đa 50.000 dòng, chống CSV formula injection). Nhận định báo cáo chạy theo quy tắc; OpenAI là tùy chọn, mặc định tắt, chỉ nhận số liệu tổng hợp và có fallback khi lỗi.":
            "Hệ thống báo biểu có KPI theo scope, settlement reconciliation, CSV UTF-8 chống formula injection và workbook XLSX an toàn với các sheet Summary, Items, Departments, Trend và TopProducts. Hộp thư N01_Notification có dedupe/idempotency; N02_EmailOutbox enqueue bền vững, retry exponential và SMTP/Mailpit disabled-by-default. Nhận định AI vẫn tùy chọn và mặc định tắt.",
        "Hai bộ test tự động hiện tại có tổng cộng 176 test (147 backend và 29 frontend), đều đạt ngày 13/07/2026. Solution Release build sạch; 12 UI test được discovery thành công nhưng chưa chạy E2E do thiếu môi trường/tài khoản/dữ liệu test cố định. Snapshot audit 54/54 ngày 28/05/2026 được ghi riêng, không tính vào kết quả hiện tại.":
            "Đến 17/07/2026, backend Release có 384 test đạt, frontend Release có 96 test đạt, integration có 14 test đạt và 6 skip theo LocalDB opt-in. UI Playwright chạy trên Aspire/LocalDB fixture cô lập hai rehearsal độc lập: lần một 16/16 pass trong 427.0 giây, lần hai 16/16 pass trong 406.9 giây; không dùng credential/provider production.",
        "Release build hiện tại đạt 0 lỗi, 0 cảnh báo; audit NuGet không phát hiện package dễ tổn thương theo các nguồn cấu hình. Các project xUnit v3 và Playwright đều build/discovery được, nhưng E2E xác thực vẫn cần môi trường tích hợp có dữ liệu kiểm thử.":
            "Solution Release build đạt 0 warning/0 error; EF pending-model clean; gitleaks current-tree clean; UI syntax smoke 2/2 pass. Hai isolated Playwright rehearsal đều có summary hoàn chỉnh 16 pass, 0 fail, 0 skip.",
        "Đơn bổ sung và đóng kỳ: Hệ thống chặn settlement khi còn đơn Pending hoặc bảng giá thiếu mặt hàng; giao dịch không được xác nhận khi validation thất bại.":
            "Đơn bổ sung và settlement: hệ thống chặn confirm khi còn Pending, thiếu giá/nhà cung cấp, input hash thay đổi hoặc allocation không reconcile; correction chỉ đi qua revision mới có reason và four-eyes.",
        "Giới hạn bằng chứng UI: Kết quả 54/54 là snapshot ngày 28/05/2026, không phải lần chạy ngày 13/07/2026. Cần chạy lại audit và Playwright E2E trên bản sắp bàn giao với môi trường test cố định.":
            "Bằng chứng UI hiện tại gồm hai lần full isolated Playwright suite 16/16 pass trên fixture Aspire/LocalDB. Snapshot 54/54 ngày 28/05/2026 vẫn được giữ như bằng chứng lịch sử, không dùng thay cho hai lần rehearsal mới.",
        "Phạm vi dữ liệu: Backend test hiện chủ yếu dùng EF Core InMemory/SQLite; chưa có kiểm thử tải dài hạn và bộ integration test đầy đủ trên SQL Server với dữ liệu gần quy mô thực tế.":
            "Phạm vi dữ liệu: integration Release có 14 pass và 6 skip theo gate LocalDB hiện hữu; kiểm thử tải dài hạn, dữ liệu 10k và backup/restore server vẫn nằm ngoài bằng chứng local hiện tại.",
        "Giới hạn báo cáo: CSV đã hoạt động nhưng chưa có mẫu Excel/PDF theo biểu mẫu doanh nghiệp, lịch gửi định kỳ hoặc lưu phiên bản báo cáo. Nhận định AI là hỗ trợ đọc số liệu, không tham gia phê duyệt hay thay đổi dữ liệu nghiệp vụ.":
            "Giới hạn báo cáo: workbook XLSX và email outbox/retry đã có trong local RC; PDF, lịch gửi định kỳ và provider email thật vẫn deferred. Nhận định AI chỉ hỗ trợ đọc số liệu, không tham gia phê duyệt hay thay đổi dữ liệu nghiệp vụ.",
        "Như vậy, các luồng nghiệp vụ cốt lõi, phân quyền theo action, thông báo bền vững và báo cáo theo phạm vi đã được hiện thực. Report có KPI, biểu đồ, top vật tư, CSV và nhận định tùy chọn; kết quả kiểm thử đủ chứng minh mức đề tài nhưng chưa thay thế kiểm thử tải, E2E tích hợp và đánh giá vận hành chính thức.":
            "Như vậy, các luồng cốt lõi từ kỳ–đơn–bổ sung đến catalog–bảng giá–settlement bất biến, báo cáo reconciliation, workbook, inbox và email sandbox đã có source/test/evidence tương ứng. Hai rehearsal UI đã pass; backup/restore owner và Word handoff vẫn là ranh giới bàn giao.",
        "Báo cáo và xuất dữ liệu: Trang Report và CSV đã hoàn thành ở mức đề tài. Phần còn lại là mẫu Excel/PDF chính thức, lịch gửi, lưu snapshot báo cáo và kiểm thử với dữ liệu gần quy mô thật.":
            "Báo cáo và xuất dữ liệu: Report, CSV và workbook XLSX đã hoàn thành ở mức local RC; PDF, lịch gửi và kiểm thử dữ liệu gần quy mô thật vẫn deferred.",
        "Bằng chứng kiểm thử giao diện: Project Playwright build thành công nhưng chưa chạy E2E ngày 13/07/2026 vì cần frontend, backend, tài khoản và dữ liệu test đồng bộ. Kết quả 54/54 lượt audit là snapshot ngày 28/05/2026, không phải lần chạy hiện tại.":
            "Bằng chứng kiểm thử giao diện: hai lần full isolated Playwright suite đều 16/16 pass, 0 fail, 0 skip; snapshot 54/54 ngày 28/05/2026 chỉ là historical audit.",
        "Tích hợp và thông báo: Hộp thư trong ứng dụng đã lưu bền vững và cập nhật qua SignalR cho đơn bổ sung/kết quả duyệt. Chưa có email, push notification, nhắc hạn gửi đơn/đóng kỳ hoặc tích hợp chính thức với hệ thống nhân sự/danh bạ.":
            "Tích hợp và thông báo: inbox bền vững/idempotent và email outbox/retry sandbox đã có; provider email thật, push, nhắc hạn và tích hợp HR/danh bạ vẫn deferred.",
        "Mở rộng báo cáo: Bổ sung mẫu Excel/PDF theo chuẩn doanh nghiệp, lịch gửi, lưu snapshot, so sánh kỳ và dashboard ngân sách; đánh giá nhận định AI bằng bộ dữ liệu/prompt test trước khi bật production.":
            "Mở rộng báo cáo: hoàn thiện PDF/lịch gửi/snapshot và probe dữ liệu lớn; tiếp tục giữ AI mặc định tắt và đánh giá bằng bộ dữ liệu/prompt test trước khi bật production.",
        "Nâng độ tin cậy kiểm thử: Xây dựng môi trường test cố định, chạy Playwright E2E và UI audit trong CI; bổ sung integration test trên SQL Server, kiểm thử tải, dữ liệu lớn và các tình huống đồng thời.":
            "Nâng độ tin cậy kiểm thử: giữ fixture cô lập, đưa hai rehearsal vào CI định kỳ, bổ sung probe dữ liệu lớn và hoàn tất integration/backup-restore theo quyền owner.",
        "Gia cố bảo mật và vận hành: Duy trì audit dependency/secret trong CI, quản lý secret theo môi trường, bổ sung giám sát log, cảnh báo, sao lưu và hướng dẫn rollback migration; kiểm thử định kỳ việc thu hồi quyền trên phiên đang mở.":
            "Gia cố bảo mật và vận hành: duy trì audit dependency/secret, correlation/log scrubbing, forward correction cho migration; backup/restore server cần owner-authorized rehearsal trước khi bàn giao.",
    }
    for old, new in paragraph_updates.items():
        count = replace_paragraph_text(doc, old, new)
        if count:
            changed[old[:40]] = count

    auth_claim_updates = {
        "- Backend mã hóa mật khẩu và gọi thủ tục lưu trữ (Stored Procedure) để xác thực tài khoản trong cơ sở dữ liệu.":
            "- Backend dùng ASP.NET Core Identity/UserManager để kiểm tra mật khẩu, trạng thái tài khoản và khóa đăng nhập; sau đó tạo JWT cùng PermissionSnapshot hiện hành.",
        "Xác thực qua stored procedure, tạo cookie, tải PermissionSnapshot và kết nối SignalR để giao diện cùng API phản ánh quyền hiện hành.":
            "Xác thực qua ASP.NET Core Identity/UserManager, phát hành JWT, tải PermissionSnapshot và kết nối SignalR để giao diện cùng API phản ánh quyền hiện hành.",
    }
    for old, new in auth_claim_updates.items():
        count = replace_paragraph_text(doc, old, new)
        if count:
            changed[old[:40]] = count
    count = replace_paragraph_prefix(
        doc,
        "Xác thực qua stored procedure",
        "Xác thực qua ASP.NET Core Identity/UserManager, phát hành JWT, tải PermissionSnapshot và kết nối SignalR để giao diện cùng API phản ánh quyền hiện hành.",
    )
    if count:
        changed["Xác thực qua stored procedure"] = count

    tables = doc.tables
    set_cell(
        tables[7],
        2,
        1,
        "Xác thực qua ASP.NET Core Identity/UserManager, phát hành JWT, tải PermissionSnapshot và kết nối SignalR để giao diện cùng API phản ánh quyền hiện hành.",
    )
    # Physical model and settlement use-case tables.
    set_cell(tables[6], 1, 0, "VPP00_Period / VPP01_RequestHeader")
    set_cell(tables[6], 1, 2, "PeriodId; Supersedes/Current revision; SettlementId")
    set_cell(tables[6], 1, 3, "Vietnam period; series/revision/lineage; rowversion; idempotency/hash; state transition guards.")
    set_cell(tables[6], 2, 3, "Qty; requested/snapshot values; immutable revision lineage; no destructive overwrite.")
    if len(tables[6].rows) < 8:
        for values in [
            ("VPP04_Settlement", "Id uniqueidentifier", "VPP00_Period; Supplier; PriceList", "Revision/current marker; input hash; calculation version; totals; confirmed/corrected timestamps."),
            ("VPP05–VPP07 snapshots", "Id uniqueidentifier", "VPP04_Settlement; request/department", "Item net/VAT/gross; charges; allocations; reconciliation and immutable correction lineage."),
        ]:
            cells = tables[6].add_row().cells
            for cell, value in zip(cells, values):
                cell.text = value

    settlement = tables[14]
    set_cell(settlement, 0, 1, "Xác nhận settlement bất biến và hiệu chỉnh revision")
    set_cell(settlement, 2, 1, "Preview toàn công ty, chọn nhà cung cấp/bảng giá, confirm snapshot VPP04–VPP07; correction tạo revision mới.")
    set_cell(settlement, 3, 1, "Đã đăng nhập và có PERIOD_SETTLE; preview hash còn nguyên; không còn Pending; giá/nhà cung cấp đủ; totals/allocation reconcile.")
    set_cell(settlement, 4, 1, "VPP04–VPP07 bất biến; period Pricing → Settled; correction giữ revision cũ, có reason và four-eyes.")
    set_cell(settlement, 5, 1, "Preview → confirm với idempotency; service mở transaction, kiểm tra hash và reconcile rồi commit snapshot.")
    set_cell(settlement, 6, 1, "Thiếu quyền/hash/giá/supplier, còn Pending hoặc variance khác 0: không commit. Correction sai reason/four-eyes: 400/403.")

    layers = tables[17]
    set_cell(layers, 1, 3, "384/384 Release; integration 14 pass/6 skip; settlement/report/notification/outbox tests.")
    set_cell(layers, 2, 3, "96/96 Release; localization/route/async/report UI helpers.")
    set_cell(layers, 3, 2, "16 authenticated Playwright tests trên fixture cô lập; hai rehearsal serial, mutation opt-in QA-only.")
    set_cell(layers, 3, 3, "Rehearsal #1: 16/16 pass; rehearsal #2: 16/16 pass; 0 fail/skip mỗi lần.")

    cases = tables[20]
    set_cell(cases, 3, 4, "Hai full isolated rehearsal 16/16 pass; snapshot cũ chỉ giữ làm historical reference.")
    set_cell(cases, 3, 5, "Đạt")
    set_cell(cases, 5, 4, "Report/workbook/inbox/outbox tests đạt; XLSX ZIP có đủ 5 sheet và dedupe idempotency.")
    set_cell(cases, 6, 4, "AI tắt mặc định; fallback rule-based; không dùng API key thật.")

    results = tables[21]
    rows = [
        ("Backend Release test", "384", "384", "0 / 0", "local", "Release run sau LEAN-07/08."),
        ("Frontend Release test", "96", "96", "0 / 0", "local", "Release run sau LEAN-07/08."),
        ("Integration Release", "20", "14", "0 / 6 skip", "local", "Existing LocalDB opt-in gate."),
        ("Solution Release build", "-", "0 lỗi", "0 cảnh báo", "local", "Build solution sạch."),
        ("UI rehearsal #1", "16", "16", "0 / 0", "427.0 s", "Aspire/LocalDB isolated fixture."),
        ("UI rehearsal #2", "16", "16", "0 / 0", "406.9 s", "Clean isolated fixture, mutation QA-only."),
    ]
    for row, values in zip(results.rows[1:], rows):
        for cell, value in zip(row.cells, values):
            cell.text = value
    if len(results.rows) < 8:
        cells = results.add_row().cells
        for cell, value in zip(cells, ("Release recovery gates", "-", "EF clean", "gitleaks clean", "local", "Protected hashes preserved; owner backup pending.")):
            cell.text = value

    goals = tables[22]
    set_cell(goals, 5, 2, "Dashboard/Report có KPI, settlement reconciliation, CSV và workbook XLSX theo cá nhân/phòng ban/toàn công ty.")
    set_cell(goals, 6, 2, "Confirm tạo VPP04–VPP07 snapshots; correction là revision mới có reason/four-eyes; totals/allocation reconcile.")
    set_cell(goals, 7, 2, "Backend 384, frontend 96, integration 14/6; UI 16/16 pass ở hai rehearsal; build/EF/gitleaks clean. Word source freeze còn owner action.")

    doc.save(path)
    return changed


def replace_settlement_media(path: Path, source_dir: Path) -> None:
    members: dict[str, bytes]
    infos: dict[str, zipfile.ZipInfo]
    with zipfile.ZipFile(path, "r") as archive:
        members = {item.filename: archive.read(item.filename) for item in archive.infolist()}
        infos = {item.filename: item for item in archive.infolist()}

    parser = etree.XMLParser(remove_blank_text=False)
    document = etree.fromstring(members["word/document.xml"], parser)
    rels = etree.fromstring(members["word/_rels/document.xml.rels"], parser)
    relation_targets = {
        rel.get("Id"): (rel.get("Target") or "")
        for rel in rels.findall("pr:Relationship", NS)
    }
    paragraphs = document.xpath(".//w:p", namespaces=NS)
    matching_captions = []
    for paragraph in paragraphs:
        text = "".join(paragraph.xpath(".//w:t/text()", namespaces=NS))
        if "Hình 3-19:" in text and "Sơ đồ tuần tự" in text:
            matching_captions.append(paragraph)
    target_paragraph = matching_captions[-1] if matching_captions else None
    if target_paragraph is None:
        raise ValueError("Hình 3-19 caption not found")

    index = paragraphs.index(target_paragraph)
    image_paragraph = None
    embeds: list[str] = []
    for paragraph in reversed(paragraphs[:index]):
        embeds = list(dict.fromkeys(paragraph.xpath(".//*[@r:embed]/@r:embed", namespaces=NS)))
        if embeds:
            image_paragraph = paragraph
            break
    if image_paragraph is None or len(embeds) != 2:
        raise ValueError(f"Expected two media relationships before Hình 3-19, got {embeds}")

    for rid in embeds:
        target = relation_targets.get(rid, "")
        target_name = target.lstrip("/").replace("\\", "/")
        if not target_name.startswith("word/"):
            target_name = "word/" + target_name
        suffix = Path(target_name).suffix.lower()
        source = source_dir / f"sequence-settle-period{suffix}"
        if suffix not in {".svg", ".png"} or not source.exists():
            raise ValueError(f"Missing settlement diagram source for {target_name}")
        members[target_name] = source.read_bytes()

    old_caption = "Hình 3-19: Sơ đồ tuần tự chức năng đóng kỳ và chụp giá"
    new_caption = "Hình 3-19: Sơ đồ tuần tự xác nhận settlement bất biến và hiệu chỉnh revision"
    for node in document.xpath(".//w:t", namespaces=NS):
        if node.text and old_caption in node.text:
            node.text = node.text.replace(old_caption, new_caption)

    old_figure_23 = "Hình 3-23: Sơ đồ hoạt động quy trình đóng kỳ"
    new_figure_23 = "Hình 3-23: Sơ đồ hoạt động quy trình xác nhận settlement và hiệu chỉnh"
    for node in document.xpath(".//w:t", namespaces=NS):
        if node.text and old_figure_23 in node.text:
            node.text = node.text.replace(old_figure_23, new_figure_23)

    for node in document.xpath(".//wp:docPr", namespaces={**NS, "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"}):
        if not (node.get("descr") or node.get("title")):
            node.set("descr", "Decorative corner on the acknowledgment page")

    temp = path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(temp, "w") as output:
        for name, data in members.items():
            if name == "word/document.xml":
                data = etree.tostring(document, xml_declaration=True, encoding="UTF-8", standalone="yes")
            output.writestr(infos[name], data)
    with zipfile.ZipFile(temp, "r") as check:
        if check.testzip():
            raise ValueError("Candidate DOCX ZIP failed integrity check")
    os.replace(temp, path)


def replace_thesis_screenshot_media(path: Path, source_dir: Path) -> None:
    screenshot_by_bookmark = {
        "fig_3_26": "ui-login.png",
        "fig_3_27": "ui-dashboard-my-orders.png",
        "fig_3_28": "ui-order-history.png",
        "fig_3_29": "ui-order-create.png",
        "fig_3_30": "ui-period-operations.png",
        "fig_3_31": "ui-price-lists.png",
        "fig_3_32": "ui-library-items.png",
        "fig_3_33": "ui-permission-groups.png",
        "fig_3_34": "ui-all-orders-summary.png",
    }
    members: dict[str, bytes]
    infos: dict[str, zipfile.ZipInfo]
    with zipfile.ZipFile(path, "r") as archive:
        members = {item.filename: archive.read(item.filename) for item in archive.infolist()}
        infos = {item.filename: item for item in archive.infolist()}

    parser = etree.XMLParser(remove_blank_text=False)
    document = etree.fromstring(members["word/document.xml"], parser)
    rels = etree.fromstring(members["word/_rels/document.xml.rels"], parser)
    relation_targets = {
        rel.get("Id"): rel.get("Target") or ""
        for rel in rels.findall("pr:Relationship", NS)
    }
    paragraphs = document.xpath(".//w:p", namespaces=NS)
    replaced = 0
    for bookmark, filename in screenshot_by_bookmark.items():
        caption_paragraph = next(
            (
                paragraph
                for paragraph in paragraphs
                if bookmark
                in paragraph.xpath(".//w:bookmarkStart/@w:name", namespaces=NS)
            ),
            None,
        )
        if caption_paragraph is None:
            raise ValueError(f"Missing thesis screenshot bookmark: {bookmark}")

        index = paragraphs.index(caption_paragraph)
        embeds: list[str] = []
        for paragraph in reversed(paragraphs[:index]):
            embeds = list(dict.fromkeys(paragraph.xpath(".//*[@r:embed]/@r:embed", namespaces=NS)))
            if embeds:
                break
        if len(embeds) != 1:
            raise ValueError(f"Expected one image relationship for {bookmark}, got {embeds}")

        target_name = relation_targets.get(embeds[0], "").lstrip("/").replace("\\", "/")
        if not target_name.startswith("word/"):
            target_name = "word/" + target_name
        source = source_dir / filename
        if not source.exists():
            raise ValueError(f"Missing thesis screenshot source: {source}")
        members[target_name] = source.read_bytes()
        replaced += 1

    temp = path.with_suffix(".screenshots.tmp.docx")
    with zipfile.ZipFile(temp, "w") as output:
        for name, data in members.items():
            if name == "word/document.xml":
                data = etree.tostring(document, xml_declaration=True, encoding="UTF-8", standalone="yes")
            output.writestr(infos[name], data)
    with zipfile.ZipFile(temp, "r") as check:
        if check.testzip():
            raise ValueError("Candidate DOCX ZIP failed screenshot replacement integrity check")
    os.replace(temp, path)
    print(f"thesis_screenshot_media_replaced={replaced}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_bytes(args.source.read_bytes())
    changed = update_docx_text(args.output)
    replace_settlement_media(args.output, Path(__file__).resolve().parent.parent / "diagrams" / "ch03")
    replace_thesis_screenshot_media(args.output, Path(__file__).resolve().parent.parent / "screenshots" / "ch03")
    print(f"output={args.output}")
    print(f"paragraph_updates={len(changed)}")
    print(f"updated_prefixes={sum(changed.values())}")


if __name__ == "__main__":
    main()
