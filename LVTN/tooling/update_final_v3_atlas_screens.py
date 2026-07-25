"""Rewrite only thesis section 3.3 with Atlas-backed screen designs.

The current DOCX remains the business-content authority. Backend code is used
only to validate permissions, states and invariants; the legacy frontend is not
used as the visual authority. The section follows the document's standard A4
portrait layout; screenshots are scaled to the usable page width.
"""

from __future__ import annotations

import argparse
import copy
import os
import shutil
import tempfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from PIL import Image


DEFAULT_DOCX = Path("LVTN/checkpoints/NguyenAnNam_DH52201078_final_v3.docx")
SCREEN_ROOT = Path("LVTN/screenshots/ch03/atlas")
MAX_IMAGE_WIDTH_CM = 15.8
MAX_IMAGE_HEIGHT_CM = 9.0


GROUPS = [
    (
        "3.3.1 Giao diện chung",
        [
            (
                "3.3.1.1 Đăng nhập hệ thống",
                "Màn hình đăng nhập thu thập tên đăng nhập và mật khẩu, hỗ trợ khôi phục mật khẩu và chỉ hiển thị thông báo lỗi an toàn. Sau khi xác thực thành công, hệ thống tải nhóm quyền và phạm vi dữ liệu tương ứng với tài khoản.",
                "01-login.png",
                "Thiết kế giao diện đăng nhập hệ thống",
            ),
            (
                "3.3.1.2 Bảng điều khiển cá nhân (dashboard)",
                "Bảng điều khiển cá nhân trình bày kỳ đang nhận đơn, hạn gửi, đơn thông thường, đơn bổ sung và đơn kỳ trước. Các thao tác sửa, hủy, xem lịch sử hoặc tạo đơn chỉ xuất hiện khi người dùng đáp ứng điều kiện nghiệp vụ.",
                "02-my-orders.png",
                "Thiết kế giao diện bảng điều khiển cá nhân",
            ),
        ],
    ),
    (
        "3.3.2 Giao diện nghiệp vụ nhân viên",
        [
            (
                "3.3.2.1 Tạo đơn thông thường",
                "Luồng tạo đơn được tổ chức thành hai bước: chọn mặt hàng và kiểm tra trước khi gửi. Người dùng có thể tìm kiếm, lọc danh mục, nhập số lượng và ghi chú; hệ thống kiểm tra kỳ, mặt hàng được phép đặt và dữ liệu bắt buộc trước khi ghi nhận đơn.",
                "03-order-create.png",
                "Thiết kế giao diện tạo đơn thông thường",
            ),
            (
                "3.3.2.2 Lịch sử và theo dõi đơn",
                "Màn hình lịch sử hỗ trợ lọc theo kỳ, loại đơn và trạng thái; đồng thời trình bày số liệu tổng hợp, chi tiết đơn và các phiên bản đã phát sinh. Người dùng có thể theo dõi quá trình sửa, hủy hoặc xử lý đơn bổ sung mà không làm mất dữ liệu lịch sử.",
                "04-history.png",
                "Thiết kế giao diện lịch sử và theo dõi đơn",
            ),
            (
                "3.3.2.3 Tra cứu danh mục mặt hàng",
                "Danh mục mặt hàng chỉ hiển thị dữ liệu người dùng được phép đặt. Bộ lọc theo tên, mã, danh mục và đơn vị tính giúp tra cứu nhanh; trạng thái áp dụng giúp phân biệt mặt hàng đang sử dụng với mặt hàng bị hạn chế.",
                "05-catalog.png",
                "Thiết kế giao diện tra cứu danh mục mặt hàng",
            ),
        ],
    ),
    (
        "3.3.3 Giao diện nghiệp vụ quản lý và vận hành kỳ",
        [
            (
                "3.3.3.1 Tổng hợp đơn theo phòng ban",
                "Quản lý theo dõi mức độ hoàn tất đơn trong phòng ban, xem số liệu theo nhân sự và mở chi tiết từng đơn. Phạm vi hiển thị được giới hạn theo quyền của tài khoản, không phụ thuộc vào thao tác lọc trên giao diện.",
                "06-department-summary.png",
                "Thiết kế giao diện tổng hợp đơn theo phòng ban",
            ),
            (
                "3.3.3.2 Duyệt hoặc từ chối đơn bổ sung",
                "Hàng chờ đơn bổ sung trình bày người đặt, phòng ban, lý do, số mặt hàng và tổng số lượng. Trước khi quyết định, quản lý xem đầy đủ chi tiết đơn; khi từ chối phải nhập lý do để người đặt theo dõi trong lịch sử.",
                "07-supplement-approval.png",
                "Thiết kế giao diện duyệt đơn bổ sung",
            ),
            (
                "3.3.3.3 Rà soát điều kiện chốt kỳ",
                "Màn hình rà soát tổng hợp các điều kiện cần thiết trước khi chuyển sang gom nhu cầu và chốt kỳ. Hệ thống nêu rõ điều kiện đã đạt, cảnh báo và điều kiện ngăn chốt kỳ, giúp người quản lý xử lý đúng nguyên nhân thay vì chỉ nhận một thông báo lỗi chung.",
                "08-period-review.png",
                "Thiết kế giao diện rà soát điều kiện chốt kỳ",
            ),
            (
                "3.3.3.4 Chọn nguồn cung và dữ liệu phân bổ",
                "Sau khi gom nhu cầu, quản lý chọn nhà cung cấp chính và bảng giá còn hiệu lực. Những mặt hàng chưa có giá phải được chọn nguồn thay thế và ghi lý do; người thực hiện, thời điểm và thay đổi được lưu trong nhật ký nghiệp vụ.",
                "09-supply-allocation.png",
                "Thiết kế giao diện chọn nguồn cung và dữ liệu phân bổ",
            ),
            (
                "3.3.3.5 Xem trước, xác nhận và hiệu chỉnh kết quả chốt kỳ",
                "Màn hình chốt kỳ cho phép xem trước số mặt hàng, tổng số lượng, nguồn cung, bảng giá, ngoại lệ và tổng tiền trước khi xác nhận. Dữ liệu được lưu tại thời điểm chốt kỳ không bị ghi đè; khi cần điều chỉnh, người có quyền tạo phiên bản hiệu chỉnh mới, nhập lý do và thực hiện theo nguyên tắc bốn mắt trên cùng không gian làm việc.",
                "10-settlement-flow.png",
                "Thiết kế giao diện xem trước và xác nhận chốt kỳ",
            ),
        ],
    ),
    (
        "3.3.4 Giao diện quản trị dữ liệu và quyền truy cập",
        [
            (
                "3.3.4.1 Quản lý mặt hàng",
                "Khu vực quản lý mặt hàng sử dụng danh sách kết hợp vùng chi tiết. Thông tin thường dùng được giữ trên lưới; quan hệ, bản dịch, nhật ký và dữ liệu kỹ thuật được đặt trong vùng chi tiết để tránh làm bảng quá rộng.",
                "11-items.png",
                "Thiết kế giao diện quản lý mặt hàng",
            ),
            (
                "3.3.4.2 Quản lý bảng giá",
                "Danh sách bảng giá thể hiện nhà cung cấp, thời gian hiệu lực và trạng thái bản nháp, đã công bố hoặc hết hiệu lực. Bảng giá đã công bố được bảo vệ khỏi thao tác sửa làm thay đổi lịch sử; khi ngừng sử dụng, hệ thống chuyển sang trạng thái hết hiệu lực.",
                "12-price-lists.png",
                "Thiết kế giao diện quản lý bảng giá",
            ),
            (
                "3.3.4.3 Quản lý người dùng và nhóm quyền",
                "Quản trị hệ thống (DEV) quản lý tài khoản, phòng ban và nhóm quyền đang áp dụng. Mỗi tài khoản chỉ có một phân công quyền đang hiệu lực; các thao tác đặt lại mật khẩu, đổi nhóm quyền hoặc vô hiệu hóa đều được ghi nhận vào nhật ký bảo mật.",
                "13-users.png",
                "Thiết kế giao diện quản lý người dùng và nhóm quyền",
            ),
            (
                "3.3.4.4 Ma trận quyền thao tác",
                "Ma trận quyền trình bày ba vai trò thống nhất: Nhân viên (EMPLOYEE), Quản lý (MANAGER) và Quản trị hệ thống (DEV). Quản lý kế thừa các thao tác của Nhân viên và được bổ sung quyền quản lý; DEV có đầy đủ quyền nhưng vẫn phải tuân theo các điều kiện nghiệp vụ của hệ thống.",
                "14-permissions.png",
                "Thiết kế giao diện ma trận quyền thao tác",
            ),
        ],
    ),
    (
        "3.3.5 Giao diện báo cáo và thông báo",
        [
            (
                "3.3.5.1 Báo cáo và xuất dữ liệu",
                "Màn hình báo cáo cho phép chọn kỳ và phạm vi được cấp, theo dõi chi phí, số lượng, xu hướng và số liệu theo phòng ban. Dữ liệu của kỳ đã chốt được lấy từ kết quả chốt kỳ; người dùng có quyền có thể xuất PDF hoặc Excel.",
                "15-reports.png",
                "Thiết kế giao diện báo cáo và xuất dữ liệu",
            ),
            (
                "3.3.5.2 Hộp thư thông báo và trạng thái hệ thống",
                "Hộp thư lưu thông báo theo tài khoản, phân biệt đã đọc và chưa đọc, đồng thời dẫn người dùng tới màn hình liên quan. Các trạng thái mất kết nối, thiếu quyền, lỗi tải dữ liệu, chưa có dữ liệu và đang tải đều có giải thích cùng hành động tiếp theo phù hợp.",
                "16-system-states.png",
                "Thiết kế giao diện hộp thư thông báo và trạng thái hệ thống",
            ),
        ],
    ),
]


def normalized(text: str) -> str:
    return " ".join(text.split())


def find_heading(document: Document, style_name: str, prefix: str):
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph.style.name == style_name and normalized(paragraph.text).startswith(prefix)
    ]
    if len(matches) != 1:
        raise RuntimeError(f"Expected one {style_name} starting with {prefix!r}, found {len(matches)}")
    return matches[0]


def replace_ppr(paragraph, template_ppr) -> None:
    current = paragraph._p.pPr
    if current is not None:
        paragraph._p.remove(current)
    paragraph._p.insert(0, copy.deepcopy(template_ppr))


def set_page_break_before(paragraph, enabled: bool) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    current = ppr.find(qn("w:pageBreakBefore"))
    if enabled and current is None:
        ppr.append(OxmlElement("w:pageBreakBefore"))
    elif not enabled and current is not None:
        ppr.remove(current)


def set_keep_with_next(paragraph, enabled: bool) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    current = ppr.find(qn("w:keepNext"))
    if enabled and current is None:
        ppr.append(OxmlElement("w:keepNext"))
    elif not enabled and current is not None:
        ppr.remove(current)


def effective_section_properties(paragraph):
    sibling = paragraph._p.getprevious()
    while sibling is not None:
        if sibling.tag == qn("w:p"):
            section = sibling.find("./w:pPr/w:sectPr", namespaces=sibling.nsmap)
            if section is not None:
                return copy.deepcopy(section)
        sibling = sibling.getprevious()
    body_section = paragraph._p.getparent().find(qn("w:sectPr"))
    if body_section is None:
        raise RuntimeError("Could not locate effective section properties before 3.3")
    return copy.deepcopy(body_section)


def prepare_section_properties(section, *, landscape: bool, preserve_first_page: bool) -> None:
    section_type = section.find(qn("w:type"))
    if section_type is None:
        section_type = OxmlElement("w:type")
        section.insert(0, section_type)
    section_type.set(qn("w:val"), "nextPage")

    page_size = section.find(qn("w:pgSz"))
    if page_size is None:
        raise RuntimeError("Section properties do not contain page size")
    width = page_size.get(qn("w:w"))
    height = page_size.get(qn("w:h"))
    if landscape:
        if int(width) < int(height):
            page_size.set(qn("w:w"), height)
            page_size.set(qn("w:h"), width)
        page_size.set(qn("w:orient"), "landscape")
        for node in list(section.findall(qn("w:pgNumType"))):
            section.remove(node)
    else:
        if int(width) > int(height):
            page_size.set(qn("w:w"), height)
            page_size.set(qn("w:h"), width)
        page_size.attrib.pop(qn("w:orient"), None)

    if not preserve_first_page:
        for node in list(section.findall(qn("w:titlePg"))):
            section.remove(node)


def insert_section_break_before(paragraph, section_properties) -> None:
    break_paragraph = paragraph.insert_paragraph_before()
    ppr = break_paragraph._p.get_or_add_pPr()
    ppr.append(copy.deepcopy(section_properties))


def add_alt_text(image_paragraph, description: str) -> None:
    for doc_property in image_paragraph._p.xpath(".//wp:docPr"):
        doc_property.set("descr", description)


def image_dimensions(image_path: Path) -> tuple[float, float]:
    with Image.open(image_path) as image:
        width_px, height_px = image.size
    ratio = width_px / height_px
    width_cm = min(MAX_IMAGE_WIDTH_CM, MAX_IMAGE_HEIGHT_CM * ratio)
    height_cm = width_cm / ratio
    return width_cm, height_cm


def body_text_outside_range(document: Document, start, end) -> tuple[str, str]:
    paragraphs = document.paragraphs
    start_index = next(index for index, paragraph in enumerate(paragraphs) if paragraph._p is start._p)
    end_index = next(index for index, paragraph in enumerate(paragraphs) if paragraph._p is end._p)
    before = "\n".join(
        text
        for paragraph in paragraphs[: start_index + 1]
        if (text := normalized(paragraph.text))
    )
    after = "\n".join(
        text
        for paragraph in paragraphs[end_index:]
        if (text := normalized(paragraph.text))
    )
    return before, after


def update_document(input_path: Path, output_path: Path) -> None:
    document = Document(input_path)
    start = find_heading(document, "Heading 2", "3.3 ")
    end = find_heading(document, "Heading 2", "3.4 ")
    before_text, after_text = body_text_outside_range(document, start, end)

    current_paragraphs = document.paragraphs
    start_index = next(index for index, paragraph in enumerate(current_paragraphs) if paragraph._p is start._p)
    end_index = next(index for index, paragraph in enumerate(current_paragraphs) if paragraph._p is end._p)
    section_paragraphs = current_paragraphs[start_index + 1 : end_index]
    heading3_template = next(p._p.pPr for p in section_paragraphs if p.style.name == "Heading 3")
    heading4_template = next(p._p.pPr for p in section_paragraphs if p.style.name == "Heading 4")
    normal_template = next(
        p._p.pPr
        for p in section_paragraphs
        if p.style.name == "Normal" and normalized(p.text) and not normalized(p.text).startswith("Hình ")
    )
    image_template = next(p._p.pPr for p in section_paragraphs if p._p.xpath(".//a:blip"))
    caption_template = next(
        p._p.pPr
        for p in section_paragraphs
        if p.style.name == "Normal" and normalized(p.text).startswith("Hình ")
    )

    portrait_properties = effective_section_properties(start)
    screen_properties = copy.deepcopy(portrait_properties)
    prepare_section_properties(portrait_properties, landscape=False, preserve_first_page=True)
    prepare_section_properties(screen_properties, landscape=False, preserve_first_page=False)

    for paragraph in section_paragraphs:
        paragraph._p.getparent().remove(paragraph._p)

    insert_section_break_before(start, portrait_properties)

    intro = end.insert_paragraph_before()
    replace_ppr(intro, normal_template)
    intro.add_run(
        "Phần này trình bày các màn hình được chuẩn hóa bằng Atlas. Nội dung nghiệp vụ kế thừa bản luận văn hiện hành và được đối chiếu với cơ chế xử lý phía máy chủ; giao diện cũ chỉ được dùng để nhận diện các chức năng đã có. Các hình sau là thiết kế giao diện phục vụ mô tả hệ thống, không được dùng làm bằng chứng kiểm thử ở Chương 4."
    )
    set_keep_with_next(intro, True)

    figure_number = 28
    for group_index, (group_heading, screens) in enumerate(GROUPS):
        group_paragraph = end.insert_paragraph_before()
        replace_ppr(group_paragraph, heading3_template)
        group_paragraph.add_run(group_heading)
        # Keep the introduction as a short orientation page, then start every
        # screen group on a clean portrait page. This prevents a heading or
        # caption from being stranded on a nearly blank page.
        set_page_break_before(group_paragraph, True)
        set_keep_with_next(group_paragraph, True)

        for screen_index, (screen_heading, description, file_name, caption_text) in enumerate(screens):
            heading = end.insert_paragraph_before()
            replace_ppr(heading, heading4_template)
            heading.add_run(screen_heading)
            set_page_break_before(heading, screen_index > 0)
            set_keep_with_next(heading, True)

            paragraph = end.insert_paragraph_before()
            replace_ppr(paragraph, normal_template)
            paragraph.add_run(description)
            set_keep_with_next(paragraph, True)

            image_path = SCREEN_ROOT / file_name
            if not image_path.is_file():
                raise FileNotFoundError(image_path)
            width_cm, height_cm = image_dimensions(image_path)
            image_paragraph = end.insert_paragraph_before()
            replace_ppr(image_paragraph, image_template)
            image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            image_paragraph.add_run().add_picture(
                str(image_path), width=Cm(width_cm), height=Cm(height_cm)
            )
            set_keep_with_next(image_paragraph, True)
            add_alt_text(image_paragraph, f"Hình 3-{figure_number}: {caption_text}")

            caption = end.insert_paragraph_before()
            replace_ppr(caption, caption_template)
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            label = caption.add_run(f"Hình 3-{figure_number}:")
            label.bold = True
            label.italic = True
            label.underline = True
            caption.add_run(f" {caption_text}")
            set_keep_with_next(caption, False)
            figure_number += 1

    if figure_number != 44:
        raise RuntimeError(f"Expected figures 3-28 through 3-43, ended at {figure_number - 1}")

    # End the dedicated screen section immediately before 3.4 so all following
    # content reuses the document's existing portrait section settings.
    insert_section_break_before(end, screen_properties)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)

    verified = Document(output_path)
    verified_start = find_heading(verified, "Heading 2", "3.3 ")
    verified_end = find_heading(verified, "Heading 2", "3.4 ")
    verified_before, verified_after = body_text_outside_range(verified, verified_start, verified_end)
    if before_text != verified_before or after_text != verified_after:
        raise RuntimeError("Text outside section 3.3 changed unexpectedly")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--docx", type=Path, default=DEFAULT_DOCX)
    parser.add_argument("--no-backup", action="store_true")
    args = parser.parse_args()

    docx_path = args.docx.resolve()
    if not docx_path.is_file():
        raise FileNotFoundError(docx_path)
    if not args.no_backup:
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        backup = docx_path.with_name(f"{docx_path.stem}.pre-atlas-3-3-{stamp}{docx_path.suffix}")
        shutil.copy2(docx_path, backup)
        print(f"BACKUP={backup}")

    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx", dir=docx_path.parent) as handle:
        temporary_path = Path(handle.name)
    try:
        update_document(docx_path, temporary_path)
        os.replace(temporary_path, docx_path)
    finally:
        temporary_path.unlink(missing_ok=True)
    print(f"OUTPUT={docx_path}")
    print("FIGURES=16")
    print("FIGURE_RANGE=3-28..3-43")


if __name__ == "__main__":
    main()
