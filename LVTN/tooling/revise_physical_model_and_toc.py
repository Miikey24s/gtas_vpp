from __future__ import annotations

import argparse
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table


NO_TOC_STYLES = {
    "Heading 1": "Heading 1 No TOC",
    "Heading 2": "Heading 2 No TOC",
    "Heading 3": "Heading 3 No TOC",
}


@dataclass(frozen=True)
class PhysicalTableRow:
    role: str
    table_names: str
    columns: str
    constraints: str


PHYSICAL_GROUPS: dict[str, list[PhysicalTableRow]] = {
    "Bảng 3-1:": [
        PhysicalTableRow(
            "Tài khoản người dùng",
            "AspNetUsers",
            "Id: int, NOT NULL, PK; UserName và NormalizedUserName: nvarchar(256), NULL; "
            "EmployeeCode: nvarchar(50), NULL; AccountStatus: nvarchar(32), NOT NULL; "
            "RowVersion: rowversion, NOT NULL.",
            "Tên đăng nhập, email và mã nhân viên có chỉ mục duy nhất có điều kiện. "
            "RowVersion kiểm soát cập nhật đồng thời.",
        ),
        PhysicalTableRow(
            "Phòng ban",
            "Departments",
            "Id: uniqueidentifier, NOT NULL, PK; Code: nvarchar(50), NOT NULL; "
            "Name: nvarchar(200), NOT NULL; ParentDepartmentId: uniqueidentifier, NULL, FK.",
            "ParentDepartmentId tự tham chiếu và dùng Restrict khi xóa. Code là duy nhất đối với bản ghi còn hiệu lực.",
        ),
        PhysicalTableRow(
            "Nhóm quyền",
            "PermissionGroups",
            "Id: uniqueidentifier, NOT NULL, PK; GroupCode và GroupName: nvarchar(50), NOT NULL; "
            "IsDeleted: bit, NOT NULL.",
            "GroupCode là duy nhất đối với nhóm quyền còn hoạt động.",
        ),
        PhysicalTableRow(
            "Phân công người dùng",
            "UserGroupMemberships",
            "Id: uniqueidentifier, NOT NULL, PK; AccountId: int, NULL, FK; DepartmentId và "
            "PermissionGroupId: uniqueidentifier, NOT NULL, FK; RowVersion: rowversion, NOT NULL.",
            "Các khóa ngoại dùng Restrict. Mỗi tài khoản chỉ có một phân công hoạt động; "
            "ràng buộc CHECK bảo đảm phòng ban chính hợp lệ.",
        ),
        PhysicalTableRow(
            "Ánh xạ quyền truy cập",
            "PageComponentMappings; GroupPageComponentMappings",
            "PageComponentMappingId và PermissionGroupId: uniqueidentifier, NOT NULL; "
            "MemberCompanyCode: bigint, NOT NULL; IsVisible và IsEnable: bit, NOT NULL.",
            "Khóa ghép xác định duy nhất quyền của nhóm đối với thành phần trang và công ty. "
            "Các khóa ngoại dùng Restrict.",
        ),
        PhysicalTableRow(
            "Nhật ký bảo mật",
            "SecurityAudits",
            "Id: uniqueidentifier, NOT NULL, PK; Action: nvarchar(100), NOT NULL; ActorUserId: int, NULL; "
            "OccurredAtUtc: datetime2, NOT NULL; Outcome: nvarchar(40), NOT NULL.",
            "Chỉ mục kết hợp hỗ trợ tra cứu theo hành động, đối tượng và thời điểm phát sinh.",
        ),
    ],
    "Bảng 3-2:": [
        PhysicalTableRow(
            "Danh mục tra cứu",
            "LookupCategories; LookupValues",
            "Id của mỗi bảng: uniqueidentifier, NOT NULL, PK; LookupCategoryId: uniqueidentifier, NULL, FK; "
            "Code và Value: nvarchar(max).",
            "Khóa ngoại dùng Restrict; LookupValues có chỉ mục theo LookupCategoryId.",
        ),
        PhysicalTableRow(
            "Danh mục văn phòng phẩm",
            "VppCategories",
            "Id: uniqueidentifier, NOT NULL, PK; VppCategoryCode và VppCategoryName: nvarchar(max), NULL.",
            "Dữ liệu hỗ trợ xóa mềm và lưu dấu thời điểm tạo, cập nhật.",
        ),
        PhysicalTableRow(
            "Văn phòng phẩm",
            "VppItems",
            "Id: uniqueidentifier, NOT NULL, PK; VppCode: nvarchar(64), NOT NULL; "
            "VppName: nvarchar(250), NOT NULL; UomId và VppCategoryId: uniqueidentifier, NOT NULL, FK.",
            "VppCode là duy nhất; các chỉ mục hỗ trợ lọc theo danh mục, đơn vị tính và trạng thái xóa.",
        ),
        PhysicalTableRow(
            "Nhà cung cấp",
            "Suppliers",
            "Id: uniqueidentifier, NOT NULL, PK; SupplierName và SupplierShortName: nvarchar(max), NULL; "
            "các trường địa chỉ cho phép NULL.",
            "Bảng hỗ trợ xóa mềm; nội dung đa ngôn ngữ được lưu trong nhóm bảng dịch riêng.",
        ),
        PhysicalTableRow(
            "Bảng giá",
            "PriceLists",
            "Id: uniqueidentifier, NOT NULL, PK; SupplierId: uniqueidentifier, NULL, FK; "
            "PriceListCode: nvarchar(50), NULL; Version: int, NOT NULL; EffectiveFromUtc: datetime2, NOT NULL; "
            "EffectiveToUtc: datetime2, NULL; RowVersion: rowversion, NOT NULL.",
            "Mã và phiên bản bảng giá là duy nhất theo nhà cung cấp. Ràng buộc CHECK kiểm tra thời gian hiệu lực, "
            "phiên bản, chiết khấu và số tiền không âm.",
        ),
        PhysicalTableRow(
            "Đơn giá theo nhà cung cấp",
            "SupplierProductMappings",
            "Id: uniqueidentifier, NOT NULL, PK; PriceListId, SupplierId và VppItemId: uniqueidentifier, NOT NULL, FK; "
            "NetPrice và MinimumOrderQuantity: decimal(19,4); VatRate: decimal(5,2); RowVersion: rowversion, NOT NULL.",
            "Mỗi mặt hàng chỉ có một dòng giá mặc định trong một bảng giá. Ràng buộc CHECK kiểm tra giá, "
            "thuế suất, số lượng đặt tối thiểu và thời gian cung ứng.",
        ),
    ],
    "Bảng 3-3:": [
        PhysicalTableRow(
            "Kỳ yêu cầu",
            "Periods",
            "Id: uniqueidentifier, NOT NULL, PK; Year, Month và State: int, NOT NULL; StartAtUtc, "
            "SubmissionDeadlineUtc và SupplementApprovalDeadlineUtc: datetime2, NOT NULL; RowVersion: rowversion, NULL.",
            "Mỗi công ty chỉ có một kỳ còn hiệu lực cho từng tháng và năm. Ràng buộc CHECK kiểm tra khoảng thời gian "
            "và trạng thái kỳ thuộc tập giá trị cho phép.",
        ),
        PhysicalTableRow(
            "Đơn yêu cầu",
            "Requests",
            "Id: uniqueidentifier, NOT NULL, PK; PeriodId: uniqueidentifier, NULL, FK; RequestSeriesId: uniqueidentifier, "
            "NOT NULL; RevisionNumber và Status: int, NOT NULL; IdempotencyKey: nvarchar(128), NULL; "
            "RowVersion: rowversion, NULL.",
            "Các chỉ mục duy nhất có điều kiện bảo đảm chỉ có một phiên bản hiện hành, một đơn thông thường trong kỳ, "
            "một đơn bổ sung chờ duyệt và một khóa chống gửi lặp.",
        ),
        PhysicalTableRow(
            "Chi tiết đơn",
            "RequestDetails",
            "Id: uniqueidentifier, NOT NULL, PK; RequestId và VppId: uniqueidentifier, NOT NULL, FK; "
            "Qty: int, NOT NULL; CurrentSinglePrice: bigint, NOT NULL.",
            "Khóa ngoại đến Requests và VppItems dùng Restrict; chỉ mục hỗ trợ tra cứu theo đơn và mặt hàng.",
        ),
        PhysicalTableRow(
            "Nhật ký đơn",
            "RequestLogs",
            "Id: uniqueidentifier, NOT NULL, PK; RequestId: uniqueidentifier, NOT NULL, FK; Action: nvarchar(64), NULL; "
            "LogDate: datetime2, NOT NULL; Reason: nvarchar(500), NULL; RevisionNumber: int, NULL.",
            "Khóa ngoại đến Requests dùng Restrict; chỉ mục theo RequestId hỗ trợ truy vết lịch sử đơn.",
        ),
    ],
    "Bảng 3-4:": [
        PhysicalTableRow(
            "Kết quả chốt kỳ",
            "Settlements",
            "Id: uniqueidentifier, NOT NULL, PK; PeriodId: uniqueidentifier, NOT NULL, FK; RevisionNumber: int, NOT NULL; "
            "IdempotencyKey: nvarchar(128), NOT NULL; các trường tiền: decimal(19,4); RowVersion: rowversion, NULL.",
            "Chỉ có một phiên bản kết quả hiện hành. Khóa chống gửi lặp và số phiên bản là duy nhất; "
            "ràng buộc CHECK bảo đảm kỳ và các số tiền hợp lệ.",
        ),
        PhysicalTableRow(
            "Dòng kết quả chốt kỳ",
            "SettlementItems",
            "Id và SettlementId: uniqueidentifier, NOT NULL; Quantity, NetUnitPrice, NetAmount, VatAmount và GrossAmount: "
            "decimal(19,4), NOT NULL; VatRate: decimal(5,2), NOT NULL.",
            "Khóa ngoại dùng Restrict; mỗi mặt hàng chỉ xuất hiện một lần trong một kết quả chốt kỳ. "
            "Ràng buộc CHECK kiểm tra số lượng, giá trị và thuế suất.",
        ),
        PhysicalTableRow(
            "Chi phí chốt kỳ",
            "SettlementCharges",
            "Id và SettlementId: uniqueidentifier, NOT NULL; ChargeType và AllocationBasis: nvarchar(32), NOT NULL; "
            "Amount: decimal(19,4), NOT NULL.",
            "Khóa ngoại dùng Restrict; mỗi loại chi phí chỉ có một dòng trong một kết quả chốt kỳ.",
        ),
        PhysicalTableRow(
            "Dữ liệu phân bổ",
            "SettlementAllocations",
            "Id, SettlementId, SettlementItemId và RequestDetailId: uniqueidentifier, NOT NULL; Quantity và các trường tiền: "
            "decimal(19,4), NOT NULL; DepartmentCode: nvarchar(64), NULL.",
            "Khóa ngoại dùng Restrict; mỗi chi tiết đơn chỉ có một dòng phân bổ trong kết quả chốt kỳ. "
            "Ràng buộc CHECK yêu cầu Quantity lớn hơn 0.",
        ),
        PhysicalTableRow(
            "Thông báo trong hệ thống",
            "Notifications",
            "Id: uniqueidentifier, NOT NULL, PK; UserId: int, NOT NULL; Type: nvarchar(50), NOT NULL; "
            "CorrelationId: nvarchar(100), NULL; ReadAt: datetime2, NULL.",
            "Chỉ mục duy nhất có điều kiện chống tạo thông báo trùng; chỉ mục hộp thư hỗ trợ lọc theo trạng thái đọc và thời điểm.",
        ),
        PhysicalTableRow(
            "Hàng đợi gửi thư",
            "EmailOutboxMessages",
            "Id: uniqueidentifier, NOT NULL, PK; DeduplicationKey: nvarchar(128), NOT NULL; Status: nvarchar(24), NOT NULL; "
            "NextAttemptAtUtc: datetime2, NOT NULL; AttemptCount: int, NOT NULL.",
            "DeduplicationKey là duy nhất; chỉ mục theo Status và NextAttemptAtUtc hỗ trợ tiến trình gửi thư và thử lại.",
        ),
    ],
}


def set_outline_level(style, level: int) -> None:
    p_pr = style.element.get_or_add_pPr()
    outline = p_pr.find(qn("w:outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        p_pr.append(outline)
    outline.set(qn("w:val"), str(level))


def ensure_no_toc_styles(document: Document) -> None:
    for source_name, target_name in NO_TOC_STYLES.items():
        if target_name in document.styles:
            target = document.styles[target_name]
        else:
            target = document.styles.add_style(target_name, WD_STYLE_TYPE.PARAGRAPH)
        target.base_style = document.styles[source_name]
        set_outline_level(target, 9)


def exclude_appendices_and_references_from_toc(document: Document) -> int:
    ensure_no_toc_styles(document)
    in_back_matter = False
    changed = 0

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text == "PHỤ LỤC":
            in_back_matter = True
        if not in_back_matter:
            continue

        source_style = paragraph.style.name
        target_style = NO_TOC_STYLES.get(source_style)
        if target_style:
            paragraph.style = document.styles[target_style]
            changed += 1

    return changed


def find_table_after_caption(document: Document, caption_prefix: str) -> Table:
    body = document.element.body
    children = list(body)
    for index, child in enumerate(children):
        if child.tag != qn("w:p"):
            continue
        text = "".join(child.xpath(".//w:t/text()"))
        if not text.strip().startswith(caption_prefix):
            continue
        for following in children[index + 1 :]:
            if following.tag == qn("w:tbl"):
                return Table(following, document)
            if following.tag == qn("w:p") and "".join(following.xpath(".//w:t/text()" )).strip():
                break
    raise ValueError(f"Could not locate table after caption {caption_prefix!r}")


def clear_cell(cell) -> None:
    tc = cell._tc
    for paragraph in list(tc.p_lst):
        tc.remove(paragraph)
    tc.add_p()


def set_cell_text(cell, parts: list[tuple[str, bool, bool]]) -> None:
    clear_cell(cell)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for index, (text, bold, italic) in enumerate(parts):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(text)
        run.bold = bold
        run.italic = italic


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def rewrite_physical_table(table: Table, rows: list[PhysicalTableRow]) -> None:
    if len(table.columns) != 3:
        raise ValueError("Expected the existing physical-model table to have three columns")
    if len(table.rows) != len(rows) + 1:
        raise ValueError("Physical-model table row count does not match the approved data map")

    header = table.cell(0, 1).merge(table.cell(0, 2))
    table.cell(0, 0).text = "Nhóm dữ liệu"
    header.text = "Cấu trúc vật lý và ràng buộc chính"
    for cell in (table.cell(0, 0), header):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
    set_repeat_header(table.rows[0])

    for row_index, data in enumerate(rows, start=1):
        details = table.cell(row_index, 1).merge(table.cell(row_index, 2))
        role_cell = table.cell(row_index, 0)
        set_cell_text(role_cell, [(data.role, True, False)])
        set_cell_text(
            details,
            [
                (f"Bảng: {data.table_names}", True, False),
                (f"Cột chính: {data.columns}", False, False),
                (f"Ràng buộc chính: {data.constraints}", False, False),
            ],
        )
        role_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        details.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cant_split(table.rows[row_index])

    table.autofit = False


def replace_physical_model_intro(document: Document) -> int:
    replacements = {
        "Mô hình vật lý được triển khai trên SQL Server.": (
            "Mô hình dữ liệu vật lý mô tả cách các đối tượng nghiệp vụ được lưu trữ trên SQL Server. "
            "Các bảng dưới đây sử dụng đúng tên bảng và tên cột trong hệ thống; chỉ những cột quyết định cấu trúc, "
            "quan hệ và tính toàn vẹn dữ liệu được trình bày trong phần chính. Khóa chính (PK), khóa ngoại (FK), "
            "khả năng để trống (NULL), kiểu rowversion và các ràng buộc quan trọng được ghi ngay tại từng nhóm dữ liệu."
        ),
        "Các bảng nghiệp vụ dùng xóa mềm qua IsDeleted": (
            "Để nội dung dễ đọc trên khổ A4, mỗi nhóm được trình bày theo hai phần: cấu trúc cột chính và ràng buộc chính. "
            "Các bảng nghiệp vụ sử dụng IsDeleted để xóa mềm; phần lớn khóa ngoại dùng Restrict nhằm tránh xóa lan dữ liệu lịch sử. "
            "Tên đầy đủ của các chỉ mục và ràng buộc kỹ thuật dài được tập hợp tại Phụ lục D."
        ),
    }
    changed = 0
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        for prefix, replacement in replacements.items():
            if text.startswith(prefix):
                paragraph.text = replacement
                changed += 1
                break
    return changed


def normalize_physical_model_wording(document: Document) -> int:
    replacements = {
        "3.1.3.3 Kỳ, đơn yêu cầu, revision và nhật ký": "3.1.3.3 Kỳ, đơn yêu cầu, phiên bản và nhật ký",
        "3.1.3.5 Identity, bảng dịch, view và keyless entity": "3.1.3.5 Identity, bảng dịch và đối tượng chỉ đọc",
        "Bảng 3-3: Cấu trúc vật lý nhóm kỳ, đơn yêu cầu, revision và nhật ký":
            "Bảng 3-3: Cấu trúc vật lý nhóm kỳ, đơn yêu cầu, phiên bản và nhật ký",
    }
    changed = 0
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text in replacements:
            paragraph.text = replacements[text]
            changed += 1
    return changed


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()

    document = Document(args.input)
    back_matter_headings = exclude_appendices_and_references_from_toc(document)
    intro_paragraphs = replace_physical_model_intro(document)
    wording_paragraphs = normalize_physical_model_wording(document)

    for caption_prefix, rows in PHYSICAL_GROUPS.items():
        rewrite_physical_table(find_table_after_caption(document, caption_prefix), rows)

    document.save(args.out)
    print(f"Back-matter headings excluded from TOC: {back_matter_headings}")
    print(f"Physical-model introductory paragraphs updated: {intro_paragraphs}")
    print(f"Physical-model headings/captions normalized: {wording_paragraphs}")
    print(f"Physical-model tables rewritten: {len(PHYSICAL_GROUPS)}")
    print(f"Output: {args.out.resolve()}")


if __name__ == "__main__":
    main()
