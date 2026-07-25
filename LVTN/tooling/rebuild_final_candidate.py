#!/usr/bin/env python3
"""Rebuild the editable thesis sections from the current GTAS VPP evidence.

The script locates ranges by stable headings, never overwrites the input DOCX,
keeps locked sections untouched, inserts figures inline, and writes a separate
candidate that Microsoft Word can update and export to PDF.
"""

from __future__ import annotations

import argparse
import hashlib
import re
import subprocess
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from rebuild_chapters_4_5 import (
    add_body,
    add_caption,
    add_heading,
    add_labeled_body,
    add_table,
    apply_font,
    copy_paragraph_properties,
    find_body_sample,
    find_caption_sample,
    find_paragraph,
    find_table_sample,
    mark_fields_dirty,
    normalized,
    preceding_section_properties,
    add_section_break,
    remove_range,
)


ROOT = Path(__file__).resolve().parents[2]
DIAGRAMS = ROOT / "LVTN" / "diagrams"
SCREENSHOTS = ROOT / "LVTN" / "screenshots" / "ch03"
GENERATED = ROOT / "LVTN" / "generated"
DOC: Document | None = None
SECTION_PROPS: dict[str, object] = {}


def add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    insertion_index = 1 if paragraph._p.pPr is not None else 0
    paragraph._p.insert(insertion_index, start)
    paragraph._p.append(end)


def add_list_item(anchor, text: str, sample, number: int | None = None):
    prefix = f"{number}. " if number is not None else "- "
    return add_body(anchor, prefix + text, sample)


def ensure_png(path: Path) -> Path:
    if path.suffix.lower() != ".svg":
        return path
    output = path.with_suffix(".png")
    if not output.exists() or output.stat().st_mtime < path.stat().st_mtime:
        source = path.read_text(encoding="utf-8")
        view_box = re.search(r'viewBox="[^"]*?\s([0-9.]+)\s([0-9.]+)"', source)
        width_attr = re.search(r'<svg[^>]+width="([0-9.]+)', source)
        height_attr = re.search(r'<svg[^>]+height="([0-9.]+)', source)
        width = int(float(view_box.group(1) if view_box else width_attr.group(1) if width_attr else 1400))
        height = int(float(view_box.group(2) if view_box else height_attr.group(1) if height_attr else 1000))
        chrome = Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe")
        if not chrome.exists():
            raise RuntimeError("Google Chrome is required to rasterize SVG diagrams.")
        subprocess.run(
            [
                str(chrome),
                "--headless=new",
                "--disable-gpu",
                "--hide-scrollbars",
                "--force-device-scale-factor=2",
                f"--window-size={width},{height}",
                f"--screenshot={output.resolve()}",
                path.resolve().as_uri(),
            ],
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
    return output


def add_image(anchor, source: Path, *, max_width=6.25, max_height=7.7):
    source = ensure_png(source)
    if not source.exists():
        raise FileNotFoundError(source)
    with Image.open(source) as image:
        width_px, height_px = image.size
    ratio = min(max_width / width_px, max_height / height_px)
    width = max(1.0, width_px * ratio)
    height = max(0.8, height_px * ratio)
    paragraph = anchor.insert_paragraph_before()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(source), width=Inches(width), height=Inches(height))
    return paragraph


def add_figure(anchor, source: Path, number: str, title: str, caption_sample, bookmark_id: int):
    add_image(anchor, source)
    caption = add_caption(anchor, f"Hình {number}: {title}", caption_sample)
    caption.paragraph_format.keep_with_next = False
    add_bookmark(caption, f"fig_{number.replace('-', '_')}", bookmark_id)
    return bookmark_id + 1


def add_table_caption(anchor, number: str, title: str, caption_sample):
    paragraph = add_caption(anchor, f"Bảng {number}: {title}", caption_sample)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def replace_range(document, start_text: str, end_text: str, builder) -> None:
    start = find_paragraph(document, start_text)
    end = find_paragraph(document, end_text)
    remove_range(start, end)
    builder(end)


def build_231(anchor, body, *_):
    add_heading(anchor, "2.3.1 Các quy trình, nghiệp vụ", "Heading 3")
    add_body(anchor, "Các quy trình được mô tả theo hành vi nghiệp vụ mà người dùng quan sát được. Việc xác thực, trạng thái, phạm vi dữ liệu và quyền thao tác đều do backend kiểm tra; giao diện chỉ hiển thị các chức năng phù hợp với snapshot quyền hiện hành.", body)

    sections = [
        ("2.3.1.1 Đăng nhập và tải quyền", [
            "Người dùng nhập tên đăng nhập và mật khẩu trên trang đăng nhập.",
            "ASP.NET Core Identity kiểm tra mật khẩu đã băm, trạng thái tài khoản, yêu cầu đổi mật khẩu và phiên đăng nhập.",
            "Khi hợp lệ, hệ thống phát hành thông tin xác thực theo cấu hình, đổi ticket một lần thành cookie và tải snapshot quyền trang, thành phần và action.",
            "Frontend khởi tạo kết nối SignalR để nhận thay đổi quyền; backend tiếp tục kiểm tra policy ở từng API nên việc ẩn nút không thay thế kiểm soát truy cập.",
        ]),
        ("2.3.1.2 Tạo đơn yêu cầu thông thường", [
            "Hệ thống xác định kỳ đang nhận đơn từ thực thể kỳ và kiểm tra hạn gửi, trạng thái kỳ cùng điều kiện mỗi người dùng chỉ có một đơn thường hiện hành trong kỳ.",
            "Người dùng tìm văn phòng phẩm trong danh mục được phép đặt, nhập số lượng, ghi chú hoặc sao chép dữ liệu từ kỳ trước.",
            "Đơn phải có ít nhất một dòng, số lượng dương, không trùng văn phòng phẩm và chỉ dùng danh mục đang hoạt động.",
            "Khi gửi, hệ thống tạo đơn, chi tiết, revision hiện hành và nhật ký trong transaction; ràng buộc duy nhất ngăn hai yêu cầu đồng thời tạo trùng đơn.",
        ]),
        ("2.3.1.3 Sao chép, chỉnh sửa và hủy đơn", [
            "Sao chép chỉ lấy các văn phòng phẩm còn hiệu lực từ đơn gần nhất và tạo dữ liệu nháp cho kỳ hiện tại.",
            "Chỉnh sửa hợp lệ tạo revision mới, giữ nguyên lịch sử revision cũ và yêu cầu rowversion để phát hiện xung đột cập nhật.",
            "Hủy đơn chuyển trạng thái theo state machine, ghi người thực hiện, thời điểm, lý do và nhật ký; dữ liệu không bị xóa vật lý khỏi lịch sử.",
        ]),
        ("2.3.1.4 Tạo và xử lý đơn bổ sung", [
            "Đơn bổ sung chỉ được tạo trong cửa sổ nghiệp vụ cho phép, phải nêu lý do, gắn với đơn gốc và không vượt giới hạn số lần hoặc tồn tại đồng thời một đơn chờ xử lý.",
            "Người có quyền REQUEST_APPROVE hoặc REQUEST_REJECT xem hàng chờ theo phạm vi, sau đó duyệt hoặc từ chối đơn Pending của người khác.",
            "Quyết định, lý do, người xử lý, thời điểm và trạng thái cuối được lưu trong đơn và nhật ký; người tạo nhận thông báo tương ứng.",
        ]),
        ("2.3.1.5 Quản lý kỳ", [
            "Kỳ là thực thể độc lập, xác định năm, tháng, múi giờ, hạn gửi, hạn xử lý đơn bổ sung và trạng thái Open, Closed, Pricing hoặc Settled.",
            "Quản lý theo dõi tiến độ, chuyển trạng thái theo điều kiện và lý do hợp lệ; rowversion ngăn ghi đè khi nhiều người thao tác đồng thời.",
        ]),
        ("2.3.1.6 Quản lý danh mục văn phòng phẩm và đơn vị tính", [
            "Quản lý tra cứu, tìm kiếm, phân trang, thêm, sửa, vô hiệu hóa hoặc khôi phục danh mục, văn phòng phẩm và đơn vị tính.",
            "Mã văn phòng phẩm là duy nhất; văn phòng phẩm liên kết đúng một danh mục và một đơn vị tính. Xóa mềm bảo toàn dữ liệu đơn và báo cáo lịch sử.",
        ]),
        ("2.3.1.7 Quản lý nhà cung cấp và bảng giá", [
            "Bảng giá thuộc nhà cung cấp, có mã, phiên bản, thời gian hiệu lực, trạng thái Draft, Published hoặc Expired và các điều kiện thương mại.",
            "Dòng giá liên kết bảng giá, nhà cung cấp và văn phòng phẩm; hệ thống kiểm tra VAT, giá, số lượng đặt tối thiểu, thời gian giao và ràng buộc không chồng lấn trước khi công bố.",
            "Bảng giá đã công bố bị khóa các thay đổi làm mất tính nhất quán; hết hiệu lực được ghi nhận bằng trạng thái và thời điểm, không xóa lịch sử.",
        ]),
        ("2.3.1.8 Xem trước và xác nhận chốt kỳ", [
            "Quản lý chọn nhà cung cấp chính rồi chọn bảng giá phù hợp, tải bản xem trước và rà soát độ phủ giá, đơn bổ sung chờ xử lý, mặt hàng thiếu giá, ngoại lệ cùng tổng tiền.",
            "Khi xác nhận, backend tính lại hash đầu vào để chặn dữ liệu cũ, tạo settlement revision bất biến gồm dòng giá, phí và phân bổ, sau đó chuyển kỳ sang Settled.",
            "Yêu cầu xác nhận lặp lại với cùng idempotency key trả về cùng kết quả thay vì tạo bản chốt trùng.",
        ]),
        ("2.3.1.9 Hiệu chỉnh kết quả chốt kỳ", [
            "Hiệu chỉnh yêu cầu lý do và nguyên tắc bốn mắt: người xác nhận hiệu chỉnh không được là người tạo revision đang hiện hành.",
            "Hệ thống tạo revision settlement mới, đánh dấu revision cũ không còn hiện hành nhưng giữ nguyên toàn bộ snapshot cũ để đối chiếu và audit.",
        ]),
        ("2.3.1.10 Tổng hợp, báo cáo, thông báo và lịch sử", [
            "Dữ liệu được giới hạn theo phạm vi cá nhân, phòng ban hoặc toàn công ty dựa trên quyền và claim của phiên đăng nhập.",
            "Báo cáo cung cấp KPI, xu hướng, trạng thái, phòng ban, văn phòng phẩm nổi bật và xuất CSV/XLSX trong đúng phạm vi được cấp.",
            "Thông báo trong ứng dụng là nguồn theo dõi chính; email outbox lưu bền vững, chống gửi trùng và thử lại khi SMTP gặp lỗi. Nhật ký đơn và security audit phục vụ truy vết thao tác.",
        ]),
    ]
    for heading, bullets in sections:
        add_heading(anchor, heading, "Heading 4")
        for idx, text in enumerate(bullets, 1):
            add_list_item(anchor, text, body, idx)


def build_233(anchor, body, caption, table_sample):
    add_heading(anchor, "2.3.3 Sơ đồ Use case tổng quát", "Heading 3")
    add_body(anchor, "Hệ thống sử dụng ba persona chuẩn. Quản lý kế thừa các chức năng cá nhân của Nhân viên và được cấp thêm nghiệp vụ theo dõi, phê duyệt, thư viện, báo cáo và vận hành kỳ. DEV là vai trò quản trị kỹ thuật; vai trò này quản lý tài khoản, nhóm quyền, quyền truy cập và nhật ký bảo mật, không mặc định thay người dùng thực hiện nghiệp vụ.", body)
    bookmark = 420
    bookmark = add_figure(anchor, DIAGRAMS / "ch02" / "use-case-overview.png", "2-3", "Sơ đồ use case tổng quát hệ thống", caption, bookmark)
    add_table_caption(anchor, "2-1", "Mô tả tác nhân của hệ thống", caption)
    add_table(
        DOC,
        anchor,
        ["Tác nhân", "Mã persona", "Phạm vi trách nhiệm"],
        [
            ["Nhân viên", "EMPLOYEE", "Lập và theo dõi đơn của mình, xem danh mục, lịch sử, báo cáo cá nhân và thông báo."],
            ["Quản lý", "MANAGER", "Kế thừa Nhân viên; xem đơn theo phạm vi được cấp, xử lý đơn bổ sung, quản lý thư viện, báo cáo và chốt kỳ."],
            ["DEV", "DEV", "Quản trị người dùng, membership, nhóm quyền, mapping truy cập và nhật ký bảo mật; không mặc định thực hiện nghiệp vụ thay người dùng."],
        ],
        [1700, 1450, 5922],
        table_sample=table_sample,
        center_columns=[1],
    )


def markdown_table(lines, start):
    headers = [cell.strip() for cell in lines[start].strip().strip("|").split("|")]
    rows = []
    index = start + 2
    while index < len(lines) and lines[index].lstrip().startswith("|"):
        cells = [re.sub(r"`([^`]*)`", r"\1", cell.strip()).replace("<br>", "\n") for cell in lines[index].strip().strip("|").split("|")]
        if not any("StoredProcedureResultDTO" in cell for cell in cells):
            rows.append(cells)
        index += 1
    return headers, rows, index


def build_31(anchor, body, caption, table_sample):
    lines = (GENERATED / "chapter-3-1-draft.md").read_text(encoding="utf-8").splitlines()
    figure_map = {
        "01-conceptual-overview.svg": ("3-1", "Mô hình dữ liệu ý niệm của GTAS VPP"),
        "02-logical-organization-access.svg": ("3-2", "Mô hình dữ liệu luận lý về tổ chức và phân quyền"),
        "03-logical-catalog-pricing.svg": ("3-3", "Mô hình dữ liệu luận lý về danh mục, nhà cung cấp và bảng giá"),
        "05-logical-period-requests.svg": ("3-4", "Mô hình dữ liệu luận lý về kỳ, đơn yêu cầu, phiên bản và nhật ký"),
        "06-logical-settlement.svg": ("3-5", "Mô hình dữ liệu luận lý về kết quả chốt kỳ, chi phí và phân bổ"),
        "07-logical-identity-views-technical.svg": ("3-6", "Mô hình dữ liệu luận lý về Identity, view keyless và bảng kỹ thuật"),
    }
    add_heading(anchor, "3.1 MÔ HÌNH DỮ LIỆU", "Heading 2")
    bookmark_id = 430
    table_number = 1
    physical_table_titles = {
        1: "Cấu trúc vật lý nhóm tổ chức và phân quyền",
        2: "Cấu trúc vật lý nhóm danh mục, nhà cung cấp và bảng giá",
        3: "Cấu trúc vật lý nhóm kỳ, đơn yêu cầu, revision và nhật ký",
        4: "Cấu trúc vật lý nhóm settlement, chi phí, phân bổ và thông báo",
        5: "Cấu trúc vật lý bảng kỹ thuật, view và keyless entity",
    }
    index = 1
    skip_translation_paragraph = False
    while index < len(lines):
        raw = lines[index].strip()
        if not raw:
            index += 1
            continue
        if raw.startswith("## 3.1.1"):
            add_heading(anchor, "3.1.1 Mô hình dữ liệu ý niệm", "Heading 3")
        elif raw.startswith("## 3.1.2"):
            add_heading(anchor, "3.1.2 Mô hình dữ liệu luận lý", "Heading 3")
        elif raw.startswith("## 3.1.3"):
            add_heading(anchor, "3.1.3 Mô hình dữ liệu vật lý", "Heading 3")
        elif raw.startswith("### "):
            add_heading(anchor, raw[4:], "Heading 4")
        elif raw.startswith("[[FIGURE:"):
            filename = raw.split(":", 1)[1].split("|", 1)[0]
            if filename in figure_map:
                number, title = figure_map[filename]
                bookmark_id = add_figure(anchor, DIAGRAMS / "ch03" / "data-models" / filename, number, title, caption, bookmark_id)
        elif raw.startswith("|") and index + 1 < len(lines) and set(lines[index + 1].replace("|", "").replace("-", "").replace(" ", "")) == set():
            headers, rows, index = markdown_table(lines, index)
            title = physical_table_titles[table_number]
            add_table_caption(anchor, f"3-{table_number}", title, caption)
            add_table(DOC, anchor, headers, rows, [1800, 4300, 2350, 2822][: len(headers)], table_sample=table_sample)
            table_number += 1
            continue
        elif raw.startswith("# "):
            pass
        elif "Nhóm bản dịch sử dụng" in raw:
            add_body(anchor, "Các bảng bản dịch hỗ trợ dữ liệu nghiệp vụ đa ngôn ngữ theo cặp thực thể - ngôn ngữ và dùng chung các trường nguồn, trạng thái cùng audit. Nhóm này không được trình bày thành sơ đồ chi tiết trong phần chính để tập trung vào dữ liệu cốt lõi.", body)
        elif "Hai view cùng keyless DTO" in raw:
            add_body(anchor, "Các bảng phụ của ASP.NET Core Identity lưu claim, thông tin đăng nhập ngoài và token. Các view keyless chỉ phục vụ truy vấn đọc; DTO kết quả thủ tục không phải entity hoặc view dữ liệu nên không được đưa vào ERD.", body)
        elif "04a-logical" not in raw and "04b-logical" not in raw and "StoredProcedureResultDTO" not in raw:
            cleaned = re.sub(r"`([^`]*)`", r"\1", raw)
            add_body(anchor, cleaned, body)
        index += 1


USE_CASES = [
    ("3.2.1.1 Chức năng đăng nhập và tải quyền", "use-case-login.svg", "Use case đăng nhập và tải quyền", "Nhân viên nhập thông tin xác thực; hệ thống kiểm tra Identity, trạng thái tài khoản, tạo phiên và tải snapshot quyền. Luồng thay thế xử lý sai mật khẩu, tài khoản bị khóa, yêu cầu đổi mật khẩu hoặc không có membership hợp lệ."),
    ("3.2.1.2 Chức năng tạo đơn yêu cầu thông thường", "use-case-create-regular-request.svg", "Use case tạo đơn yêu cầu thông thường", "Người dùng chọn hoặc sao chép văn phòng phẩm, nhập số lượng, kiểm tra kỳ và gửi đơn. Backend bảo đảm một đơn thường hiện hành cho mỗi người dùng trong kỳ."),
    ("3.2.1.3 Chức năng chỉnh sửa và hủy đơn", "use-case-edit-cancel-request.svg", "Use case chỉnh sửa và hủy đơn", "Chỉ chủ sở hữu được sửa hoặc hủy đơn ở trạng thái cho phép. Chỉnh sửa tạo revision mới; hủy giữ lịch sử và ghi audit."),
    ("3.2.1.4 Chức năng tạo đơn bổ sung", "use-case-create-additional-request.svg", "Use case tạo đơn bổ sung", "Đơn bổ sung phải gắn với đơn gốc, có lý do và đáp ứng cửa sổ nghiệp vụ, quota cùng điều kiện không có đơn Pending khác."),
    ("3.2.1.5 Chức năng duyệt hoặc từ chối đơn bổ sung", "use-case-approve-additional-request.svg", "Use case duyệt hoặc từ chối đơn bổ sung", "Quản lý kiểm tra nội dung đơn Pending trong phạm vi được cấp, duyệt hoặc từ chối và phát thông báo cho người tạo."),
    ("3.2.1.6 Chức năng quản lý danh mục văn phòng phẩm", "use-case-manage-catalog.svg", "Use case quản lý danh mục văn phòng phẩm", "Quản lý tìm kiếm, phân trang, thêm, sửa, vô hiệu hóa hoặc khôi phục danh mục, văn phòng phẩm và đơn vị tính."),
    ("3.2.1.7 Chức năng quản lý bảng giá", "use-case-manage-pricing.svg", "Use case quản lý bảng giá", "Bảng giá được quản lý theo nhà cung cấp, phiên bản, hiệu lực và trạng thái; việc công bố hoặc hết hiệu lực phải đáp ứng các ràng buộc thương mại."),
    ("3.2.1.8 Chức năng xem trước và chốt kỳ", "use-case-settle-period.svg", "Use case xem trước và chốt kỳ", "Quản lý chọn nhà cung cấp và bảng giá, rà soát blocker, xem trước tổng tiền rồi xác nhận tạo settlement revision bất biến."),
    ("3.2.1.9 Chức năng quản lý người dùng và phân quyền", "use-case-manage-permissions.svg", "Use case quản lý người dùng và phân quyền", "DEV quản lý membership, persona và mapping trong giới hạn an toàn; backend ngăn tự nâng quyền hoặc loại bỏ quyền quản trị cuối cùng."),
    ("3.2.1.10 Chức năng xem dashboard và dữ liệu tổng hợp", "use-case-dashboard-summary.svg", "Use case xem dashboard và dữ liệu tổng hợp", "Dữ liệu dashboard và báo cáo được giới hạn theo cá nhân, phòng ban hoặc toàn công ty dựa trên quyền hiện hành."),
]

SEQUENCES = [
    ("3.2.2.1 Đăng nhập và tải quyền", "sequence-login-permissions.svg", "Sơ đồ tuần tự đăng nhập và tải quyền"),
    ("3.2.2.2 Tạo đơn yêu cầu thông thường", "sequence-create-regular-request.svg", "Sơ đồ tuần tự tạo đơn yêu cầu thông thường"),
    ("3.2.2.3 Tạo đơn bổ sung", "sequence-create-additional-request.svg", "Sơ đồ tuần tự tạo đơn bổ sung"),
    ("3.2.2.4 Duyệt đơn bổ sung", "sequence-approve-additional-request.svg", "Sơ đồ tuần tự duyệt đơn bổ sung"),
    ("3.2.2.5 Xem trước và xác nhận chốt kỳ", "sequence-settle-period.svg", "Sơ đồ tuần tự xem trước và xác nhận chốt kỳ"),
    ("3.2.2.6 Quản lý bảng giá", "sequence-manage-pricing.svg", "Sơ đồ tuần tự quản lý bảng giá"),
]

ACTIVITIES = [
    ("3.2.3.1 Tạo đơn yêu cầu thông thường", "activity-create-regular-request.svg", "Sơ đồ hoạt động tạo đơn yêu cầu thông thường"),
    ("3.2.3.2 Xử lý đơn bổ sung", "activity-additional-request.svg", "Sơ đồ hoạt động xử lý đơn bổ sung"),
    ("3.2.3.3 Xem trước và chốt kỳ", "activity-settle-period.svg", "Sơ đồ hoạt động xem trước và chốt kỳ"),
    ("3.2.3.4 Quản lý bảng giá", "activity-manage-pricing.svg", "Sơ đồ hoạt động quản lý bảng giá"),
    ("3.2.3.5 Phân quyền", "activity-manage-permissions.svg", "Sơ đồ hoạt động quản lý phân quyền"),
]


def build_32(anchor, body, caption, *_):
    add_heading(anchor, "3.2 MÔ HÌNH XỬ LÝ", "Heading 2")
    add_body(anchor, "Mô hình xử lý được mô tả bằng use case chi tiết, sơ đồ tuần tự và sơ đồ hoạt động. Các sơ đồ thống nhất với cơ chế xác thực Identity, revision đơn yêu cầu, workflow đơn bổ sung, vòng đời bảng giá, settlement revision và phân quyền theo action policy.", body)
    number = 7
    bookmark = 450
    add_heading(anchor, "3.2.1 Use case chi tiết", "Heading 3")
    for heading, filename, title, text in USE_CASES:
        add_heading(anchor, heading, "Heading 4")
        add_body(anchor, text, body)
        bookmark = add_figure(anchor, DIAGRAMS / "ch03" / filename, f"3-{number}", title, caption, bookmark)
        number += 1
    add_heading(anchor, "3.2.2 Sơ đồ tuần tự", "Heading 3")
    for heading, filename, title in SEQUENCES:
        add_heading(anchor, heading, "Heading 4")
        add_body(anchor, "Sơ đồ thể hiện thứ tự trao đổi giữa người dùng, giao diện, API, dịch vụ nghiệp vụ và cơ sở dữ liệu; các nhánh lỗi dừng trước khi ghi dữ liệu không hợp lệ.", body)
        bookmark = add_figure(anchor, DIAGRAMS / "ch03" / filename, f"3-{number}", title, caption, bookmark)
        number += 1
    add_heading(anchor, "3.2.3 Sơ đồ hoạt động", "Heading 3")
    for heading, filename, title in ACTIVITIES:
        add_heading(anchor, heading, "Heading 4")
        add_body(anchor, "Sơ đồ làm rõ các điều kiện rẽ nhánh, trạng thái chờ, bước kiểm tra và kết quả cuối của quy trình nghiệp vụ.", body)
        bookmark = add_figure(anchor, DIAGRAMS / "ch03" / filename, f"3-{number}", title, caption, bookmark)
        number += 1


UI_FIGURES = [
    ("3.3.1.1 Đăng nhập hệ thống", "ui-login.png", "Thiết kế giao diện đăng nhập hệ thống", "Thiết kế đăng nhập cung cấp biểu mẫu tên đăng nhập, mật khẩu, trạng thái xử lý và thông báo lỗi an toàn. Hình được dùng để minh họa bố cục; cơ chế xác thực thực tế theo ASP.NET Core Identity."),
    ("3.3.1.2 Dashboard đơn hàng cá nhân", "ui-dashboard-my-orders.png", "Thiết kế giao diện dashboard đơn hàng cá nhân", "Dashboard trình bày kỳ đang mở, hạn gửi, KPI và lối tắt tạo, sao chép, theo dõi đơn của người dùng."),
    ("3.3.2.1 Lịch sử đơn yêu cầu", "ui-order-history.png", "Thiết kế giao diện lịch sử đơn yêu cầu", "Màn hình hỗ trợ lọc theo kỳ và trạng thái, xem chi tiết, revision cùng thao tác sửa hoặc hủy khi đủ điều kiện."),
    ("3.3.2.2 Tạo đơn yêu cầu văn phòng phẩm", "ui-order-create.png", "Thiết kế giao diện tạo đơn yêu cầu văn phòng phẩm", "Luồng tạo đơn dùng hai bước: chọn văn phòng phẩm và rà soát trước khi gửi; các trạng thái lỗi, rỗng và disabled phản ánh điều kiện nghiệp vụ."),
    ("3.3.2.3 Rà soát kỳ và đơn bổ sung", "ui-period-operations.png", "Thiết kế giao diện rà soát kỳ và đơn bổ sung", "Khu vực vận hành kỳ hiển thị tiến độ, blocker, nhà cung cấp, bảng giá, bản xem trước và hàng chờ đơn bổ sung theo quyền quản lý."),
    ("3.3.3.1 Quản lý bảng giá", "ui-price-lists.png", "Thiết kế giao diện quản lý bảng giá", "Workspace bảng giá thể hiện nhà cung cấp, phiên bản, hiệu lực, trạng thái và panel chi tiết để hạn chế đưa quá nhiều cột lên lưới chính."),
    ("3.3.3.2 Quản lý danh mục văn phòng phẩm", "ui-library-items.png", "Thiết kế giao diện quản lý danh mục văn phòng phẩm", "Lưới danh mục hỗ trợ tìm kiếm, phân trang, chọn cột và panel chi tiết; thao tác thay đổi trạng thái dùng xóa mềm và khôi phục."),
    ("3.3.3.3 Phân quyền theo nhóm", "ui-permission-groups.png", "Thiết kế giao diện phân quyền theo nhóm", "Màn hình DEV quản lý persona, membership và mapping quyền với thông tin giải thích, giới hạn an toàn và audit."),
    ("3.3.4.1 Tổng hợp toàn công ty theo kỳ", "ui-all-orders-summary.png", "Thiết kế giao diện tổng hợp toàn công ty theo kỳ", "Màn hình tổng hợp hiển thị KPI, bộ lọc và danh sách đơn trong phạm vi công ty; dữ liệu thực tế vẫn do backend giới hạn theo claim và action policy."),
]


def build_33_34(anchor, body, caption, *_):
    add_heading(anchor, "3.3 HỆ THỐNG MÀN HÌNH", "Heading 2")
    add_body(anchor, "Frontend Blazor/Radzen xác định route và capability đang có; Atlas được dùng làm nguồn thiết kế bố cục sau khi đối chiếu với API, trạng thái và quyền của backend. Vì một số màn hình còn đang hoàn thiện, các hình dưới đây được gọi là thiết kế giao diện, không được dùng làm bằng chứng kiểm thử ở Chương 4.", body)
    groups = {
        "3.3.1": "3.3.1 Giao diện chung",
        "3.3.2": "3.3.2 Giao diện nghiệp vụ đơn yêu cầu",
        "3.3.3": "3.3.3 Giao diện quản trị",
        "3.3.4": "3.3.4 Giao diện tổng hợp và báo biểu",
    }
    inserted = set()
    number = 28
    bookmark = 490
    for heading, filename, title, text in UI_FIGURES:
        group = heading[:5]
        if group not in inserted:
            add_heading(anchor, groups[group], "Heading 3")
            inserted.add(group)
        add_heading(anchor, heading, "Heading 4")
        add_body(anchor, text, body)
        bookmark = add_figure(anchor, SCREENSHOTS / filename, f"3-{number}", title, caption, bookmark)
        number += 1
    add_heading(anchor, "3.4 HỆ THỐNG BÁO BIỂU", "Heading 2")
    add_body(anchor, "Hệ thống cung cấp báo cáo theo ba phạm vi: cá nhân, phòng ban và toàn công ty. Phạm vi hiển thị được suy ra từ permission của phiên đăng nhập; client không thể tự truyền mã phòng ban hoặc công ty để mở rộng dữ liệu.", body)
    add_body(anchor, "Báo cáo gồm KPI tổng đơn, tổng dòng, tổng số lượng và tổng giá trị; xu hướng theo kỳ; phân bố trạng thái; tổng hợp theo phòng ban; văn phòng phẩm nổi bật. Khi kỳ đã chốt, dữ liệu tiền ưu tiên settlement snapshot và allocation để bảo toàn khả năng đối chiếu.", body)
    add_body(anchor, "Người có REPORT_EXPORT có thể xuất CSV UTF-8 hoặc workbook XLSX gồm các sheet Summary, Items, Departments, Trend và TopProducts. Chức năng xuất kiểm soát phạm vi, giới hạn dữ liệu và xử lý giá trị có nguy cơ trở thành công thức trong bảng tính.", body)
    add_body(anchor, "Thông báo trong ứng dụng hiển thị số chưa đọc, cho phép đánh dấu đã đọc và mở đúng route liên quan. Email outbox là kênh bổ sung có trạng thái, số lần thử, lịch thử lại và khóa chống trùng; khi email bị tắt hoặc gửi lỗi, thông báo trong ứng dụng vẫn là nguồn theo dõi chính.", body)


def build_ch45(anchor, body, caption, table_sample):
    add_heading(anchor, "Chương 4. THỬ NGHIỆM", "Heading 1")
    add_body(anchor, "Chương này trình bày các cổng kiểm chứng tại commit 5f24fda ngày 25/07/2026. Kết quả được ghi đúng theo lệnh chạy thực tế; lỗi còn lại được báo cáo trung thực và không được thay bằng số liệu lịch sử. Log ứng dụng được cấu trúc hóa bằng Serilog để hỗ trợ truy vết trong quá trình thử nghiệm [10].", body)
    add_heading(anchor, "4.1 CÁC KỊCH BẢN THỬ NGHIỆM", "Heading 2")
    add_heading(anchor, "4.1.1 Mốc mã nguồn và môi trường", "Heading 3")
    add_body(anchor, "Mốc kiểm tra là commit 5f24fda286799d4b9869c0d381b8042fbc2c2120 trên Windows, .NET 10, cấu hình Release. Repository sạch trước khi dựng ứng viên luận văn. Integration test mặc định chạy các kiểm tra an toàn; sáu kiểm tra SQL Server LocalDB đặc thù chỉ chạy khi bật opt-in cô lập.", body)
    add_table_caption(anchor, "4-1", "Phạm vi lệnh kiểm chứng", caption)
    add_table(DOC, anchor, ["Hạng mục", "Lệnh", "Mục đích"], [
        ["Build", "dotnet build gtas_vpp.sln -c Release", "Biên dịch toàn solution."],
        ["Backend unit test", "dotnet test gtas_vpp_be.Tests -c Release --no-build", "Kiểm tra service, controller, RBAC và quy tắc nghiệp vụ."],
        ["Frontend unit test", "dotnet test gtas_vpp_fe.Tests -c Release --no-build", "Kiểm tra helper, localization và hợp đồng kiến trúc UI."],
        ["Integration test", "dotnet test gtas_vpp.IntegrationTests -c Release --no-build", "Kiểm tra safety contract và fixture SQL opt-in."],
        ["EF model", "dotnet ef migrations has-pending-model-changes", "Đối chiếu model với migration snapshot."],
        ["Dependency audit", "dotnet list gtas_vpp.sln package --vulnerable --include-transitive", "Tìm package có lỗ hổng đã biết."],
    ], [1800, 3800, 3472], table_sample=table_sample)
    add_heading(anchor, "4.1.2 Kịch bản nghiệp vụ đại diện", "Heading 3")
    add_table_caption(anchor, "4-2", "Ma trận kiểm thử nghiệp vụ", caption)
    add_table(DOC, anchor, ["Mã", "Kịch bản", "Kết quả mong đợi", "Bằng chứng"], [
        ["TC-01", "Đăng nhập và tải quyền", "Identity xác thực đúng, quyền bị cưỡng chế ở API.", "Backend/frontend unit test và Playwright E2E cô lập."],
        ["TC-02", "Tạo, sửa, hủy và revision đơn thường", "Không tạo trùng đơn hiện hành; lịch sử và rowversion được bảo toàn.", "Backend unit và integration safety."],
        ["TC-03", "Đơn bổ sung", "Kiểm tra lý do, quota, Pending và quyết định duyệt/từ chối.", "Backend service/controller test."],
        ["TC-04", "Danh mục, nhà cung cấp và bảng giá", "Ràng buộc mã, hiệu lực, publish/expire và dòng giá đúng.", "Backend unit và EF model."],
        ["TC-05", "Xem trước, xác nhận và hiệu chỉnh chốt kỳ", "Blocker được phát hiện; snapshot bất biến, idempotent và four-eyes.", "Settlement service tests."],
        ["TC-06", "Báo cáo, thông báo và email outbox", "Đúng phạm vi, xuất dữ liệu an toàn, chống gửi trùng và retry.", "Report/notification/outbox tests."],
    ], [900, 2300, 2900, 2972], table_sample=table_sample, center_columns=[0])
    add_heading(anchor, "4.2 KẾT QUẢ THỬ NGHIỆM", "Heading 2")
    add_table_caption(anchor, "4-3", "Tổng hợp kết quả ngày 25/07/2026", caption)
    add_table(DOC, anchor, ["Hạng mục", "Tổng", "Đạt", "Không đạt/bỏ qua", "Kết luận"], [
        ["Build Release", "Toàn solution", "0 lỗi", "0 cảnh báo", "Đạt"],
        ["Backend unit test", "414", "412", "2 / 0", "Chưa đạt"],
        ["Frontend unit test", "151", "150", "1 / 0", "Chưa đạt"],
        ["Integration test", "20", "14", "0 / 6", "Đạt trong phạm vi mặc định"],
        ["EF pending model changes", "1", "1", "0", "Đạt: model khớp snapshot"],
        ["NuGet vulnerability audit", "Toàn solution", "Không phát hiện", "0", "Đạt tại thời điểm kiểm tra"],
        ["Playwright E2E cô lập", "27", "12", "15 / 0", "Chưa đạt"],
    ], [2100, 1200, 1600, 1900, 2272], table_sample=table_sample, center_columns=[1,2,3,4])
    add_heading(anchor, "4.2.1 Phân tích kết quả", "Heading 3")
    add_body(anchor, "Hai backend test không đạt cho thấy hợp đồng kiểm thử chưa thống nhất với mô hình ba persona hiện hành: một test còn chờ bốn persona, một test chờ giới hạn permission mapping khác với response thực tế. Đây là lỗi cổng chất lượng cần xử lý trước khi kết luận sẵn sàng bàn giao.", body)
    add_body(anchor, "Frontend còn một test không đạt do thiếu 17 khóa tài nguyên VI/EN của luồng tạo đơn. Integration test mặc định đạt 14/14; sáu test LocalDB được bỏ qua đúng thiết kế opt-in. EF xác nhận model không có thay đổi chưa scaffold, nhưng trạng thái migration đã áp dụng cần một database đích cụ thể mới kiểm tra được.", body)
    add_body(anchor, "Playwright E2E chạy trên Aspire và LocalDB cô lập đạt 12/27 ca. Mười lăm ca không đạt tập trung ở luồng tạo và bổ sung đơn, đăng ký tài khoản, hộp thông báo, phân quyền, catalog, bảng dữ liệu dài, bố cục sticky/header, nhận diện thương hiệu và accessibility của trang đăng nhập. Kết quả này cho thấy chức năng nền đã có nhưng giao diện hiện hành chưa đủ ổn định để kết luận sẵn sàng bàn giao.", body)
    add_heading(anchor, "4.3 XỬ LÝ NGOẠI LỆ VÀ TỒN ĐỌNG KỸ THUẬT", "Heading 2")
    add_heading(anchor, "4.3.1 Xử lý ngoại lệ", "Heading 3")
    for label, text in [
        ("Xác thực và quyền: ", "Tài khoản không hợp lệ, bị khóa hoặc thiếu action policy bị từ chối; lỗi an toàn được hiển thị thay cho chi tiết nội bộ."),
        ("Đơn yêu cầu: ", "Danh sách rỗng, số lượng không dương, trùng văn phòng phẩm, quá hạn hoặc sai trạng thái bị chặn trước khi ghi."),
        ("Đồng thời: ", "Unique index, rowversion, transaction và idempotency key ngăn tạo trùng hoặc ghi đè dữ liệu cũ."),
        ("Chốt kỳ: ", "Thiếu giá, còn đơn Pending hoặc hash preview cũ làm quy trình dừng; revision cũ không bị ghi đè khi hiệu chỉnh."),
        ("Thông báo: ", "Outbox giữ trạng thái và lịch retry; lỗi SMTP không làm mất thông báo trong ứng dụng."),
    ]:
        add_labeled_body(anchor, label, text, body)
    add_heading(anchor, "4.3.2 Tồn đọng kỹ thuật", "Heading 3")
    add_body(anchor, "Cổng unit test còn ba lỗi; Playwright E2E còn 15/27 ca không đạt, chủ yếu do selector hoặc thành phần tương tác không xuất hiện, sai khác bố cục sticky/header, mô hình persona cũ trong test và thiếu tiêu đề tài liệu ở trang đăng nhập. Chưa có kiểm thử tải dài hạn, dữ liệu gần quy mô vận hành, nhà cung cấp SMTP thật hoặc bằng chứng giám sát production cho phiên bản này.", body)

    add_heading(anchor, "Chương 5. KẾT LUẬN", "Heading 1")
    add_body(anchor, "Chương này đối chiếu kết quả với sáu mục tiêu tại Bảng 1.4 và phân biệt rõ chức năng đã hiện thực với mức bằng chứng kiểm thử tại mốc hiện hành.", body)
    add_heading(anchor, "5.1 KẾT QUẢ ĐỐI CHIẾU VỚI MỤC TIÊU", "Heading 2")
    add_table_caption(anchor, "5-1", "Đối chiếu kết quả thực hiện với mục tiêu đề tài", caption)
    add_table(DOC, anchor, ["STT", "Mục tiêu", "Kết quả", "Đánh giá"], [
        ["1", "Website quản lý yêu cầu văn phòng phẩm theo kỳ", "Đã có kỳ độc lập, đơn thường, đơn bổ sung, revision, lịch sử và trạng thái.", "Đạt về chức năng"],
        ["2", "Quản lý danh mục và bảng giá", "Đã có danh mục, đơn vị tính, nhà cung cấp, bảng giá, hiệu lực và dòng giá.", "Đạt về chức năng"],
        ["3", "Tổng hợp nhu cầu và chốt kỳ", "Đã có preview, blocker, settlement snapshot, allocation và correction revision.", "Đạt về chức năng"],
        ["4", "Phân quyền theo nhóm và thành phần", "Đã có ba persona, membership, page/component/action mapping; hai test RBAC còn lỗi.", "Chưa đạt cổng kiểm thử"],
        ["5", "Kiểm thử và triển khai", "Build và integration mặc định đạt; unit test còn ba lỗi và Playwright E2E đạt 12/27 ca.", "Chưa đạt"],
        ["6", "Bảo mật và toàn vẹn dữ liệu", "Identity, policy, rowversion, unique constraint, audit, idempotency và snapshot đã được triển khai.", "Đạt về chức năng"],
    ], [800, 2500, 3972, 1800], table_sample=table_sample, center_columns=[0,3])
    add_heading(anchor, "5.2 CÁC VẤN ĐỀ CÒN TỒN ĐỌNG", "Heading 2")
    add_body(anchor, "Ba unit test cần được đồng bộ với mô hình persona và tài nguyên bản địa hóa trước khi bàn giao. Mười lăm lỗi Playwright E2E cần được phân loại thành lỗi sản phẩm và lỗi hợp đồng kiểm thử, sau đó sửa và chạy lại toàn bộ trên đúng commit ứng viên. Các test LocalDB opt-in cần kết quả riêng nếu được dùng làm bằng chứng migration, seed hoặc concurrency.", body)
    add_body(anchor, "Hệ thống chưa có kiểm thử tải dài hạn, mẫu báo cáo PDF chính thức, tích hợp nhân sự, nhà cung cấp email thật và bộ bằng chứng monitoring/backup/restore cho môi trường vận hành của phiên bản hiện hành.", body)
    add_heading(anchor, "5.3 HƯỚNG PHÁT TRIỂN", "Heading 2")
    add_body(anchor, "Ưu tiên gần nhất là đưa toàn bộ unit, integration opt-in và Playwright E2E về trạng thái xanh; sau đó chuẩn hóa dữ liệu gần thực tế, kiểm thử tải, mẫu XLSX/PDF, SMTP production, đồng bộ tổ chức và giám sát vận hành.", body)
    add_body(anchor, "GTAS VPP đã hình thành nền tảng nghiệp vụ đầy đủ cho quản lý yêu cầu văn phòng phẩm theo kỳ. Giá trị tiếp theo nằm ở việc hoàn thiện cổng kiểm thử và bằng chứng vận hành để chuyển từ sản phẩm học thuật có chức năng sang hệ thống có thể bàn giao ổn định.", body)


def build_appendix_references(anchor, body, *_):
    add_heading(anchor, "PHỤ LỤC", "Heading 1")
    add_heading(anchor, "PHỤ LỤC A. HƯỚNG DẪN DÀNH CHO NHÂN VIÊN", "Heading 2")
    add_heading(anchor, "A.1 Đăng nhập", "Heading 3")
    add_body(anchor, "Truy cập trang GTAS VPP, nhập tên đăng nhập và mật khẩu. Sau khi xác thực, hệ thống tải quyền và mở route đầu tiên được cấp. Không có lựa chọn môi trường trên giao diện sử dụng thông thường.", body)
    add_heading(anchor, "A.2 Tạo, sao chép, sửa và hủy đơn", "Heading 3")
    add_body(anchor, "Tại Đơn hàng của tôi, chọn tạo đơn mới hoặc sao chép kỳ trước. Tìm văn phòng phẩm, nhập số lượng dương, thêm ghi chú và chuyển sang bước rà soát trước khi gửi. Mỗi người dùng chỉ có một đơn thường hiện hành trong kỳ.", body)
    add_body(anchor, "Mở lịch sử để xem chi tiết. Nếu đơn và kỳ còn cho phép, chọn sửa để tạo revision mới hoặc hủy và nhập lý do. Đơn đã duyệt, từ chối, hủy hoặc thuộc kỳ đã chốt không được chỉnh sửa trái trạng thái.", body)
    add_heading(anchor, "A.3 Đơn bổ sung và thông báo", "Heading 3")
    add_body(anchor, "Chọn tạo đơn bổ sung khi hệ thống hiển thị chức năng, nhập lý do bắt buộc và gửi chờ xử lý. Theo dõi trạng thái trong lịch sử và hộp thư thông báo; chọn thông báo để mở chức năng liên quan.", body)
    add_heading(anchor, "PHỤ LỤC B. HƯỚNG DẪN DÀNH CHO QUẢN LÝ", "Heading 2")
    add_heading(anchor, "B.1 Duyệt hoặc từ chối đơn bổ sung", "Heading 3")
    add_body(anchor, "Mở danh sách đơn bổ sung chờ xử lý, kiểm tra kỳ, người tạo, lý do và chi tiết. Chọn Duyệt hoặc Từ chối; thao tác chỉ áp dụng cho đơn Pending trong phạm vi được cấp.", body)
    add_heading(anchor, "B.2 Xem trước và xác nhận chốt kỳ", "Heading 3")
    add_body(anchor, "Chọn kỳ, nhà cung cấp chính và bảng giá phù hợp. Tải bản xem trước, xử lý toàn bộ blocker, mặt hàng thiếu giá và đơn Pending. Khi dữ liệu hợp lệ, xác nhận chốt để tạo settlement snapshot; không sửa trực tiếp dữ liệu snapshot trong cơ sở dữ liệu.", body)
    add_heading(anchor, "B.3 Báo cáo và xuất dữ liệu", "Heading 3")
    add_body(anchor, "Chọn phạm vi được cấp, năm và tháng để xem KPI, xu hướng, trạng thái, phòng ban và văn phòng phẩm nổi bật. Người có REPORT_EXPORT có thể tải CSV hoặc XLSX; nếu dữ liệu lớn, thu hẹp bộ lọc trước khi xuất.", body)
    add_heading(anchor, "PHỤ LỤC C. HƯỚNG DẪN DÀNH CHO DEV", "Heading 2")
    add_body(anchor, "DEV quản lý tài khoản, membership, persona và mapping quyền trong trang Phân quyền; kiểm tra audit trước và sau thay đổi. Không cấp quyền nghiệp vụ vượt quá nhu cầu, không tự loại bỏ quyền quản trị cuối cùng và không chỉnh trực tiếp PasswordHash hoặc token trong cơ sở dữ liệu.", body)

    add_heading(anchor, "TÀI LIỆU THAM KHẢO", "Heading 1")
    refs = [
        '[1] Microsoft, “ASP.NET Core documentation,” Microsoft Learn. [Trực tuyến]. Địa chỉ: https://learn.microsoft.com/aspnet/core/. [Truy cập: 25/07/2026].',
        '[2] Microsoft, “ASP.NET Core Identity,” Microsoft Learn. [Trực tuyến]. Địa chỉ: https://learn.microsoft.com/aspnet/core/security/authentication/identity. [Truy cập: 25/07/2026].',
        '[3] Microsoft, “ASP.NET Core Blazor,” Microsoft Learn. [Trực tuyến]. Địa chỉ: https://learn.microsoft.com/aspnet/core/blazor/. [Truy cập: 25/07/2026].',
        '[4] Microsoft, “Entity Framework Core documentation,” Microsoft Learn. [Trực tuyến]. Địa chỉ: https://learn.microsoft.com/ef/core/. [Truy cập: 25/07/2026].',
        '[5] Microsoft, “SQL Server documentation,” Microsoft Learn. [Trực tuyến]. Địa chỉ: https://learn.microsoft.com/sql/sql-server/. [Truy cập: 25/07/2026].',
        '[6] Microsoft, “.NET Aspire documentation,” Microsoft Learn. [Trực tuyến]. Địa chỉ: https://learn.microsoft.com/dotnet/aspire/. [Truy cập: 25/07/2026].',
        '[7] Radzen, “Radzen Blazor Components,” Radzen. [Trực tuyến]. Địa chỉ: https://blazor.radzen.com/. [Truy cập: 25/07/2026].',
        '[8] Microsoft, “Playwright for .NET,” Playwright. [Trực tuyến]. Địa chỉ: https://playwright.dev/dotnet/. [Truy cập: 25/07/2026].',
        '[9] Docker, “Docker Compose,” Docker Docs. [Trực tuyến]. Địa chỉ: https://docs.docker.com/compose/. [Truy cập: 25/07/2026].',
        '[10] Serilog Project, “Serilog,” [Trực tuyến]. Địa chỉ: https://serilog.net/. [Truy cập: 25/07/2026].',
        '[11] Mapster Project, “Mapster,” GitHub. [Trực tuyến]. Địa chỉ: https://github.com/MapsterMapper/Mapster. [Truy cập: 25/07/2026].',
        '[12] Odoo, “Purchase,” Odoo documentation. [Trực tuyến]. Địa chỉ: https://www.odoo.com/documentation/18.0/applications/inventory_and_mrp/purchase.html. [Truy cập: 25/07/2026].',
        '[13] Atlassian, “Service request management,” Atlassian. [Trực tuyến]. Địa chỉ: https://www.atlassian.com/software/jira/service-management/product-guide/getting-started/service-request-management. [Truy cập: 25/07/2026].',
        '[14] Nguyễn An Nam, “Mã nguồn và hồ sơ kiểm chứng đề tài GTAS VPP,” 2026.',
        '[15] Zoho Creator, “Approval workflows,” Zoho. [Trực tuyến]. Địa chỉ: https://www.zoho.com/creator/approval-workflow/. [Truy cập: 25/07/2026].',
    ]
    for ref in refs:
        paragraph = add_body(anchor, ref, body)
        paragraph.paragraph_format.space_after = Pt(0)


def digest_locked(document: Document) -> dict[str, str]:
    ranges = [
        ("locked_chapter_1", "Chương 1. GIỚI THIỆU", "Chương 2. PHƯƠNG PHÁP THỰC HIỆN"),
        ("locked_2_1", "2.1 CÁC HỆ THỐNG TƯƠNG TỰ", "2.2 CÔNG NGHỆ SỬ DỤNG"),
        ("locked_2_2", "2.2 CÔNG NGHỆ SỬ DỤNG", "2.3 PHÂN TÍCH YÊU CẦU"),
        ("locked_2_3_2", "2.3.2 Sơ đồ chức năng", "2.3.3 Sơ đồ Use case tổng quát"),
    ]
    result = {}
    for name, start_text, end_text in ranges:
        start = find_paragraph(document, start_text)._p
        end = find_paragraph(document, end_text)._p
        chunks = []
        current = start
        while current is not None and current is not end:
            chunks.append(current.xml)
            current = current.getnext()
        result[name] = hashlib.sha256("".join(chunks).encode("utf-8")).hexdigest()
    return result


def rebuild(input_path: Path, output_path: Path) -> None:
    global DOC, SECTION_PROPS
    if input_path.resolve() == output_path.resolve():
        raise ValueError("Input and output DOCX must be different.")
    document = Document(input_path)
    DOC = document
    body = find_body_sample(document)
    caption = find_caption_sample(document)
    table_sample = find_table_sample(document)
    locked_before = digest_locked(document)
    SECTION_PROPS = {
        "chapter_3": preceding_section_properties(find_paragraph(document, "CHƯƠNG 3. THIẾT KẾ")),
        "chapter_4": preceding_section_properties(find_paragraph(document, "Chương 4. THỬ NGHIỆM")),
        "chapter_5": preceding_section_properties(find_paragraph(document, "Chương 5. KẾT LUẬN")),
        "appendix": preceding_section_properties(find_paragraph(document, "PHỤ LỤC")),
        "references": preceding_section_properties(find_paragraph(document, "TÀI LIỆU THAM KHẢO")),
    }

    # Work from the end of the document toward the beginning so stable anchors remain available.
    replace_range(document, "PHỤ LỤC", "TÀI LIỆU THAM KHẢO", lambda anchor: None)
    references = find_paragraph(document, "TÀI LIỆU THAM KHẢO")
    end = document.add_paragraph("__END_OF_DOCUMENT__")
    remove_range(references, end)
    build_appendix_references(end, body)
    end._element.getparent().remove(end._element)

    replace_range(document, "Chương 4. THỬ NGHIỆM", "PHỤ LỤC", lambda anchor: build_ch45(anchor, body, caption, table_sample))
    replace_range(document, "3.3 HỆ THỐNG MÀN HÌNH", "Chương 4. THỬ NGHIỆM", lambda anchor: build_33_34(anchor, body, caption, table_sample))
    replace_range(document, "3.2 MÔ HÌNH XỬ LÝ", "3.3 HỆ THỐNG MÀN HÌNH", lambda anchor: build_32(anchor, body, caption, table_sample))
    replace_range(document, "3.1 MÔ HÌNH DỮ LIỆU", "3.2 MÔ HÌNH XỬ LÝ", lambda anchor: build_31(anchor, body, caption, table_sample))
    replace_range(document, "2.3.3 Sơ đồ Use case tổng quát", "CHƯƠNG 3. THIẾT KẾ", lambda anchor: build_233(anchor, body, caption, table_sample))
    replace_range(document, "2.3.1 Các quy trình, nghiệp vụ", "2.3.2 Sơ đồ chức năng", lambda anchor: build_231(anchor, body))

    # Restore the exact section boundaries only after all content ranges have
    # been replaced; otherwise a later range replacement can consume the empty
    # paragraph that carries w:sectPr immediately before its end heading.
    for heading, key in [
        ("CHƯƠNG 3. THIẾT KẾ", "chapter_3"),
        ("Chương 4. THỬ NGHIỆM", "chapter_4"),
        ("Chương 5. KẾT LUẬN", "chapter_5"),
        ("PHỤ LỤC", "appendix"),
        ("TÀI LIỆU THAM KHẢO", "references"),
    ]:
        add_section_break(find_paragraph(document, heading), SECTION_PROPS[key])

    locked_after = digest_locked(document)
    if locked_before != locked_after:
        changed = [name for name in locked_before if locked_before[name] != locked_after[name]]
        raise RuntimeError(f"Locked ranges changed unexpectedly: {changed}")
    mark_fields_dirty(document)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    rebuild(args.input, args.output)
    print(args.output.resolve())


if __name__ == "__main__":
    main()
