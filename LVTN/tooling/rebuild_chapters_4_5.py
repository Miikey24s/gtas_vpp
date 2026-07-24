#!/usr/bin/env python3
"""Rebuild Chapters 4 and 5 of the GTAS VPP thesis.

The script locates the replacement range by stable headings, keeps the input
document untouched, and writes a separate DOCX output.
"""

from __future__ import annotations

import argparse
import copy
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


FONT_NAME = "Times New Roman"
FONT_SIZE = Pt(13)
CHAPTER_4_HEADING = "Chương 4. THỬ NGHIỆM"
CHAPTER_5_HEADING = "Chương 5. KẾT LUẬN"
APPENDIX_HEADING = "PHỤ LỤC"

PRESERVED_HEADING_BOOKMARKS = {
    "CHƯƠNG 4. THỬ NGHIỆM": "Chương 4. THỬ NGHIỆM",
    "4.1 CÁC KỊCH BẢN THỬ NGHIỆM": "4.1 CÁC KỊCH BẢN THỬ NGHIỆM",
    "4.1.1 Phạm vi và phương pháp ghi nhận": "4.1.1 Mốc mã nguồn, môi trường và phương pháp ghi nhận",
    "4.1.2 Ma trận test case đại diện": "4.1.2 Kịch bản nghiệp vụ đại diện",
    "4.2 KẾT QUẢ THỬ NGHIỆM CÁC KỊCH BẢN": "4.2 KẾT QUẢ THỬ NGHIỆM",
    "4.3 XỬ LÝ NGOẠI LỆ VÀ TỒN ĐỌNG KỸ THUẬT": "4.3 XỬ LÝ NGOẠI LỆ VÀ TỒN ĐỌNG KỸ THUẬT",
    "4.3.1 Xử lý các trường hợp ngoại lệ": "4.3.1 Xử lý các trường hợp ngoại lệ",
    "4.3.2 Tồn đọng kỹ thuật": "4.3.2 Tồn đọng kỹ thuật",
    "CHƯƠNG 5. KẾT LUẬN": "Chương 5. KẾT LUẬN",
    "5.1 KẾT QUẢ ĐỐI CHIẾU VỚI MỤC TIÊU": "5.1 KẾT QUẢ ĐỐI CHIẾU VỚI MỤC TIÊU",
    "5.2 CÁC VẤN ĐỀ CÒN TỒN ĐỌNG": "5.2 CÁC VẤN ĐỀ CÒN TỒN ĐỌNG",
    "5.3 HƯỚNG PHÁT TRIỂN": "5.3 HƯỚNG PHÁT TRIỂN",
}


def normalized(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip().casefold()


def find_paragraph(document: Document, exact_text: str):
    wanted = normalized(exact_text)
    matches = [p for p in document.paragraphs if normalized(p.text) == wanted]
    if len(matches) != 1:
        raise RuntimeError(
            f"Expected exactly one paragraph '{exact_text}', found {len(matches)}."
        )
    return matches[0]


def find_body_sample(document: Document):
    sample_prefix = "Trong hoạt động của doanh nghiệp"
    for paragraph in document.paragraphs:
        if normalized(paragraph.text).startswith(normalized(sample_prefix)):
            return paragraph
    raise RuntimeError("Cannot locate the body-format sample in Chapter 1.")


def find_caption_sample(document: Document):
    for paragraph in document.paragraphs:
        if normalized(paragraph.text).startswith(normalized("Bảng 3-1:")):
            return paragraph
    raise RuntimeError("Cannot locate an existing table-caption sample.")


def find_table_sample(document: Document):
    """Use the approved Chapter 1 result table as the table-format authority."""
    for table in document.tables:
        headers = [normalized(cell.text) for cell in table.rows[0].cells]
        if headers[:2] == [normalized("STT"), normalized("Kết quả cần đạt")]:
            return table
    raise RuntimeError("Cannot locate the Chapter 1 table-format sample.")


def remove_range(start_paragraph, end_paragraph) -> None:
    """Remove body elements from start (inclusive) to end (exclusive)."""
    element = start_paragraph._p
    stop = end_paragraph._p
    while element is not None and element is not stop:
        next_element = element.getnext()
        element.getparent().remove(element)
        element = next_element
    if element is None:
        raise RuntimeError("The end heading was not reached while replacing chapters 4–5.")


def preceding_section_properties(paragraph):
    element = paragraph._p.getprevious()
    while element is not None:
        if element.tag == qn("w:p"):
            p_pr = element.find(qn("w:pPr"))
            if p_pr is not None:
                sect_pr = p_pr.find(qn("w:sectPr"))
                if sect_pr is not None:
                    return copy.deepcopy(sect_pr)
        element = element.getprevious()
    raise RuntimeError(f"Cannot find the section break before '{paragraph.text}'.")


def add_section_break(anchor, section_properties) -> None:
    paragraph = anchor.insert_paragraph_before()
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.append(copy.deepcopy(section_properties))


def apply_font(run, *, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = FONT_NAME
    run.font.size = FONT_SIZE
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_NAME)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def copy_paragraph_properties(target, sample) -> None:
    current = target._p.pPr
    if current is not None:
        target._p.remove(current)
    if sample._p.pPr is not None:
        target._p.insert(0, copy.deepcopy(sample._p.pPr))


def copy_run_properties(target, sample) -> None:
    current = target._r.rPr
    if current is not None:
        target._r.remove(current)
    if sample._r.rPr is not None:
        target._r.insert(0, copy.deepcopy(sample._r.rPr))


def add_heading(anchor, text: str, style: str):
    # Heading styles in the approved Chapters 1–2 already carry the template
    # scale (24/15/14 pt), spacing and alignment. Direct formatting here would
    # override that hierarchy and make later chapters look smaller.
    return anchor.insert_paragraph_before(text, style=style)


def add_body(anchor, text: str, sample):
    paragraph = anchor.insert_paragraph_before()
    copy_paragraph_properties(paragraph, sample)
    run = paragraph.add_run(text)
    apply_font(run, bold=False, italic=False)
    return paragraph


def add_labeled_body(anchor, label: str, text: str, sample):
    paragraph = anchor.insert_paragraph_before()
    copy_paragraph_properties(paragraph, sample)
    label_run = paragraph.add_run(label)
    apply_font(label_run, bold=True, italic=False)
    value_run = paragraph.add_run(text)
    apply_font(value_run, bold=False, italic=False)
    return paragraph


def add_caption(anchor, text: str, sample):
    paragraph = anchor.insert_paragraph_before()
    copy_paragraph_properties(paragraph, sample)
    if ":" not in text:
        raise ValueError(f"Caption must contain a colon: {text}")
    number, description = text.split(":", 1)
    number_run = paragraph.add_run(number + ":")
    apply_font(number_run, bold=True, italic=True)
    number_run.font.underline = True
    description_run = paragraph.add_run(" " + description.strip())
    apply_font(description_run, bold=False, italic=False)
    description_run.font.underline = False
    return paragraph


def set_cell_width(cell, width_twips: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    # Chapter 1–2 tables use the legacy left/right OOXML names. Keeping those
    # names makes the rebuilt tables structurally and visually identical.
    for tag, value in (("top", top), ("left", start), ("bottom", bottom), ("right", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def mark_row_repeat(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_table_geometry(table, widths: list[int]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    # Chapters 1–2 use the full 9072-DXA content width for data tables.
    target_width = 9072
    scale = target_width / sum(widths)
    widths = [round(width * scale) for width in widths]
    widths[-1] += target_width - sum(widths)

    tbl_w.set(qn("w:w"), str(target_width))
    tbl_w.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        prevent_row_split(row)
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def format_table_text(
    table,
    sample_table,
    center_columns: set[int] | None = None,
) -> None:
    center_columns = center_columns or set()
    header_sample = sample_table.rows[0].cells[0]
    centered_body_sample = sample_table.rows[1].cells[0]
    left_body_sample = sample_table.rows[1].cells[1]

    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            for paragraph in cell.paragraphs:
                sample_paragraph = (
                    header_sample.paragraphs[0]
                    if row_index == 0
                    else centered_body_sample.paragraphs[0]
                    if column_index in center_columns
                    else left_body_sample.paragraphs[0]
                )
                copy_paragraph_properties(paragraph, sample_paragraph)
                for run in paragraph.runs:
                    sample_run = sample_paragraph.runs[0]
                    copy_run_properties(run, sample_run)
            if row_index == 0:
                shade_cell(cell, "EDEDED")
    mark_row_repeat(table.rows[0])


def add_table(
    document,
    anchor,
    headers,
    rows,
    widths,
    *,
    table_sample,
    center_columns=None,
):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = table_sample.style
    for index, value in enumerate(headers):
        table.rows[0].cells[index].text = value
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
    set_table_geometry(table, widths)
    format_table_text(table, table_sample, set(center_columns or []))
    anchor._p.addprevious(table._tbl)
    return table


def mark_fields_dirty(document: Document) -> None:
    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def add_heading_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))

    insertion_index = 1 if paragraph._p.pPr is not None else 0
    paragraph._p.insert(insertion_index, start)
    paragraph._p.append(end)


def capture_heading_bookmarks(document: Document):
    captured = {}
    for old_heading, rebuilt_heading in PRESERVED_HEADING_BOOKMARKS.items():
        paragraph = next(
            (
                candidate
                for candidate in document.paragraphs
                if candidate.style.name.startswith("Heading")
                and normalized(candidate.text)
                in {normalized(old_heading), normalized(rebuilt_heading)}
            ),
            None,
        )
        if paragraph is None:
            raise RuntimeError(f"Cannot find original heading bookmark: {old_heading}")
        bookmark = paragraph._p.find(".//" + qn("w:bookmarkStart"))
        if bookmark is None:
            raise RuntimeError(f"Original heading has no bookmark: {old_heading}")
        captured[old_heading] = (
            bookmark.get(qn("w:name")),
            int(bookmark.get(qn("w:id"))),
        )
    return captured


def reattach_heading_bookmarks(document: Document, captured) -> None:
    for old_heading, new_heading in PRESERVED_HEADING_BOOKMARKS.items():
        paragraph = next(
            (
                candidate
                for candidate in document.paragraphs
                if candidate.style.name.startswith("Heading")
                and normalized(candidate.text) == normalized(new_heading)
            ),
            None,
        )
        if paragraph is None:
            raise RuntimeError(f"Cannot find rebuilt heading for bookmark: {new_heading}")
        name, bookmark_id = captured[old_heading]
        add_heading_bookmark(paragraph, name, bookmark_id)


def rebuild(input_path: Path, output_path: Path) -> None:
    if input_path.resolve() == output_path.resolve():
        raise ValueError("Input and output DOCX paths must be different.")
    if not input_path.exists():
        raise FileNotFoundError(input_path)

    document = Document(input_path)
    start = find_paragraph(document, CHAPTER_4_HEADING)
    original_chapter_5 = find_paragraph(document, CHAPTER_5_HEADING)
    appendix = find_paragraph(document, APPENDIX_HEADING)
    body_sample = find_body_sample(document)
    caption_sample = find_caption_sample(document)
    table_sample = find_table_sample(document)
    preserved_bookmarks = capture_heading_bookmarks(document)
    chapter_4_section_properties = preceding_section_properties(original_chapter_5)
    chapter_5_section_properties = preceding_section_properties(appendix)

    remove_range(start, appendix)

    # Chapter 4
    add_heading(appendix, CHAPTER_4_HEADING, "Heading 1")
    add_body(
        appendix,
        "Chương này trình bày phương pháp kiểm chứng, các kịch bản đại diện và kết quả thu được tại một mốc mã nguồn xác định. Số liệu hiện hành được tách khỏi bằng chứng lịch sử để bảo đảm khả năng truy vết và tránh cộng gộp các lần chạy khác nhau.",
        body_sample,
    )
    add_heading(appendix, "4.1 CÁC KỊCH BẢN THỬ NGHIỆM", "Heading 2")
    add_heading(appendix, "4.1.1 Mốc mã nguồn, môi trường và phương pháp ghi nhận", "Heading 3")
    add_body(
        appendix,
        "Kết quả hiện hành được ghi nhận ngày 25/07/2026 tại mốc WT-2026-07-25, trên nền commit fbbaa49. Mốc này gồm các thay đổi chưa commit của đợt hoàn thiện vai trò, tài nguyên bản địa hóa và giao diện; do đó mọi kết luận chỉ áp dụng cho đúng trạng thái làm việc tại thời điểm chạy.",
        body_sample,
    )
    add_body(
        appendix,
        "Môi trường kiểm chứng là Windows, .NET 10 và cấu hình Release. Kiểm thử tích hợp sử dụng SQL Server LocalDB cô lập; kiểm tra gói phụ thuộc dùng các nguồn NuGet đang cấu hình. Toàn bộ Playwright E2E không được chạy tại mốc này vì cổng kiểm thử đơn vị và tích hợp còn lỗi.",
        body_sample,
    )
    add_body(
        appendix,
        "Bên cạnh kết quả hiện hành, luận văn sử dụng riêng hồ sơ LEAN-09 ngày 17/07/2026 để chứng minh hệ thống từng hoàn thành các lượt E2E xác thực cô lập. Bằng chứng lịch sử được ghi thành bảng riêng và không cộng vào tổng số test của ngày 25/07/2026.",
        body_sample,
    )
    add_caption(appendix, "Bảng 4-1: Phạm vi và bằng chứng kiểm thử tại mốc WT-2026-07-25", caption_sample)
    add_table(
        document,
        appendix,
        ["Hạng mục", "Mốc/ngày", "Môi trường", "Bằng chứng", "Mục đích"],
        [
            ["Build", "WT-2026-07-25\n25/07/2026", ".NET 10, Release", "dotnet build solution -c Release --no-restore", "Kiểm tra khả năng biên dịch và liên kết toàn solution."],
            ["Backend test", "WT-2026-07-25\n25/07/2026", "xUnit, Release", "dotnet test bộ phía máy chủ --no-build", "Kiểm tra dịch vụ, bộ điều khiển, phân quyền, kiến trúc và quy tắc nghiệp vụ."],
            ["Frontend test", "WT-2026-07-25\n25/07/2026", "xUnit, Release", "dotnet test bộ giao diện --no-build", "Kiểm tra helper, bản địa hóa, kiến trúc và hợp đồng giao diện."],
            ["Integration test", "WT-2026-07-25\n25/07/2026", "SQL Server LocalDB cô lập", "dotnet test bộ tích hợp với opt-in LocalDB", "Kiểm tra migration, seed, reseed, reset và ràng buộc dữ liệu thật."],
            ["Playwright E2E", "LEAN-09\n17/07/2026", "Aspire và dữ liệu cô lập", "Hồ sơ docs/execution/LEAN-09.md", "Kiểm tra luồng xác thực, responsive và hành vi người dùng trên trình duyệt."],
            ["Dependency audit", "WT-2026-07-25\n25/07/2026", "NuGet", "dotnet package list --vulnerable", "Kiểm tra lỗ hổng đã biết của gói trực tiếp và bắc cầu."],
        ],
        [1600, 1650, 1450, 2250, 2050],
        table_sample=table_sample,
    )

    add_heading(appendix, "4.1.2 Kịch bản nghiệp vụ đại diện", "Heading 3")
    add_body(
        appendix,
        "Các kịch bản được chọn theo mức độ rủi ro nghiệp vụ: xác thực và phân quyền; kỳ đặt hàng; đơn thường và đơn bổ sung; danh mục, nhà cung cấp và bảng giá; tổng hợp, chốt kỳ; báo cáo, thông báo và email outbox. Kết quả thực tế được liên kết với đúng nhóm test hoặc hồ sơ kiểm chứng.",
        body_sample,
    )
    add_caption(appendix, "Bảng 4-2: Ma trận kiểm thử nghiệp vụ đại diện", caption_sample)
    add_table(
        document,
        appendix,
        ["Mã", "Kịch bản", "Kết quả mong đợi", "Bằng chứng tại mốc kiểm tra"],
        [
            ["TC-01", "Đăng nhập và kiểm tra phạm vi hành động theo vai trò.", "Tài khoản hợp lệ được xác thực; API chỉ chấp nhận hành động nằm trong phạm vi quyền.", "Phần lớn kiểm thử phân quyền đạt, nhưng còn 2 kiểm thử phía máy chủ không đạt do hợp đồng vai trò đang được tái cấu trúc."],
            ["TC-02", "Tạo, sửa, hủy, sao chép và theo dõi đơn thường trong kỳ.", "Chỉ một đơn thường hiện hành cho mỗi người dùng/kỳ; phiên bản chỉnh sửa và lịch sử được lưu; thao tác bị giới hạn bởi trạng thái và hạn gửi.", "Các trường hợp thuộc 412 kiểm thử phía máy chủ đạt; luồng trình duyệt từng đạt trong hồ sơ LEAN-09."],
            ["TC-03", "Tạo và xử lý đơn bổ sung.", "Kiểm tra điều kiện, lý do, hạn mức và trạng thái; chuyên viên quản lý văn phòng phẩm được phê duyệt hoặc từ chối.", "Các kiểm thử service/controller liên quan đạt trong bộ phía máy chủ; luồng UI có trong danh sách 27 kịch bản Playwright."],
            ["TC-04", "Quản lý danh mục văn phòng phẩm, đơn vị tính, nhà cung cấp và bảng giá.", "Dữ liệu tuân thủ trạng thái hiệu lực, thời gian áp dụng và ràng buộc duy nhất; thao tác sai bị từ chối.", "Các test danh mục/bảng giá đạt trong bộ backend; XLSX đã có bằng chứng trong LEAN-07."],
            ["TC-05", "Tổng hợp nhu cầu, chọn nhà cung cấp/bảng giá và chốt kỳ.", "Hệ thống chặn khi còn điều kiện chưa đạt; khi hợp lệ tạo bản chụp dữ liệu giá, chi phí và phân bổ bất biến để đối chiếu.", "Các kiểm thử chốt kỳ và tính nhất quán dữ liệu đạt trong bộ phía máy chủ hiện hành."],
            ["TC-06", "Xem báo cáo, phát thông báo và xử lý hàng đợi email.", "Dữ liệu đúng phạm vi; XLSX đối chiếu được; thông báo đúng người nhận; hàng đợi tránh gửi trùng và hỗ trợ thử lại.", "Các kiểm thử báo cáo, thông báo và outbox đạt; môi trường cục bộ hỗ trợ SMTP tương thích Mailpit theo LEAN-07."],
        ],
        [1000, 2400, 2900, 2700],
        table_sample=table_sample,
        center_columns=[0],
    )

    add_heading(appendix, "4.1.3 Kịch bản chất lượng và vận hành", "Heading 3")
    add_caption(appendix, "Bảng 4-3: Ma trận kiểm thử chất lượng", caption_sample)
    add_table(
        document,
        appendix,
        ["Mã", "Kịch bản", "Tiêu chí", "Kết quả"],
        [
            ["TC-07", "Build toàn solution ở cấu hình Release.", "Không có lỗi hoặc cảnh báo biên dịch.", "Đạt: 0 cảnh báo, 0 lỗi."],
            ["TC-08", "Chạy toàn bộ backend test.", "Tất cả test đạt trước khi bàn giao.", "Chưa đạt: 412/414 đạt; còn 2 lỗi về vai trò và giới hạn cấp quyền."],
            ["TC-09", "Chạy toàn bộ frontend test.", "Tất cả test đạt và tài nguyên VI/EN đồng bộ.", "Chưa đạt: 150/151 đạt; thiếu 17 khóa bản địa hóa của luồng tạo đơn."],
            ["TC-10", "Chạy integration test trên LocalDB cô lập.", "Migration, seed, reseed, reset và ràng buộc dữ liệu đều đạt.", "Chưa đạt: 19/20 đạt; fixture vai trò chưa khớp refactor."],
            ["TC-11", "Khám phá và chạy Playwright E2E.", "Các luồng xác thực quan trọng chạy trên môi trường cô lập.", "Phát hiện 27 kịch bản; chưa chạy toàn bộ ở mốc hiện tại vì cổng kiểm thử đang lỗi."],
            ["TC-12", "Audit gói phụ thuộc.", "Không có lỗ hổng đã biết trong kết quả NuGet hiện hành.", "Đạt: không phát hiện package dễ tổn thương từ các nguồn cấu hình."],
        ],
        [1000, 2500, 2900, 2600],
        table_sample=table_sample,
        center_columns=[0],
    )

    add_heading(appendix, "4.2 KẾT QUẢ THỬ NGHIỆM", "Heading 2")
    add_body(
        appendix,
        "Kết quả ngày 25/07/2026 cho thấy solution vẫn biên dịch ổn định, nhưng cổng kiểm thử chưa đạt hoàn toàn. Bốn test không đạt đều tập trung ở đợt thay đổi vai trò và tài nguyên bản địa hóa, không phải lỗi build. Vì vậy mốc WT-2026-07-25 chưa được xem là mốc sẵn sàng bàn giao.",
        body_sample,
    )
    add_caption(appendix, "Bảng 4-4: Tổng hợp kết quả hiện hành ngày 25/07/2026", caption_sample)
    add_table(
        document,
        appendix,
        ["Hạng mục", "Tổng", "Đạt", "Không đạt/bỏ qua", "Kết luận"],
        [
            ["Build Release", "Toàn solution", "0 lỗi", "0 cảnh báo", "Đạt"],
            ["Backend test", "414", "412", "2 / 0", "Chưa đạt"],
            ["Frontend test", "151", "150", "1 / 0", "Chưa đạt"],
            ["Integration test", "20", "19", "1 / 0", "Chưa đạt"],
            ["Playwright discovery", "27", "27 được phát hiện", "Chưa chạy toàn bộ", "Chưa đủ bằng chứng hiện hành"],
            ["NuGet vulnerability audit", "13 project", "Không phát hiện", "0", "Đạt tại thời điểm kiểm tra"],
        ],
        [2100, 1200, 1700, 1900, 2450],
        table_sample=table_sample,
        center_columns=[1, 2, 3, 4],
    )

    add_heading(appendix, "4.2.1 Bằng chứng E2E lịch sử", "Heading 3")
    add_body(
        appendix,
        "Hồ sơ LEAN-09 tại commit b681c556, ngày 17/07/2026, ghi nhận hai lượt Playwright E2E xác thực cô lập đều đạt 16/16 kịch bản và một lượt kiểm tra responsive shell đạt 1/1. Đây là bằng chứng rằng hạ tầng E2E và các luồng chính từng hoạt động ổn định, nhưng không thay thế việc chạy lại sau khi hoàn tất refactor hiện nay.",
        body_sample,
    )
    add_caption(appendix, "Bảng 4-5: Kết quả E2E tại mốc LEAN-09 ngày 17/07/2026", caption_sample)
    add_table(
        document,
        appendix,
        ["Lượt", "Mốc mã nguồn", "Phạm vi", "Kết quả", "Thời gian"],
        [
            ["Diễn tập 1", "b681c556", "16 luồng UI xác thực cô lập", "16/16 đạt", "427,0 giây"],
            ["Diễn tập 2", "b681c556", "16 luồng UI xác thực cô lập", "16/16 đạt", "406,9 giây"],
            ["Responsive shell", "b681c556", "Kiểm tra shell có mục tiêu", "1/1 đạt", "49,0 giây"],
        ],
        [1600, 1400, 2672, 1600, 1800],
        table_sample=table_sample,
        center_columns=[0, 1, 3, 4],
    )

    add_heading(appendix, "4.2.2 Phân tích kết quả", "Heading 3")
    add_body(
        appendix,
        "Hai backend test không đạt phản ánh sự không thống nhất giữa mô hình vai trò đang được điều chỉnh và hợp đồng kiểm thử về số vai trò chuẩn cùng giới hạn cấp quyền. Lỗi integration test có cùng nguyên nhân khi dữ liệu seed trong fixture chưa tạo đúng tập vai trò được mong đợi.",
        body_sample,
    )
    add_body(
        appendix,
        "Frontend test không đạt do 17 khóa bản địa hóa của luồng tạo đơn chưa có đủ cặp tài nguyên tiếng Việt và tiếng Anh. Đây là lỗi chất lượng giao diện có thể xác định rõ, cần được sửa trước khi chạy lại Playwright E2E và audit responsive trên mốc mã nguồn mới.",
        body_sample,
    )

    add_heading(appendix, "4.3 XỬ LÝ NGOẠI LỆ VÀ TỒN ĐỌNG KỸ THUẬT", "Heading 2")
    add_heading(appendix, "4.3.1 Xử lý các trường hợp ngoại lệ", "Heading 3")
    add_labeled_body(appendix, "Xác thực và phân quyền: ", "Yêu cầu không có thông tin xác thực hoặc không đủ quyền bị từ chối tại API; thay đổi quyền được kiểm tra lại trước hành động nghiệp vụ.", body_sample)
    add_labeled_body(appendix, "Dữ liệu đơn: ", "Danh sách rỗng, số lượng không dương, mặt hàng trùng hoặc thao tác ngoài trạng thái cho phép bị chặn trước khi ghi dữ liệu.", body_sample)
    add_labeled_body(appendix, "Đồng thời và revision: ", "Ràng buộc duy nhất và transaction ngăn tạo nhiều đơn thường hiện hành cho cùng người dùng/kỳ; thay đổi hợp lệ tạo revision và lưu nhật ký.", body_sample)
    add_labeled_body(appendix, "Đơn bổ sung: ", "Hệ thống kiểm tra kỳ, lý do, quota và trạng thái chờ; quyết định phê duyệt hoặc từ chối được lưu cùng người thực hiện và thời điểm.", body_sample)
    add_labeled_body(appendix, "Chốt kỳ: ", "Quy trình dừng khi còn blocker, thiếu giá hoặc dữ liệu không nhất quán; khi thành công tạo snapshot giá, chi phí và phân bổ để bảo toàn kết quả đối chiếu.", body_sample)
    add_labeled_body(appendix, "Thông báo và email outbox: ", "Bản ghi có trạng thái xử lý, số lần thử và khóa chống trùng; lỗi gửi không làm mất sự kiện nghiệp vụ và có thể được thử lại.", body_sample)

    add_heading(appendix, "4.3.2 Tồn đọng kỹ thuật", "Heading 3")
    add_body(appendix, "Cổng kiểm thử của mốc WT-2026-07-25 còn bốn test không đạt. Cần hoàn tất hợp đồng vai trò, seed dữ liệu và tài nguyên bản địa hóa trước khi xem đây là mốc ứng viên bàn giao.", body_sample)
    add_body(appendix, "Chưa có kết quả kiểm thử tải dài hạn, kiểm thử sức chịu tải và bộ dữ liệu gần quy mô vận hành. Các kết quả unit/integration hiện hành chưa đủ để suy ra hiệu năng production.", body_sample)
    add_body(appendix, "Hệ thống đã xuất XLSX, nhưng chưa có mẫu XLSX/PDF chính thức theo biểu mẫu doanh nghiệp và chưa kiểm chứng đầy đủ bố cục in với dữ liệu lớn.", body_sample)
    add_body(appendix, "Notification và email outbox đã có; môi trường cục bộ hỗ trợ SMTP tương thích Mailpit. Nhà cung cấp email thật, push notification production và quy trình theo dõi gửi thất bại chưa được hoàn thiện.", body_sample)
    add_body(appendix, "Tích hợp chính thức với hệ thống nhân sự chưa có. Quy trình backup, restore, giám sát và cảnh báo trên máy chủ thật cũng chưa có đủ bằng chứng diễn tập.", body_sample)

    add_section_break(appendix, chapter_4_section_properties)

    # Chapter 5
    add_heading(appendix, CHAPTER_5_HEADING, "Heading 1")
    add_body(
        appendix,
        "Chương này đối chiếu kết quả thực hiện với sáu mục tiêu đã xác lập, nêu rõ giới hạn còn lại và đề xuất hướng phát triển. Đánh giá dựa trên bằng chứng Chương 4, trong đó tách riêng chức năng đã hiện thực với trạng thái cổng kiểm thử tại mốc kiểm tra.",
        body_sample,
    )
    add_heading(appendix, "5.1 KẾT QUẢ ĐỐI CHIẾU VỚI MỤC TIÊU", "Heading 2")
    add_caption(appendix, "Bảng 5-1: Đối chiếu kết quả thực hiện với mục tiêu đề tài", caption_sample)
    add_table(
        document,
        appendix,
        ["STT", "Mục tiêu", "Kết quả và bằng chứng", "Đánh giá"],
        [
            ["1", "Website quản lý yêu cầu theo kỳ", "Đã có luồng đơn thường, đơn bổ sung, phiên bản chỉnh sửa, lịch sử, hạn gửi và trạng thái; được bao phủ bởi kiểm thử phía máy chủ và bằng chứng E2E lịch sử.", "Đạt"],
            ["2", "Quản lý danh mục và bảng giá", "Đã quản lý văn phòng phẩm, đơn vị tính, nhà cung cấp, bảng giá, giá hiệu lực và ràng buộc liên quan; có xuất XLSX.", "Đạt"],
            ["3", "Tổng hợp nhu cầu và chốt kỳ", "Đã tổng hợp theo phạm vi, kiểm tra điều kiện chặn, chọn giá và tạo bản chụp kết quả chốt kỳ gồm giá, chi phí và phân bổ.", "Đạt"],
            ["4", "Phân quyền theo nhóm và thành phần", "Nền tảng phân quyền đã có, nhưng mốc kiểm tra còn 2 kiểm thử phía máy chủ và 1 kiểm thử tích hợp không đạt do hợp đồng vai trò đang được tái cấu trúc.", "Chưa đạt tại mốc kiểm tra"],
            ["5", "Kiểm thử và triển khai", "Build đạt và có bằng chứng E2E lịch sử; tuy nhiên bộ test hiện hành còn lỗi và toàn bộ E2E chưa được chạy lại sau tái cấu trúc.", "Chưa đạt tại mốc kiểm tra"],
            ["6", "Bảo mật và toàn vẹn dữ liệu", "Đã có kiểm tra quyền tại API, kiểm tra dữ liệu đầu vào, giao dịch, ràng buộc dữ liệu, phiên bản chỉnh sửa, nhật ký và outbox; kiểm tra NuGet không phát hiện lỗ hổng đã biết.", "Đạt"],
        ],
        [850, 2300, 4150, 1700],
        table_sample=table_sample,
        center_columns=[0, 3],
    )
    add_body(
        appendix,
        "Bốn trong sáu mục tiêu đạt bằng chứng hiện hành. Hai mục tiêu về phân quyền và kiểm thử/triển khai chưa đạt tại mốc WT-2026-07-25 vì đợt tái cấu trúc chưa vượt qua cổng kiểm thử. Kết luận này không phủ nhận các chức năng đã có, mà xác định rõ điều kiện cần hoàn thành trước khi bàn giao.",
        body_sample,
    )

    add_heading(appendix, "5.2 CÁC VẤN ĐỀ CÒN TỒN ĐỌNG", "Heading 2")
    add_labeled_body(appendix, "Cổng kiểm thử sau tái cấu trúc: ", "Hợp đồng vai trò, dữ liệu seed và tài nguyên bản địa hóa chưa đồng bộ hoàn toàn, làm bốn test không đạt. Toàn bộ Playwright E2E và UI audit chưa được chạy lại trên mốc mã nguồn này.", body_sample)
    add_labeled_body(appendix, "Kiểm thử tải và dữ liệu quy mô lớn: ", "Hệ thống chưa có kết quả kiểm thử tải dài hạn, đo tải đồng thời hoặc đánh giá với dữ liệu gần quy mô vận hành. Vì vậy chưa đủ cơ sở kết luận về ngưỡng đáp ứng và tài nguyên máy chủ cần thiết.", body_sample)
    add_labeled_body(appendix, "Biểu mẫu báo cáo chính thức: ", "Chức năng XLSX đã tồn tại, nhưng mẫu XLSX/PDF theo quy định doanh nghiệp, bố cục in thống nhất và kiểm chứng với đơn có nhiều dòng chưa được hoàn thiện.", body_sample)
    add_labeled_body(appendix, "Tích hợp thông báo và dữ liệu tổ chức: ", "Notification, email outbox và SMTP cục bộ đã có. Phần còn thiếu là nhà cung cấp email thật, push notification production và tích hợp chính thức với hệ thống nhân sự để đồng bộ người dùng, phòng ban và trạng thái làm việc.", body_sample)
    add_labeled_body(appendix, "Vận hành trên máy chủ thật: ", "Quy trình backup, restore, monitoring, cảnh báo và diễn tập phục hồi trên môi trường máy chủ thật chưa có đủ bằng chứng. Đây là khoảng cách chính giữa mức hoàn thành đề tài và khả năng vận hành production ổn định.", body_sample)

    add_heading(appendix, "5.3 HƯỚNG PHÁT TRIỂN", "Heading 2")
    add_labeled_body(appendix, "Hoàn tất cổng kiểm thử: ", "Chuẩn hóa mô hình vai trò và seed dữ liệu, bổ sung đầy đủ tài nguyên bản địa hóa, chạy lại toàn bộ unit test, integration test, Playwright E2E và UI audit trước mỗi mốc bàn giao.", body_sample)
    add_labeled_body(appendix, "Đánh giá hiệu năng: ", "Xây dựng kịch bản tải theo số người dùng, số phòng ban, số đơn và số dòng mặt hàng; đo thời gian đáp ứng, thông lượng, bộ nhớ và truy vấn chậm trên SQL Server.", body_sample)
    add_labeled_body(appendix, "Hoàn thiện tài liệu xuất: ", "Chuẩn hóa mẫu PDF/XLSX theo biểu mẫu được phê duyệt, hỗ trợ dữ liệu dài, kiểm tra ngắt trang và lưu phiên bản tài liệu đã phát hành.", body_sample)
    add_labeled_body(appendix, "Mở rộng tích hợp: ", "Kết nối nhà cung cấp email và push phù hợp, đồng bộ dữ liệu tổ chức từ hệ thống nhân sự, đồng thời thiết lập giám sát trạng thái outbox và cơ chế xử lý thất bại.", body_sample)
    add_labeled_body(appendix, "Nâng cao vận hành: ", "Thiết lập backup tự động, kiểm thử restore định kỳ, dashboard giám sát, cảnh báo và tài liệu xử lý sự cố; chỉ mở production sau khi các diễn tập đạt yêu cầu.", body_sample)
    add_body(
        appendix,
        "GTAS VPP đã hình thành đầy đủ nền tảng nghiệp vụ cho quản lý yêu cầu văn phòng phẩm theo kỳ. Giá trị tiếp theo của đề tài nằm ở việc hoàn tất cổng kiểm thử, kiểm chứng quy mô và chuẩn hóa vận hành để chuyển từ sản phẩm học thuật có chức năng hoàn chỉnh sang hệ thống có thể triển khai ổn định trong thực tế.",
        body_sample,
    )

    add_section_break(appendix, chapter_5_section_properties)

    reattach_heading_bookmarks(document, preserved_bookmarks)
    mark_fields_dirty(document)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path, help="Input DOCX path")
    parser.add_argument("--output", required=True, type=Path, help="Output DOCX path")
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    rebuild(args.input, args.output)
    print(args.output.resolve())
