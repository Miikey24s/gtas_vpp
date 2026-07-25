#!/usr/bin/env python3
"""Refine the GTAS VPP thesis candidate without rewriting stable sections.

The script uses headings and stable text as locators, preserves the cover and
locked ranges, restructures the physical data model, adds focused use-case
specifications, and rebuilds the reference list in first-citation order.
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree

import rebuild_final_candidate as base
from rebuild_chapters_4_5 import (
    add_body,
    add_caption,
    add_heading,
    add_labeled_body,
    add_table,
    apply_font,
    find_body_sample,
    find_caption_sample,
    find_paragraph,
    find_table_sample,
    mark_fields_dirty,
    normalized,
    remove_range,
    set_table_geometry,
)


ROOT = Path(__file__).resolve().parents[2]
DIAGRAMS = ROOT / "LVTN" / "diagrams"


REFERENCES = [
    '[1] Google Workspace Learning Center, “Share & collaborate on a spreadsheet,” Google. [Trực tuyến]. Địa chỉ: https://support.google.com/a/users/answer/13309904. [Truy cập: 25/07/2026].',
    '[2] Microsoft Support, “Quick tips: Share and collaborate with Excel for the web,” Microsoft. [Trực tuyến]. Địa chỉ: https://support.microsoft.com/office/c72e0df7-d999-4225-b839-6fd79fc97145. [Truy cập: 25/07/2026].',
    '[3] Odoo, “Purchase,” Odoo 18.0 Documentation. [Trực tuyến]. Địa chỉ: https://www.odoo.com/documentation/18.0/applications/inventory_and_mrp/purchase.html. [Truy cập: 25/07/2026].',
    '[4] Atlassian, “Get started with service requests in Jira Service Management,” Atlassian. [Trực tuyến]. Địa chỉ: https://www.atlassian.com/software/jira/service-management/product-guide/getting-started/service-request-management. [Truy cập: 25/07/2026].',
    '[5] Zoho Creator Help, “Understanding workflows,” Zoho. [Trực tuyến]. Địa chỉ: https://help.zoho.com/portal/en/kb/creator/developer-guide/workflows/understand-workflows/articles/understand-workflows. [Truy cập: 25/07/2026].',
    '[6] Microsoft, “ASP.NET Core documentation,” Microsoft Learn. [Trực tuyến]. Địa chỉ: https://learn.microsoft.com/aspnet/core/?view=aspnetcore-10.0. [Truy cập: 25/07/2026].',
    '[7] Microsoft, “Introduction to Identity on ASP.NET Core,” Microsoft Learn. [Trực tuyến]. Địa chỉ: https://learn.microsoft.com/aspnet/core/security/authentication/identity?view=aspnetcore-10.0. [Truy cập: 25/07/2026].',
    '[8] Microsoft, “ASP.NET Core Blazor,” Microsoft Learn. [Trực tuyến]. Địa chỉ: https://learn.microsoft.com/aspnet/core/blazor/?view=aspnetcore-10.0. [Truy cập: 25/07/2026].',
    '[9] Microsoft, “Overview of Entity Framework Core,” Microsoft Learn. [Trực tuyến]. Địa chỉ: https://learn.microsoft.com/ef/core/. [Truy cập: 25/07/2026].',
    '[10] Microsoft, “Microsoft SQL documentation,” Microsoft Learn. [Trực tuyến]. Địa chỉ: https://learn.microsoft.com/sql/?view=sql-server-ver17. [Truy cập: 25/07/2026].',
    '[11] Radzen, “Radzen Blazor Components,” Radzen. [Trực tuyến]. Địa chỉ: https://blazor.radzen.com/. [Truy cập: 25/07/2026].',
    '[12] Mapster, “Mapster - The Mapper of Your Domain,” Mapster. [Trực tuyến]. Địa chỉ: https://mapstermapper.github.io/Mapster/. [Truy cập: 25/07/2026].',
    '[13] Serilog Project, “Serilog - simple .NET logging with fully-structured events,” Serilog. [Trực tuyến]. Địa chỉ: https://serilog.net/. [Truy cập: 25/07/2026].',
    '[14] Docker, “Docker Compose,” Docker Docs. [Trực tuyến]. Địa chỉ: https://docs.docker.com/compose/. [Truy cập: 25/07/2026].',
    '[15] NGINX, “Module ngx_http_proxy_module,” NGINX Documentation. [Trực tuyến]. Địa chỉ: https://nginx.org/en/docs/http/ngx_http_proxy_module.html. [Truy cập: 25/07/2026].',
    '[16] Microsoft, “Aspire documentation,” Aspire. [Trực tuyến]. Địa chỉ: https://aspire.dev/. [Truy cập: 25/07/2026].',
    '[17] xUnit.net, “Getting Started with xUnit.net v3,” xUnit.net. [Trực tuyến]. Địa chỉ: https://xunit.net/docs/getting-started/v3/getting-started. [Truy cập: 25/07/2026].',
    '[18] Microsoft, “Installation - Playwright .NET,” Playwright. [Trực tuyến]. Địa chỉ: https://playwright.dev/dotnet/docs/intro. [Truy cập: 25/07/2026].',
]


PHYSICAL_GROUPS = [
    (
        "Tổ chức và phân quyền",
        "3-1",
        [
            ["Người dùng (AspNetUsers)", "Id int NOT NULL (PK); UserName, NormalizedUserName nvarchar(256) NULL; EmployeeCode nvarchar(50) NULL; AccountStatus nvarchar(32) NOT NULL; RowVersion rowversion NOT NULL.", "Filtered UNIQUE cho tên đăng nhập, email và mã nhân viên; rowversion kiểm soát đồng thời."],
            ["Phòng ban (Departments)", "Id uniqueidentifier NOT NULL (PK); Code nvarchar(50) NOT NULL; Name nvarchar(200) NOT NULL; ParentDepartmentId uniqueidentifier NULL (FK).", "FK tự tham chiếu dùng Restrict; filtered UNIQUE cho Code của bản ghi còn hiệu lực."],
            ["Nhóm quyền (PermissionGroups)", "Id uniqueidentifier NOT NULL (PK); GroupCode, GroupName nvarchar(50) NOT NULL; IsDeleted bit NOT NULL.", "Filtered UNIQUE cho GroupCode đang hoạt động."],
            ["Thành viên nhóm", "Id uniqueidentifier NOT NULL (PK); AccountId int NULL (FK); DepartmentId, PermissionGroupId uniqueidentifier NOT NULL (FK); RowVersion rowversion NOT NULL.", "Các FK dùng Restrict; một membership hoạt động cho mỗi tài khoản; CHECK phòng ban chính hợp lệ."],
            ["Ánh xạ truy cập", "PageComponentMappingId, PermissionGroupId uniqueidentifier NOT NULL; MemberCompanyCode bigint NOT NULL; IsVisible, IsEnable bit NOT NULL.", "Khóa ghép theo nhóm quyền, thành phần trang và công ty; các FK dùng Restrict."],
            ["Nhật ký bảo mật (SecurityAudits)", "Id uniqueidentifier NOT NULL (PK); Action nvarchar(100) NOT NULL; ActorUserId int NULL; OccurredAtUtc datetime2 NOT NULL; Outcome nvarchar(40) NOT NULL.", "Index kết hợp theo hành động, đối tượng và thời điểm phát sinh."],
        ],
    ),
    (
        "Danh mục, nhà cung cấp và bảng giá",
        "3-2",
        [
            ["Danh mục tra cứu", "LookupCategories.Id và LookupValues.Id: uniqueidentifier NOT NULL (PK); LookupCategoryId uniqueidentifier NULL (FK); Code, Value nvarchar(max).", "FK dùng Restrict; index theo LookupCategoryId."],
            ["Danh mục VPP", "VppCategories.Id uniqueidentifier NOT NULL (PK); VppCategoryCode, VppCategoryName nvarchar(max) NULL.", "Xóa mềm và audit bằng IsDeleted, CreatedAtUtc, UpdatedAtUtc."],
            ["Văn phòng phẩm (VppItems)", "Id uniqueidentifier NOT NULL (PK); VppCode nvarchar(64) NOT NULL; VppName nvarchar(250) NOT NULL; UomId, VppCategoryId uniqueidentifier NOT NULL (FK).", "UNIQUE VppCode; index theo danh mục, đơn vị tính và trạng thái xóa."],
            ["Nhà cung cấp (Suppliers)", "Id uniqueidentifier NOT NULL (PK); SupplierName, SupplierShortName nvarchar(max) NULL; các trường địa chỉ nullable.", "Xóa mềm; dữ liệu dịch tách theo ngôn ngữ."],
            ["Bảng giá (PriceLists)", "Id uniqueidentifier NOT NULL (PK); SupplierId uniqueidentifier NULL (FK); PriceListCode nvarchar(50) NULL; Version int NOT NULL; EffectiveFromUtc datetime2 NOT NULL; EffectiveToUtc datetime2 NULL; RowVersion rowversion NOT NULL.", "Filtered UNIQUE bảng giá mặc định và mã/phiên bản theo NCC; CHECK hiệu lực, version, chiết khấu và số tiền không âm."],
            ["Dòng giá nhà cung cấp", "Id uniqueidentifier NOT NULL (PK); PriceListId, SupplierId, VppItemId uniqueidentifier NOT NULL (FK); NetPrice decimal(19,4); VatRate decimal(5,2); MinimumOrderQuantity decimal(19,4); RowVersion rowversion NOT NULL.", "Filtered UNIQUE một dòng mặc định/mặt hàng/bảng giá; CHECK giá, VAT, MOQ và lead time."],
        ],
    ),
    (
        "Kỳ, đơn yêu cầu, revision và nhật ký",
        "3-3",
        [
            ["Kỳ (Periods)", "Id uniqueidentifier NOT NULL (PK); Year, Month, State int NOT NULL; StartAtUtc, SubmissionDeadlineUtc, SupplementApprovalDeadlineUtc datetime2 NOT NULL; RowVersion rowversion NULL.", "Filtered UNIQUE công ty/năm/tháng đang hoạt động; CHECK khoảng thời gian và State IN (0,1,2,3)."],
            ["Đơn yêu cầu (Requests)", "Id uniqueidentifier NOT NULL (PK); PeriodId uniqueidentifier NULL (FK); RequestSeriesId uniqueidentifier NOT NULL; RevisionNumber int NOT NULL; Status int NOT NULL; IdempotencyKey nvarchar(128) NULL; RowVersion rowversion NULL.", "Filtered UNIQUE revision hiện hành, đơn thường/kỳ, đơn bổ sung Pending, lần bổ sung và idempotency key."],
            ["Chi tiết đơn (RequestDetails)", "Id uniqueidentifier NOT NULL (PK); RequestId, VppId uniqueidentifier NOT NULL (FK); Qty int NOT NULL; CurrentSinglePrice bigint NOT NULL.", "FK đến Requests và VppItems dùng Restrict; index theo RequestId và VppId."],
            ["Nhật ký đơn (RequestLogs)", "Id uniqueidentifier NOT NULL (PK); RequestId uniqueidentifier NOT NULL (FK); Action nvarchar(64) NULL; LogDate datetime2 NOT NULL; Reason nvarchar(500) NULL; RevisionNumber int NULL.", "FK đến Requests dùng Restrict; index theo RequestId."],
        ],
    ),
    (
        "Kết quả chốt kỳ, chi phí, phân bổ và thông báo",
        "3-4",
        [
            ["Kết quả chốt kỳ (Settlements)", "Id uniqueidentifier NOT NULL (PK); PeriodId uniqueidentifier NOT NULL (FK); RevisionNumber int NOT NULL; IdempotencyKey nvarchar(128) NOT NULL; các số tiền decimal(19,4); RowVersion rowversion NULL.", "Filtered UNIQUE revision hiện hành; UNIQUE idempotency và số revision; CHECK kỳ và các số tiền không âm."],
            ["Dòng chốt kỳ (SettlementItems)", "Id, SettlementId uniqueidentifier NOT NULL; Quantity, NetUnitPrice, NetAmount, VatAmount, GrossAmount decimal(19,4) NOT NULL; VatRate decimal(5,2) NOT NULL.", "FK dùng Restrict; UNIQUE SettlementId/VppId; CHECK số lượng dương, giá trị và VAT hợp lệ."],
            ["Phí chốt kỳ (SettlementCharges)", "Id, SettlementId uniqueidentifier NOT NULL; ChargeType, AllocationBasis nvarchar(32) NOT NULL; Amount decimal(19,4) NOT NULL.", "FK dùng Restrict; UNIQUE SettlementId/ChargeType."],
            ["Phân bổ (SettlementAllocations)", "Id, SettlementId, SettlementItemId, RequestDetailId uniqueidentifier NOT NULL; Quantity và các số tiền decimal(19,4) NOT NULL; DepartmentCode nvarchar(64) NULL.", "FK dùng Restrict; UNIQUE một phân bổ cho mỗi chi tiết đơn trong settlement; CHECK Quantity > 0."],
            ["Thông báo (Notifications)", "Id uniqueidentifier NOT NULL (PK); UserId int NOT NULL; Type nvarchar(50) NOT NULL; CorrelationId nvarchar(100) NULL; ReadAt datetime2 NULL.", "Filtered UNIQUE người dùng/công ty/loại/correlation; index hộp thư theo trạng thái đọc và thời điểm."],
            ["Email outbox", "Id uniqueidentifier NOT NULL (PK); DeduplicationKey nvarchar(128) NOT NULL; Status nvarchar(24) NOT NULL; NextAttemptAtUtc datetime2 NOT NULL; AttemptCount int NOT NULL.", "UNIQUE DeduplicationKey; index hàng đợi theo Status/NextAttemptAtUtc."],
        ],
    ),
]


CONSTRAINT_APPENDIX = [
    ("UserNameIndex", "Filtered UNIQUE trên AspNetUsers.NormalizedUserName khi giá trị khác NULL."),
    ("UX_AspNetUsers_EmployeeCode", "Filtered UNIQUE mã nhân viên khi EmployeeCode khác NULL."),
    ("UX_Departments_Code_Active", "Filtered UNIQUE mã phòng ban khi IsDeleted = 0."),
    ("UX_UserGroupMemberships_OneActivePerUser", "Mỗi tài khoản chỉ có một membership đang hoạt động."),
    ("UX_PriceLists_OneDefault", "Mỗi phạm vi chỉ có một bảng giá mặc định còn hiệu lực."),
    ("UX_PriceLists_SupplierCodeVersion", "Không trùng mã và phiên bản bảng giá trong cùng nhà cung cấp."),
    ("UX_SupplierProductMappings_OneDefaultPerItemAndList", "Mỗi mặt hàng chỉ có một dòng giá mặc định trong một bảng giá."),
    ("UX_Periods_Company_Year_Month_Active", "Mỗi công ty chỉ có một kỳ hoạt động cho cùng năm và tháng."),
    ("UX_Requests_CurrentRevisionSeries", "Mỗi chuỗi đơn chỉ có một revision hiện hành."),
    ("UX_Requests_OneRegularPerUserPeriod", "Mỗi người dùng chỉ có một đơn thường hiện hành trong một kỳ."),
    ("UX_Requests_OnePendingSupplement", "Không tồn tại đồng thời nhiều đơn bổ sung Pending cho cùng đơn gốc."),
    ("UX_Requests_IdempotencyKey", "Chống lặp lệnh tạo đơn theo người dùng và idempotency key."),
    ("UX_Settlements_Current", "Mỗi công ty, năm và tháng chỉ có một settlement revision hiện hành."),
    ("UX_Settlements_Idempotency", "Chống tạo trùng settlement theo idempotency key."),
    ("UX_Notifications_UserCompanyTypeCorrelation", "Chống tạo trùng thông báo theo người dùng, công ty, loại và correlation."),
    ("UX_EmailOutboxMessages_DeduplicationKey", "Chống xếp trùng email vào outbox."),
]


CHECK_APPENDIX = [
    ("CK_UserGroupMemberships_ActivePrimaryDepartment", "Membership hoạt động phải gắn tài khoản và phòng ban chính hợp lệ."),
    ("CK_PriceLists_EffectiveWindow", "EffectiveToUtc phải lớn hơn EffectiveFromUtc khi có ngày kết thúc."),
    ("CK_PriceLists_VersionPositive", "Version phải lớn hơn 0."),
    ("CK_PriceLists_DiscountRateRange", "DiscountRate nằm trong khoảng 0 đến 100."),
    ("CK_SupplierProductMappings_VatRateRange", "VatRate nằm trong khoảng 0 đến 100."),
    ("CK_SupplierProductMappings_NetPriceNonNegative", "NetPrice không âm."),
    ("CK_Periods_ValidRange", "Năm, tháng, thời hạn gửi, thời hạn bổ sung và state phải hợp lệ."),
    ("CK_Settlements_AmountsNonNegative", "Các thành phần tiền của settlement không âm."),
    ("CK_Settlements_Period", "Năm, tháng và revision number của settlement hợp lệ."),
    ("CK_SettlementItems_AmountsNonNegative", "Số lượng dương; giá, thuế và thành tiền không âm."),
    ("CK_SettlementAllocations_QuantityPositive", "Số lượng phân bổ phải lớn hơn 0."),
]


PHYSICAL_TABLE_NAMES = {
    "Thành viên nhóm": "UserGroupMemberships",
    "Ánh xạ truy cập": "PageComponentMappings; GroupPageComponentMappings",
    "Danh mục tra cứu": "LookupCategories; LookupValues",
    "Danh mục VPP": "VppCategories",
    "Dòng giá nhà cung cấp": "SupplierProductMappings",
    "Email outbox": "EmailOutboxMessages",
}


USE_CASE_SPECS = {
    "login": {
        "name": "Đăng nhập và tải quyền",
        "actor": "Nhân viên, Quản lý hoặc DEV",
        "description": "Xác thực tài khoản bằng ASP.NET Core Identity và tải snapshot quyền hiện hành.",
        "pre": "Tài khoản tồn tại; người dùng chưa có phiên hợp lệ.",
        "post": "Phiên đăng nhập được tạo, cookie được thiết lập và quyền trang/thành phần/action được tải.",
        "main": "1. Nhập tên đăng nhập và mật khẩu.\n2. Hệ thống kiểm tra mật khẩu, trạng thái tài khoản và membership.\n3. Hệ thống tạo phiên, tải quyền và điều hướng đến route được phép.",
        "alt": "Sai thông tin, tài khoản bị khóa/chưa duyệt, bắt buộc đổi mật khẩu hoặc không có membership: từ chối đăng nhập và hiển thị lỗi an toàn.",
    },
    "create": {
        "name": "Tạo đơn thông thường",
        "actor": "Nhân viên hoặc Quản lý khi tạo đơn của mình",
        "description": "Tạo một đơn yêu cầu thông thường trong kỳ đang nhận đơn.",
        "pre": "Kỳ ở trạng thái đang nhận đơn; chưa quá hạn; người dùng chưa có đơn thường hiện hành trong kỳ.",
        "post": "Đơn và chi tiết được lưu; revision hiện hành và nhật ký được tạo.",
        "main": "1. Mở tạo đơn hoặc sao chép kỳ trước.\n2. Chọn văn phòng phẩm, nhập số lượng và ghi chú.\n3. Rà soát và gửi.\n4. Backend kiểm tra điều kiện, lưu đơn và trả kết quả.",
        "alt": "Danh sách rỗng, số lượng không dương, trùng mặt hàng, kỳ hết hạn hoặc lệnh trùng: hệ thống từ chối hoặc trả kết quả idempotent phù hợp.",
    },
    "edit": {
        "name": "Chỉnh sửa hoặc hủy đơn",
        "actor": "Chủ sở hữu đơn",
        "description": "Thay đổi đơn hợp lệ bằng revision mới hoặc hủy đơn nhưng vẫn giữ lịch sử.",
        "pre": "Đơn thuộc người dùng, còn ở trạng thái cho phép và rowversion còn hiệu lực.",
        "post": "Chỉnh sửa tạo revision mới; hủy đơn cập nhật trạng thái, lý do và nhật ký.",
        "main": "1. Mở đơn trong lịch sử.\n2. Chọn sửa hoặc hủy.\n3. Hệ thống kiểm tra quyền, trạng thái và rowversion.\n4. Lưu revision mới hoặc trạng thái hủy.",
        "alt": "Không phải chủ sở hữu, kỳ đã định giá/chốt, trạng thái không hợp lệ hoặc rowversion cũ: từ chối thao tác.",
    },
    "supplement": {
        "name": "Tạo đơn bổ sung",
        "actor": "Nhân viên hoặc Quản lý khi tạo đơn của mình",
        "description": "Tạo yêu cầu bổ sung gắn với đơn gốc sau cửa sổ đơn thông thường.",
        "pre": "Có đơn gốc hợp lệ; còn cửa sổ bổ sung; có lý do; chưa vượt quota và không có đơn bổ sung Pending khác.",
        "post": "Đơn bổ sung ở trạng thái Pending, gắn chuỗi đơn gốc và có nhật ký.",
        "main": "1. Chọn tạo đơn bổ sung.\n2. Nhập lý do, mặt hàng và số lượng.\n3. Hệ thống kiểm tra quota/cửa sổ/trạng thái.\n4. Lưu đơn Pending và thông báo cho người xử lý.",
        "alt": "Thiếu lý do, quá hạn, vượt quota, không có đơn gốc hoặc đã có đơn Pending: không tạo đơn.",
    },
    "approve": {
        "name": "Duyệt hoặc từ chối đơn bổ sung",
        "actor": "Quản lý có REQUEST_APPROVE hoặc REQUEST_REJECT",
        "description": "Ra quyết định đối với đơn bổ sung Pending trong phạm vi được cấp.",
        "pre": "Đơn ở trạng thái Pending, thuộc phạm vi cho phép và không do chính người xử lý tạo.",
        "post": "Đơn chuyển Approved hoặc Rejected; quyết định, lý do, thời điểm, nhật ký và thông báo được lưu.",
        "main": "1. Mở hàng chờ.\n2. Xem lý do và chi tiết.\n3. Chọn duyệt hoặc từ chối; nhập lý do khi cần.\n4. Backend kiểm tra lại và ghi quyết định.",
        "alt": "Đơn đã được xử lý, ngoài phạm vi, do chính người xử lý tạo hoặc trạng thái cũ: từ chối thao tác.",
    },
    "pricing": {
        "name": "Quản lý bảng giá",
        "actor": "Quản lý có LIBRARY_MANAGE",
        "description": "Quản lý vòng đời bảng giá theo nhà cung cấp, phiên bản, hiệu lực và dòng giá.",
        "pre": "Nhà cung cấp và văn phòng phẩm tồn tại; người dùng có quyền quản lý thư viện.",
        "post": "Bảng giá Draft được lưu hoặc chuyển Published/Expired cùng audit và rowversion mới.",
        "main": "1. Tạo hoặc sửa bảng giá Draft.\n2. Khai báo hiệu lực, VAT, phí và dòng giá.\n3. Kiểm tra độ phủ và ràng buộc.\n4. Công bố hoặc cho hết hiệu lực.",
        "alt": "Version/hiệu lực chồng lấn, giá hoặc VAT không hợp lệ, thiếu dòng giá hay rowversion cũ: không công bố/cập nhật.",
    },
    "settle": {
        "name": "Xem trước và xác nhận chốt kỳ",
        "actor": "Quản lý có quyền vận hành kỳ",
        "description": "Chọn nhà cung cấp/bảng giá, rà soát blocker và tạo settlement snapshot bất biến.",
        "pre": "Kỳ đã đóng nhận đơn hoặc đang định giá; bảng giá đã công bố và còn hiệu lực; không còn đơn Pending.",
        "post": "Settlement revision hiện hành được tạo, phân bổ khớp tổng tiền và kỳ chuyển sang đã chốt.",
        "main": "1. Chọn kỳ, nhà cung cấp và bảng giá.\n2. Tải preview, kiểm tra độ phủ, ngoại lệ và tổng tiền.\n3. Xử lý toàn bộ blocker.\n4. Xác nhận bằng hash preview và idempotency key.",
        "alt": "Thiếu giá, còn đơn Pending, preview cũ, bảng giá không hợp lệ hoặc lệnh trùng: dừng hoặc trả lại kết quả idempotent.",
    },
    "correct": {
        "name": "Hiệu chỉnh kết quả chốt kỳ",
        "actor": "Quản lý có quyền hiệu chỉnh",
        "description": "Tạo revision settlement mới mà không ghi đè snapshot lịch sử.",
        "pre": "Kỳ đã chốt; có settlement hiện hành; người xác nhận khác người xác nhận revision hiện tại; có lý do.",
        "post": "Revision mới trở thành hiện hành; revision cũ và toàn bộ dòng/phí/phân bổ được giữ nguyên để đối chiếu.",
        "main": "1. Mở kết quả chốt hiện hành.\n2. Chọn hiệu chỉnh và nhập lý do.\n3. Tải lại preview, rà soát thay đổi.\n4. Người đủ điều kiện xác nhận revision mới.",
        "alt": "Cùng người xác nhận, thiếu lý do, preview cũ hoặc dữ liệu không hợp lệ: từ chối hiệu chỉnh.",
    },
    "permissions": {
        "name": "Quản lý người dùng và phân quyền",
        "actor": "DEV",
        "description": "Quản lý tài khoản, membership, persona và mapping truy cập trong giới hạn an toàn.",
        "pre": "DEV đã đăng nhập và có permission quản trị tương ứng.",
        "post": "Thay đổi hợp lệ được lưu, session version/quyền được cập nhật và security audit được ghi.",
        "main": "1. Tìm tài khoản hoặc persona.\n2. Cập nhật trạng thái, membership hoặc mapping.\n3. Backend kiểm tra giới hạn quyền.\n4. Lưu thay đổi và phát tín hiệu cập nhật quyền.",
        "alt": "Tự nâng quyền trái phép, loại bỏ quyền quản trị cuối cùng, mapping vượt trần persona hoặc rowversion cũ: từ chối thao tác.",
    },
}


USE_CASE_LAYOUT = [
    ("3.2.1.1 Chức năng đăng nhập và tải quyền", "use-case-login.svg", "Use case đăng nhập và tải quyền", "login"),
    ("3.2.1.2 Chức năng tạo đơn yêu cầu thông thường", "use-case-create-regular-request.svg", "Use case tạo đơn yêu cầu thông thường", "create"),
    ("3.2.1.3 Chức năng chỉnh sửa và hủy đơn", "use-case-edit-cancel-request.svg", "Use case chỉnh sửa và hủy đơn", "edit"),
    ("3.2.1.4 Chức năng tạo đơn bổ sung", "use-case-create-additional-request.svg", "Use case tạo đơn bổ sung", "supplement"),
    ("3.2.1.5 Chức năng duyệt hoặc từ chối đơn bổ sung", "use-case-approve-additional-request.svg", "Use case duyệt hoặc từ chối đơn bổ sung", "approve"),
    ("3.2.1.6 Chức năng quản lý danh mục văn phòng phẩm", "use-case-manage-catalog.svg", "Use case quản lý danh mục văn phòng phẩm", None),
    ("3.2.1.7 Chức năng quản lý bảng giá", "use-case-manage-pricing.svg", "Use case quản lý bảng giá", "pricing"),
    ("3.2.1.8 Chức năng xem trước và chốt kỳ", "use-case-settle-period.svg", "Use case xem trước và chốt kỳ", "settle"),
    ("3.2.1.9 Chức năng hiệu chỉnh kết quả chốt kỳ", None, None, "correct"),
    ("3.2.1.10 Chức năng quản lý người dùng và phân quyền", "use-case-manage-permissions.svg", "Use case quản lý người dùng và phân quyền", "permissions"),
    ("3.2.1.11 Chức năng xem dashboard và dữ liệu tổng hợp", "use-case-dashboard-summary.svg", "Use case xem dashboard và dữ liệu tổng hợp", None),
]


def paragraph_text(element) -> str:
    return "".join(node.text or "" for node in element.iter(qn("w:t"))).strip()


def xml_range(document: Document, start_text: str | None, end_text: str, *, normalize_citations=False) -> bytes:
    body = document._element.body
    started = start_text is None
    parts: list[bytes] = []
    for child in body.iterchildren():
        text = paragraph_text(child)
        if not started and normalized(text) == normalized(start_text or ""):
            started = True
        if started and normalized(text) == normalized(end_text):
            break
        if started:
            data = etree.tostring(child, with_tail=False)
            if normalize_citations:
                data = re.sub(rb"\[\d+\]", b"[#]", data)
            parts.append(data)
    if not parts:
        raise RuntimeError(f"Cannot capture protected range: {start_text!r} -> {end_text!r}")
    return b"".join(parts)


def semantic_text_range(document: Document, start_text: str, end_text: str) -> str:
    body = document._element.body
    started = False
    values: list[str] = []
    for child in body.iterchildren():
        text = paragraph_text(child)
        if not started and normalized(text) == normalized(start_text):
            started = True
        if started and normalized(text) == normalized(end_text):
            break
        if started:
            # Join runs inside a paragraph/table block without introducing
            # artificial spaces. Replacing a citation may collapse several
            # Word runs into one while leaving the visible text unchanged.
            values.append("".join(node.text or "" for node in child.iter(qn("w:t"))))
    if not values:
        raise RuntimeError(f"Cannot capture semantic range: {start_text!r} -> {end_text!r}")
    text = "\n".join(values)
    text = re.sub(r"\[\d+\]", "[#]", text)
    return re.sub(r"\s+", " ", text).strip()


def protected_digests(document: Document) -> dict[str, str]:
    section_21_text = semantic_text_range(
        document,
        "2.1 CÁC HỆ THỐNG TƯƠNG TỰ",
        "2.2 CÔNG NGHỆ SỬ DỤNG",
    )
    ranges = {
        "cover": xml_range(document, None, "LỜI CẢM ƠN"),
        "chapter_1": xml_range(document, "Chương 1. GIỚI THIỆU", "Chương 2. PHƯƠNG PHÁP THỰC HIỆN"),
        "section_2_1": section_21_text.encode("utf-8"),
        "section_2_3_2": xml_range(document, "2.3.2 Sơ đồ chức năng", "2.3.3 Sơ đồ Use case tổng quát"),
    }
    return {name: hashlib.sha256(value).hexdigest() for name, value in ranges.items()}


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_paragraph_text(paragraph, text: str) -> None:
    sample_run = paragraph.runs[0] if paragraph.runs else None
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    if sample_run is not None and sample_run._r.rPr is not None:
        run._r.insert(0, copy.deepcopy(sample_run._r.rPr))
    else:
        apply_font(run, bold=False, italic=False)


def set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    set_paragraph_text(paragraph, text)
    for extra in list(cell.paragraphs[1:]):
        cell._tc.remove(extra._p)


def replace_text_everywhere(document: Document, old: str, new: str) -> int:
    count = 0
    containers = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                containers.extend(cell.paragraphs)
    for paragraph in containers:
        if old in paragraph.text:
            set_paragraph_text(paragraph, paragraph.text.replace(old, new))
            count += 1
    return count


def replace_range(document: Document, start_text: str, end_text: str, builder) -> None:
    start = find_paragraph(document, start_text)
    end = find_paragraph(document, end_text)
    remove_range(start, end)
    builder(end)


def set_cell_no_wrap(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    if tc_pr.find(qn("w:noWrap")) is None:
        tc_pr.append(OxmlElement("w:noWrap"))


def add_specification_table(document, anchor, number: int, spec: dict[str, str], caption_sample, table_sample):
    add_caption(anchor, f"Bảng 3-{number}: Đặc tả use case {spec['name'].lower()}", caption_sample)
    rows = [
        ["Tên use case", spec["name"]],
        ["Tác nhân", spec["actor"]],
        ["Mô tả", spec["description"]],
        ["Tiền điều kiện", spec["pre"]],
        ["Hậu điều kiện", spec["post"]],
        ["Luồng chính", spec["main"]],
        ["Luồng thay thế/ngoại lệ", spec["alt"]],
    ]
    table = add_table(
        document,
        anchor,
        ["Thành phần", "Nội dung"],
        rows,
        [1900, 7172],
        table_sample=table_sample,
    )
    for row in table.rows[1:]:
        set_cell_no_wrap(row.cells[0])
    return table


def build_personas(document, anchor, body, caption_sample, table_sample) -> None:
    add_heading(anchor, "2.3.3 Sơ đồ Use case tổng quát", "Heading 3")
    add_body(
        anchor,
        "Hệ thống sử dụng ba persona chuẩn EMPLOYEE, MANAGER và DEV. Trên sơ đồ UML, Quản lý được nối với Nhân viên bằng quan hệ generalization có đầu tam giác rỗng, biểu thị Quản lý sử dụng lại các use case cá nhân và có thêm nghiệp vụ quản lý. Đây là quan hệ mô hình hóa actor; các persona lưu trong RBAC vẫn là ba nhóm phẳng độc lập. DEV là vai trò quản trị kỹ thuật, không thay người dùng thực hiện nghiệp vụ thường ngày.",
        body,
    )
    base.add_figure(anchor, DIAGRAMS / "ch02" / "use-case-overview.png", "2-3", "Sơ đồ use case tổng quát hệ thống", caption_sample, 420)
    add_caption(anchor, "Bảng 2-1: Mô tả tác nhân của hệ thống", caption_sample)
    table = add_table(
        document,
        anchor,
        ["Tác nhân", "Mã persona", "Phạm vi trách nhiệm"],
        [
            ["Nhân viên", "EMPLOYEE", "Lập và theo dõi đơn của mình, xem danh mục, lịch sử, báo cáo cá nhân và thông báo."],
            ["Quản lý", "MANAGER", "Thực hiện use case của Nhân viên; xem đơn theo phạm vi được cấp, xử lý đơn bổ sung, quản lý thư viện, báo cáo và chốt kỳ."],
            ["DEV", "DEV", "Quản trị tài khoản, membership, persona, mapping truy cập và nhật ký bảo mật; không mặc định thực hiện nghiệp vụ thay người dùng."],
        ],
        [1700, 1850, 5522],
        table_sample=table_sample,
        center_columns=[1],
    )
    for row in table.rows[1:]:
        set_cell_no_wrap(row.cells[1])


def build_physical_model(document, anchor, body, caption_sample, table_sample) -> None:
    add_heading(anchor, "3.1.3 Mô hình dữ liệu vật lý", "Heading 3")
    add_body(
        anchor,
        "Mô hình vật lý được triển khai trên SQL Server. Khóa nghiệp vụ chủ yếu dùng uniqueidentifier, khóa tài khoản dùng int, thời điểm dùng datetime2, số tiền dùng decimal(19,4), thuế suất dùng decimal(5,2). Ký hiệu NULL/NOT NULL trong các bảng dưới đây thể hiện đúng tính bắt buộc của cột; PK, FK, filtered UNIQUE, CHECK và rowversion được nêu tại các điểm chi phối trực tiếp tính toàn vẹn nghiệp vụ.",
        body,
    )
    add_body(
        anchor,
        "Các bảng nghiệp vụ dùng xóa mềm qua IsDeleted và phần lớn quan hệ dùng DeleteBehavior.Restrict để tránh xóa lan dữ liệu lịch sử. Ba bảng phụ Identity AspNetUserClaims, AspNetUserLogins và AspNetUserTokens giữ Cascade theo ASP.NET Core Identity. Tên đầy đủ của các index và constraint quan trọng được tập hợp tại Phụ lục D để tránh xé identifier trong phần chính.",
        body,
    )
    for title, number, rows in PHYSICAL_GROUPS:
        add_heading(anchor, title, "Heading 4")
        add_caption(anchor, f"Bảng {number}: Cấu trúc vật lý nhóm {title.lower()}", caption_sample)
        display_rows = []
        for label, columns, constraints in rows:
            match = re.fullmatch(r"(.+?)\s+\(([^()]+)\)", label)
            if match:
                label, table_names = match.groups()
            else:
                table_names = PHYSICAL_TABLE_NAMES.get(label)
            if table_names:
                columns = f"{table_names}: {columns}"
            display_rows.append([label, columns, constraints])
        table = add_table(
            document,
            anchor,
            ["Bảng/nhóm dữ liệu", "Cột và kiểu dữ liệu chính", "Khóa, index và constraint"],
            display_rows,
            [2250, 3800, 3022],
            table_sample=table_sample,
        )
        for row in table.rows[1:]:
            set_cell_no_wrap(row.cells[0])
    add_heading(anchor, "Identity, bảng dịch, view và keyless entity", "Heading 4")
    add_body(
        anchor,
        "AspNetUsers là bảng tài khoản chính; AspNetUserClaims, AspNetUserLogins và AspNetUserTokens là các bảng phụ của Identity. Bảy bảng dịch lưu bản dịch theo cặp thực thể - LanguageCode bằng filtered unique index và không được trình bày thành data dictionary chi tiết trong phần chính. BusinessDataTranslationBase dùng chiến lược TPC nên không có bảng riêng.",
        body,
    )
    add_body(
        anchor,
        "Hai view v_Users và v_WFXCompany là keyless, chỉ đọc và không thuộc migration model. StoredProcedureResultDTO chỉ là DTO keyless nhận kết quả thủ tục, không phải entity hoặc view nên không được đưa vào ERD. Chi tiết tên bảng kỹ thuật, rowversion, delete behavior và các ràng buộc được đối chiếu tại Phụ lục D.",
        body,
    )


def build_use_case_specs(document, anchor, body, caption_sample, table_sample) -> None:
    add_heading(anchor, "3.2.1 Use case chi tiết", "Heading 3")
    figure_number = 7
    bookmark_id = 450
    table_number = 5
    for heading, filename, title, spec_key in USE_CASE_LAYOUT:
        add_heading(anchor, heading, "Heading 4")
        if spec_key is not None:
            spec = USE_CASE_SPECS[spec_key]
            add_body(anchor, spec["description"], body)
            add_specification_table(document, anchor, table_number, spec, caption_sample, table_sample)
            table_number += 1
        elif "danh mục" in heading:
            add_body(anchor, "Quản lý tìm kiếm, phân trang, thêm, sửa, vô hiệu hóa hoặc khôi phục danh mục, văn phòng phẩm và đơn vị tính trong phạm vi được cấp.", body)
        else:
            add_body(anchor, "Dashboard và dữ liệu tổng hợp được giới hạn theo cá nhân, phòng ban hoặc toàn công ty dựa trên action policy và claim hiện hành.", body)
        if filename and title:
            bookmark_id = base.add_figure(anchor, DIAGRAMS / "ch03" / filename, f"3-{figure_number}", title, caption_sample, bookmark_id)
            figure_number += 1


def build_appendix_d(anchor, body) -> None:
    add_heading(anchor, "PHỤ LỤC D. TỪ ĐIỂN RÀNG BUỘC DỮ LIỆU", "Heading 2")
    add_heading(anchor, "D.1 Chỉ mục duy nhất và chỉ mục có điều kiện", "Heading 3")
    for name, description in CONSTRAINT_APPENDIX:
        add_labeled_body(anchor, f"{name}: ", description, body)
    add_heading(anchor, "D.2 Check constraint", "Heading 3")
    for name, description in CHECK_APPENDIX:
        add_labeled_body(anchor, f"{name}: ", description, body)
    add_heading(anchor, "D.3 Rowversion và hành vi xóa", "Heading 3")
    add_body(anchor, "RowVersion được dùng tại AspNetUsers, UserGroupMemberships, PriceLists, SupplierProductMappings, Periods, Requests và Settlements để phát hiện ghi đè đồng thời.", body)
    add_body(anchor, "DeleteBehavior.Restrict áp dụng cho phần lớn quan hệ nghiệp vụ. AspNetUserClaims, AspNetUserLogins và AspNetUserTokens dùng Cascade theo Identity; xóa mềm qua IsDeleted bảo toàn lịch sử nghiệp vụ.", body)
    add_heading(anchor, "D.4 Bảng dịch và đối tượng chỉ đọc", "Heading 3")
    add_body(anchor, "Các bảng DepartmentTranslations, LookupCategoryTranslations, LookupValueTranslations, VppCategoryTranslations, VppItemTranslations, SupplierTranslations và PriceListTranslations dùng khóa chính Id, FK đến thực thể gốc và filtered UNIQUE theo thực thể - LanguageCode khi IsDeleted = 0.", body)
    add_body(anchor, "v_Users và v_WFXCompany là view keyless chỉ đọc. BusinessDataTranslationBase không có store riêng; StoredProcedureResultDTO không phải entity hoặc view dữ liệu.", body)


def rebuild_references(document: Document, body) -> None:
    start = find_paragraph(document, "TÀI LIỆU THAM KHẢO")
    end = document.add_paragraph("__END_REFERENCES__")
    remove_range(start, end)
    add_heading(end, "TÀI LIỆU THAM KHẢO", "Heading 1")
    for reference in REFERENCES:
        paragraph = add_body(end, reference, body)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.left_indent = Cm(0.7)
        paragraph.paragraph_format.first_line_indent = Cm(-0.7)
        paragraph.paragraph_format.space_after = Pt(0)
    end._element.getparent().remove(end._element)


def update_citations_and_chapter_2(document: Document, body) -> None:
    paragraph = next(p for p in document.paragraphs if p.text.startswith("Tài liệu chính thức cho thấy Google Sheets"))
    set_paragraph_text(
        paragraph,
        "Tài liệu chính thức cho thấy Google Sheets và Excel for the web hỗ trợ chia sẻ, đồng biên soạn; Odoo Purchase hỗ trợ quản lý báo giá và đơn mua hàng; Jira Service Management và Zoho Creator hỗ trợ tiếp nhận yêu cầu, luồng xử lý hoặc phê duyệt [1], [2], [3], [4], [5]. Bảng so sánh dưới đây đánh giá các nhóm giải pháp theo mức độ phù hợp với nghiệp vụ riêng của GTAS VPP, gồm quy tắc ngày 5, đơn bổ sung, chụp giá và phân quyền chi tiết.",
    )
    technology = next(p for p in document.paragraphs if p.text.startswith("Hệ thống GTAS VPP được hiện thực"))
    set_paragraph_text(
        technology,
        "Hệ thống GTAS VPP được hiện thực trên .NET 10 và ASP.NET Core [6], dùng ASP.NET Core Identity [7], frontend Blazor Web App theo chế độ Interactive Server [8], Entity Framework Core [9] và SQL Server [10]. Giao diện sử dụng Radzen Blazor [11], ánh xạ DTO dùng Mapster [12], nhật ký cấu trúc dùng Serilog [13]. Môi trường triển khai dùng Docker Compose [14], NGINX reverse proxy [15] và .NET Aspire cho điều phối local [16].",
    )
    if replace_text_everywhere(document, "Serilog, bảng VPP03_Log", "Serilog và nhóm dữ liệu nhật ký nghiệp vụ") != 1:
        raise RuntimeError("Expected exactly one legacy Serilog/schema phrase.")
    for table in document.tables:
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        if headers[:2] != ["Nhóm", "Công nghệ"]:
            continue
        technology_citations = {
            "Backend": "ASP.NET Core Web API (.NET 10) [6]",
            "Frontend": "Blazor Server, Radzen Blazor [8], [11]",
            "Cơ sở dữ liệu": "SQL Server, Entity Framework Core [10], [9]",
            "Bảo mật": "JWT Bearer, cookie frontend, policy-based authorization [6], [7]",
            "Ánh xạ dữ liệu": "Mapster [12]",
            "Ghi log": "Serilog và nhóm dữ liệu nhật ký nghiệp vụ [13]",
            "Triển khai": "Docker Compose, Nginx reverse proxy [14], [15]",
            "Kiểm thử": "xUnit, Playwright, FluentAssertions, Moq, EF Core InMemory/SQLite [17], [18]",
        }
        for row in table.rows[1:]:
            group = row.cells[0].text.strip()
            if group in technology_citations:
                set_cell_text(row.cells[1], technology_citations[group])
        break
    period_line = next(p for p in document.paragraphs if "trạng thái Open, Closed, Pricing hoặc Settled" in p.text)
    set_paragraph_text(
        period_line,
        "1. Kỳ là thực thể độc lập, xác định năm, tháng, múi giờ, hạn gửi, hạn xử lý đơn bổ sung và bốn trạng thái: đang nhận đơn (Open), đã đóng nhận đơn (SubmissionClosed), đang định giá (Pricing) và đã chốt (Settled).",
    )


def update_chapter_4_and_tables(document: Document) -> None:
    chapter_intro = next(p for p in document.paragraphs if p.text.startswith("Chương này trình bày các cổng kiểm chứng"))
    set_paragraph_text(
        chapter_intro,
        "Chương này trình bày các cổng kiểm chứng tại commit 5f24fda ngày 25/07/2026. Kết quả được ghi đúng theo lệnh chạy thực tế; lỗi còn lại được báo cáo trung thực và không được thay bằng số liệu lịch sử. Log ứng dụng được cấu trúc hóa bằng Serilog [13]; unit/integration test chạy bằng xUnit.net v3 [17] và E2E chạy bằng Playwright .NET [18].",
    )
    replace_text_everywhere(document, "21,9828 phút", "khoảng 21,98 phút")
    replace_text_everywhere(document, "21.9828 Minutes", "khoảng 21,98 phút")
    for table in document.tables:
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        if headers[:2] == ["Mã", "Kịch bản"]:
            set_table_geometry(table, [1150, 2150, 2822, 2950])
            for row in table.rows[1:]:
                set_cell_no_wrap(row.cells[0])
        if headers[:2] == ["Hạng mục", "Tổng"] and len(headers) == 5:
            for row in table.rows[1:]:
                if row.cells[0].text.strip() == "Integration test":
                    row.cells[1].text = "20"
                    row.cells[2].text = "14"
                    row.cells[3].text = "6 chưa chạy (LocalDB opt-in)"
                    row.cells[4].text = "Đạt trong phạm vi mặc định"
                if row.cells[0].text.strip() == "Playwright E2E cô lập":
                    row.cells[1].text = "27"
                    row.cells[2].text = "12"
                    row.cells[3].text = "15 không đạt"
                    row.cells[4].text = "Chưa đạt"


def refine(input_path: Path, output_path: Path) -> None:
    if input_path.resolve() == output_path.resolve():
        raise ValueError("Input and output paths must be different.")
    document = Document(input_path)
    locked_before = protected_digests(document)
    section_21_before = semantic_text_range(document, "2.1 CÁC HỆ THỐNG TƯƠNG TỰ", "2.2 CÔNG NGHỆ SỬ DỤNG")
    body = find_body_sample(document)
    caption_sample = find_caption_sample(document)
    table_sample = find_table_sample(document)
    base.DOC = document

    update_citations_and_chapter_2(document, body)
    replace_range(
        document,
        "2.3.3 Sơ đồ Use case tổng quát",
        "CHƯƠNG 3. THIẾT KẾ",
        lambda anchor: build_personas(document, anchor, body, caption_sample, table_sample),
    )
    find_paragraph(document, "CHƯƠNG 3. THIẾT KẾ").paragraph_format.page_break_before = True
    replace_range(
        document,
        "3.1.3 Mô hình dữ liệu vật lý",
        "3.2 MÔ HÌNH XỬ LÝ",
        lambda anchor: build_physical_model(document, anchor, body, caption_sample, table_sample),
    )
    replace_range(
        document,
        "3.2.1 Use case chi tiết",
        "3.2.2 Sơ đồ tuần tự",
        lambda anchor: build_use_case_specs(document, anchor, body, caption_sample, table_sample),
    )
    references_heading = find_paragraph(document, "TÀI LIỆU THAM KHẢO")
    build_appendix_d(references_heading, body)
    update_chapter_4_and_tables(document)
    rebuild_references(document, body)
    mark_fields_dirty(document)

    locked_after = protected_digests(document)
    changed = [name for name in locked_before if locked_before[name] != locked_after[name]]
    if changed:
        detail = ""
        if "section_2_1" in changed:
            section_21_after = semantic_text_range(document, "2.1 CÁC HỆ THỐNG TƯƠNG TỰ", "2.2 CÔNG NGHỆ SỬ DỤNG")
            detail = f"\nBEFORE={section_21_before!r}\nAFTER={section_21_after!r}"
        raise RuntimeError(f"Protected content changed unexpectedly: {changed}{detail}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    refine(args.input, args.output)
    print(args.output.resolve())


if __name__ == "__main__":
    main()
