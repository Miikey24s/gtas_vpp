"""
edit_deck.py - Programmatic utilities and CLI for editing PPTX presentations.

Usage:
    python edit_deck.py <input.pptx> --output <output.pptx> [operations]

Operations:
    --replace-text <old> <new> [--slide <num>]
    --set-notes <slide_num> <notes_text_or_file>
    --replace-image <slide_num> <shape_name_or_id> <new_image_file>
"""

import argparse
import os
import sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt

# Ensure UTF-8 output on Windows
if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if sys.stderr and hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")


def replace_text_in_text_frame(tf, old_text, new_text, match_case=True):
    count = 0
    for p in tf.paragraphs:
        # If the entire paragraph matches or substring matches
        if old_text in p.text:
            # Check individual runs
            found_in_runs = False
            for r in p.runs:
                if old_text in r.text:
                    r.text = r.text.replace(old_text, new_text)
                    count += 1
                    found_in_runs = True
            # If text was split across runs, replace at paragraph level
            if not found_in_runs and old_text in p.text:
                p.text = p.text.replace(old_text, new_text)
                count += 1
    return count


def replace_text_in_deck(prs, old_text, new_text, target_slide=None):
    count = 0
    for idx, slide in enumerate(prs.slides, start=1):
        if target_slide is not None and idx != target_slide:
            continue
        for shape in slide.shapes:
            if shape.has_text_frame:
                count += replace_text_in_text_frame(shape.text_frame, old_text, new_text)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text_frame:
                            count += replace_text_in_text_frame(cell.text_frame, old_text, new_text)
    return count


def set_speaker_notes(prs, slide_number, notes_text):
    if slide_number < 1 or slide_number > len(prs.slides):
        raise ValueError(f"Slide number {slide_number} out of range (1..{len(prs.slides)})")
    slide = prs.slides[slide_number - 1]
    if not slide.has_notes_slide:
        slide.notes_slide  # accessing it creates notes slide
    slide.notes_slide.notes_text_frame.text = notes_text
    return True


def replace_image_in_slide(prs, slide_number, shape_identifier, new_image_path):
    if not os.path.exists(new_image_path):
        raise FileNotFoundError(f"New image not found: {new_image_path}")

    if slide_number < 1 or slide_number > len(prs.slides):
        raise ValueError(f"Slide number {slide_number} out of range (1..{len(prs.slides)})")

    slide = prs.slides[slide_number - 1]
    target_shape = None

    for shape in slide.shapes:
        if str(shape.shape_id) == str(shape_identifier) or shape.name.lower() == str(shape_identifier).lower():
            target_shape = shape
            break

    if target_shape is None:
        # Try matching by picture index
        pictures = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
        if str(shape_identifier).isdigit() and int(shape_identifier) <= len(pictures):
            target_shape = pictures[int(shape_identifier) - 1]

    if target_shape is None:
        raise ValueError(f"Shape '{shape_identifier}' not found in slide {slide_number}")

    left = target_shape.left
    top = target_shape.top
    width = target_shape.width
    height = target_shape.height

    # Delete old shape element
    sp_elem = target_shape.element
    sp_elem.getparent().remove(sp_elem)

    # Insert new picture at same position
    new_pic = slide.shapes.add_picture(new_image_path, left, top, width, height)
    return new_pic


def main():
    parser = argparse.ArgumentParser(description="Edit PPTX presentation files")
    parser.add_argument("pptx", help="Path to input PPTX file")
    parser.add_argument("--output", "-o", required=True, help="Path to output PPTX file")
    parser.add_argument("--replace-text", nargs=2, metavar=("OLD", "NEW"), help="Find and replace text")
    parser.add_argument("--slide", type=int, help="Target specific slide number for operation")
    parser.add_argument("--set-notes", nargs=2, metavar=("SLIDE_NUM", "TEXT_OR_FILE"), help="Set speaker notes for slide")
    parser.add_argument("--replace-image", nargs=3, metavar=("SLIDE_NUM", "SHAPE_ID", "IMG_PATH"), help="Replace image shape with new image")

    args = parser.parse_args()

    input_path = os.path.abspath(args.pptx)
    output_path = os.path.abspath(args.output)

    if not os.path.exists(input_path):
        print(f"File not found: {input_path}", file=sys.stderr)
        sys.exit(1)

    prs = Presentation(input_path)
    modified = False

    if args.replace_text:
        old_txt, new_txt = args.replace_text
        count = replace_text_in_deck(prs, old_txt, new_txt, target_slide=args.slide)
        print(f"[Edit] Replaced {count} occurrences of '{old_txt}' -> '{new_txt}'")
        modified = True

    if args.set_notes:
        slide_num = int(args.set_notes[0])
        notes_input = args.set_notes[1]
        if os.path.exists(notes_input):
            with open(notes_input, "r", encoding="utf-8") as f:
                notes_text = f.read()
        else:
            notes_text = notes_input
        set_speaker_notes(prs, slide_num, notes_text)
        print(f"[Edit] Updated speaker notes for slide {slide_num}")
        modified = True

    if args.replace_image:
        slide_num = int(args.replace_image[0])
        shape_id = args.replace_image[1]
        img_path = args.replace_image[2]
        replace_image_in_slide(prs, slide_num, shape_id, img_path)
        print(f"[Edit] Replaced image '{shape_id}' in slide {slide_num} with '{img_path}'")
        modified = True

    if modified:
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        prs.save(output_path)
        print(f"[Done] Saved modified presentation to: {output_path}")
    else:
        print("[Warning] No edit operations performed.")


if __name__ == "__main__":
    main()
