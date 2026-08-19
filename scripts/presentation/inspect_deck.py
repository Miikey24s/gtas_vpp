"""
inspect_deck.py - Inspect and dump PPTX presentation structure and contents.

Usage:
    python inspect_deck.py <input.pptx> [--format markdown|json|summary] [--slide <index>] [--output <file>]
"""

import argparse
import json
import os
import sys

# Ensure UTF-8 output on Windows
if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if sys.stderr and hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt


def pt_to_str(pt_val):
    if pt_val is None:
        return "N/A"
    return f"{pt_val.pt:.1f}pt"


def color_to_str(color_format):
    try:
        if color_format.type == 1:  # RGB
            return f"#{color_format.rgb}"
        elif color_format.type == 2:  # Theme
            return f"Theme({color_format.theme_color})"
        return "Default"
    except Exception:
        return "Unknown"


def inspect_deck(pptx_path, target_slide=None):
    if not os.path.exists(pptx_path):
        raise FileNotFoundError(f"File not found: {pptx_path}")

    prs = Presentation(pptx_path)
    slide_width_in = prs.slide_width.inches
    slide_height_in = prs.slide_height.inches
    aspect_ratio = f"{slide_width_in:.2f}x{slide_height_in:.2f} in ({slide_width_in/slide_height_in:.2f}:1)"

    deck_info = {
        "file": os.path.abspath(pptx_path),
        "total_slides": len(prs.slides),
        "dimensions": {
            "width_inches": slide_width_in,
            "height_inches": slide_height_in,
            "aspect_ratio": aspect_ratio,
        },
        "slides": [],
    }

    for idx, slide in enumerate(prs.slides, start=1):
        if target_slide is not None and idx != target_slide:
            continue

        slide_data = {
            "slide_number": idx,
            "title": "",
            "layout_name": slide.slide_layout.name if slide.slide_layout else "Unknown",
            "notes": "",
            "shapes_count": len(slide.shapes),
            "text_blocks": [],
            "tables": [],
            "images": [],
            "shapes": [],
        }

        # Notes
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            slide_data["notes"] = slide.notes_slide.notes_text_frame.text.strip()

        # Iterate shapes
        for s_idx, shape in enumerate(slide.shapes, start=1):
            shape_info = {
                "id": shape.shape_id,
                "name": shape.name,
                "type": str(shape.shape_type),
                "left_in": round(shape.left.inches, 2) if shape.left else 0,
                "top_in": round(shape.top.inches, 2) if shape.top else 0,
                "width_in": round(shape.width.inches, 2) if shape.width else 0,
                "height_in": round(shape.height.inches, 2) if shape.height else 0,
            }

            # Check title
            if shape == slide.shapes.title:
                slide_data["title"] = shape.text.strip()

            # Check TextFrame
            if shape.has_text_frame:
                paragraphs_data = []
                for p in shape.text_frame.paragraphs:
                    p_text = p.text.strip()
                    if not p_text:
                        continue
                    runs_data = []
                    for run in p.runs:
                        font = run.font
                        runs_data.append({
                            "text": run.text,
                            "font_name": font.name,
                            "font_size": pt_to_str(font.size),
                            "bold": font.bold,
                            "italic": font.italic,
                            "color": color_to_str(font.color) if font.color else "Default",
                        })
                    paragraphs_data.append({
                        "level": p.level,
                        "text": p_text,
                        "runs": runs_data,
                    })
                if paragraphs_data:
                    slide_data["text_blocks"].append({
                        "shape_name": shape.name,
                        "paragraphs": paragraphs_data,
                    })

            # Check Table
            if shape.has_table:
                table = shape.table
                table_rows = []
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells]
                    table_rows.append(row_cells)
                slide_data["tables"].append({
                    "shape_name": shape.name,
                    "rows": len(table.rows),
                    "cols": len(table.columns),
                    "data": table_rows,
                })

            # Check Image
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                slide_data["images"].append({
                    "shape_name": shape.name,
                    "left_in": shape_info["left_in"],
                    "top_in": shape_info["top_in"],
                    "width_in": shape_info["width_in"],
                    "height_in": shape_info["height_in"],
                })

            slide_data["shapes"].append(shape_info)

        deck_info["slides"].append(slide_data)

    return deck_info


def format_as_markdown(deck_info):
    lines = []
    lines.append(f"# Presentation Inspection: `{os.path.basename(deck_info['file'])}`")
    lines.append(f"- **Total Slides**: {deck_info['total_slides']}")
    lines.append(f"- **Dimensions**: {deck_info['dimensions']['aspect_ratio']}")
    lines.append("")

    for s in deck_info["slides"]:
        lines.append(f"## Slide {s['slide_number']}: {s['title'] or '(No Title)'}")
        lines.append(f"- **Layout**: `{s['layout_name']}` | **Shapes**: {s['shapes_count']} | **Images**: {len(s['images'])} | **Tables**: {len(s['tables'])}")

        if s["text_blocks"]:
            lines.append("### Content:")
            for tb in s["text_blocks"]:
                lines.append(f"*{tb['shape_name']}*:")
                for p in tb["paragraphs"]:
                    indent = "  " * p["level"]
                    font_details = ""
                    if p["runs"] and p["runs"][0]["font_size"] != "N/A":
                        font_details = f" `[{p['runs'][0]['font_name']}, {p['runs'][0]['font_size']}]`"
                    lines.append(f"{indent}- {p['text']}{font_details}")

        if s["tables"]:
            lines.append("### Tables:")
            for t in s["tables"]:
                lines.append(f"- Table `{t['shape_name']}` ({t['rows']}x{t['cols']}):")
                for row in t["data"]:
                    lines.append(f"  | {' | '.join(row)} |")

        if s["notes"]:
            lines.append("### Speaker Notes:")
            lines.append(f"> {s['notes']}")

        lines.append("\n---\n")

    return "\n".join(lines)


def format_as_summary(deck_info):
    lines = []
    lines.append(f"Deck: {os.path.basename(deck_info['file'])} ({deck_info['total_slides']} slides, {deck_info['dimensions']['aspect_ratio']})")
    for s in deck_info["slides"]:
        title = s['title'] or '(No Title)'
        notes_flag = " [Notes]" if s["notes"] else ""
        imgs_flag = f" [{len(s['images'])} img]" if s["images"] else ""
        tbls_flag = f" [{len(s['tables'])} tbl]" if s["tables"] else ""
        lines.append(f"  Slide {s['slide_number']:02d}: {title}{notes_flag}{imgs_flag}{tbls_flag}")
    return "\n".join(lines)


def main():
    parser = argparse.ArgumentParser(description="Inspect PPTX presentations")
    parser.add_argument("pptx", help="Path to PPTX file")
    parser.add_argument("--format", choices=["markdown", "json", "summary"], default="markdown", help="Output format")
    parser.add_argument("--slide", type=int, help="Inspect a specific slide number only")
    parser.add_argument("--output", help="Save output to file")

    args = parser.parse_args()

    deck_info = inspect_deck(args.pptx, target_slide=args.slide)

    if args.format == "json":
        result = json.dumps(deck_info, indent=2, ensure_ascii=False)
    elif args.format == "summary":
        result = format_as_summary(deck_info)
    else:
        result = format_as_markdown(deck_info)

    if args.output:
        with open(args.output, "w", encoding="utf-8") as f:
            f.write(result)
        print(f"Inspection written to {args.output}")
    else:
        print(result)


if __name__ == "__main__":
    main()
