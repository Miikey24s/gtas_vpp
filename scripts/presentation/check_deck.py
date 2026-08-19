"""
check_deck.py - Automated Quality Assurance & Consistency Checker for PPTX presentations.

Usage:
    python check_deck.py <input.pptx> [--min-font-size <pt>] [--max-bullets <int>]
"""

import argparse
import os
import sys
from pptx import Presentation

# Ensure UTF-8 output on Windows
if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if sys.stderr and hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")


def check_presentation(pptx_path, min_font_pt=14, max_bullets=7):
    prs = Presentation(pptx_path)
    issues = []
    warnings = []
    stats = {
        "total_slides": len(prs.slides),
        "slides_with_notes": 0,
        "slides_with_title": 0,
        "total_images": 0,
        "total_tables": 0,
    }

    for idx, slide in enumerate(prs.slides, start=1):
        # 1. Title check
        has_title = False
        if slide.shapes.title and slide.shapes.title.text.strip():
            has_title = True
            stats["slides_with_title"] += 1
        else:
            # Check if any text box contains a prominent heading
            first_text = ""
            for s in slide.shapes:
                if s.has_text_frame and s.text_frame.text.strip():
                    first_text = s.text_frame.text.strip().split("\n")[0]
                    break
            if idx > 1 and not first_text:
                warnings.append(f"Slide {idx:02d}: Không có tiêu đề rõ ràng hoặc slide trống.")

        # 2. Speaker notes check
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame and slide.notes_slide.notes_text_frame.text.strip():
            stats["slides_with_notes"] += 1
        else:
            warnings.append(f"Slide {idx:02d}: Thiếu speaker notes cho phần thuyết trình.")

        # 3. Bullets / density check
        total_bullets = 0
        for s in slide.shapes:
            if s.has_text_frame:
                for p in s.text_frame.paragraphs:
                    if p.text.strip():
                        total_bullets += 1
                    # Font size check
                    for r in p.runs:
                        if r.font.size:
                            pt_val = r.font.size.pt
                            if pt_val < min_font_pt and pt_val > 0:
                                warnings.append(
                                    f"Slide {idx:02d}: Font size nhỏ ({pt_val:.1f}pt < {min_font_pt}pt) tại đoạn '{p.text[:40]}...'"
                                )

            if s.has_table:
                stats["total_tables"] += 1

            if s.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
                stats["total_images"] += 1

        if total_bullets > max_bullets and idx > 1:
            warnings.append(
                f"Slide {idx:02d}: Mật độ nội dung cao ({total_bullets} dòng/bullet > {max_bullets}). Cân nhắc tách slide hoặc rút gọn."
            )

    return stats, warnings, issues


def main():
    parser = argparse.ArgumentParser(description="Check PPTX presentation quality")
    parser.add_argument("pptx", help="Path to PPTX file")
    parser.add_argument("--min-font-size", type=float, default=12.0, help="Minimum font size in pt (default: 12.0)")
    parser.add_argument("--max-bullets", type=int, default=8, help="Maximum paragraphs/bullets per slide (default: 8)")

    args = parser.parse_args()
    pptx_path = os.path.abspath(args.pptx)

    if not os.path.exists(pptx_path):
        print(f"File not found: {pptx_path}", file=sys.stderr)
        sys.exit(1)

    print(f"=== PPTX QA Report: {os.path.basename(pptx_path)} ===")
    stats, warnings, issues = check_presentation(pptx_path, args.min_font_size, args.max_bullets)

    print(f"Tổng số slides: {stats['total_slides']}")
    print(f"Slides có tiêu đề: {stats['slides_with_title']}/{stats['total_slides']}")
    print(f"Slides có Speaker Notes: {stats['slides_with_notes']}/{stats['total_slides']}")
    print(f"Tổng số hình ảnh: {stats['total_images']}")
    print(f"Tổng số bảng biểu: {stats['total_tables']}")
    print("-" * 50)

    if issues:
        print(f"\n[LỖI NGHIÊM TRỌNG] ({len(issues)}):")
        for err in issues:
            print(f"  ❌ {err}")

    if warnings:
        print(f"\n[CẢNH BÁO/GỢI Ý TỐI ƯU] ({len(warnings)}):")
        for warn in warnings:
            print(f"  ⚠️  {warn}")
    else:
        print("\n✅ Không có cảnh báo. Slide deck đạt chuẩn trình chiếu!")

    print("=" * 50)


if __name__ == "__main__":
    main()
