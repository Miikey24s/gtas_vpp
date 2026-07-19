"""Normalize the thesis VPP workbook into deterministic demo seed datasets.

ETL means Extract-Transform-Load: read the workbook, clean/reconcile it, then
write small TSV files that the .NET seed can load without Excel at runtime.
The source workbook is never modified and personal notes are not exported.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from statistics import median
from typing import Any

from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_SOURCE = REPO_ROOT / "LVTN/data/DANG KY VPP - CAC DON VI.xlsx"
OUTPUT_DIR = REPO_ROOT / "gtas_vpp_be/gtas_vpp_be.Service/Helpers/Data/Demo"

NON_ORDER_SHEETS = {"Danh mục", "List"}
LATEST_SOURCE_MONTH = 3

CATEGORY_SPECS = {
    "bang keo bam kim bam lo": ("TAPE_FASTENING", "Băng keo, bấm kim, bấm lỗ"),
    "bia cong bia phan trang trinh ky": ("BINDING_FILING", "Bìa còng, bìa phân trang, trình ký"),
    "kep buom kim bam keo": ("CLIPS_STAPLES_SCISSORS", "Kẹp bướm, kim bấm, kéo"),
    "but viet": ("WRITING", "Bút viết"),
    "tap viet thuoc ke": ("NOTEBOOK_RULER", "Tập viết, Thước kẻ"),
    "giay in cac loai": ("PAPER", "Giấy in các loại"),
    "khac": ("OTHER", "Khác"),
    "phuc vu van phong": ("OFFICE_SUPPORT", "Phục vụ Văn phòng"),
}

UOM_SPECS = {
    "bo": ("BUNDLE", "Bó"),
    "bich": ("BAG", "Bịch"),
    "bo set": ("SET", "Bộ"),
    "can": ("CANISTER", "Can"),
    "chai": ("BOTTLE", "Chai"),
    "chiec": ("PIECE", "Chiếc"),
    "cuon volume": ("VOLUME", "Cuốn"),
    "cuon": ("ROLL", "Cuộn"),
    "cai": ("ITEM", "Cái"),
    "cay": ("STICK", "Cây"),
    "cap": ("PAIR_PACK", "Cặp"),
    "cuc": ("BLOCK", "Cục"),
    "goi": ("PACK", "Gói"),
    "hop": ("BOX", "Hộp"),
    "kg": ("KILOGRAM", "Kg"),
    "loc": ("MULTIPACK", "Lốc"),
    "quyen": ("BOOK", "Quyển"),
    "ram": ("REAM", "Ram"),
    "soi": ("STRAND", "Sợi"),
    "thung": ("CARTON", "Thùng"),
    "tam": ("SHEET", "Tấm"),
    "to": ("LEAF", "Tờ"),
    "tuyp": ("TUBE", "Tuýp"),
    "vien": ("TABLET", "Viên"),
    "vi": ("BLISTER", "Vỉ"),
    "xap": ("STACK", "Xấp"),
    "doi": ("PAIR", "Đôi"),
}

DEPARTMENT_SPECS = {
    "CongNgheMay": ("CONGNGHEMAY", "Công nghệ may", ""),
    "LIÊN VÙNG 1": ("LIENVUNG1", "Liên vùng 1", ""),
    "P.HCQT": ("HCQT", "Hành chính Quản trị", ""),
    "PhapChe": ("PHAPCHE", "Pháp chế", ""),
    "KeHoach": ("KEHOACH", "Kế hoạch", ""),
    "KD28": ("KD28", "KD28", ""),
    "CongNgheWash": ("CONGNGHEWASH", "Công nghệ Wash", ""),
    "CBSX": ("CBSX", "CBSX", ""),
    "NSTL": ("NSTL", "NSTL", ""),
    "Sourcing": ("SOURCING", "Sourcing", ""),
    "CPD": ("CPD", "CPD", ""),
    "FD": ("FD", "FD", ""),
    "DauTu": ("DAUTU", "Đầu tư phát triển", ""),
    "FQM": ("FQM", "FQM", ""),
    "TTHTPhuocLong": ("HOANTATPHUOCLONG", "Hoàn tất Phước Long", ""),
    "KD17": ("KD17", "KD17", ""),
    "KD16": ("KD16", "KD16", ""),
    "DinhMuc": ("DINHMUC", "Định mức", ""),
    "GiamDinh": ("GIAMDINH", "Giám định", ""),
    "KD19": ("KD19", "KD19", ""),
    "IT": ("IT", "Công nghệ thông tin", ""),
    "QA": ("QA", "QA", ""),
    "KD3": ("KD3", "KD3", ""),
    "KD1": ("KD1", "KD1", ""),
    "MayMauLongAn": ("MAULONGAN", "May mẫu Long An", ""),
    "P.KT LA": ("PKTLA", "Phòng Kỹ thuật Long An", ""),
    "KD2": ("KD2", "KD2", ""),
    "KD5+6+18": ("KD5+6+18", "KD5 + KD6 + KD18", "KD5+6"),
    "KD4": ("KD4", "KD4", ""),
    "KD26": ("KD26", "KD26", ""),
    "Sale7": ("KD7", "KD7", ""),
    "KD8": ("KD8", "KD8", ""),
    "KD25": ("KD25", "KD25", ""),
    "KD27": ("KD27", "KD27", ""),
    "KhoTong": ("KHOTONG", "Kho tổng", ""),
    "LongAn": ("LONGAN", "Long An", ""),
    "KHO LA": ("KHOLA", "Kho Long An", ""),
    "PPJ-WISER": ("PPJWISER", "PPJ-WISER", ""),
    "FittechRap": ("NHOMFITTECHRAP", "Nhóm Fittech Rập", "NHOMFITTECH Rap"),
    "MayMauHCM": ("XUONGMAYMAU", "Xưởng may mẫu HCM", ""),
    "Purchasing": ("PURCHASING", "Purchasing", ""),
    "KeHoachPNC": ("KEHOACHPNC", "Kế hoạch PNC", ""),
    "QLTBMay": ("QLTBMAY", "Quản lý thiết bị may", ""),
    "RandD": ("RD", "R&D", "R&D"),
    "TCKT": ("TCKT", "TCKT", ""),
    "TCKTKho": ("TCKTKHO", "TCKT Kho", ""),
    "TCKTVTJ": ("TCKTVTJ", "TCKT VTJ", ""),
    "TQM": ("TQM", "TQM", ""),
    "TheuMau": ("THEUMAU", "Thêu mẫu", "c"),
    "TTHTLinhTrung": ("TTHTLINHTRUNG", "TTHT Linh Trung", ""),
    "WashLinhTrung": ("WASHLINHTRUNG", "Wash Linh Trung", ""),
    "XNK": ("XNK", "Xuất nhập khẩu", ""),
    "KD7": ("KD7", "KD7", ""),
}

SYNTHETIC_NAMES = [
    "Nguyễn Minh Anh", "Trần Hoàng Nam", "Lê Thu Hà", "Phạm Quốc Bảo",
    "Võ Ngọc Mai", "Đặng Thành Đạt", "Bùi Khánh Linh", "Đỗ Đức Huy",
    "Hồ Gia Hân", "Ngô Tuấn Kiệt", "Dương Thảo Vy", "Lý Nhật Minh",
    "Nguyễn Hải Yến", "Trần Quang Vinh", "Lê Phương Thảo", "Phạm Minh Khang",
    "Võ Thanh Trúc", "Đặng Anh Dũng", "Bùi Mỹ Linh", "Đỗ Hoài Phong",
    "Hồ Bảo Ngọc", "Ngô Đức Anh", "Dương Kim Oanh", "Lý Thanh Tùng",
    "Nguyễn Quỳnh Như", "Trần Gia Bảo", "Lê Hồng Nhung", "Phạm Tuấn Anh",
    "Võ Minh Châu", "Đặng Quốc Khánh", "Bùi Thanh Hằng", "Đỗ Trung Hiếu",
    "Hồ Ngọc Trâm", "Ngô Minh Triết", "Dương Mai Phương", "Lý Hoàng Long",
    "Nguyễn Thùy Dương", "Trần Thanh Sơn", "Lê Ngọc Diệp", "Phạm Anh Tuấn",
    "Võ Khánh Vy", "Đặng Minh Quân", "Bùi Thu Trang", "Đỗ Thành Công",
    "Hồ Phương Linh", "Ngô Quốc Việt", "Dương Ngọc Ánh", "Lý Đức Thịnh",
    "Nguyễn Thanh Tâm", "Trần Minh Nhật", "Lê Bảo Châu", "Phạm Hải Đăng",
]

ITEM_ALIASES = {
    "bao nilon 20x35cm khong quai": "Bao PE 20x35cm không quai",
    "bao nilon 40x50cm khong quai": "Bao PE 40x50cm không quai",
    "bao nilon 45x60cm khong quai": "Bao PE 45x60cm không quai",
    "bao pe 37x47cm": "Bao PE ( 37 x 47 cm)",
    "bao ziper 14x10cm": "Bao zipper 14x10",
    "bao ziper 5x8cm": "Bao zipper đựng mẫu 5*8cm",
    "bao zipper 12x17cm": "Bao zipper 12x17 cm",
    "but gel b 01 master xanh": "Bút Gel B-01 Master",
    "giay bia day a4 250gsm": "Giấy bìa dày - A4 250gam",
    "ho kho hq 8g": "Hồ khô 8gr G-36S",
    "kim sung ban gia": "Kim súng bắn giá (N4-P)",
}

ADDITIONAL_ITEMS = [
    ("OFFICE_SUPPORT", "Giấy hộp Pulppy", "BOX", Decimal("0.08"), Decimal("23500"), "inferred_from_giay_hop_japani"),
    ("OTHER", "Hộp bút xoay Xukiva - 174", "ITEM", Decimal("0.08"), Decimal("45000"), "inferred_from_hop_but_xoay_ttm_3006"),
    ("OFFICE_SUPPORT", "Tủ y tế", "ITEM", Decimal("0.08"), Decimal("450000"), "inferred_missing_price"),
]

IT_BASKET_ITEMS = [
    "Giấy A4 trắng 80 Excell",
    "Bút bi xanh 027",
    "Bìa lỗ TQ GM",
    "Giấy note vàng 3*3 - UNC",
    "Băng keo trong 5cm 100y",
    "Pin đồng hồ Pin 2A Energizer",
]


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFC", str(value))
    text = text.replace("’", "'").replace("`", "'")
    return re.sub(r"\s+", " ", text).strip()


def match_key(value: Any) -> str:
    text = clean_text(value).casefold().replace("đ", "d").replace("_", ",")
    text = unicodedata.normalize("NFD", text)
    text = "".join(character for character in text if not unicodedata.combining(character))
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def decimal_value(value: Any) -> Decimal | None:
    if value is None or value == "" or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            return None
        return Decimal(str(value))
    try:
        return Decimal(clean_text(value).replace(".", "").replace(",", ""))
    except InvalidOperation:
        return None


def month_value(value: Any) -> int | None:
    number = decimal_value(value)
    if number is None or number != number.to_integral_value():
        return None
    month = int(number)
    return month if 1 <= month <= 12 else None


def category_spec(value: str) -> tuple[str, str]:
    spec = CATEGORY_SPECS.get(match_key(value))
    if spec is None:
        raise ValueError(f"Unknown category: {value!r}")
    return spec


def uom_spec(value: str) -> tuple[str, str]:
    key = match_key(value)
    if key == "bo":
        key = "bo set" if clean_text(value).casefold() == "bộ" else "bo"
    elif key == "cuon":
        key = "cuon volume" if clean_text(value).casefold() == "cuốn" else "cuon"
    elif key in {"vi", "vi blister"}:
        key = "vi"
    spec = UOM_SPECS.get(key)
    if spec is None:
        raise ValueError(f"Unknown UOM: {value!r} ({key})")
    return spec


def stable_item_code(category_code: str, item_name: str, uom_code: str) -> str:
    source = f"{category_code}|{match_key(item_name)}|{uom_code}".encode("utf-8")
    return "VPP_" + hashlib.sha256(source).hexdigest()[:24].upper()


@dataclass
class CatalogItem:
    category_code: str
    category_name: str
    item_code: str
    item_name: str
    uom_code: str
    uom_name: str
    vat_rate: Decimal
    unit_price: Decimal
    source_rows: list[int]
    resolution: str


def write_tsv(path: Path, fieldnames: list[str], rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        for row in rows:
            writer.writerow({key: row.get(key, "") for key in fieldnames})


def normalize(source: Path) -> dict[str, Any]:
    workbook = load_workbook(source, read_only=True, data_only=True, keep_links=False)
    catalog_sheet = workbook["Danh mục"]

    raw_catalog: dict[str, list[dict[str, Any]]] = defaultdict(list)
    current_category = ""
    for row_number, row in enumerate(
        catalog_sheet.iter_rows(min_row=4, max_row=catalog_sheet.max_row, min_col=1, max_col=5, values_only=True),
        start=4,
    ):
        _, item_raw, uom_raw, vat_raw, price_raw = row
        item_name = clean_text(item_raw)
        uom_name = clean_text(uom_raw)
        vat_rate = decimal_value(vat_raw)
        unit_price = decimal_value(price_raw)
        if item_name and not uom_name and vat_rate is None and unit_price is None:
            current_category = item_name
            continue
        if not item_name:
            continue
        if not current_category or not uom_name or vat_rate is None or unit_price is None:
            raise ValueError(f"Incomplete catalog row at Danh mục!{row_number}")
        raw_catalog[match_key(item_name)].append(
            {
                "source_row": row_number,
                "category_name": current_category,
                "item_name": item_name,
                "uom_name": uom_name,
                "vat_rate": vat_rate,
                "unit_price": unit_price,
            }
        )

    catalog: dict[str, CatalogItem] = {}
    duplicate_resolutions = []
    for item_key, rows in raw_catalog.items():
        chosen = rows[0]
        category_code, category_name = category_spec(chosen["category_name"])
        uom_code, canonical_uom_name = uom_spec(chosen["uom_name"])
        resolution = "source"
        if len(rows) > 1:
            resolution = "duplicate_first_source_row"
            duplicate_resolutions.append(
                {
                    "item_key": item_key,
                    "chosen_source_row": chosen["source_row"],
                    "all_source_rows": [row["source_row"] for row in rows],
                }
            )
        catalog[item_key] = CatalogItem(
            category_code=category_code,
            category_name=category_name,
            item_code=stable_item_code(category_code, chosen["item_name"], uom_code),
            item_name=chosen["item_name"],
            uom_code=uom_code,
            uom_name=canonical_uom_name,
            vat_rate=chosen["vat_rate"],
            unit_price=chosen["unit_price"],
            source_rows=[row["source_row"] for row in rows],
            resolution=resolution,
        )

    for category_code, item_name, uom_code, vat_rate, unit_price, resolution in ADDITIONAL_ITEMS:
        category_name = next(name for code, name in CATEGORY_SPECS.values() if code == category_code)
        uom_name = next(name for code, name in UOM_SPECS.values() if code == uom_code)
        item_key = match_key(item_name)
        catalog[item_key] = CatalogItem(
            category_code=category_code,
            category_name=category_name,
            item_code=stable_item_code(category_code, item_name, uom_code),
            item_name=item_name,
            uom_code=uom_code,
            uom_name=uom_name,
            vat_rate=vat_rate,
            unit_price=unit_price,
            source_rows=[],
            resolution=resolution,
        )

    alias_key_to_target_key = {
        alias_key: match_key(target_name) for alias_key, target_name in ITEM_ALIASES.items()
    }
    invalid_item_keys = {match_key(name) for _, name in CATEGORY_SPECS.values()}
    raw_orders = []
    skipped_blank_quantities = 0
    note_rows_removed = 0
    invalid_category_rows_removed = 0

    for sheet in workbook.worksheets:
        if sheet.title in NON_ORDER_SHEETS:
            continue
        if sheet.title not in DEPARTMENT_SPECS:
            raise ValueError(f"Missing department mapping for sheet {sheet.title!r}")
        department_code, _, _ = DEPARTMENT_SPECS[sheet.title]
        data_started = False
        consecutive_empty = 0
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=5, max_row=min(sheet.max_row, 20_000), min_col=1, max_col=10, values_only=True),
            start=5,
        ):
            _, category_raw, item_raw, _, _, price_raw, month_raw, quantity_raw, _, note_raw = row
            item_name = clean_text(item_raw)
            source_month = month_value(month_raw)
            quantity = decimal_value(quantity_raw)
            if not item_name and source_month is None and quantity is None:
                consecutive_empty += 1
                if data_started and consecutive_empty >= 500:
                    break
                continue
            consecutive_empty = 0
            if not item_name:
                continue
            data_started = True
            item_key = match_key(item_name)
            if item_key in invalid_item_keys:
                if source_month is not None and quantity is not None and quantity > 0:
                    invalid_category_rows_removed += 1
                continue
            if source_month is None or quantity is None or quantity <= 0:
                skipped_blank_quantities += 1
                continue
            if quantity != quantity.to_integral_value():
                raise ValueError(f"Non-integer quantity at {sheet.title}!{row_number}: {quantity}")
            target_key = alias_key_to_target_key.get(item_key, item_key)
            catalog_item = catalog.get(target_key)
            if catalog_item is None:
                raise ValueError(f"Unresolved item at {sheet.title}!{row_number}: {item_name!r}")
            source_price = decimal_value(price_raw)
            if source_price is None or source_price <= 0:
                source_price = catalog_item.unit_price
            if clean_text(note_raw):
                note_rows_removed += 1
            raw_orders.append(
                {
                    "department_code": department_code,
                    "source_month": source_month,
                    "item_key": target_key,
                    "quantity": int(quantity),
                    "unit_price": int(source_price),
                    "source_ref": f"{sheet.title}!{row_number}",
                    "is_synthesized": False,
                }
            )

    grouped_orders: dict[tuple[str, int, str], list[dict[str, Any]]] = defaultdict(list)
    for row in raw_orders:
        grouped_orders[(row["department_code"], row["source_month"], row["item_key"])].append(row)

    for item_name in IT_BASKET_ITEMS:
        item_key = match_key(item_name)
        source_quantities = [
            row["quantity"]
            for row in raw_orders
            if row["source_month"] == LATEST_SOURCE_MONTH and row["item_key"] == item_key
        ]
        if not source_quantities:
            raise ValueError(f"Cannot synthesize IT basket; no source quantities for {item_name!r}")
        grouped_orders[("IT", LATEST_SOURCE_MONTH, item_key)].append(
            {
                "department_code": "IT",
                "source_month": LATEST_SOURCE_MONTH,
                "item_key": item_key,
                "quantity": int(median(source_quantities)),
                "unit_price": int(catalog[item_key].unit_price),
                "source_ref": f"SYNTHETIC_MEDIAN_MONTH_{LATEST_SOURCE_MONTH}",
                "is_synthesized": True,
            }
        )

    normalized_orders = []
    for (department_code, source_month, item_key), rows in sorted(grouped_orders.items()):
        prices = Counter(row["unit_price"] for row in rows)
        unit_price = prices.most_common(1)[0][0]
        catalog_item = catalog[item_key]
        normalized_orders.append(
            {
                "DepartmentCode": department_code,
                "SourceMonth": source_month,
                "ItemCode": catalog_item.item_code,
                "Quantity": sum(row["quantity"] for row in rows),
                "UnitPrice": unit_price,
                "SourceRowCount": len(rows),
                "SourceRefs": ";".join(row["source_ref"] for row in rows),
                "IsSynthesized": str(any(row["is_synthesized"] for row in rows)).lower(),
            }
        )

    department_groups: dict[str, dict[str, Any]] = {}
    for sheet_name, (code, name, legacy_codes) in DEPARTMENT_SPECS.items():
        current = department_groups.setdefault(
            code,
            {
                "DepartmentCode": code,
                "DepartmentName": name,
                "LegacyCodes": set(),
                "SourceSheets": [],
            },
        )
        current["SourceSheets"].append(sheet_name)
        if legacy_codes:
            current["LegacyCodes"].update(code.strip() for code in legacy_codes.split(";") if code.strip())

    departments = []
    for value in sorted(department_groups.values(), key=lambda row: row["DepartmentCode"]):
        departments.append(
            {
                "DepartmentCode": value["DepartmentCode"],
                "DepartmentName": value["DepartmentName"],
                "LegacyCodes": ";".join(sorted(value["LegacyCodes"])),
                "SourceSheets": ";".join(sorted(value["SourceSheets"])),
            }
        )

    if len(departments) > len(SYNTHETIC_NAMES):
        raise ValueError("Synthetic name catalog is smaller than the department catalog")
    users = []
    for index, department in enumerate(departments):
        code = department["DepartmentCode"]
        normalized_username_code = re.sub(r"[^a-z0-9]+", ".", code.casefold()).strip(".")
        users.append(
            {
                "DepartmentCode": code,
                "Username": f"vpp.{normalized_username_code}",
                "FullName": SYNTHETIC_NAMES[index],
                "Email": f"vpp.{normalized_username_code}@demo.gtas.local",
                "EmployeeCode": f"VPP{index + 1:03d}",
            }
        )

    catalog_rows = [
        {
            "CategoryCode": item.category_code,
            "CategoryName": item.category_name,
            "ItemCode": item.item_code,
            "ItemName": item.item_name,
            "UomCode": item.uom_code,
            "UomName": item.uom_name,
            "VatPercent": format(item.vat_rate * 100, "f"),
            "UnitPrice": int(item.unit_price),
            "SourceRows": ";".join(str(row) for row in item.source_rows),
            "Resolution": item.resolution,
        }
        for item in sorted(catalog.values(), key=lambda value: (value.category_code, match_key(value.item_name)))
    ]

    write_tsv(
        OUTPUT_DIR / "demo-catalog.tsv",
        ["CategoryCode", "CategoryName", "ItemCode", "ItemName", "UomCode", "UomName", "VatPercent", "UnitPrice", "SourceRows", "Resolution"],
        catalog_rows,
    )
    write_tsv(
        OUTPUT_DIR / "demo-departments.tsv",
        ["DepartmentCode", "DepartmentName", "LegacyCodes", "SourceSheets"],
        departments,
    )
    write_tsv(
        OUTPUT_DIR / "demo-users.tsv",
        ["DepartmentCode", "Username", "FullName", "Email", "EmployeeCode"],
        users,
    )
    write_tsv(
        OUTPUT_DIR / "demo-orders.tsv",
        ["DepartmentCode", "SourceMonth", "ItemCode", "Quantity", "UnitPrice", "SourceRowCount", "SourceRefs", "IsSynthesized"],
        normalized_orders,
    )

    source_sha256 = hashlib.sha256(source.read_bytes()).hexdigest()
    source_label = source.name
    try:
        source_label = source.relative_to(REPO_ROOT).as_posix()
    except ValueError:
        pass

    summary = {
        "source": source_label,
        "source_sha256": source_sha256,
        "source_sheet_count": len(workbook.sheetnames),
        "source_catalog_rows": sum(len(rows) for rows in raw_catalog.values()),
        "canonical_catalog_items": len(catalog_rows),
        "canonical_categories": len({row["CategoryCode"] for row in catalog_rows}),
        "canonical_uoms": len({row["UomCode"] for row in catalog_rows}),
        "vat_percent_distribution": dict(sorted(Counter(row["VatPercent"] for row in catalog_rows).items())),
        "duplicate_catalog_groups_resolved": duplicate_resolutions,
        "item_aliases_applied": [
            {"source_key": source_key, "target_item": target_name}
            for source_key, target_name in sorted(ITEM_ALIASES.items())
        ],
        "inferred_catalog_items": [
            {"item_name": item_name, "resolution": resolution, "unit_price": int(unit_price)}
            for _, item_name, _, _, unit_price, resolution in ADDITIONAL_ITEMS
        ],
        "canonical_departments": len(departments),
        "department_aliases": [
            {
                "department_code": row["DepartmentCode"],
                "legacy_codes": row["LegacyCodes"],
                "source_sheets": row["SourceSheets"],
            }
            for row in departments
            if row["LegacyCodes"] or ";" in row["SourceSheets"]
        ],
        "raw_positive_quantity_rows": len(raw_orders) + invalid_category_rows_removed,
        "source_valid_order_rows": len(raw_orders),
        "invalid_category_rows_removed": invalid_category_rows_removed,
        "normalized_order_lines": len(normalized_orders),
        "duplicate_order_rows_merged": len(raw_orders) + len(IT_BASKET_ITEMS) - len(normalized_orders),
        "source_months": sorted({row["SourceMonth"] for row in normalized_orders}),
        "skipped_blank_quantity_rows": skipped_blank_quantities,
        "personal_note_rows_removed": note_rows_removed,
        "synthetic_it_lines": len(IT_BASKET_ITEMS),
        "latest_source_month": LATEST_SOURCE_MONTH,
        "period_mapping": "rolling_12_months_ending_at_current_business_period",
    }
    (OUTPUT_DIR / "demo-source-audit.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return summary


if __name__ == "__main__":
    import argparse
    import sys

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    args = parser.parse_args()
    result = normalize(args.source.resolve())
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))
